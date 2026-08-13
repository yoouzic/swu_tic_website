import os
import tempfile
import unittest
from datetime import datetime

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash


TEST_DB_DIR = tempfile.TemporaryDirectory()
os.environ['SQLITE_DB_PATH'] = os.path.join(TEST_DB_DIR.name, 'ui_third_pass.db')
os.environ['INSTANCE_DIR'] = TEST_DB_DIR.name
os.environ['SECRET_KEY'] = 'test-secret-key'
os.environ['STORAGE_CLEANUP_ENABLED'] = '0'

from app.app import app
from app.models import LectureForm, User, db
from tests.app_test_utils import cleanup_sqlite_database, configure_sqlite_database


class ThirdPassRouteContractTest(unittest.TestCase):
    @classmethod
    def tearDownClass(cls):
        with app.app_context():
            cleanup_sqlite_database(db)
        TEST_DB_DIR.cleanup()

    def setUp(self):
        configure_sqlite_database(app, db, os.environ['SQLITE_DB_PATH'])
        app.config.update(TESTING=True)
        self.context = app.app_context()
        self.context.push()
        db.drop_all()
        db.create_all()
        self.client = app.test_client()

        self.super_admin = self._create_user('super', 'SA001', '中心管理员', '超级管理员', '中心')
        self.department_admin = self._create_user('department-admin', 'DA001', '甲部门管理员', '管理员', '甲部门')
        self.listener_a = self._create_user('listener-a', 'U001', '甲用户', '信息员', '甲部门')
        self.listener_b = self._create_user('listener-b', 'U002', '乙用户', '信息员', '乙部门')
        db.session.flush()
        self.pending = self._create_form(
            self.listener_a,
            status='待审核',
            lecture_date='2026-08-01',
            teacher_name='教师甲',
            teacher_college='学院甲',
            created_at=datetime(2026, 8, 1, 9, 0),
        )
        self.department_reviewed = self._create_form(
            self.listener_a,
            status='部门已审核',
            lecture_date='2026-08-02',
            teacher_name='教师乙',
            teacher_college='学院甲',
            created_at=datetime(2026, 8, 2, 9, 0),
        )
        self.rejected = self._create_form(
            self.listener_b,
            status='已驳回',
            lecture_date='2026-08-03',
            teacher_name='教师丙',
            teacher_college='学院乙',
            created_at=datetime(2026, 8, 3, 9, 0),
        )
        db.session.commit()
        self._login(self.super_admin)

    def tearDown(self):
        cleanup_sqlite_database(db, drop_all=True)
        self.context.pop()

    def _create_user(self, student_id, number, name, role, department):
        user = User(
            number=number,
            department=department,
            name=name,
            gender='-',
            grade='-',
            college='-',
            major='-',
            dormitory='-',
            phone='-',
            qq='-',
            student_id=student_id,
            password_hash=generate_password_hash('password'),
            role=role,
            group='一组',
            is_active=True,
        )
        db.session.add(user)
        return user

    def _create_form(self, listener, status, lecture_date, teacher_name, teacher_college, created_at):
        form = LectureForm(
            listener_name=listener.name,
            listener_number=listener.number,
            lecture_date=lecture_date,
            class_period='1-2节',
            lecture_location='教室A',
            teacher_name=teacher_name,
            teacher_college=teacher_college,
            course_title=f'{teacher_name}课程',
            student_grade_class='2024级',
            teaching_method='讲授',
            classroom_discipline='良好',
            classroom_atmosphere='良好',
            courseware_quality='良好',
            overall_effect='良好',
            quality_case='无',
            course_feedback='清晰',
            student_signature1='学生甲',
            contact_phone1='10000000000',
            status=status,
            created_at=created_at,
            updated_at=created_at,
        )
        db.session.add(form)
        db.session.flush()
        form.unique_id = form.id
        return form

    def _login(self, user):
        with self.client.session_transaction() as sess:
            sess['user_id'] = user.id
            sess['user_role'] = user.role
            sess['user_name'] = user.name

    def test_statistics_page_applies_filters_and_visualization_state(self):
        response = self.client.get(
            '/admin/statistics?department=甲部门&status=已审核&dimension=user'
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('甲用户'.encode('utf-8'), response.data)
        self.assertNotIn('乙用户'.encode('utf-8'), response.data)
        self.assertIn('暂不足以形成趋势'.encode('utf-8'), response.data)

        response = self.client.get('/admin/statistics?start_date=2026-09-01')
        self.assertEqual(response.status_code, 200)
        self.assertIn('暂无提交趋势数据'.encode('utf-8'), response.data)

    def test_statistics_detail_and_export_reuse_filter_scope(self):
        detail = self.client.get(
            '/admin/api/statistics_detail?department=甲部门&detail_dimension=user&detail_value=U001'
        )
        self.assertEqual(detail.status_code, 200)
        payload = detail.get_json()
        self.assertEqual(payload['pagination']['total'], 2)
        self.assertEqual({item['listener_department'] for item in payload['forms']}, {'甲部门'})

        exported = self.client.get(
            '/admin/export_statistics?department=甲部门&status=已审核&dimension=user'
        )
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(
            exported.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = load_workbook(filename=__import__('io').BytesIO(exported.data), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:2], ('用户', '部门'))
        self.assertEqual(rows[1][0:6], ('甲用户', '甲部门', 1, 0, 1, 0))

        self.listener_a.name = '=HYPERLINK("https://invalid.example")'
        db.session.commit()
        exported = self.client.get(
            '/admin/export_statistics?department=甲部门&status=已审核&dimension=user'
        )
        workbook = load_workbook(filename=__import__('io').BytesIO(exported.data), read_only=True)
        self.assertTrue(workbook.active['A2'].value.startswith("'="))

    def test_statistics_department_admin_cannot_expand_scope_with_query_param(self):
        self._login(self.department_admin)
        response = self.client.get('/admin/api/statistics_detail?department=乙部门')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload['pagination']['total'], 2)
        self.assertEqual({item['listener_department'] for item in payload['forms']}, {'甲部门'})

    def test_review_queue_serializes_server_authoritative_reviewability(self):
        newer = self._create_form(
            self.listener_a,
            status='部门已审核',
            lecture_date='2026-08-04',
            teacher_name='教师丁',
            teacher_college='学院甲',
            created_at=datetime(2026, 8, 4, 9, 0),
        )
        newer.unique_id = self.department_reviewed.unique_id
        db.session.commit()
        response = self.client.get('/admin/api/review/forms')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        by_id = {
            form['id']: form['can_review']
            for group in payload['forms']
            for form in group['forms']
        }
        self.assertFalse(by_id[self.pending.id])
        self.assertFalse(by_id[self.department_reviewed.id])
        self.assertTrue(by_id[newer.id])
        self.assertFalse(by_id[self.rejected.id])

        detail = self.client.get(f'/admin/api/review/form/{self.department_reviewed.id}')
        self.assertEqual(detail.status_code, 200)
        self.assertFalse(detail.get_json()['form']['can_review'])


if __name__ == '__main__':
    unittest.main()
