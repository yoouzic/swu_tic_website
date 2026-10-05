"""Administrator safety contracts for the no-current-schedule launch."""
import io
import os
import tempfile
import unittest
from pathlib import Path

import openpyxl
from werkzeug.datastructures import FileStorage
from werkzeug.security import generate_password_hash

from app.app import app
from app.blueprints.admin.forms_io import _parse_form_import_rows
from app.models import Department, Group, LectureForm, PasswordAuditLog, Permission, RolePermission, SystemSetting, User, db
from app.services import import_state
from tests.app_test_utils import cleanup_sqlite_database, configure_sqlite_database


CONTACT_COLUMNS = ['编号', '部门/组别', '姓名', '性别', '年级', '学院', '专业', '宿舍', '手机号码', 'QQ号码', '学号']
PERSISTENCE_EVIDENCE = []


def workbook_stream(book):
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


class LaunchAdminSafetyTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix='launch-admin-', dir=tempfile.gettempdir())
        self.root = Path(self.temp_dir.name)
        configure_sqlite_database(app, db, self.root / 'case.sqlite')
        app.config.update(TESTING=False, WTF_CSRF_ENABLED=False)
        self.ctx = app.app_context()
        self.ctx.push()
        db.drop_all()
        db.create_all()
        self.department = Department(name='审计部')
        self.group = Group(name='一组', department='审计部', max_members=10)
        db.session.add_all([self.department, self.group])
        self.super = self.make_user('0000', '999000', '超级管理员')
        self.member = self.make_user('0001', '999001')
        self.admin = self.make_user('0002', '999002', '管理员')
        db.session.commit()
        self.client = app.test_client()
        with self.client.session_transaction() as sess:
            sess.update(user_id=self.super.id, user_role=self.super.role, user_name=self.super.name)
        self.previous_env = {key: os.environ.get(key) for key in ('CONTACT_TEMPLATE_PATH', 'EXPORT_DIR')}
        os.environ['CONTACT_TEMPLATE_PATH'] = str(self.root / 'contacts.xlsx')
        os.environ['EXPORT_DIR'] = str(self.root / 'exports')
        book = openpyxl.Workbook()
        book.active.append(CONTACT_COLUMNS)
        book.save(os.environ['CONTACT_TEMPLATE_PATH'])

    def tearDown(self):
        cleanup_sqlite_database(db, drop_all=True)
        self.ctx.pop()
        for key, previous in self.previous_env.items():
            if previous is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = previous
        self.temp_dir.cleanup()

    def make_user(self, number, sid, role='信息员'):
        user = User(number=number, student_id=sid, role=role, name=f'人员{number}', gender='男', grade='2026',
                    department='审计部', group='一组', college='测试学院', major='测试专业', dormitory='测试宿舍',
                    phone='13800000000', qq='123456', password_hash=generate_password_hash('OldPass-2026'), is_active=True)
        db.session.add(user)
        db.session.flush()
        return user

    def row(self, number, sid, name='更新姓名'):
        return {'number': number, 'student_id': sid, 'name': name, 'department': '审计部', 'group': '一组',
                'role': '信息员', 'gender': '男', 'grade': '2026', 'college': '测试学院', 'major': '测试专业',
                'dormitory': '测试宿舍', 'phone': '13800000000', 'qq': '123456', 'has_error': False}

    def preview(self, rows):
        book = openpyxl.Workbook()
        book.active.append(CONTACT_COLUMNS)
        for row in rows:
            book.active.append([row['number'], row['department'] + '/' + row['group'], row['name'], row['gender'],
                                row['grade'], row['college'], row['major'], row['dormitory'], row['phone'], row['qq'], row['student_id']])
        return self.client.post('/admin/preview_import', data={'file': (workbook_stream(book), 'contacts.xlsx')}).get_json()

    def confirm(self, rows, overwrite=True):
        import_id = import_state.save_import_preview(rows, CONTACT_COLUMNS)
        return self.client.post('/admin/confirm_import', json={'import_id': import_id, 'overwrite': overwrite})

    def assert_invalid(self, response):
        self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
        self.assertIs(response.get_json()['success'], False)

    def test_preview_rejects_number_belonging_to_another_student(self):
        result = self.preview([self.row(self.super.number, '888888')])
        self.assertTrue(result['success'])
        self.assertEqual(result['valid_rows'], 0)
        self.assertTrue(result['preview'][0]['has_error'])
        self.assertIn('编号', ' '.join(result['errors']))

    def test_preview_rejects_existing_student_reusing_another_number(self):
        result = self.preview([self.row(self.super.number, self.member.student_id)])
        self.assertEqual(result['valid_rows'], 0)

    def test_confirm_rejects_cross_student_number_overwrite(self):
        before = (self.super.student_id, self.super.role, self.super.password_hash)
        self.assert_invalid(self.confirm([self.row(self.super.number, '888888')]))
        db.session.expire_all()
        self.assertEqual((self.super.student_id, self.super.role, self.super.password_hash), before)
        self.assertEqual(User.query.filter_by(role='超级管理员').count(), 1)
        PERSISTENCE_EVIDENCE.append({'case': 'contacts-number-conflict', 'status': 400, 'student_id': self.super.student_id,
                                     'role': self.super.role, 'password_unchanged': self.super.password_hash == before[2], 'super_count': 1})

    def test_confirm_rechecks_conflicts_after_preview_and_rolls_back_whole_batch(self):
        rows = [self.row('0010', '999010'), self.row('0011', '999011')]
        preview = self.preview(rows)
        self.assertEqual(preview['valid_rows'], 2)
        self.make_user('0011', '888888')
        db.session.commit()
        response = self.client.post('/admin/confirm_import', json={'import_id': preview['import_id'], 'overwrite': True})
        self.assert_invalid(response)
        self.assertIsNone(User.query.filter_by(student_id='999010').first())
        self.assertIsNotNone(import_state.load_import_preview(preview['import_id']))
        PERSISTENCE_EVIDENCE.append({'case': 'contacts-preview-stale-conflict', 'status': response.status_code,
                                     'first_batch_student_created': False, 'preview_retryable': True})

    def test_preview_flags_duplicate_identity_inside_upload(self):
        result = self.preview([self.row('0010', '999010'), self.row('0010', '999011')])
        self.assertEqual(result['valid_rows'], 1)
        self.assertTrue(result['preview'][1]['has_error'])

    def test_contacts_update_preserves_admin_role_password_and_custom_permissions(self):
        permission = Permission(name='管理部门', description='test')
        db.session.add(permission)
        db.session.flush()
        db.session.add(RolePermission(role=f'特殊角色_{self.admin.id}', permission_id=permission.id))
        db.session.commit()
        old_hash = self.admin.password_hash
        response = self.confirm([self.row(self.admin.number, self.admin.student_id)])
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()['updated_count'], 1)
        db.session.expire_all()
        self.assertEqual(self.admin.name, '更新姓名')
        self.assertEqual(self.admin.role, '管理员')
        self.assertEqual(self.admin.password_hash, old_hash)
        self.assertEqual(RolePermission.query.filter_by(role=f'特殊角色_{self.admin.id}').count(), 1)
        self.assertEqual(PasswordAuditLog.query.filter_by(target_user_id=self.admin.id).count(), 0)
        self.assertIsNone(response.get_json()['password_file_url'])
        PERSISTENCE_EVIDENCE.append({'case': 'contacts-admin-preserved', 'status': response.status_code,
                                     'name': self.admin.name, 'role': self.admin.role, 'password_unchanged': self.admin.password_hash == old_hash,
                                     'custom_permissions': 1, 'password_audits_added': 0, 'password_file_url': None})

    def test_contacts_update_preserves_only_super_and_session_access(self):
        old_hash = self.super.password_hash
        response = self.confirm([self.row(self.super.number, self.super.student_id)])
        self.assertEqual(response.status_code, 200)
        db.session.expire_all()
        self.assertEqual(self.super.role, '超级管理员')
        self.assertEqual(self.super.password_hash, old_hash)
        self.assertEqual(self.client.get('/admin/api/settings/teaching').status_code, 200)
        PERSISTENCE_EVIDENCE.append({'case': 'contacts-last-super-preserved', 'role': self.super.role,
                                     'password_unchanged': self.super.password_hash == old_hash, 'settings_access_status': 200})

    def test_confirm_rejects_string_overwrite_without_claiming_preview(self):
        rows = [self.row(self.member.number, self.member.student_id)]
        import_id = import_state.save_import_preview(rows, CONTACT_COLUMNS)
        response = self.client.post('/admin/confirm_import', json={'import_id': import_id, 'overwrite': 'false'})
        self.assert_invalid(response)
        self.assertIsNotNone(import_state.load_import_preview(import_id))

    def teaching(self):
        return {'first_week_monday': '2026-09-07', 'total_weeks': 16, 'required_submission_count': 1,
                'check_dept_review': False, 'check_center_review': False, 'show_auto_review_details': True,
                'enable_typos_check': True, 'course_weekly_limit_enabled': False}

    def test_teaching_settings_reject_non_boolean_flags_before_writes(self):
        for key in ('check_dept_review', 'check_center_review', 'show_auto_review_details', 'enable_typos_check', 'course_weekly_limit_enabled'):
            for invalid in ('false', 0, 1, None, []):
                with self.subTest(key=key, invalid=invalid):
                    data = self.teaching()
                    data[key] = invalid
                    self.assert_invalid(self.client.post('/admin/api/settings/teaching', json=data))
                    self.assertEqual(SystemSetting.query.count(), 0)

    def test_boolean_flags_keep_false_on_normal_settings_path(self):
        response = self.client.post('/admin/api/settings/teaching', json=self.teaching())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(SystemSetting.query.filter_by(key='teaching_check_dept_review').one().value, 'false')
        PERSISTENCE_EVIDENCE.append({'case': 'boolean-false-persisted', 'value': SystemSetting.query.filter_by(key='teaching_check_dept_review').one().value})

    def user_payload(self):
        return {key: getattr(self.member, key) for key in ('name', 'department', 'gender', 'grade', 'college', 'major', 'dormitory', 'phone', 'qq', 'student_id', 'role')}

    def test_user_create_rejects_unknown_role_without_creating_account(self):
        data = self.user_payload()
        data.update(student_id='888888', role='不存在的角色')
        self.assert_invalid(self.client.post('/admin/api/users', json=data))
        self.assertIsNone(User.query.filter_by(student_id='888888').first())

    def test_user_update_rejects_unknown_role(self):
        self.assert_invalid(self.client.put(f'/admin/api/users/{self.member.id}', json={'role': '不存在的角色'}))
        self.assertEqual(self.member.role, '信息员')

    def test_user_update_rejects_empty_or_non_text_name(self):
        old_name = self.member.name
        for name in ('', '   ', None, 123):
            with self.subTest(name=name):
                self.assert_invalid(self.client.put(f'/admin/api/users/{self.member.id}', json={'name': name}))
                self.assertEqual(self.member.name, old_name)

    def test_last_active_super_cannot_be_demoted(self):
        inactive = self.make_user('0003', '999003', '超级管理员')
        inactive.is_active = False
        db.session.commit()
        self.assert_invalid(self.client.put(f'/admin/api/users/{self.super.id}', json={'role': '管理员'}))
        self.assertEqual(self.super.role, '超级管理员')
        PERSISTENCE_EVIDENCE.append({'case': 'last-active-super-demotion-rejected', 'remaining_role': self.super.role,
                                     'inactive_super_does_not_satisfy_guard': True})

    def test_super_can_be_demoted_when_another_active_super_exists(self):
        self.make_user('0003', '999003', '超级管理员')
        db.session.commit()
        response = self.client.put(f'/admin/api/users/{self.super.id}', json={'role': '管理员'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.super.role, '管理员')

    def test_group_create_and_update_reject_invalid_capacities_without_mutation(self):
        for method, path in (('POST', '/admin/api/groups'), ('PUT', f'/admin/api/groups/{self.group.id}')):
            for capacity in (-1, 0, True, 1.5, '2.5', None):
                with self.subTest(method=method, capacity=capacity):
                    payload = {'name': '更新组', 'department': '审计部', 'max_members': capacity}
                    self.assert_invalid(self.client.open(path, method=method, json=payload))
                    self.assertEqual(Group.query.count(), 1)
                    self.assertEqual(self.group.max_members, 10)

    def test_group_positive_capacity_remains_supported(self):
        response = self.client.post('/admin/api/groups', json={'name': '二组', 'department': '审计部', 'max_members': 2})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Group.query.filter_by(name='二组').one().max_members, 2)

    def test_group_form_data_string_capacity_remains_supported(self):
        response = self.client.post('/admin/api/groups', json={'name': '二组', 'department': '审计部', 'max_members': '10'})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(Group.query.filter_by(name='二组').one().max_members, 10)

    def test_group_form_data_string_capacity_can_be_updated(self):
        response = self.client.put(f'/admin/api/groups/{self.group.id}', json={'name': '一组', 'department': '审计部', 'max_members': '2'})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(self.group.max_members, 2)

    def test_mutating_admin_apis_reject_non_object_json(self):
        routes = [
            ('POST', '/admin/api/departments'), ('PUT', f'/admin/api/departments/{self.department.id}'),
            ('POST', '/admin/api/groups'), ('PUT', f'/admin/api/groups/{self.group.id}'),
            ('POST', f'/admin/api/departments/{self.department.id}/disband'),
            ('POST', f'/admin/api/groups/{self.group.id}/disband'),
            ('POST', f'/admin/api/groups/{self.group.id}/move_members'),
            ('POST', f'/admin/api/departments/{self.department.id}/move_members'),
            ('POST', '/admin/api/users'), ('PUT', f'/admin/api/users/{self.member.id}'),
            ('PUT', f'/admin/api/users/{self.member.id}/permissions'),
            ('POST', '/admin/confirm_import'),
        ]
        for method, path in routes:
            for value in ([1], 'false', True, None):
                with self.subTest(method=method, path=path, value=value):
                    self.assert_invalid(self.client.open(path, method=method, data=__import__('json').dumps(value), content_type='application/json'))
        self.assertEqual(User.query.count(), 3)
        self.assertEqual(Department.query.count(), 1)
        self.assertEqual(Group.query.count(), 1)

    def test_system_form_template_preserves_padded_number_through_import(self):
        for number in ('0001', '1001'):
            with self.subTest(number=number):
                if number == '1001':
                    self.make_user(number, '999101')
                    db.session.commit()
                template = self.client.get('/admin/api/forms/import/template')
                self.assertEqual(template.status_code, 200)
                book = openpyxl.load_workbook(io.BytesIO(template.data))
                sheet = book.active
                sheet.cell(2, 1, number)
                sheet.cell(2, 4, '2026/09/21星期一')
                sheet.cell(2, 27, '2026-09-21 12:00:00')
                sheet.cell(2, 28, '2026-09-21 12:00:00')
                preview = self.client.post('/admin/api/forms/import/preview', data={'file': (workbook_stream(book), 'forms.xlsx')}).get_json()
                self.assertEqual(preview['summary']['valid_rows'], 1, preview)
                self.assertEqual(preview['rows'][0]['listener_number'], number)
                imported = self.client.post('/admin/api/forms/import', data={'file': (workbook_stream(book), 'forms.xlsx')}).get_json()
                self.assertEqual(imported['imported_count'], 1, imported)
                self.assertEqual(LectureForm.query.filter_by(listener_number=number).count(), 1)
                PERSISTENCE_EVIDENCE.append({'case': 'system-form-template-import', 'number': number,
                                             'preview_number': preview['rows'][0]['listener_number'], 'valid_rows': preview['summary']['valid_rows'],
                                             'imported_count': imported['imported_count'], 'persisted_count_for_number': 1})

    def test_legacy_import_preserves_text_padded_number(self):
        book = openpyxl.Workbook()
        values = [None] * 25
        values[2], values[7], values[10] = '0001', '教师', '课程'
        book.active.append(values)
        rows, error = _parse_form_import_rows(FileStorage(stream=workbook_stream(book), filename='legacy.xlsx'))
        self.assertIsNone(error)
        self.assertEqual(rows[0]['listener_number'], '0001')


if __name__ == '__main__':
    unittest.main()
