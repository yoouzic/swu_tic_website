# -*- coding: utf-8 -*-
"""Tests for the shared workbook infrastructure extracted in Step 1."""
import io
import unittest

from openpyxl import Workbook, load_workbook

from urllib.parse import unquote

from app.services.workbook import (
    autosize_worksheet,
    build_export_filename,
    workbook_response,
)


class BuildExportFilenameTests(unittest.TestCase):
    def test_chinese_filename_kept(self):
        self.assertEqual(build_export_filename('考评统计', '默认'), '考评统计.xlsx')

    def test_illegal_characters_replaced(self):
        self.assertEqual(
            build_export_filename('a<b>c:"d/e\\f|g?h*i', '默认'),
            'a_b_c__d_e_f_g_h_i.xlsx',
        )

    def test_empty_title_falls_back_to_default(self):
        self.assertEqual(build_export_filename('', '默认标题'), '默认标题.xlsx')
        self.assertEqual(build_export_filename('   ', '默认标题'), '默认标题.xlsx')

    def test_existing_xlsx_suffix_not_duplicated(self):
        self.assertEqual(build_export_filename('数据', '默认'), '数据.xlsx')
        self.assertEqual(build_export_filename('数据.xlsx', '默认'), '数据.xlsx')
        self.assertEqual(build_export_filename('数据.XLSX', '默认'), '数据.XLSX')


class AutosizeWorksheetTests(unittest.TestCase):
    def test_autosize_respects_min_and_max(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['short', 'x' * 100])
        ws.append(['a', 'b'])
        autosize_worksheet(ws, min_width=12, max_width=40)
        widths = {ws.column_dimensions[letter].width for letter in ('A', 'B')}
        self.assertGreaterEqual(min(widths), 12)
        self.assertLessEqual(max(widths), 40)

    def test_autosize_expands_longer_text_within_cap(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['hello world'])
        autosize_worksheet(ws, min_width=1, max_width=40)
        self.assertEqual(ws.column_dimensions['A'].width, 13)


class WorkbookResponseTests(unittest.TestCase):
    def test_response_mime_and_attachment_filename(self):
        from app.app import create_app
        app = create_app({'TESTING': True})

        @app.route('/_test_workbook_response')
        def _respond():
            wb = Workbook()
            ws = wb.active
            ws.append(['x'])
            return workbook_response(wb, '测试文件.xlsx')

        response = app.test_client().get('/_test_workbook_response')
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        disposition = response.headers['Content-Disposition']
        self.assertIn('attachment', disposition)
        self.assertIn('filename*=UTF-8', disposition)
        self.assertIn('测试文件.xlsx', unquote(disposition))

    def test_response_body_is_openable_workbook(self):
        from app.app import create_app
        app = create_app({'TESTING': True})

        @app.route('/_test_workbook_body')
        def _respond():
            wb = Workbook()
            ws = wb.active
            ws.append(['姓名'])
            ws.append(['张三'])
            return workbook_response(wb, '数据.xlsx')

        response = app.test_client().get('/_test_workbook_body')
        reloaded = load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertEqual(reloaded.active['A1'].value, '姓名')


if __name__ == '__main__':
    unittest.main()
