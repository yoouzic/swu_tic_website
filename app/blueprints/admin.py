from flask import Blueprint, render_template, request, redirect, url_for, flash, session, jsonify, send_file, current_app
from ..models import User, Department, Group, LectureForm, Permission, RolePermission, Teacher, Venue, Course, ListeningBan, CourseRegistration, db, SystemSetting, ScoreRecord, ScoreItem, StatisticsSnapshot, PersonnelMovementRecord, AssessmentOverride
from sqlalchemy import func
from datetime import datetime, timedelta
from ..utils.auto_review import AutoReviewEngine, SETTING_KEY_SEMESTER_MONDAY, SETTING_KEY_SCHEDULE_PATH, SETTING_KEY_CONTACTS_PATH, SETTING_KEY_FEEDBACK_PATH
from .auth import login_required, role_required
from ..utils.review_permissions import (
    get_user_review_permission, 
    can_review_status, 
    get_user_structure_for_review, 
    get_reviewable_users,
    get_reviewable_status_list,
    get_next_status_after_review,
    get_review_permission_presentation,
)
from ..utils.permission_feedback import build_forbidden_payload, forbidden_json, flash_forbidden
from ..utils.manage_permissions import (
    get_user_manage_permission,
    check_manage_permission
)
from ..utils.password_audit import record_password_audit
from ..utils.audit_tags import (
    REVIEW_TAG_OPTIONS,
    LATE_TAG_OPTIONS,
    LATE_TAG_LATE,
    REVIEW_TAG_REQUIRED,
    build_audit_tag,
    parse_audit_tag,
    validate_audit_tag,
    is_auto_review_allowed,
)
from ..utils.leave_management import (
    ASSESSMENT_EXEMPT_OVERRIDE_TYPES,
    LEAVE_OVERRIDE_TYPE,
    build_leave_status_payload,
    get_current_teaching_week as get_leave_current_teaching_week,
    get_form_effective_week_no as get_leave_form_effective_week_no,
    get_leave_makeup_forms,
    get_teaching_settings as get_leave_teaching_settings,
    parse_lecture_date_value,
    set_leave_makeup_forms,
)
from ..utils.user_status import (
    UNASSIGNED_DEPARTMENT_NAME,
    UNASSIGNED_GROUP_NAME,
    active_user_filter,
    is_user_active,
)
from ..utils.profile_settings import (
    PROFILE_EDITABLE_FIELD_OPTIONS,
    SETTING_KEY_PROFILE_EDITABLE_FIELDS,
    get_profile_editable_fields,
    normalize_profile_editable_fields,
)
from ..utils.course_registration_limits import (
    SETTING_KEY_COURSE_WEEKLY_LIMIT_COUNT,
    SETTING_KEY_COURSE_WEEKLY_LIMIT_ENABLED,
    get_course_weekly_limit_settings,
    normalize_course_weekly_limit_count,
)
from ..utils.review_drafts import (
    delete_review_form_draft,
    load_review_form_draft,
    normalize_review_draft_payload,
    parse_review_form_draft,
    save_review_form_draft,
)
from ..utils.env_config import env_path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
import os
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
import secrets
import string
from datetime import datetime
import json
import numpy as np
import re
from collections import defaultdict

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')

# 允许上传的文件扩展名
ALLOWED_EXTENSIONS = {'xlsx', 'xls'}
DEFAULT_CONTACT_TEMPLATE_PATH = os.path.join('data', 'storage', 'templates', 'contacts', '通讯录.xlsx')
DEFAULT_SCHEDULE_TEMPLATE_PATH = os.path.join('data', 'storage', 'templates', 'schedule', '全校课表.xls')
DEFAULT_EXPORT_DIR = os.path.join('data', 'storage', 'exports', 'contacts')
DEFAULT_AUTO_REVIEW_UPLOAD_DIR = os.path.join('data', 'storage', 'uploads', 'auto_review')
DEFAULT_AUTO_REVIEW_REPORT_DIR = os.path.join('data', 'storage', 'exports', 'auto_review')

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


@admin_bp.route('/download_template')
@role_required('超级管理员')
def download_template():
    """下载通讯录模板文件"""
    template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
    return send_file(template_path, as_attachment=True)


@admin_bp.route('/view_sample')
@role_required('超级管理员')
def view_sample():
    """在线查看模板示例"""
    template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
    return send_file(template_path, as_attachment=False)


@admin_bp.route('/download_passwords/<filename>')
@role_required('超级管理员')
def download_passwords(filename):
    """下载密码文件"""
    file_path = os.path.join(env_path('EXPORT_DIR', DEFAULT_EXPORT_DIR), filename)
    if os.path.exists(file_path):
        return send_file(file_path, as_attachment=True)
    else:
        flash('文件不存在', 'error')
        return redirect(url_for('admin.super_admin_dashboard'))


@admin_bp.route('/auto_review')
@role_required('超级管理员')
def auto_review_page():
    """自动审核旧入口，兼容书签并定位到设置中心。"""
    return redirect(url_for('admin.system_management', tab='automation'))


@admin_bp.route('/api/auto_review/settings', methods=['GET', 'POST'])
@role_required('超级管理员')
def auto_review_settings():
    """自动审核基础设置：获取/设置第一周星期一，并检查文件存在"""
    if request.method == 'GET':
        engine = AutoReviewEngine()
        return jsonify({'success': True, 'status': engine.files_status()})

    data = request.get_json() or {}
    monday = data.get('semester_monday')
    if not monday:
        return jsonify({'success': False, 'message': '缺少第一周星期一日期（YYYY-MM-DD）'}), 400
    try:
        datetime.strptime(monday, '%Y-%m-%d')
    except Exception:
        return jsonify({'success': False, 'message': '日期格式错误，应为YYYY-MM-DD'}), 400
    SystemSetting.set(SETTING_KEY_SEMESTER_MONDAY, monday)
    engine = AutoReviewEngine()
    return jsonify({'success': True, 'status': engine.files_status()})


@admin_bp.route('/api/auto_review/upload', methods=['POST'])
@role_required('超级管理员')
def auto_review_upload():
    """上传/替换课表、通讯录或反馈文件，并保存路径到系统设置"""
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '未选择文件'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': '未选择文件'}), 400
    file_type = request.form.get('file_type') or request.form.get('type')
    if file_type not in ('schedule', 'contacts', 'feedback'):
        return jsonify({'success': False, 'message': '缺少或错误的文件类型（schedule/contacts/feedback）'}), 400
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'message': '文件格式不支持，请上传.xls或.xlsx文件'}), 400

    upload_dir = env_path('AUTO_REVIEW_UPLOAD_DIR', DEFAULT_AUTO_REVIEW_UPLOAD_DIR)
    os.makedirs(upload_dir, exist_ok=True)

    ext = file.filename.rsplit('.', 1)[1].lower()
    prefix = 'schedule' if file_type == 'schedule' else ('contacts' if file_type == 'contacts' else 'feedback')
    save_name = f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"
    save_path = os.path.join(upload_dir, save_name)
    file.save(save_path)

    if file_type == 'schedule':
        SystemSetting.set(SETTING_KEY_SCHEDULE_PATH, save_path)
    elif file_type == 'contacts':
        SystemSetting.set(SETTING_KEY_CONTACTS_PATH, save_path)
    else:
        SystemSetting.set(SETTING_KEY_FEEDBACK_PATH, save_path)

    engine = AutoReviewEngine()
    return jsonify({'success': True, 'message': '上传成功', 'path': save_path, 'status': engine.files_status()})


@admin_bp.route('/api/auto_review/feedback_run', methods=['POST'])
@role_required('超级管理员')
def auto_review_feedback_run():
    """执行反馈文件审核，并可导出报告"""
    data = request.get_json() or {}
    export = data.get('export', True)

    engine = AutoReviewEngine()
    result = engine.review_feedback()

    report_path = None
    if export:
        report_path = engine.export_report(result)

    download_url = None
    if report_path:
        basename = os.path.basename(report_path)
        download_url = url_for('admin.auto_review_download', filename=basename)

    return jsonify({'success': True, 'data': result, 'report_path': report_path, 'download_url': download_url})


@admin_bp.route('/auto_review/download')
@role_required('超级管理员')
def auto_review_download():
    """下载自动审核报告"""
    filename = request.args.get('filename')
    if not filename:
        return jsonify({'success': False, 'message': '缺少文件名'}), 400
    reports_dir = env_path('AUTO_REVIEW_REPORT_DIR', DEFAULT_AUTO_REVIEW_REPORT_DIR)
    real_path = os.path.join(reports_dir, os.path.basename(filename))
    if not os.path.exists(real_path):
        return jsonify({'success': False, 'message': '报告不存在'}), 404
    return send_file(real_path, as_attachment=True)


@admin_bp.route('/api/settings/teaching', methods=['GET'])
@role_required('超级管理员')
def get_teaching_settings():
    """获取听课制度设置"""
    try:
        settings = {
            'first_week_monday': SystemSetting.query.filter_by(key='teaching_first_week_monday').first(),
            'week_start_day': SystemSetting.query.filter_by(key='teaching_week_start_day').first(),
            'total_weeks': SystemSetting.query.filter_by(key='teaching_total_weeks').first(),
            'required_submission_count': SystemSetting.query.filter_by(key='teaching_required_submission').first(),
            'check_dept_review': SystemSetting.query.filter_by(key='teaching_check_dept_review').first(),
            'check_center_review': SystemSetting.query.filter_by(key='teaching_check_center_review').first(),
            'show_auto_review_details': SystemSetting.query.filter_by(
                key='teaching_show_auto_review_details'
            ).first(),
            'enable_typos_check': SystemSetting.query.filter_by(key='teaching_enable_typos_check').first(),
            'reviewer_display_mode': SystemSetting.query.filter_by(key='teaching_reviewer_display_mode').first(),
        }
        profile_editable_fields = get_profile_editable_fields()
        course_weekly_limit = get_course_weekly_limit_settings()

        data = {
            'first_week_monday': settings['first_week_monday'].value if settings['first_week_monday'] else None,
            'week_start_day': int(settings['week_start_day'].value) if settings['week_start_day'] else 0,
            'total_weeks': int(settings['total_weeks'].value) if settings['total_weeks'] else 20,
            'required_submission_count': (
                int(settings['required_submission_count'].value)
                if settings['required_submission_count']
                else 1
            ),
            'required_listening_count': (
                int(settings['required_submission_count'].value)
                if settings['required_submission_count']
                else 1
            ),
            'check_dept_review': settings['check_dept_review'].value == 'true'
            if settings['check_dept_review']
            else False,
            'check_center_review': settings['check_center_review'].value == 'true'
            if settings['check_center_review']
            else False,
            'show_auto_review_details': settings['show_auto_review_details'].value == 'true'
            if settings['show_auto_review_details']
            else True,
            'enable_typos_check': settings['enable_typos_check'].value == 'true'
            if settings['enable_typos_check']
            else True,
            'reviewer_display_mode': (
                settings['reviewer_display_mode'].value
                if settings['reviewer_display_mode']
                else 'name'
            ),
            'profile_editable_fields': profile_editable_fields,
            'profile_editable_field_options': PROFILE_EDITABLE_FIELD_OPTIONS,
            'course_weekly_limit_enabled': course_weekly_limit['enabled'],
            'course_weekly_limit_count': course_weekly_limit['limit_count'],
        }

        return jsonify({'success': True, 'data': data})
    except Exception as exc:
        return jsonify({'success': False, 'message': str(exc)}), 500


@admin_bp.route('/api/settings/teaching', methods=['POST'])
@role_required('超级管理员')
def update_teaching_settings():
    """更新听课制度设置"""
    try:
        data = request.get_json()
        required_fields = ['first_week_monday', 'total_weeks', 'required_submission_count']
        for field in required_fields:
            if field not in data:
                return jsonify({'success': False, 'message': f'缺少必填字段: {field}'}), 400

        settings_map = {
            'teaching_first_week_monday': str(data['first_week_monday']),
            'teaching_week_start_day': str(data.get('week_start_day', 0)),
            'teaching_total_weeks': str(data['total_weeks']),
            'teaching_required_submission': str(data['required_submission_count']),
            'teaching_check_dept_review': 'true' if data.get('check_dept_review') else 'false',
            'teaching_check_center_review': 'true' if data.get('check_center_review') else 'false',
            'teaching_show_auto_review_details': 'true'
            if data.get('show_auto_review_details')
            else 'false',
            'teaching_enable_typos_check': 'true' if data.get('enable_typos_check') else 'false',
            'teaching_reviewer_display_mode': data.get('reviewer_display_mode', 'name')
            if data.get('reviewer_display_mode', 'name') in ['name', 'number']
            else 'name',
            SETTING_KEY_COURSE_WEEKLY_LIMIT_ENABLED: 'true'
            if data.get('course_weekly_limit_enabled')
            else 'false',
            SETTING_KEY_COURSE_WEEKLY_LIMIT_COUNT: str(
                normalize_course_weekly_limit_count(data.get('course_weekly_limit_count'))
            ),
            SETTING_KEY_PROFILE_EDITABLE_FIELDS: json.dumps(
                normalize_profile_editable_fields(data.get('profile_editable_fields')),
                ensure_ascii=False,
            ),
        }

        for key, value in settings_map.items():
            setting = SystemSetting.query.filter_by(key=key).first()
            if not setting:
                setting = SystemSetting(key=key, value=value)
                db.session.add(setting)
            else:
                setting.value = value

        db.session.commit()
        return jsonify({'success': True, 'message': '设置已更新'})
    except Exception as exc:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(exc)}), 500

def generate_random_password(length=8):
    """生成随机密码"""
    characters = string.ascii_letters + string.digits
    return ''.join(secrets.choice(characters) for _ in range(length))

def get_reviewer_display_mode():
    mode = SystemSetting.get('teaching_reviewer_display_mode', 'name')
    if mode not in ['name', 'number']:
        mode = 'name'
    return mode


def _serialize_audit_tag(audit_tag):
    parsed = parse_audit_tag(audit_tag)
    return {
        'raw': parsed['raw'],
        'review_tag': parsed['review_tag'],
        'week_correction_tag': parsed.get('week_correction_tag'),
        'week_correction_week_no': parsed.get('week_correction_week_no'),
        'legacy_late_tag': parsed.get('legacy_late_tag'),
        # 兼容旧前端字段
        'late_tag': parsed.get('late_tag'),
        'review_tag_defined': parsed['review_tag'] in REVIEW_TAG_OPTIONS,
        'week_correction_defined': parsed.get('week_correction_week_no') is not None,
        'second_tag_defined': parsed.get('second_tag') is not None,
        'late_tag_defined': parsed['late_tag'] in LATE_TAG_OPTIONS if parsed.get('late_tag') else False,
        'has_late_tag': parsed.get('late_tag') is not None,
        'extra_tags': parsed['extra_tags'],
    }


def _load_review_forms_for_operation(form_ids):
    normalized_form_ids = []
    invalid_form_ids = []
    for form_id in form_ids or []:
        try:
            value = int(form_id)
        except (TypeError, ValueError):
            invalid_form_ids.append(form_id)
            continue
        if value not in normalized_form_ids:
            normalized_form_ids.append(value)

    if invalid_form_ids:
        return None, jsonify({'success': False, 'message': f'存在非法表单ID：{invalid_form_ids[:5]}'}), 400
    if not normalized_form_ids:
        return None, jsonify({'success': False, 'message': '请选择至少一个表单'}), 400

    forms = LectureForm.query.filter(LectureForm.id.in_(normalized_form_ids)).all()
    form_map = {form.id: form for form in forms}
    listener_numbers = {form.listener_number for form in forms if form.listener_number}
    users_by_number = {
        user.number: user
        for user in _active_user_query().filter(User.number.in_(listener_numbers)).all()
    } if listener_numbers else {}
    reviewable_user_ids = set(get_reviewable_users(session['user_id']))

    ordered_forms = []
    missing_ids = []
    unauthorized_ids = []
    for form_id in normalized_form_ids:
        form = form_map.get(form_id)
        if not form:
            missing_ids.append(form_id)
            continue
        form_user = users_by_number.get(form.listener_number)
        if not form_user or form_user.id not in reviewable_user_ids:
            unauthorized_ids.append(form_id)
            continue
        ordered_forms.append(form)

    if missing_ids:
        return None, jsonify({'success': False, 'message': f'以下表单不存在：{missing_ids[:5]}'}), 404
    if unauthorized_ids:
        return None, jsonify({'success': False, 'message': f'您无权操作以下表单：{unauthorized_ids[:5]}'}), 403
    return ordered_forms, None, None


def _snapshot_user_for_movement(user):
    return {
        'id': user.id,
        'number': user.number,
        'name': user.name,
        'department': user.department,
        'group': user.group,
        'group_id': user.group_id,
        'gender': user.gender,
        'grade': user.grade,
        'college': user.college,
        'major': user.major,
        'dormitory': user.dormitory,
        'phone': user.phone,
        'qq': user.qq,
        'student_id': user.student_id,
        'role': user.role,
        'is_active': user.is_active,
    }


def _create_personnel_movement_record(operator_user_id, target_snapshot, action_type, summary, changes=None):
    department_name = (
        target_snapshot.get('department_after')
        or target_snapshot.get('department')
        or target_snapshot.get('department_before')
    )
    group_name = (
        target_snapshot.get('group_after')
        or target_snapshot.get('group')
        or target_snapshot.get('group_before')
    )
    record = PersonnelMovementRecord(
        target_user_id=target_snapshot.get('id'),
        target_user_number=target_snapshot.get('number'),
        target_user_name=target_snapshot.get('name') or '未知用户',
        department_name=department_name,
        group_name=group_name,
        action_type=action_type,
        summary=summary,
        details_json=json.dumps({'changes': changes or []}, ensure_ascii=False) if changes else None,
        operator_user_id=operator_user_id
    )
    db.session.add(record)
    return record


def _active_user_query():
    return User.query.filter(active_user_filter())


@admin_bp.route('/api/review/form/<int:form_id>/draft', methods=['GET', 'PUT', 'DELETE'])
@login_required
def review_form_draft(form_id):
    """Save, load, or delete the current reviewer's draft for one form."""
    try:
        user_id = session['user_id']
        permission = get_user_review_permission(user_id)
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限，如有疑问，请联系管理员'}), 403

        form = LectureForm.query.get_or_404(form_id)
        reviewable_user_ids = get_reviewable_users(user_id)
        form_user = _active_user_query().filter_by(number=form.listener_number).first()
        if not form_user or form_user.id not in reviewable_user_ids:
            return jsonify({'success': False, 'message': '您没有权限审核此表单'}), 403

        if request.method == 'GET':
            draft = load_review_form_draft(user_id, form_id)
            return jsonify({
                'success': True,
                'exists': draft is not None,
                'data': parse_review_form_draft(draft),
                'updated_at': draft.updated_at.strftime('%Y-%m-%d %H:%M:%S') if draft else None,
            })

        if request.method == 'DELETE':
            deleted = delete_review_form_draft(user_id, form_id)
            db.session.commit()
            return jsonify({'success': True, 'deleted': deleted})

        body = request.get_json(silent=True) or {}
        raw_payload = body.get('data', body)
        payload = normalize_review_draft_payload(raw_payload)
        if payload is None:
            return jsonify({'success': False, 'message': 'Draft payload must be a JSON object'}), 400

        draft = save_review_form_draft(user_id, form_id, payload)
        db.session.commit()
        return jsonify({
            'success': True,
            'updated_at': draft.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500


def _serialize_user_basic(user):
    return {
        'id': user.id,
        'number': user.number,
        'name': user.name,
        'gender': user.gender,
        'grade': user.grade,
        'college': user.college,
        'major': user.major,
        'dormitory': user.dormitory,
        'phone': user.phone,
        'qq': user.qq,
        'student_id': user.student_id,
        'role': user.role,
        'department': user.department,
        'group': user.group,
        'group_id': user.group_id,
        'is_active': user.is_active is not False,
    }


def _set_user_unassigned(user):
    user.department = UNASSIGNED_DEPARTMENT_NAME
    user.group_id = None
    user.group = UNASSIGNED_GROUP_NAME


def _build_group_payload(group, users, *, name=None, department=None, description=None, max_members=None, is_virtual=False):
    group_name = name if name is not None else group.name
    group_department = department if department is not None else group.department
    payload = {
        'id': None if is_virtual else group.id,
        'name': group_name,
        'department': group_department,
        'leader': None if is_virtual else group.leader,
        'description': description if description is not None else (None if is_virtual else group.description),
        'max_members': max_members if max_members is not None else ('无限制' if is_virtual else group.max_members),
        'member_count': len(users),
        'current_members': len(users),
        'members': [],
        'users': [],
        'is_virtual': is_virtual,
    }
    for user in users:
        user_data = _serialize_user_basic(user)
        payload['members'].append(user_data)
        payload['users'].append(user_data)
    return payload


def _build_review_form_filter_datetime(group_data, time_filter_type):
    latest_form = (group_data or {}).get('latest_form')
    if not latest_form:
        return None
    if time_filter_type == 'lecture':
        lecture_date = _parse_lecture_date_value(latest_form.lecture_date)
        if not lecture_date:
            return None
        return datetime.combine(lecture_date, datetime.min.time())
    if time_filter_type == 'updated':
        return _get_form_latest_timestamp(latest_form)

    created_candidates = [
        form.created_at for form in ((group_data or {}).get('forms') or [])
        if getattr(form, 'created_at', None)
    ]
    if created_candidates:
        return min(created_candidates)
    return latest_form.created_at


def _normalize_review_form_time_filter(time_filter_type):
    if time_filter_type in ['lecture', 'updated', 'created']:
        return time_filter_type
    return 'created'


def _get_form_latest_timestamp(form):
    if not form:
        return None
    return form.updated_at or form.created_at


def _load_review_form_groups_for_definition(form_ids):
    requested_forms, error_response, status_code = _load_review_forms_for_operation(form_ids)
    if error_response:
        return None, error_response, status_code

    ordered_unique_ids = []
    requested_form_ids_by_unique_id = {}
    for form in requested_forms:
        unique_id = form.unique_id or form.id
        if unique_id not in ordered_unique_ids:
            ordered_unique_ids.append(unique_id)
        requested_form_ids_by_unique_id.setdefault(unique_id, []).append(form.id)

    if not ordered_unique_ids:
        return [], None, None

    group_forms = LectureForm.query.filter(
        (LectureForm.unique_id.in_(ordered_unique_ids)) | (LectureForm.id.in_(ordered_unique_ids))
    ).all()
    grouped_forms = {}
    for form in group_forms:
        unique_id = form.unique_id or form.id
        grouped_forms.setdefault(unique_id, []).append(form)

    entries = []
    for unique_id in ordered_unique_ids:
        form_list = grouped_forms.get(unique_id, [])
        if not form_list:
            continue
        sorted_forms = sorted(
            form_list,
            key=lambda form: (_get_form_latest_timestamp(form) or datetime.min, form.id),
            reverse=True
        )
        entries.append({
            'unique_id': unique_id,
            'latest_form': sorted_forms[0],
            'forms': sorted_forms,
            'requested_form_ids': requested_form_ids_by_unique_id.get(unique_id, []),
        })
    return entries, None, None

def _can_view_managed_user(current_user, target_user, manage_permission):
    if not current_user or not target_user or not manage_permission:
        return False
    if manage_permission == '超级管理员':
        return True
    if target_user.department != current_user.department:
        return False
    if manage_permission == '管理部门':
        return target_user.role != '超级管理员'
    if manage_permission == '管理部门小组':
        current_group = current_user.group_id or current_user.group
        target_group = target_user.group_id or target_user.group
        return current_group == target_group
    return False

def _build_user_profile_stats(user):
    stats = None
    if not user or user.role not in ['信息员', '管理员', '超级管理员']:
        return stats

    first_week_setting = SystemSetting.query.filter_by(key='teaching_first_week_monday').first()
    if first_week_setting and first_week_setting.value:
        try:
            first_week_date = datetime.strptime(first_week_setting.value, '%Y-%m-%d').date()
        except ValueError:
            first_week_date = None
    else:
        first_week_date = None

    week_start_day_setting = SystemSetting.query.filter_by(key='teaching_week_start_day').first()
    week_start_day = int(week_start_day_setting.value) if week_start_day_setting and week_start_day_setting.value else 0

    required_submission_setting = SystemSetting.query.filter_by(key='teaching_required_submission').first()
    required_submission = int(required_submission_setting.value) if required_submission_setting and required_submission_setting.value else 1
    required_listening = required_submission

    check_dept_setting = SystemSetting.query.filter_by(key='teaching_check_dept_review').first()
    check_dept = check_dept_setting.value == 'true' if check_dept_setting else False

    check_center_setting = SystemSetting.query.filter_by(key='teaching_check_center_review').first()
    check_center = check_center_setting.value == 'true' if check_center_setting else False

    def get_week_num(date_obj):
        if not first_week_date or not date_obj:
            return -1
        current_date = date_obj.date() if isinstance(date_obj, datetime) else date_obj
        days_to_subtract = (first_week_date.weekday() - week_start_day) % 7
        actual_start_date = first_week_date - timedelta(days=days_to_subtract)
        diff = (current_date - actual_start_date).days
        if diff < 0:
            return -1
        return (diff // 7) + 1

    def get_form_group_week_num(versions):
        if not versions:
            return -1
        latest_version = versions[-1]
        parsed_tag = parse_audit_tag(latest_version.audit_tag)
        if parsed_tag.get('week_correction_week_no') is not None:
            return parsed_tag['week_correction_week_no']
        return get_week_num(parse_lecture_date_value(latest_version.lecture_date))

    def count_feedback_chars(form):
        if not form:
            return 0
        text = f"{form.course_feedback or ''}{form.suggestions or ''}"
        return len(''.join(str(text).split()))

    current_week_num = None
    current_week_label = '当前不在教学周内'
    if first_week_date:
        resolved_week_num = get_week_num(datetime.now().date())
        if resolved_week_num > 0:
            current_week_num = resolved_week_num
            current_week_label = f'当前教学周（第 {resolved_week_num} 周）'
        else:
            current_week_label = '本学期教学周尚未开始'

    all_forms = LectureForm.query.filter_by(listener_number=user.number).order_by(LectureForm.created_at.asc()).all()
    form_groups = {}
    for form in all_forms:
        uid = form.unique_id or form.id
        form_groups.setdefault(uid, []).append(form)

    total_submitted_count = len(form_groups)
    total_approved_count = 0
    total_reward_count = 0
    total_feedback_chars = 0
    week_submitted_count = 0
    week_approved_count = 0
    weekly_error_free_counts = {}

    for versions in form_groups.values():
        versions.sort(key=lambda item: item.id)
        submit_week = get_form_group_week_num(versions)

        if submit_week == current_week_num:
            week_submitted_count += 1

        approved_versions = [version for version in versions if version.status == '中心已审核']
        approved_form = approved_versions[-1] if approved_versions else None
        if approved_form:
            total_approved_count += 1
            total_feedback_chars += count_feedback_chars(approved_form)
            if submit_week == current_week_num:
                week_approved_count += 1

        has_dept_review = False
        dept_review_ok = True
        has_center_review = False
        center_review_ok = True

        for version in versions:
            if version.status == '部门已审核':
                has_dept_review = True
                score_record = ScoreRecord.query.filter_by(form_id=version.id).first()
                if score_record:
                    auto_items = ScoreItem.query.filter_by(score_record_id=score_record.id, is_auto_generated=True).count()
                    if auto_items > 0:
                        dept_review_ok = False

            if version.status == '中心已审核':
                has_center_review = True
                score_record = ScoreRecord.query.filter_by(form_id=version.id).first()
                if score_record:
                    auto_items = ScoreItem.query.filter_by(score_record_id=score_record.id, is_auto_generated=True).count()
                    if auto_items > 0:
                        center_review_ok = False

        criteria_met = True
        if check_dept and (not has_dept_review or not dept_review_ok):
            criteria_met = False
        if check_center and (not has_center_review or not center_review_ok):
            criteria_met = False
        if not check_dept and not check_center:
            criteria_met = False

        if criteria_met and submit_week > 0:
            weekly_error_free_counts[submit_week] = weekly_error_free_counts.get(submit_week, 0) + 1

    week_error_free_count = weekly_error_free_counts.get(current_week_num, 0)
    week_reward_count = max(0, week_error_free_count - required_listening)
    for count in weekly_error_free_counts.values():
        total_reward_count += max(0, count - required_listening)

    days_since_last = '无'
    last_form_obj = LectureForm.query.filter_by(listener_number=user.number).order_by(LectureForm.created_at.desc()).first()
    if last_form_obj and last_form_obj.created_at:
        days_since_last = (datetime.now() - last_form_obj.created_at).days

    this_month_start = datetime(datetime.now().year, datetime.now().month, 1)
    this_month = 0
    for versions in form_groups.values():
        if versions and versions[0].created_at and versions[0].created_at >= this_month_start:
            this_month += 1

    last_submit = last_form_obj.created_at.strftime('%Y-%m-%d') if last_form_obj and last_form_obj.created_at else '无'
    average_feedback_chars = total_feedback_chars / total_approved_count if total_approved_count else 0

    user_form_ids = [form.id for form in all_forms]
    if user_form_ids:
        total_deduction = db.session.query(func.sum(ScoreRecord.total_personal_score)).filter(
            ScoreRecord.form_id.in_(user_form_ids)
        ).scalar() or 0.0
    else:
        total_deduction = 0.0

    assessment_items = []
    if user_form_ids:
        item_rows = db.session.query(ScoreItem, ScoreRecord, LectureForm)\
            .join(ScoreRecord, ScoreItem.score_record_id == ScoreRecord.id)\
            .join(LectureForm, ScoreRecord.form_id == LectureForm.id)\
            .filter(
                LectureForm.listener_number == user.number,
                db.or_(ScoreItem.personal_score > 0, ScoreItem.department_score > 0)
            ).all()

        for score_item, score_record, form in item_rows:
            assessment_time = (
                score_record.updated_at
                or score_record.created_at
                or score_item.created_at
                or form.review_time
                or form.updated_at
                or form.created_at
            )
            personal_score = float(score_item.personal_score or 0.0)
            department_score = float(score_item.department_score or 0.0)
            assessment_items.append({
                'reason': score_item.reason or '未填写考评原因',
                'personal_score': personal_score,
                'department_score': department_score,
                'score_sort': personal_score,
                'assessment_time': assessment_time,
                'assessment_timestamp': assessment_time.timestamp() if assessment_time else 0,
                'assessment_time_text': assessment_time.strftime('%Y-%m-%d %H:%M') if assessment_time else '-',
                'course_title': form.course_title or '-',
                'teacher_name': form.teacher_name or '-',
                'form_id': form.id
            })

    assessment_items.sort(
        key=lambda item: (
            item['score_sort'],
            item['department_score'],
            item['assessment_timestamp'],
            item['form_id']
        ),
        reverse=True
    )

    results = db.session.query(
        LectureForm.listener_number,
        func.sum(ScoreRecord.total_personal_score).label('total')
    ).join(ScoreRecord, LectureForm.id == ScoreRecord.form_id).group_by(LectureForm.listener_number).all()

    scores_map = {row[0]: float(row[1] or 0.0) for row in results}
    all_scores = list(scores_map.values())
    all_scores.sort(reverse=True)

    has_evaluation_sample = user.number in scores_map
    user_score = scores_map.get(user.number)
    percentile = None
    if has_evaluation_sample and all_scores:
        try:
            rank_index = all_scores.index(user_score)
            percentile = (rank_index + 1) / len(all_scores)
        except ValueError:
            percentile = None

    if percentile is None:
        rating_label = '暂无评级'
    elif percentile <= 0.2:
        rating_label = '需要继续努力'
    elif percentile <= 0.7:
        rating_label = '表现良好'
    else:
        rating_label = '行为很好'

    numeric_deduction = float(total_deduction or 0.0)
    deduction_display = '0.00' if abs(numeric_deduction) < 0.005 else f'-{abs(numeric_deduction):.2f}'

    stats = {
        'total_forms': total_submitted_count,
        'this_month': this_month,
        'last_submit': last_submit,
        'total_deduction': total_deduction,
        'deduction_display': deduction_display,
        'has_evaluation_sample': has_evaluation_sample,
        'rating_label': rating_label,
        'percentile': percentile,
        'current_week': current_week_num,
        'current_week_label': current_week_label,
        'week_submitted': week_submitted_count,
        'week_approved': week_approved_count,
        'week_reward': week_reward_count,
        'total_submitted': total_submitted_count,
        'total_approved': total_approved_count,
        'total_reward': total_reward_count,
        'average_feedback_chars': average_feedback_chars,
        'total_feedback_chars': total_feedback_chars,
        'days_since_last': days_since_last,
        'assessment_items': assessment_items
    }
    return stats

def _build_user_form_groups(user, search='', date_from='', date_to=''):
    query = LectureForm.query.filter_by(listener_number=user.number)

    if search:
        query = query.filter(
            db.or_(
                LectureForm.course_title.contains(search),
                LectureForm.teacher_name.contains(search),
                LectureForm.lecture_location.contains(search)
            )
        )

    if date_from:
        query = query.filter(LectureForm.lecture_date >= date_from)
    if date_to:
        query = query.filter(LectureForm.lecture_date <= date_to)

    all_forms = query.order_by(LectureForm.updated_at.desc(), LectureForm.created_at.desc(), LectureForm.id.desc()).all()
    grouped_forms = {}
    for form in all_forms:
        unique_key = form.unique_id if form.unique_id else f"single_{form.id}"
        grouped_forms.setdefault(unique_key, []).append(form)

    form_groups = []
    for unique_id, versions in grouped_forms.items():
        versions.sort(key=lambda item: item.updated_at or item.created_at, reverse=True)
        form_groups.append({
            'unique_id': unique_id,
            'latest_form': versions[0],
            'versions': versions,
            'version_count': len(versions)
        })

    form_groups.sort(
        key=lambda item: item['latest_form'].updated_at or item['latest_form'].created_at,
        reverse=True
    )
    return form_groups

def _build_user_reservations(user, search='', date_from='', date_to=''):
    query = CourseRegistration.query.filter_by(user_id=user.id)

    if date_from:
        start_dt = datetime.strptime(date_from, '%Y-%m-%d')
        query = query.filter(CourseRegistration.created_at >= start_dt)
    if date_to:
        end_dt = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1)
        query = query.filter(CourseRegistration.created_at < end_dt)

    registrations = query.order_by(CourseRegistration.created_at.desc()).all()
    reservations = []
    normalized_search = (search or '').strip().lower()

    for registration in registrations:
        course = Course.query.filter_by(
            course_code=registration.course_code,
            selection_code=registration.selection_code
        ).first()
        bind_count = LectureForm.query.filter_by(registration_id=registration.id).count()
        item = {
            'id': registration.id,
            'course_code': registration.course_code,
            'selection_code': registration.selection_code,
            'course_name': course.course_name if course else '课程已删除',
            'teacher_name': course.teacher.name if course and course.teacher else '未知',
            'class_time': course.class_time if course else '',
            'class_location': course.class_location if course else '',
            'listening_info': registration.listening_info,
            'created_at': registration.created_at,
            'is_bound': bind_count > 0,
            'bind_count': bind_count,
            'can_edit': False,
            'can_delete': False
        }

        if normalized_search:
            haystack = ' '.join([
                item['course_name'] or '',
                item['teacher_name'] or '',
                item['course_code'] or '',
                item['selection_code'] or '',
                item['listening_info'] or ''
            ]).lower()
            if normalized_search not in haystack:
                continue

        reservations.append(item)

    return reservations

FORM_IMPORT_TEMPLATE_HEADERS = [
    '听课人编号', '听课人姓名+学院', '课程信息变化', '听课时间', '第几节', '听课地点',
    '授课教师', '教师所属学院', '课程名称', '专业年级', '异常情况反映', '主要教学方法',
    '管理课堂纪律', '调动课堂气氛', '课件制作质量', '整体教学效果', '优质案例推荐',
    '课程反馈', '不足及建议', '听课班级同学签名1', '联系电话1', '听课班级同学签名2', '联系电话2',
    '状态', '审核标签', '审核意见', '创建时间', '更新时间'
]

FORM_IMPORT_COLUMN_ALIASES = {
    'listener_number': ['听课人编号', '听课人(填写编号)', '编号', '信息员编号', '学号'],
    'listener_name': ['听课人姓名+学院', '听课人', '听课人姓名'],
    'course_changes': ['课程信息变化', '课程信息变化(一般三个方面:老师、教室、时间,没有变化填“无”)'],
    'lecture_date': ['听课时间', '听课时间(如:2023/10/19星期四)', '日期', '上课时间'],
    'class_period': ['第几节', '第几节(如:第1-3节)', '节次', '上课节次'],
    'lecture_location': ['听课地点', '听课地点(如:32-302)', '上课地点', '教室'],
    'teacher_name': ['授课教师', '授课教师(谨防错别字)', '教师', '教师姓名'],
    'teacher_college': ['教师所属学院', '教师所属学院(对照全校课表填写)', '教师学院', '学院'],
    'course_title': ['课程名称', '课程(总标题)', '课程', '课程名'],
    'student_grade_class': ['专业年级', '专业年级(如:2018级植物生产类05、06班)', '班级信息'],
    'abnormal_situation': ['异常情况反映', '异常情况反映(根据事实,没有则填"无")'],
    'teaching_method': ['主要教学方法'],
    'classroom_discipline': ['管理课堂纪律', '课堂纪律'],
    'classroom_atmosphere': ['调动课堂气氛', '课堂气氛'],
    'courseware_quality': ['课件制作质量', '课件质量'],
    'overall_effect': ['整体教学效果'],
    'quality_case': ['优质案例推荐', '是否推荐优质案例'],
    'course_feedback': ['课程反馈', '课程反馈(优点,五十字以上,评价的内容实在且有针对性,结尾不需要句号)'],
    'suggestions': ['不足及建议', '不足及建议(根据事实,没有则填"无")', '建议'],
    'student_signature1': ['听课班级同学签名1', '听课班级同学签名'],
    'contact_phone1': ['联系电话1', '联系电话'],
    'student_signature2': ['听课班级同学签名2'],
    'contact_phone2': ['联系电话2'],
    'status': ['status', '状态'],
    'audit_tag': ['audit_tag', '审核标签'],
    'review_comment': ['review_comment', '审核意见'],
    'created_at': ['created_at', '创建时间'],
    'updated_at': ['updated_at', '更新时间']
}

LEGACY_IMPORT_INDEX_MAP = {
    'listener_name': 1,
    'listener_number': 2,
    'course_changes': 3,
    'lecture_date': 4,
    'class_period': 5,
    'lecture_location': 6,
    'teacher_name': 7,
    'teacher_college': 8,
    'course_title': 9,
    'student_grade_class': 10,
    'abnormal_situation': 11,
    'teaching_method': 12,
    'classroom_discipline': 13,
    'classroom_atmosphere': 14,
    'courseware_quality': 15,
    'overall_effect': 16,
    'quality_case': 17,
    'course_feedback': 18,
    'suggestions': 19,
    'student_signature1': 21,
    'contact_phone1': 22,
    'student_signature2': 23,
    'contact_phone2': 24
}

def _excel_cell_to_text(value):
    if value is None or pd.isna(value):
        return ''
    text = str(value).strip()
    if text.lower() == 'nan':
        return ''
    return text

def _normalize_listener_number(value):
    text = _excel_cell_to_text(value).replace(' ', '')
    if text.endswith('.0'):
        candidate = text[:-2]
        if candidate.isdigit():
            return candidate
    return text

def _extract_named_form_rows(df):
    columns = [str(c).strip() for c in df.columns]
    alias_to_column = {}
    for key, aliases in FORM_IMPORT_COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in columns:
                alias_to_column[key] = alias
                break
    if not all(k in alias_to_column for k in ['listener_number', 'teacher_name', 'course_title']):
        return []
    rows = []
    for _, row in df.iterrows():
        if row.isna().all():
            continue
        item = {}
        for key, col_name in alias_to_column.items():
            item[key] = _excel_cell_to_text(row.get(col_name))
        if item.get('listener_number') or item.get('teacher_name') or item.get('course_title'):
            rows.append(item)
    return rows

def _extract_legacy_form_rows(df):
    rows = []
    for _, row in df.iterrows():
        item = {}
        has_content = False
        for key, idx in LEGACY_IMPORT_INDEX_MAP.items():
            value = _excel_cell_to_text(row[idx]) if idx < len(row) else ''
            item[key] = value
            if key in ['listener_number', 'teacher_name', 'course_title'] and value:
                has_content = True
        if has_content:
            rows.append(item)
    return rows

def _parse_form_import_rows(file_storage):
    try:
        df_named = pd.read_excel(file_storage, sheet_name=0)
    except Exception:
        return [], 'Excel读取失败，请检查文件格式'
    rows = _extract_named_form_rows(df_named)
    if rows:
        return rows, None
    try:
        file_storage.stream.seek(0)
        df_legacy = pd.read_excel(file_storage, sheet_name=0, header=None)
    except Exception:
        return [], 'Excel读取失败，请检查文件格式'
    legacy_rows = _extract_legacy_form_rows(df_legacy)
    if legacy_rows:
        return legacy_rows, None
    return [], '未识别到可导入的数据，请使用系统模板或检查表头'

def _to_int_or_none(value):
    text = _excel_cell_to_text(value)
    if not text:
        return None
    if text.endswith('.0'):
        text = text[:-2]
    try:
        return int(text)
    except Exception:
        return None

def _parse_optional_datetime(value, fallback):
    if value is None:
        return fallback
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    text = _excel_cell_to_text(value)
    if not text:
        return fallback
    normalized = text.replace('T', ' ').replace('/', '-').strip()
    formats = [
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%Y-%m-%d'
    ]
    for fmt in formats:
        try:
            return datetime.strptime(normalized, fmt)
        except Exception:
            continue
    try:
        return datetime.fromisoformat(normalized)
    except Exception:
        return fallback

def _prepare_import_form_row(row, row_no, importer_id, import_time):
    errors = []
    listener_number = _normalize_listener_number(row.get('listener_number', ''))
    if not listener_number:
        errors.append('听课人编号为空')
        return {'ok': False, 'errors': errors, 'row_no': row_no}
    listener_user = _active_user_query().filter_by(number=listener_number).first()
    if not listener_user:
        errors.append(f'听课人编号 {listener_number} 不存在')
        return {'ok': False, 'errors': errors, 'row_no': row_no, 'listener_number': listener_number}

    lecture_date = _excel_cell_to_text(row.get('lecture_date'))
    class_period = _excel_cell_to_text(row.get('class_period'))
    lecture_location = _excel_cell_to_text(row.get('lecture_location'))
    teacher_name = _excel_cell_to_text(row.get('teacher_name'))
    teacher_college = _excel_cell_to_text(row.get('teacher_college'))
    course_title = _excel_cell_to_text(row.get('course_title'))
    student_grade_class = _excel_cell_to_text(row.get('student_grade_class'))
    teaching_method = _excel_cell_to_text(row.get('teaching_method'))
    classroom_discipline = _excel_cell_to_text(row.get('classroom_discipline'))
    classroom_atmosphere = _excel_cell_to_text(row.get('classroom_atmosphere'))
    overall_effect = _excel_cell_to_text(row.get('overall_effect'))
    quality_case = _excel_cell_to_text(row.get('quality_case'))
    course_feedback = _excel_cell_to_text(row.get('course_feedback'))
    student_signature1 = _excel_cell_to_text(row.get('student_signature1'))
    contact_phone1 = _excel_cell_to_text(row.get('contact_phone1'))

    required_pairs = [
        ('听课时间', lecture_date),
        ('第几节', class_period),
        ('听课地点', lecture_location),
        ('授课教师', teacher_name),
        ('教师所属学院', teacher_college),
        ('课程名称', course_title),
        ('专业年级', student_grade_class),
        ('主要教学方法', teaching_method),
        ('管理课堂纪律', classroom_discipline),
        ('调动课堂气氛', classroom_atmosphere),
        ('整体教学效果', overall_effect),
        ('优质案例推荐', quality_case),
        ('课程反馈', course_feedback),
        ('听课班级同学签名1', student_signature1),
        ('联系电话1', contact_phone1)
    ]
    missing_fields = [name for name, value in required_pairs if not value]
    if missing_fields:
        errors.append(f'缺少必填字段 {",".join(missing_fields[:5])}')

    status = _excel_cell_to_text(row.get('status')) or '部门已审核'
    allowed_status = {'待审核', '部门已审核', '中心已审核', '已驳回'}
    if status not in allowed_status:
        errors.append('status 不合法')
    audit_tag = _excel_cell_to_text(row.get('audit_tag')) or REVIEW_TAG_REQUIRED
    errors.extend(validate_audit_tag(audit_tag))

    courseware_quality = _excel_cell_to_text(row.get('courseware_quality')) or '无'
    if 'PPT演示法' not in teaching_method and not _excel_cell_to_text(row.get('courseware_quality')):
        courseware_quality = '无'

    listener_name = _excel_cell_to_text(row.get('listener_name')) or f'{listener_user.name}（{listener_user.college}）'
    created_at = _parse_optional_datetime(row.get('created_at'), import_time)
    updated_at = _parse_optional_datetime(row.get('updated_at'), import_time)
    payload = {
        'listener_name': listener_name,
        'listener_number': listener_number,
        'course_changes': _excel_cell_to_text(row.get('course_changes')) or '无',
        'lecture_date': lecture_date,
        'class_period': class_period,
        'lecture_location': lecture_location,
        'teacher_name': teacher_name,
        'teacher_college': teacher_college,
        'course_title': course_title,
        'student_grade_class': student_grade_class,
        'abnormal_situation': _excel_cell_to_text(row.get('abnormal_situation')) or '无',
        'teaching_method': teaching_method,
        'classroom_discipline': classroom_discipline,
        'classroom_atmosphere': classroom_atmosphere,
        'courseware_quality': courseware_quality,
        'overall_effect': overall_effect,
        'quality_case': quality_case,
        'course_feedback': course_feedback,
        'suggestions': _excel_cell_to_text(row.get('suggestions')) or '无',
        'student_signature1': student_signature1,
        'contact_phone1': contact_phone1,
        'student_signature2': _excel_cell_to_text(row.get('student_signature2')),
        'contact_phone2': _excel_cell_to_text(row.get('contact_phone2')),
        'registration_id': None,
        'status': status,
        'audit_tag': audit_tag,
        'reviewer_id': importer_id,
        'review_time': import_time,
        'review_comment': _excel_cell_to_text(row.get('review_comment')),
        'created_at': created_at,
        'updated_at': updated_at
    }
    if errors:
        return {
            'ok': False,
            'errors': errors,
            'row_no': row_no,
            'listener_number': listener_number,
            'teacher_name': teacher_name,
            'course_title': course_title,
            'status': status
        }
    return {
        'ok': True,
        'payload': payload,
        'row_no': row_no,
        'listener_number': listener_number,
        'teacher_name': teacher_name,
        'course_title': course_title,
        'status': status
    }

@admin_bp.route('/super_admin_dashboard')
@role_required('超级管理员')
def super_admin_dashboard():
    """超级管理员仪表板"""
    return redirect(url_for('main.index'))

@admin_bp.route('/admin_dashboard')
@role_required('管理员')
def admin_dashboard():
    """管理员仪表板"""
    return redirect(url_for('main.index'))

@admin_bp.route('/manage_departments')
@role_required('管理员')
def manage_departments():
    """部门管理页面"""
    user = User.query.get(session['user_id'])
    
    # 检查用户是否有管理权限
    manage_permission = get_user_manage_permission(user.id)
    
    if not manage_permission:
        flash('您没有权限访问此页面', 'danger')
        if user.role == '管理员':
            return redirect(url_for('admin.admin_dashboard'))
        else:
            return redirect(url_for('admin.super_admin_dashboard'))
            
    return render_template('admin/manage_departments.html', 
                          manage_permission=manage_permission,
                          current_user_dept=user.department)

@admin_bp.route('/users/<int:user_id>/view')
@role_required('管理员')
def view_managed_user(user_id):
    current_user = User.query.get(session['user_id'])
    manage_permission = get_user_manage_permission(current_user.id)
    if not manage_permission:
        flash('您没有权限查看该用户', 'error')
        return redirect(url_for('admin.manage_departments'))

    target_user = User.query.get_or_404(user_id)
    if not is_user_active(target_user):
        flash('该用户已离任', 'error')
        return redirect(url_for('admin.manage_departments'))
    if not _can_view_managed_user(current_user, target_user, manage_permission):
        flash('您没有权限查看该用户', 'error')
        return redirect(url_for('admin.manage_departments'))

    search = request.args.get('search', '').strip()
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()

    detail_context = {
        'stats': _build_user_profile_stats(target_user),
        'forms': _build_user_form_groups(target_user, search=search, date_from=date_from, date_to=date_to),
        'my_reservations': _build_user_reservations(target_user, search=search, date_from=date_from, date_to=date_to),
        'search': search,
        'date_from': date_from,
        'date_to': date_to,
    }
    if request.args.get('format') == 'fragment':
        return render_template(
            'admin/_user_detail_panel.html',
            user=target_user,
            target_user=target_user,
            current_user=current_user,
            is_fragment=True,
            **detail_context
        )
    return render_template(
        'admin/user_detail.html',
        user=target_user,
        target_user=target_user,
        current_user=current_user,
        is_fragment=False,
        **detail_context
    )

@admin_bp.route('/preview_import', methods=['POST'])
@role_required('超级管理员')
def preview_import():
    """预览通讯录并进行格式校验（支持 xls/xlsx，严格对标模板列）"""
    try:
        from flask import current_app
        current_app.logger.info('通讯录预览开始')
        if 'file' not in request.files:
            return jsonify({'success': False, 'message': '未收到文件'}), 400
        file = request.files['file']
        if file.filename == '':
            return jsonify({'success': False, 'message': '请选择文件'}), 400
        if not allowed_file(file.filename):
            return jsonify({'success': False, 'message': '文件格式不支持，仅支持xls/xlsx'}), 400

        ext = file.filename.rsplit('.', 1)[1].lower()
        engine = 'openpyxl' if ext == 'xlsx' else 'xlrd'
        # 以字符串读入，保留学号/手机号等前导零
        df = pd.read_excel(file, dtype=str, engine=engine, keep_default_na=False)

        # 加载模板列
        template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
        template_df = pd.read_excel(template_path, dtype=str, engine='openpyxl', keep_default_na=False)
        expected_cols = [c.strip() for c in template_df.columns.tolist()]
        incoming_cols = [c.strip() for c in df.columns.tolist()]
        if incoming_cols != expected_cols:
            return jsonify({'success': False, 'message': 'Excel列不匹配，请使用模板列名与顺序'}), 400

        errors = []
        preview_rows = []
        valid_rows = 0

        for i, row in df.iterrows():
            rec = {col: str(row[col]).strip() for col in expected_cols}
            dep_group = rec['部门/组别']
            if '/' in dep_group:
                department = dep_group.split('/', 1)[0]
                group_name = dep_group.split('/', 1)[1]
            else:
                department = dep_group
                group_name = UNASSIGNED_GROUP_NAME

            row_errors = []
            # 基本校验
            if not rec['编号'] or not rec['编号'].isdigit():
                row_errors.append('编号必须为纯数字')
            if not rec['姓名']:
                row_errors.append('姓名不能为空')
            if rec['性别'] not in ('男', '女'):
                row_errors.append('性别必须为男/女')
            grade = rec['年级']
            if not grade or not (grade.isdigit() or (grade.endswith('级') and grade[:-1].isdigit())):
                row_errors.append('年级格式不正确')
            phone = rec['手机号码']
            if phone and not (phone.isdigit() and len(phone) >= 6):
                row_errors.append('手机号码必须为数字')
            qq = rec['QQ号码']
            if qq and not (qq.isdigit() and 5 <= len(qq) <= 12):
                row_errors.append('QQ号码格式不正确')
            student_id = rec['学号']
            if not student_id or not student_id.isdigit():
                row_errors.append('学号必须为纯数字')

            # 重复检查（以学号判定）
            existing_user_by_sid = User.query.filter_by(student_id=student_id).first()
            is_duplicate = existing_user_by_sid is not None

            preview_rows.append({
                'row_number': i + 1,
                'number': rec['编号'],
                'department': department,
                'name': rec['姓名'],
                'gender': rec['性别'],
                'grade': rec['年级'],
                'college': rec['学院'],
                'major': rec['专业'],
                'dormitory': rec['宿舍'],
                'phone': phone,
                'qq': qq,
                'student_id': student_id,
                'role': '信息员',
                'group': group_name,
                'has_error': bool(row_errors),
                'errors': row_errors,
                'is_duplicate': is_duplicate
            })
            if not row_errors:
                valid_rows += 1
            errors.extend([f"第{i + 1}行: {e}" for e in row_errors])

        import_id = secrets.token_hex(8)
        cache = current_app.config.setdefault('IMPORT_CACHE', {})
        cache[import_id] = {
            'rows': preview_rows,
            'expected_cols': expected_cols,
        }

        return jsonify({
            'success': True,
            'import_id': import_id,
            'columns': expected_cols,
            'total_rows': len(preview_rows),
            'valid_rows': valid_rows,
            'error_rows': len(preview_rows) - valid_rows,
            'errors': errors,
            'preview': preview_rows[:50]
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'预览失败：{str(e)}'}), 500

@admin_bp.route('/confirm_import', methods=['POST'])
@role_required('超级管理员')
def confirm_import():
    """确认导入到数据库，生成随机密码并保存哈希，返回总用户数"""
    try:
        from flask import current_app
        current_app.logger.info('通讯录导入开始')
        data = request.get_json()
        import_id = data.get('import_id')
        overwrite = bool(data.get('overwrite', False))
        cache = current_app.config.get('IMPORT_CACHE', {})
        if not import_id or import_id not in cache:
            return jsonify({'success': False, 'message': '导入会话已失效，请重新预览'}), 400

        rows = cache[import_id]['rows']
        actor_user_id = session.get('user_id')
        imported_count = 0
        updated_count = 0
        skipped_count = 0
        password_list = []

        for r in rows:
            if r['has_error']:
                skipped_count += 1
                continue

            # 以学号判定唯一
            user = User.query.filter_by(student_id=r['student_id']).first()
            if user:
                if not overwrite:
                    skipped_count += 1
                    continue
                # 更新信息并重置密码
                user.number = r['number']
                user.department = r['department']
                user.name = r['name']
                user.gender = r['gender']
                user.grade = r['grade']
                user.college = r['college']
                user.major = r['major']
                user.dormitory = r['dormitory']
                user.phone = r['phone']
                user.qq = r['qq']
                user.role = r['role']
                user.group = r['group']
                # 保障部门/小组存在
                if user.department and not Department.query.filter_by(name=user.department).first():
                    db.session.add(Department(name=user.department))
                if user.group and user.department:
                    if not Group.query.filter_by(name=user.group, department=user.department).first() and user.group not in ['待分配', UNASSIGNED_GROUP_NAME]:
                        db.session.add(Group(name=user.group, department=user.department))
                password = generate_random_password()
                user.password_hash = generate_password_hash(password)
                record_password_audit(
                    actor_user_id=actor_user_id,
                    target_user_id=user.id,
                    action='import_overwrite_reset_password',
                    details={'match_type': 'student_id', 'overwrite': True}
                )
                updated_count += 1
                password_list.append({
                    'number': user.number,
                    'name': user.name,
                    'student_id': user.student_id,
                    'password': password
                })
            else:
                # 编号重复处理
                existed_by_number = User.query.filter_by(number=r['number']).first()
                if existed_by_number:
                    if not overwrite:
                        skipped_count += 1
                        continue
                    # 按编号更新该用户
                    user = existed_by_number
                    user.department = r['department']
                    user.name = r['name']
                    user.gender = r['gender']
                    user.grade = r['grade']
                    user.college = r['college']
                    user.major = r['major']
                    user.dormitory = r['dormitory']
                    user.phone = r['phone']
                    user.qq = r['qq']
                    user.student_id = r['student_id']
                    user.role = r['role']
                    user.group = r['group']
                    # 保障部门/小组存在
                    if user.department and not Department.query.filter_by(name=user.department).first():
                        db.session.add(Department(name=user.department))
                    if user.group and user.department:
                        if not Group.query.filter_by(name=user.group, department=user.department).first() and user.group not in ['待分配', UNASSIGNED_GROUP_NAME]:
                            db.session.add(Group(name=user.group, department=user.department))
                    password = generate_random_password()
                    user.password_hash = generate_password_hash(password)
                    record_password_audit(
                        actor_user_id=actor_user_id,
                        target_user_id=user.id,
                        action='import_overwrite_reset_password',
                        details={'match_type': 'number', 'overwrite': True}
                    )
                    updated_count += 1
                    password_list.append({
                        'number': user.number,
                        'name': user.name,
                        'student_id': user.student_id,
                        'password': password
                    })
                else:
                    # 新增用户
                    password = generate_random_password()
                    user = User(
                        number=r['number'],
                        department=r['department'],
                        name=r['name'],
                        gender=r['gender'],
                        grade=r['grade'],
                        college=r['college'],
                        major=r['major'],
                        dormitory=r['dormitory'],
                        phone=r['phone'],
                        qq=r['qq'],
                        student_id=r['student_id'],
                        password_hash=generate_password_hash(password),
                        role=r['role'],
                        group=r['group']
                    )
                    # 保障部门/小组存在
                    if user.department and not Department.query.filter_by(name=user.department).first():
                        db.session.add(Department(name=user.department))
                    if user.group and user.department:
                        if not Group.query.filter_by(name=user.group, department=user.department).first() and user.group not in ['待分配', UNASSIGNED_GROUP_NAME]:
                            db.session.add(Group(name=user.group, department=user.department))
                    db.session.add(user)
                    db.session.flush()
                    record_password_audit(
                        actor_user_id=actor_user_id,
                        target_user_id=user.id,
                        action='import_create_user_password_init',
                        details={'overwrite': overwrite}
                    )
                    imported_count += 1
                    password_list.append({
                        'number': user.number,
                        'name': user.name,
                        'student_id': user.student_id,
                        'password': password
                    })

        db.session.commit()
        total_users = _active_user_query().count()

        export_filename = None
        if password_list:
            password_df = pd.DataFrame(password_list)
            export_filename = f'passwords_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
            export_dir = env_path('EXPORT_DIR', DEFAULT_EXPORT_DIR)
            os.makedirs(export_dir, exist_ok=True)
            export_path = os.path.join(export_dir, export_filename)
            password_df.to_excel(export_path, index=False)

        # 清理会话缓存
        cache.pop(import_id, None)

        current_app.logger.info(f'通讯录导入完成 新增{imported_count} 更新{updated_count} 跳过{skipped_count}')
        return jsonify({
            'success': True,
            'message': f'导入完成：新增 {imported_count}，更新 {updated_count}，跳过 {skipped_count}',
            'imported_count': imported_count,
            'updated_count': updated_count,
            'skipped_count': skipped_count,
            'total_users': total_users,
            'password_file_url': (url_for('admin.download_passwords', filename=export_filename) if export_filename else None)
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'导入失败：{str(e)}'}), 500

@admin_bp.route('/manage_groups')
@role_required('管理员')
def manage_groups():
    """小组管理"""
    user = User.query.get(session['user_id'])
    manage_permission = get_user_manage_permission(user.id)
    
    # 检查用户是否有管理权限（管理部门或管理部门小组）
    if not manage_permission and manage_permission != '超级管理员':
         flash('您没有权限访问此页面', 'danger')
         return redirect(url_for('admin.admin_dashboard'))
    
    # 确定要管理的小组范围
    if manage_permission == '超级管理员':
        groups = Group.query.all()
    elif manage_permission == '管理部门':
        # 管理员只能管理本部门的小组
        groups = Group.query.filter_by(department=user.department).all()
        # 如果通过department name关联
        if not groups and user.department:
            dept = Department.query.filter_by(name=user.department).first()
            if dept:
                # Group表使用department名称关联，不是ID
                groups = Group.query.filter_by(department=dept.name).all()
    elif manage_permission == '管理部门小组':
        # 只能管理本小组
        if user.group and user.department:
            groups = Group.query.filter_by(department=user.department, name=user.group).all()
        else:
            groups = []
    else:
        groups = []
            
    return render_template('admin/manage_groups.html', groups=groups, user=user, manage_permission=manage_permission)

@admin_bp.route('/view_forms')
@login_required
def view_forms():
    """查看听课表单 - 按unique_id分组显示"""
    user = User.query.get(session['user_id'])
    page = request.args.get('page', 1, type=int)
    per_page = 20
    
    # 搜索参数
    search = request.args.get('search', '')
    date_filter = request.args.get('date', '')
    status_filter = request.args.get('status', '')
    
    # 基础查询
    if user.role == '超级管理员':
        # 超级管理员可以查看所有表单
        base_query = LectureForm.query.join(User, LectureForm.listener_number == User.number).filter(active_user_filter())
    elif user.role == '管理员':
        # 管理员只能查看本部门的表单
        base_query = LectureForm.query.join(User, LectureForm.listener_number == User.number)\
                                     .filter(User.department == user.department, active_user_filter())
    else:
        # 信息员只能查看自己的表单
        base_query = LectureForm.query.filter_by(listener_number=user.number)
    
    # 应用搜索过滤
    if search:
        base_query = base_query.filter(
            db.or_(
                LectureForm.listener_name.contains(search),
                LectureForm.teacher_name.contains(search),
                LectureForm.course_title.contains(search)
            )
        )
    
    if date_filter:
        base_query = base_query.filter(LectureForm.lecture_date.contains(date_filter))
    
    if status_filter:
        base_query = base_query.filter(LectureForm.status == status_filter)
    
    # 获取所有符合条件的表单，按unique_id分组
    all_forms = base_query.order_by(LectureForm.unique_id, LectureForm.updated_at.desc()).all()
    
    # 按unique_id分组，每组只取最新的表单作为代表
    form_groups = {}
    for form in all_forms:
        unique_id = form.unique_id or f"single_{form.id}"  # 处理没有unique_id的旧表单
        if unique_id not in form_groups:
            form_groups[unique_id] = {
                'latest_form': form,
                'versions': []
            }
        form_groups[unique_id]['versions'].append(form)
    
    # 转换为列表并按创建时间排序（使用组内第一个版本的创建时间）
    grouped_forms = list(form_groups.values())
    
    # 获取每个组的最早版本时间作为排序依据
    for group in grouped_forms:
        # 在 form_groups 构建时，versions 是按 updated_at desc 排列的
        # 所以 versions[-1] 是最早的版本（如果按时间倒序）
        # 或者更稳妥地，重新排序 versions 找到最早的 created_at
        versions = group['versions']
        if versions:
            # 找到最早的创建时间
            earliest_time = min(v.created_at for v in versions if v.created_at)
            group['sort_time'] = earliest_time
        else:
            group['sort_time'] = datetime.now() # Fallback

    # 按最早创建时间倒序排序
    grouped_forms.sort(key=lambda x: x['sort_time'], reverse=True)
    
    # 手动实现分页
    total = len(grouped_forms)
    start = (page - 1) * per_page
    end = start + per_page
    page_forms = grouped_forms[start:end]
    
    # 创建分页对象
    class Pagination:
        def __init__(self, page, per_page, total, items):
            self.page = page
            self.per_page = per_page
            self.total = total
            self.items = items
            self.pages = (total + per_page - 1) // per_page
            self.has_prev = page > 1
            self.has_next = page < self.pages
            self.prev_num = page - 1 if self.has_prev else None
            self.next_num = page + 1 if self.has_next else None
        
        def iter_pages(self, left_edge=2, left_current=2, right_current=3, right_edge=2):
            last = self.pages
            for num in range(1, last + 1):
                if num <= left_edge or \
                   (self.page - left_current - 1 < num < self.page + right_current) or \
                   num > last - right_edge:
                    yield num
    
    forms = Pagination(page, per_page, total, page_forms)
    
    reviewer_display_mode = get_reviewer_display_mode()
    return render_template('admin/view_forms.html', forms=forms, user=user, reviewer_display_mode=reviewer_display_mode)

@admin_bp.route('/import_forms_excel')
@role_required('超级管理员')
def import_forms_excel_page():
    return render_template('admin/import_forms_excel.html')

@admin_bp.route('/api/forms/import/template', methods=['GET'])
@role_required('超级管理员')
def download_forms_import_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表单导入模板'
    ws.append(FORM_IMPORT_TEMPLATE_HEADERS)
    ws.append([
        '20230001', '张三（计算机与信息科学学院）', '无', '2026/03/15星期三', '第3-4节', '32-302',
        '李老师', '计算机与信息科学学院', '数据结构', '2023级计算机1班', '无', 'PPT演示法；案例教学法',
        '秩序良好', '互动积极', '图文清晰', '整体较好', '可推荐', '课堂目标明确，内容组织清晰，学生参与度高。',
        '建议增加课堂练习时间', '王同学', '13800000000', '', '',
        '部门已审核', '需要人工审核', '', '2026-03-16 10:00:00', '2026-03-16 10:05:00'
    ])
    for col_idx in range(1, len(FORM_IMPORT_TEMPLATE_HEADERS) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name='听课表单导入模板.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@admin_bp.route('/api/forms/import', methods=['POST'])
@role_required('超级管理员')
def import_forms_from_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'message': '仅支持xls或xlsx文件'}), 400

    rows, err = _parse_form_import_rows(file)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    if not rows:
        return jsonify({'success': False, 'message': '未读取到可导入的数据'}), 400

    operator_id = session.get('user_id')
    import_time = datetime.now()
    imported_count = 0
    skipped_count = 0
    skipped_details = []
    max_detail = 15
    imported_form_ids = []

    for idx, row in enumerate(rows, start=1):
        prepared = _prepare_import_form_row(row, idx, operator_id, import_time)
        if not prepared['ok']:
            skipped_count += 1
            if len(skipped_details) < max_detail:
                skipped_details.append(f"第{idx}行：{'；'.join(prepared['errors'][:2])}")
            continue
        payload = prepared['payload']
        form = LectureForm(**payload)
        try:
            db.session.add(form)
            db.session.flush()
            imported_updated_at = payload.get('updated_at') or import_time
            LectureForm.query.filter_by(id=form.id).update({
                'unique_id': form.id,
                'updated_at': imported_updated_at
            }, synchronize_session=False)
            db.session.commit()
            imported_count += 1
            imported_form_ids.append(form.id)
        except Exception as e:
            db.session.rollback()
            skipped_count += 1
            if len(skipped_details) < max_detail:
                skipped_details.append(f'第{idx}行：写入失败 {str(e)}')

    export_url = None
    if imported_form_ids:
        export_url = url_for('admin.export_imported_forms_excel', form_ids=','.join(str(fid) for fid in imported_form_ids))

    return jsonify({
        'success': True,
        'imported_count': imported_count,
        'skipped_count': skipped_count,
        'skipped_details': skipped_details,
        'imported_form_ids': imported_form_ids,
        'export_url': export_url
    })

@admin_bp.route('/api/forms/import/preview', methods=['POST'])
@role_required('超级管理员')
def preview_forms_from_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择预览文件'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'success': False, 'message': '请先选择预览文件'}), 400
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'message': '仅支持xls或xlsx文件'}), 400

    rows, err = _parse_form_import_rows(file)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    if not rows:
        return jsonify({'success': False, 'message': '未读取到可预览的数据'}), 400

    operator_id = session.get('user_id')
    import_time = datetime.now()
    preview_rows = []
    valid_rows = 0
    invalid_rows = 0
    for idx, row in enumerate(rows, start=1):
        prepared = _prepare_import_form_row(row, idx, operator_id, import_time)
        if prepared['ok']:
            valid_rows += 1
            preview_rows.append({
                'row_no': idx,
                'valid': True,
                'listener_number': prepared.get('listener_number', ''),
                'teacher_name': prepared.get('teacher_name', ''),
                'course_title': prepared.get('course_title', ''),
                'status': prepared.get('status', '部门已审核'),
                'errors': []
            })
        else:
            invalid_rows += 1
            preview_rows.append({
                'row_no': idx,
                'valid': False,
                'listener_number': prepared.get('listener_number', ''),
                'teacher_name': prepared.get('teacher_name', ''),
                'course_title': prepared.get('course_title', ''),
                'status': prepared.get('status', '部门已审核'),
                'errors': prepared.get('errors', [])
            })
    return jsonify({
        'success': True,
        'summary': {
            'total_rows': len(rows),
            'valid_rows': valid_rows,
            'invalid_rows': invalid_rows
        },
        'rows': preview_rows[:50]
    })

@admin_bp.route('/api/forms/import/export', methods=['GET'])
@role_required('超级管理员')
def export_imported_forms_excel():
    form_ids_raw = request.args.get('form_ids', '').strip()
    if not form_ids_raw:
        return jsonify({'success': False, 'message': '缺少form_ids参数'}), 400
    form_ids = []
    for part in form_ids_raw.split(','):
        part = part.strip()
        if not part:
            continue
        try:
            form_ids.append(int(part))
        except Exception:
            continue
    if not form_ids:
        return jsonify({'success': False, 'message': 'form_ids参数无效'}), 400
    forms = LectureForm.query.filter(LectureForm.id.in_(form_ids)).order_by(LectureForm.id.asc()).all()
    if not forms:
        return jsonify({'success': False, 'message': '未找到可导出的表单'}), 404
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '导入结果明细'
    headers = [
        '表单ID', '表单组ID', '听课人编号', '听课人姓名+学院', '课程信息变化', '听课时间', '第几节', '听课地点',
        '授课教师', '教师所属学院', '课程名称', '专业年级', '异常情况反映', '主要教学方法', '管理课堂纪律', '调动课堂气氛',
        '课件制作质量', '整体教学效果', '优质案例推荐', '课程反馈', '不足及建议', '听课班级同学签名1', '联系电话1',
        '听课班级同学签名2', '联系电话2', '状态', '审核标签', '审核人ID', '审核时间', '审核意见',
        '登记ID', '创建时间', '更新时间'
    ]
    ws.append(headers)
    for form in forms:
        ws.append([
            form.id,
            form.unique_id,
            form.listener_number or '',
            form.listener_name or '',
            form.course_changes or '',
            form.lecture_date or '',
            form.class_period or '',
            form.lecture_location or '',
            form.teacher_name or '',
            form.teacher_college or '',
            form.course_title or '',
            form.student_grade_class or '',
            form.abnormal_situation or '',
            form.teaching_method or '',
            form.classroom_discipline or '',
            form.classroom_atmosphere or '',
            form.courseware_quality or '',
            form.overall_effect or '',
            form.quality_case or '',
            form.course_feedback or '',
            form.suggestions or '',
            form.student_signature1 or '',
            form.contact_phone1 or '',
            form.student_signature2 or '',
            form.contact_phone2 or '',
            form.status or '',
            form.audit_tag or '',
            form.reviewer_id if form.reviewer_id is not None else '',
            form.review_time.strftime('%Y-%m-%d %H:%M:%S') if form.review_time else '',
            form.review_comment or '',
            form.registration_id if form.registration_id is not None else '',
            form.created_at.strftime('%Y-%m-%d %H:%M:%S') if form.created_at else '',
            form.updated_at.strftime('%Y-%m-%d %H:%M:%S') if form.updated_at else ''
        ])
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'导入表单明细_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@admin_bp.route('/review_forms')
@login_required
def review_forms():
    """表单审核页面"""
    # 检查用户权限
    permission = get_user_review_permission(session['user_id'])
    if not permission:
        flash('您不具有审表权限，如有疑问，请联系管理员', 'error')
        return redirect(url_for('admin.admin_dashboard'))
    
    # 直接返回模板，让前端JavaScript处理数据加载
    current_user = User.query.get(session['user_id'])
    is_super_admin = bool(current_user and current_user.role == '超级管理员')
    try:
        teaching_total_weeks = int(SystemSetting.get('teaching_total_weeks', '20') or 20)
    except (TypeError, ValueError):
        teaching_total_weeks = 20
    if teaching_total_weeks < 1:
        teaching_total_weeks = 20
    if teaching_total_weeks > 52:
        teaching_total_weeks = 52
    return render_template(
        'admin/review_forms.html',
        is_super_admin=is_super_admin,
        teaching_total_weeks=teaching_total_weeks
    )

@admin_bp.route('/api/review/reject/<int:form_id>', methods=['POST'])
@login_required
def reject_form_review(form_id):
    """驳回表单"""
    try:
        if 'user_id' not in session:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        # 允许所有管理员驳回，或者检查审表权限
        user = User.query.get(session['user_id'])
        if user.role not in ['管理员', '超级管理员'] and not get_user_review_permission(user.id):
             return jsonify({'success': False, 'message': '您无权驳回表单'}), 403

        original_form = LectureForm.query.get_or_404(form_id)
        
        # 获取最新版本
        unique_id = original_form.unique_id or original_form.id
        latest_form = LectureForm.query.filter_by(unique_id=unique_id).order_by(LectureForm.id.desc()).first()
        
        # 校验：确保操作的是最新版本
        if latest_form and latest_form.id != original_form.id:
            return jsonify({'success': False, 'message': '该表单已有更新版本，请刷新页面后操作'}), 400
        
        # 检查可见性权限
        reviewable_user_ids = get_reviewable_users(session['user_id'])
        form_user = _active_user_query().filter_by(number=original_form.listener_number).first()
        
        if user.role != '超级管理员':
            if not form_user or form_user.id not in reviewable_user_ids:
                if not get_user_review_permission(user.id):
                     return jsonify({'success': False, 'message': '您没有权限操作此表单'}), 403

        data = request.get_json()
        reason = data.get('reason', '')
        
        # 检查当前状态，决定是更新还是新建
        if latest_form.status == '已驳回':
            # 更新现有记录
            latest_form.review_comment = reason
            latest_form.reviewer_id = session['user_id']
            latest_form.review_time = datetime.now()
            latest_form.updated_at = datetime.now()
            
            db.session.add(latest_form)
            delete_review_form_draft(session['user_id'], form_id)
            db.session.commit()
            
            return jsonify({
                'success': True, 
                'message': '驳回意见已更新',
                'new_status': '已驳回'
            })
        else:
            # 创建新的表单记录（驳回版本）
            new_form = LectureForm(
                # 复制原表单所有字段
                listener_name=original_form.listener_name,
                listener_number=original_form.listener_number,
                course_changes=original_form.course_changes,
                lecture_date=original_form.lecture_date,
                class_period=original_form.class_period,
                lecture_location=original_form.lecture_location,
                teacher_name=original_form.teacher_name,
                teacher_college=original_form.teacher_college,
                course_title=original_form.course_title,
                student_grade_class=original_form.student_grade_class,
                abnormal_situation=original_form.abnormal_situation,
                teaching_method=original_form.teaching_method,
                classroom_discipline=original_form.classroom_discipline,
                classroom_atmosphere=original_form.classroom_atmosphere,
                courseware_quality=original_form.courseware_quality,
                overall_effect=original_form.overall_effect,
                quality_case=original_form.quality_case,
                course_feedback=original_form.course_feedback,
                suggestions=original_form.suggestions,
                student_signature1=original_form.student_signature1,
                contact_phone1=original_form.contact_phone1,
                student_signature2=original_form.student_signature2,
                contact_phone2=original_form.contact_phone2,
                registration_id=original_form.registration_id, # 保持关联
                audit_tag=original_form.audit_tag,
                
                # 设置驳回状态
                status='已驳回',
                reviewer_id=session['user_id'],
                review_time=datetime.now(),
                review_comment=reason,
                
                # 继承unique_id
                unique_id=unique_id,
                created_at=original_form.created_at,
                updated_at=datetime.now()
            )
            
            # 确保原表单有unique_id
            if not original_form.unique_id:
                original_form.unique_id = original_form.id
                db.session.add(original_form)
                
            db.session.add(new_form)
            delete_review_form_draft(session['user_id'], form_id)
            db.session.commit()
            
            return jsonify({
                'success': True, 
                'message': '表单已驳回',
                'new_status': '已驳回'
            })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/submit/<int:form_id>', methods=['POST'])
@login_required
def submit_review(form_id):
    """提交表单审核"""
    try:
        # 检查用户权限
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您没有审核权限'})
        
        # 获取原表单
        original_form = LectureForm.query.filter_by(id=form_id).first()
        if not original_form:
            return jsonify({'success': False, 'message': '表单不存在'})
        
        # 检查是否可以审核该状态的表单
        if not can_review_status(session['user_id'], original_form.status):
            return jsonify({'success': False, 'message': f'您无权审核状态为"{original_form.status}"的表单'})
        
        # 获取表单数据
        # 兼容 JSON 格式（前端使用 fetch JSON 提交）和 Form 格式
        if request.is_json:
            data = request.get_json()
            form_data = data.get('form_data', {})
            review_comment = data.get('review_comment', '无')
            score_data_list = data.get('score_data', [])
        else:
            form_data = request.form.to_dict()
            review_comment = form_data.get('review_comment', '无')
            score_data_list = []
            # 尝试从 form 字段中解析 score_data JSON 字符串（兼容旧方式或隐藏域提交）
            score_data_json = form_data.get('score_data')
            if score_data_json:
                try:
                    score_data_list = json.loads(score_data_json)
                except:
                    pass
        
        # 确定新状态
        new_status = get_next_status_after_review(session['user_id'])

        course_changes_raw = (form_data.get('course_changes') or '').strip()
        resolved_course_changes = course_changes_raw or (original_form.course_changes or '无')
        
        # 获取最新版本
        unique_id = original_form.unique_id or original_form.id
        latest_form = LectureForm.query.filter_by(unique_id=unique_id).order_by(LectureForm.id.desc()).first()
        
        # 校验：确保操作的是最新版本
        if latest_form and latest_form.id != original_form.id:
             return jsonify({'success': False, 'message': '该表单已有更新版本，请刷新页面后操作'}), 400
        
        # 准备表单数据
        new_form_data = {
            'unique_id': unique_id,
            'audit_tag': original_form.audit_tag,
            'listener_name': form_data.get('listener_name'),
            'listener_number': form_data.get('listener_number'),
            'course_changes': resolved_course_changes,
            'lecture_date': form_data.get('lecture_date'),
            'class_period': form_data.get('class_period'),
            'lecture_location': form_data.get('lecture_location'),
            'teacher_name': form_data.get('teacher_name'),
            'teacher_college': form_data.get('teacher_college'),
            'course_title': form_data.get('course_title'),
            'student_grade_class': form_data.get('student_grade_class'),
            'abnormal_situation': form_data.get('abnormal_situation', '无'),
            'teaching_method': form_data.get('teaching_method'),
            'classroom_discipline': form_data.get('classroom_discipline'),
            'classroom_atmosphere': form_data.get('classroom_atmosphere'),
            'courseware_quality': form_data.get('courseware_quality'),
            'overall_effect': form_data.get('overall_effect'),
            'quality_case': form_data.get('quality_case'),
            'course_feedback': form_data.get('course_feedback'),
            'suggestions': form_data.get('suggestions', '无'),
            'student_signature1': form_data.get('student_signature1'),
            'contact_phone1': form_data.get('contact_phone1'),
            'student_signature2': form_data.get('student_signature2'),
            'contact_phone2': form_data.get('contact_phone2'),
            'status': new_status,
            'reviewer_id': session['user_id'],
            'review_time': datetime.now(),
            'review_comment': review_comment,
            'updated_at': datetime.now()  # 更新时间为审核时间
        }

        field_names = {
            'listener_name': '听课人姓名',
            'course_changes': '课程信息变化',
            'lecture_date': '听课时间',
            'class_period': '第几节',
            'lecture_location': '听课地点',
            'teacher_name': '授课教师',
            'teacher_college': '教师所属学院',
            'course_title': '课程总标题',
            'student_grade_class': '专业年级',
            'abnormal_situation': '异常情况反映',
            'teaching_method': '主要教学方法',
            'classroom_discipline': '管理课堂纪律',
            'classroom_atmosphere': '调动课堂气氛',
            'courseware_quality': '课件制作质量',
            'overall_effect': '整体教学效果',
            'quality_case': '优质案例推荐',
            'course_feedback': '课程反馈',
            'suggestions': '不足及建议',
            'student_signature1': '听课班级同学签名1',
            'contact_phone1': '联系电话1',
            'student_signature2': '听课班级同学签名2',
            'contact_phone2': '联系电话2'
        }
        modified_fields = []
        for field, label in field_names.items():
            new_value = new_form_data.get(field)
            old_value = getattr(original_form, field, None)
            if str(new_value or '') != str(old_value or ''):
                modified_fields.append(label)

        if modified_fields:
            base_comment = (review_comment or '').strip() or '无'
            modification_note = f"\n\n[系统记录] 审核人修改了以下字段：{', '.join(modified_fields)}"
            new_form_data['review_comment'] = base_comment + modification_note
        
        target_form = None
        if latest_form.status == new_status:
             # 更新现有记录
             target_form = latest_form
             for k, v in new_form_data.items():
                 if k != 'unique_id' and hasattr(target_form, k): 
                     setattr(target_form, k, v)
             
             # 删除旧的评分记录以便重新创建
             old_score_record = ScoreRecord.query.filter_by(form_id=target_form.id).first()
             if old_score_record:
                 db.session.delete(old_score_record)
        else:
             # 创建新的表单记录
             target_form = LectureForm(**new_form_data)
             target_form.created_at = original_form.created_at # 继承创建时间
             db.session.add(target_form)
        
        db.session.flush()  # 获取新表单ID
        
        # 处理评分数据
        if score_data_list:
            try:
                # 计算总分
                total_dept_score = sum(float(item.get('department_score', 0)) for item in score_data_list)
                total_personal_score = sum(float(item.get('personal_score', 0)) for item in score_data_list)
                
                # 创建评分记录
                score_record = ScoreRecord(
                    form_id=target_form.id,
                    reviewer_id=session['user_id'],
                    total_department_score=total_dept_score,
                    total_personal_score=total_personal_score
                )
                db.session.add(score_record)
                db.session.flush() # 获取评分记录ID
                
                # 创建评分项
                for item in score_data_list:
                    score_item = ScoreItem(
                        score_record_id=score_record.id,
                        reason=item.get('reason'),
                        department_score=float(item.get('department_score', 0)),
                        personal_score=float(item.get('personal_score', 0)),
                        is_auto_generated=item.get('is_auto_generated', item.get('is_auto', False))
                    )
                    db.session.add(score_item)
            except Exception as e:
                print(f"Error processing score data: {e}")
                # 不中断主流程，只记录错误
        
        delete_review_form_draft(session['user_id'], form_id)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '审核提交成功',
            'new_status': new_status,
            'form_id': target_form.id
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'审核提交失败：{str(e)}'})


def _safe_review_return_url(value):
    candidate = (value or '').strip()
    if candidate.startswith('/admin/review_forms'):
        return candidate
    return url_for('admin.review_forms')


@admin_bp.route('/review/form/<int:form_id>')
@login_required
def review_form_page(form_id):
    """审核表单页面（独立页面）"""
    # 检查用户权限
    permission = get_user_review_permission(session['user_id'])
    if not permission:
        flash('您不具有审表权限，如有疑问，请联系管理员', 'error')
        return redirect(url_for('main.index'))
    
    return render_template(
        'admin/review_form.html',
        review_return_url=_safe_review_return_url(request.args.get('return_to'))
    )

@admin_bp.route('/form/<int:form_id>')
@login_required
def get_form_detail(form_id):
    """统一表单详情入口。"""
    try:
        # 检查用户权限
        user = User.query.get(session['user_id'])
        form = LectureForm.query.get_or_404(form_id)
        
        # 权限检查
        has_permission = False
        if user.role == '超级管理员':
            has_permission = True
        elif user.role == '管理员':
            # 管理员只能查看本部门的表单
            listener = _active_user_query().filter_by(number=form.listener_number).first()
            if listener and listener.department == user.department:
                has_permission = True
        else:
            # 信息员只能查看自己的表单
            if form.listener_number == user.number:
                has_permission = True
                
        if not has_permission:
            if request.args.get('format') != 'json':
                flash('您没有权限查看此表单', 'error')
                return redirect(url_for('main.index'))
            return jsonify({'success': False, 'message': '您没有权限查看此表单'})
            
        # 渲染表单详情模板
        reviewer_display_mode = get_reviewer_display_mode()
        score_record = ScoreRecord.query.filter_by(form_id=form.id).first()
        if request.args.get('format') != 'json':
            return render_template(
                'admin/form_detail_page.html',
                form=form,
                reviewer_display_mode=reviewer_display_mode,
                score_record=score_record
            )
        html = render_template(
            'admin/form_detail_content.html',
            form=form,
            reviewer_display_mode=reviewer_display_mode,
            score_record=score_record
        )
        return jsonify({'success': True, 'html': html})
        
    except Exception as e:
        if request.args.get('format') != 'json':
            flash(str(e), 'error')
            return redirect(url_for('main.index'))
        return jsonify({'success': False, 'message': str(e)})

# ==================== 表单审核 API 路由 ====================

@admin_bp.route('/api/review/permission', methods=['GET'])
def get_review_permission():
    """获取当前用户的审核权限"""
    try:
        if 'user_id' not in session:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        permission = get_user_review_permission(session['user_id'])
        presentation = get_review_permission_presentation(permission)
        
        return jsonify({
            'success': True,
            'permission': permission,
            'has_permission': permission is not None,
            **presentation,
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/batch_auto_check', methods=['POST'])
@login_required
def batch_auto_check():
    """Legacy adapter for the additive evidence-only batch API."""
    permission = get_user_review_permission(session['user_id'])
    if permission != '审表_中心':
        return jsonify({'success': False, 'message': '您无权限执行一键自动审核'}), 403
    from ..review_automation.routes import create_batch_response

    return create_batch_response(
        legacy_force_ignored=True,
        actor_id=session['user_id'],
    )

@admin_bp.route('/api/review/batch_auto_check/preview', methods=['POST'])
@login_required
def batch_auto_check_preview():
    """Legacy adapter for the evidence-only batch preview."""
    permission = get_user_review_permission(session['user_id'])
    if permission != '审表_中心':
        return jsonify({'success': False, 'message': '您无权限查看自动审核预览'}), 403
    from ..review_automation.routes import preview_batch_response

    return preview_batch_response(
        legacy_force_ignored=True,
        actor_id=session['user_id'],
    )

@admin_bp.route('/api/review/auto_check/status', methods=['GET'])
@login_required
def get_auto_check_status():
    """Legacy status adapter backed by the newest evidence-only batch."""
    permission = get_user_review_permission(session['user_id'])
    if permission != '审表_中心':
        return jsonify({'success': False, 'message': '您无权限查看自动审核任务状态'}), 403
    from ..review_automation.routes import status_batch_response

    return status_batch_response(
        request.args.get('batch_id'),
        actor_id=session['user_id'],
    )

@admin_bp.route('/auto_review/results')
@login_required
def auto_review_results_page():
    """自动审核结果页面（专业审查）"""
    permission = get_user_review_permission(session['user_id'])
    if permission != '审表_中心':
        return redirect(url_for('admin.review_forms'))
    return render_template('admin/auto_review_results.html')

@admin_bp.route('/api/review/user-structure', methods=['GET'])
@login_required
def get_user_structure_api():
    """获取用户结构数据（用于审核权限范围显示）"""
    try:
        permission = get_user_review_permission(session['user_id'])
        if permission:
            structure = get_user_structure_for_review(session['user_id'])
            presentation = get_review_permission_presentation(permission)
            return jsonify({
                'success': True,
                'permission': permission,
                'structure': structure,
                **presentation,
            })
        manage_permission = get_user_manage_permission(session['user_id'])
        if manage_permission == '管理部门':
            current_user = User.query.get(session['user_id'])
            users = _active_user_query().filter(
                User.role.in_(['信息员', '管理员']),
                User.department == current_user.department
            ).all()
            groups = {}
            for user in users:
                if user.id == current_user.id:
                    continue
                group = Group.query.get(user.group_id) if user.group_id else None
                group_name = group.name if group else UNASSIGNED_GROUP_NAME
                if group_name not in groups:
                    groups[group_name] = []
                groups[group_name].append({
                    'id': user.id,
                    'name': user.name,
                    'number': user.number,
                    'role': user.role
                })
            return jsonify({
                'success': True,
                'permission': manage_permission,
                'structure': groups
            })
        return jsonify({'success': False, 'message': '您不具有查看范围用户的权限'}), 403
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/statistics', methods=['GET'])
@login_required
def get_review_statistics():
    """获取审核统计数据（按筛选时间统计完成/未完成任务人数）"""
    try:
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限'}), 403

        start_date_str = request.args.get('start_date')
        end_date_str = request.args.get('end_date')
        if not start_date_str or not end_date_str:
            return jsonify({'success': False, 'message': '请先在筛选条件中选择开始和结束时间'}), 400

        start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d') + timedelta(days=1)
        time_filter_type = _normalize_review_form_time_filter(request.args.get('time_filter_type'))

        required_count = int(SystemSetting.get('teaching_required_submission', 1) or 1)
        required_count = max(required_count, 1)

        current_user = User.query.get(session['user_id'])
        query = _active_user_query().filter(User.role.in_(['信息员', '管理员']))
        if permission == '审表_小组':
            query = query.filter(User.group_id == current_user.group_id)
        elif permission == '审表_部门':
            query = query.filter(User.department == current_user.department)
        users = [u for u in query.all() if u.id != current_user.id]
        user_numbers = [u.number for u in users]

        if not user_numbers:
            return jsonify({
                'success': True,
                'range_info': {
                    'start_date': start_date_str,
                    'end_date': end_date_str
                },
                'stats': {
                    'completed': {'count': 0, 'users': []},
                    'incomplete': {'count': 0, 'users': []}
                }
            })

        user_group_map = {}
        for group_data in _latest_form_groups_for_users(user_numbers):
            latest_form = group_data.get('latest_form')
            if not latest_form:
                continue
            filter_dt = _build_review_form_filter_datetime(group_data, time_filter_type)
            if not filter_dt or filter_dt < start_date or filter_dt >= end_date:
                continue
            user_group_map.setdefault(latest_form.listener_number, []).append(group_data)

        completed_users = []
        incomplete_users = []
        valid_completed_statuses = {'待审核', '部门已审核', '中心已审核'}

        for user in users:
            submitted_groups = user_group_map.get(user.number, []) or []
            completed_count = 0
            form_summaries = []

            for group_data in submitted_groups:
                latest = group_data['latest_form']
                if latest.status in valid_completed_statuses:
                    completed_count += 1
                form_summaries.append({
                    'id': latest.id,
                    'date': latest.lecture_date,
                    'period': latest.class_period,
                    'status': latest.status
                })

            user_info = {
                'id': user.id,
                'name': user.name,
                'number': user.number,
                'department': user.department,
                'group': user.group,
                'submitted_count': completed_count,
                'required_count': required_count,
                'forms': form_summaries
            }

            if completed_count >= required_count:
                completed_users.append(user_info)
            else:
                incomplete_users.append(user_info)

        return jsonify({
            'success': True,
            'range_info': {
                'start_date': start_date_str,
                'end_date': end_date_str
            },
            'stats': {
                'completed': {
                    'count': len(completed_users),
                    'users': completed_users
                },
                'incomplete': {
                    'count': len(incomplete_users),
                    'users': incomplete_users
                }
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/forms', methods=['GET'])
def get_forms_for_review():
    """获取可审核的表单列表"""
    try:
        if 'user_id' not in session:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限，如有疑问，请联系管理员'}), 403
        
        # 获取查询参数
        start_date = request.args.get('start_date')
        end_date = request.args.get('end_date')
        status_filter = request.args.get('status')
        time_filter_type = _normalize_review_form_time_filter(request.args.get('time_filter_type'))
        user_ids = request.args.getlist('user_ids')  # 可以是多个用户ID
        
        # 构建查询
        query = LectureForm.query
        
        # 用户范围过滤
        reviewable_user_ids = get_reviewable_users(session['user_id'])
        if user_ids:
            # 取交集，确保只能查看有权限的用户
            filtered_user_ids = [int(uid) for uid in user_ids if int(uid) in reviewable_user_ids]
            if filtered_user_ids:
                # 通过listener_number关联用户
                users = _active_user_query().filter(User.id.in_(filtered_user_ids)).all()
                listener_numbers = [user.number for user in users]
                query = query.filter(LectureForm.listener_number.in_(listener_numbers))
            else:
                # 没有有效的用户ID，返回空结果
                return jsonify({'success': True, 'forms': []})
        else:
            # 没有指定用户，使用权限范围内的所有用户
            users = _active_user_query().filter(User.id.in_(reviewable_user_ids)).all()
            listener_numbers = [user.number for user in users]
            query = query.filter(LectureForm.listener_number.in_(listener_numbers))
        
        parsed_start = None
        parsed_end = None
        if start_date:
            try:
                parsed_start = datetime.strptime(start_date, '%Y-%m-%d')
            except Exception:
                return jsonify({'success': False, 'message': '开始时间格式错误'}), 400
        if end_date:
            try:
                parsed_end = datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1)
            except Exception:
                return jsonify({'success': False, 'message': '结束时间格式错误'}), 400

        # 获取所有表单，按updated_at倒序排列
        forms = query.order_by(LectureForm.updated_at.desc(), LectureForm.id.desc()).all()
        reviewer_display_mode = get_reviewer_display_mode()
        reviewer_ids = list({f.reviewer_id for f in forms if f.reviewer_id})
        reviewer_map = {}
        if reviewer_ids:
            reviewers = User.query.filter(User.id.in_(reviewer_ids)).all()
            reviewer_map = {r.id: r for r in reviewers}
        
        # 按unique_id分组，保持每组内按updated_at倒序
        grouped_forms = {}
        for form in forms:
            unique_id = form.unique_id or form.id
            if unique_id not in grouped_forms:
                grouped_forms[unique_id] = []
            grouped_forms[unique_id].append(form)
        
        # 构建返回数据，按最新更新时间排序分组
        result = []
        # 按每组最新的updated_at排序
        sorted_groups = sorted(
            grouped_forms.items(),
            key=lambda x: max((form.updated_at or form.created_at) for form in x[1]),
            reverse=True
        )
        from ..review_automation.routes import latest_assessment_summaries
        automation_summaries = latest_assessment_summaries([
            max(
                form_group,
                key=lambda item: (item.updated_at or item.created_at, item.id),
            ).id
            for _, form_group in sorted_groups
        ])
        
        for unique_id, form_group in sorted_groups:
            sorted_form_group = sorted(form_group, key=lambda f: (f.updated_at or f.created_at), reverse=True)
            latest_form = sorted_form_group[0]
            group_data_for_filter = {
                'unique_id': unique_id,
                'latest_form': latest_form,
                'forms': sorted_form_group
            }
            filter_dt = _build_review_form_filter_datetime(group_data_for_filter, time_filter_type)
            if parsed_start and (not filter_dt or filter_dt < parsed_start):
                continue
            if parsed_end and (not filter_dt or filter_dt >= parsed_end):
                continue
            if status_filter and latest_form.status != status_filter:
                continue
            group_data = {
                'unique_id': unique_id,
                'forms': []
            }
            
            for form in sorted_form_group:
                reviewer = reviewer_map.get(form.reviewer_id) if form.reviewer_id else None
                if reviewer_display_mode == 'number':
                    reviewer_display = reviewer.number if reviewer else '-'
                else:
                    reviewer_display = reviewer.name if reviewer else '-'
                group_data['forms'].append({
                    'id': form.id,
                    'unique_id': form.unique_id or form.id,
                    'listener_name': form.listener_name,
                    'listener_number': form.listener_number,
                    'course_title': form.course_title,
                    'teacher_name': form.teacher_name,
                    'lecture_date': form.lecture_date,
                    'status': form.status,
                    'can_review': (
                        form.id == latest_form.id
                        and can_review_status(session['user_id'], form.status, permission=permission)
                    ),
                    'created_at': form.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                    'updated_at': form.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
                    'review_comment': form.review_comment or '',
                    'reviewer_id': form.reviewer_id,
                    'reviewer_display': reviewer_display,
                    'audit_tag': form.audit_tag or '',
                    'audit_tag_info': _serialize_audit_tag(form.audit_tag),
                    'automation': automation_summaries.get(form.id) if form.id == latest_form.id else None,
                })
            
            result.append(group_data)
        
        return jsonify({'success': True, 'forms': result})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/form/<int:form_id>', methods=['GET'])
def get_form_for_review(form_id):
    """获取单个表单的详细信息（用于审核）"""
    try:
        if 'user_id' not in session:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限，如有疑问，请联系管理员'}), 403
        
        form = LectureForm.query.get_or_404(form_id)
        
        # 检查是否有权限审核此表单
        reviewable_user_ids = get_reviewable_users(session['user_id'])
        form_user = _active_user_query().filter_by(number=form.listener_number).first()
        if not form_user or form_user.id not in reviewable_user_ids:
            return jsonify({'success': False, 'message': '您没有权限审核此表单'}), 403
        
        # 构建表单数据
        reviewer_display_mode = get_reviewer_display_mode()
        if form.reviewer:
            reviewer_display = form.reviewer.number if reviewer_display_mode == 'number' else form.reviewer.name
        else:
            reviewer_display = None
        latest_form = form.get_latest_version()
        form_data = {
            'id': form.id,
            'unique_id': form.unique_id or form.id,
            'listener_name': form.listener_name,
            'listener_number': form.listener_number,
            'course_changes': form.course_changes,
            'lecture_date': form.lecture_date,
            'class_period': form.class_period,
            'lecture_location': form.lecture_location,
            'teacher_name': form.teacher_name,
            'teacher_college': form.teacher_college,
            'course_title': form.course_title,
            'student_grade_class': form.student_grade_class,
            'abnormal_situation': form.abnormal_situation,
            'teaching_method': form.teaching_method,
            'classroom_discipline': form.classroom_discipline,
            'classroom_atmosphere': form.classroom_atmosphere,
            'courseware_quality': form.courseware_quality,
            'overall_effect': form.overall_effect,
            'quality_case': form.quality_case,
            'course_feedback': form.course_feedback,
            'suggestions': form.suggestions,
            'student_signature1': form.student_signature1,
            'contact_phone1': form.contact_phone1,
            'student_signature2': form.student_signature2,
            'contact_phone2': form.contact_phone2,
            'status': form.status,
            'can_review': (
                form.id == latest_form.id
                and can_review_status(session['user_id'], form.status, permission=permission)
            ),
            'reviewer_id': form.reviewer_id,
            'reviewer_display': reviewer_display,
            'review_time': form.review_time.strftime('%Y-%m-%d %H:%M:%S') if form.review_time else None,
            'review_comment': form.review_comment,
            'created_at': form.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'updated_at': form.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
            'audit_tag': form.audit_tag or '',
            'audit_tag_info': _serialize_audit_tag(form.audit_tag)
        }
        from ..review_automation.routes import latest_assessment_summaries
        automation = latest_assessment_summaries([latest_form.id]).get(latest_form.id)
        form_data['automation'] = automation
        form_data['automation_evidence_url'] = automation.get('evidence_url') if automation else None
        
        score_record = ScoreRecord.query.filter_by(form_id=form.id).first()
        score_data = []
        if score_record:
            score_data = [{
                'reason': item.reason,
                'department_score': item.department_score,
                'personal_score': item.personal_score,
                'is_auto_generated': item.is_auto_generated
            } for item in score_record.items]

        return jsonify({
            'success': True,
            'form': form_data,
            'automation': automation,
            'automation_evidence_url': form_data['automation_evidence_url'],
            'permission': permission,
            'score_data': score_data
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@admin_bp.route('/api/review/form-tags/preview', methods=['POST'])
@role_required('超级管理员')
def preview_review_form_tags():
    data = request.get_json() or {}
    group_entries, error_response, status_code = _load_review_form_groups_for_definition(data.get('form_ids', []))
    if error_response:
        return error_response, status_code

    preview_rows = []
    for entry in group_entries:
        form = entry['latest_form']
        preview_rows.append({
            'form_id': form.id,
            'unique_id': form.unique_id or form.id,
            'listener_name': form.listener_name,
            'listener_number': form.listener_number,
            'teacher_name': form.teacher_name,
            'course_title': form.course_title,
            'status': form.status,
            'audit_tag': form.audit_tag or '',
            'audit_tag_info': _serialize_audit_tag(form.audit_tag),
            'version_count': len(entry['forms']),
        })

    return jsonify({
        'success': True,
        'forms': preview_rows,
    })


@admin_bp.route('/api/review/form-tags', methods=['PUT'])
@role_required('超级管理员')
def update_review_form_tags():
    data = request.get_json() or {}
    definitions = data.get('definitions', [])
    if not isinstance(definitions, list) or not definitions:
        return jsonify({'success': False, 'message': '请提交需要定义的表单标签'}), 400

    normalized_definitions = {}
    for definition in definitions:
        try:
            form_id = int(definition.get('form_id'))
        except (TypeError, ValueError):
            return jsonify({'success': False, 'message': '存在非法表单ID'}), 400
        normalized_definitions[form_id] = {
            'review_tag': str(definition.get('review_tag') or '').strip(),
            'week_correction_tag': str(
                definition.get('week_correction_tag')
                or definition.get('late_tag')
                or ''
            ).strip(),
        }

    group_entries, error_response, status_code = _load_review_form_groups_for_definition(normalized_definitions.keys())
    if error_response:
        return error_response, status_code

    updated_rows = []
    updated_version_count = 0
    for entry in group_entries:
        latest_form = entry['latest_form']
        definition = normalized_definitions.get(latest_form.id)
        if not definition:
            for requested_form_id in entry.get('requested_form_ids', []):
                definition = normalized_definitions.get(requested_form_id)
                if definition:
                    break
        if not definition:
            continue
        review_tag = definition.get('review_tag')
        week_correction_tag = definition.get('week_correction_tag') or None
        audit_tag = build_audit_tag(review_tag, week_correction_tag)
        errors = validate_audit_tag(audit_tag)
        if errors:
            return jsonify({
                'success': False,
                'message': f'表单组最新版本ID {latest_form.id} 的标签不合法：{"；".join(errors)}'
            }), 400
        for form in entry['forms']:
            form.audit_tag = audit_tag
        updated_version_count += len(entry['forms'])
        updated_rows.append({
            'form_id': latest_form.id,
            'unique_id': entry['unique_id'],
            'audit_tag': audit_tag,
            'audit_tag_info': _serialize_audit_tag(audit_tag),
            'version_count': len(entry['forms']),
        })

    db.session.commit()
    return jsonify({
        'success': True,
        'message': f'已更新 {len(updated_rows)} 个表单组，共同步 {updated_version_count} 个版本的标签定义',
        'updated_forms': updated_rows,
    })

@admin_bp.route('/api/review/form/<int:form_id>', methods=['POST'])
def submit_form_review(form_id):
    """提交表单审核"""
    try:
        if 'user_id' not in session:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限，如有疑问，请联系管理员'}), 403
        
        original_form = LectureForm.query.get_or_404(form_id)
        
        # 检查是否有权限审核此表单
        reviewable_user_ids = get_reviewable_users(session['user_id'])
        form_user = _active_user_query().filter_by(number=original_form.listener_number).first()
        if not form_user or form_user.id not in reviewable_user_ids:
            return jsonify({'success': False, 'message': '您没有权限审核此表单'}), 403
        
        # 检查是否可以审核此状态的表单
        if not can_review_status(session['user_id'], original_form.status):
            return jsonify({'success': False, 'message': f'您无权限审核处于"{original_form.status}"状态的表单'}), 403
        
        data = request.get_json()
        review_comment = data.get('review_comment', '')
        form_data = data.get('form_data', {})
        
        # 检查表单数据是否有修改
        modified_fields = []
        original_data = {
            'listener_name': original_form.listener_name,
            'course_changes': original_form.course_changes,
            'lecture_date': original_form.lecture_date,
            'class_period': original_form.class_period,
            'lecture_location': original_form.lecture_location,
            'teacher_name': original_form.teacher_name,
            'teacher_college': original_form.teacher_college,
            'course_title': original_form.course_title,
            'student_grade_class': original_form.student_grade_class,
            'abnormal_situation': original_form.abnormal_situation,
            'teaching_method': original_form.teaching_method,
            'classroom_discipline': original_form.classroom_discipline,
            'classroom_atmosphere': original_form.classroom_atmosphere,
            'courseware_quality': original_form.courseware_quality,
            'overall_effect': original_form.overall_effect,
            'quality_case': original_form.quality_case,
            'course_feedback': original_form.course_feedback,
            'suggestions': original_form.suggestions,
            'student_signature1': original_form.student_signature1,
            'contact_phone1': original_form.contact_phone1,
            'student_signature2': original_form.student_signature2,
            'contact_phone2': original_form.contact_phone2
        }
        
        field_names = {
            'listener_name': '听课人姓名',
            'course_changes': '课程信息变化',
            'lecture_date': '听课时间',
            'class_period': '第几节',
            'lecture_location': '听课地点',
            'teacher_name': '授课教师',
            'teacher_college': '教师所属学院',
            'course_title': '课程总标题',
            'student_grade_class': '专业年级',
            'abnormal_situation': '异常情况反映',
            'teaching_method': '主要教学方法',
            'classroom_discipline': '管理课堂纪律',
            'classroom_atmosphere': '调动课堂气氛',
            'courseware_quality': '课件制作质量',
            'overall_effect': '整体教学效果',
            'quality_case': '优质案例推荐',
            'course_feedback': '课程反馈',
            'suggestions': '不足及建议',
            'student_signature1': '听课班级同学签名1',
            'contact_phone1': '联系电话1',
            'student_signature2': '听课班级同学签名2',
            'contact_phone2': '联系电话2'
        }
        
        for field, original_value in original_data.items():
            new_value = form_data.get(field, original_value)
            if str(new_value or '') != str(original_value or ''):
                modified_fields.append(field_names.get(field, field))
        
        # 处理听课时间（优先使用带星期几的显示格式）
        lecture_date = form_data.get('lecture_date_display') or form_data.get('lecture_date', original_form.lecture_date)
        
        # 处理节次（优先使用自动补全的格式）
        class_period = form_data.get('class_period') or f"第{form_data.get('start_period', '')}-{form_data.get('end_period', '')}节" if form_data.get('start_period') else form_data.get('class_period', original_form.class_period)
        
        new_status = get_next_status_after_review(session['user_id'])
        unique_id = original_form.unique_id or original_form.id
        latest_form = LectureForm.query.filter_by(unique_id=unique_id).order_by(LectureForm.id.desc()).first()
        if latest_form and latest_form.id != original_form.id:
            return jsonify({'success': False, 'message': '该表单已有更新版本，请刷新页面后操作'}), 400

        target_form = None
        # 如果当前版本已经是审核产生的版本（ID != unique_id），则直接在当前版本上修改
        # 或者如果状态没有改变，也在当前版本上修改
        if latest_form and (latest_form.id != unique_id or latest_form.status == new_status):
            target_form = latest_form
        else:
            target_form = LectureForm(
                listener_name=original_form.listener_name,
                listener_number=original_form.listener_number,
                course_changes=original_form.course_changes,
                lecture_date=original_form.lecture_date,
                class_period=original_form.class_period,
                lecture_location=original_form.lecture_location,
                teacher_name=original_form.teacher_name,
                teacher_college=original_form.teacher_college,
                course_title=original_form.course_title,
                student_grade_class=original_form.student_grade_class,
                abnormal_situation=original_form.abnormal_situation,
                teaching_method=original_form.teaching_method,
                classroom_discipline=original_form.classroom_discipline,
                classroom_atmosphere=original_form.classroom_atmosphere,
                courseware_quality=original_form.courseware_quality,
                overall_effect=original_form.overall_effect,
                quality_case=original_form.quality_case,
                course_feedback=original_form.course_feedback,
                suggestions=original_form.suggestions,
                student_signature1=original_form.student_signature1,
                contact_phone1=original_form.contact_phone1,
                student_signature2=original_form.student_signature2,
                contact_phone2=original_form.contact_phone2,
                audit_tag=original_form.audit_tag,
                unique_id=unique_id,
                created_at=original_form.created_at
            )
            db.session.add(target_form)

        target_form.listener_name = form_data.get('listener_name', original_form.listener_name)
        target_form.listener_number = original_form.listener_number
        course_changes_raw = (form_data.get('course_changes') or '').strip()
        target_form.course_changes = course_changes_raw or original_form.course_changes or '无'
        target_form.lecture_date = lecture_date
        target_form.class_period = class_period
        target_form.lecture_location = form_data.get('lecture_location', original_form.lecture_location)
        target_form.teacher_name = form_data.get('teacher_name', original_form.teacher_name)
        target_form.teacher_college = form_data.get('teacher_college', original_form.teacher_college)
        target_form.course_title = form_data.get('course_title', original_form.course_title)
        target_form.student_grade_class = form_data.get('student_grade_class', original_form.student_grade_class)
        target_form.abnormal_situation = form_data.get('abnormal_situation', original_form.abnormal_situation)
        target_form.teaching_method = form_data.get('teaching_method', original_form.teaching_method)
        target_form.classroom_discipline = form_data.get('classroom_discipline', original_form.classroom_discipline)
        target_form.classroom_atmosphere = form_data.get('classroom_atmosphere', original_form.classroom_atmosphere)
        target_form.courseware_quality = form_data.get('courseware_quality', original_form.courseware_quality)
        target_form.overall_effect = form_data.get('overall_effect', original_form.overall_effect)
        target_form.quality_case = form_data.get('quality_case', original_form.quality_case)
        target_form.course_feedback = form_data.get('course_feedback', original_form.course_feedback)
        target_form.suggestions = form_data.get('suggestions', original_form.suggestions)
        target_form.student_signature1 = form_data.get('student_signature1', original_form.student_signature1)
        target_form.contact_phone1 = form_data.get('contact_phone1', original_form.contact_phone1)
        target_form.student_signature2 = form_data.get('student_signature2', original_form.student_signature2)
        target_form.contact_phone2 = form_data.get('contact_phone2', original_form.contact_phone2)
        target_form.status = new_status
        target_form.reviewer_id = session['user_id']
        target_form.review_time = datetime.now()
        target_form.review_comment = review_comment
        target_form.updated_at = datetime.now()

        if not original_form.unique_id:
            original_form.unique_id = original_form.id
            db.session.add(original_form)

        db.session.flush()
        old_score_record = ScoreRecord.query.filter_by(form_id=target_form.id).first()
        if old_score_record:
            db.session.delete(old_score_record)

        score_data = data.get('score_data', [])
        if score_data:
            total_dept = 0.0
            total_pers = 0.0
            score_record = ScoreRecord(
                form_id=target_form.id,
                reviewer_id=session['user_id']
            )
            db.session.add(score_record)
            db.session.flush() # 获取 score_record.id
            
            for item in score_data:
                reason = item.get('reason')
                dept_score = float(item.get('department_score', 0))
                pers_score = float(item.get('personal_score', 0))
                is_auto = item.get('is_auto', False)
                
                if dept_score > 0 or pers_score > 0:
                    score_item = ScoreItem(
                        score_record_id=score_record.id,
                        reason=reason,
                        department_score=dept_score,
                        personal_score=pers_score,
                        is_auto_generated=is_auto
                    )
                    db.session.add(score_item)
                    total_dept += dept_score
                    total_pers += pers_score
            
            score_record.total_department_score = total_dept
            score_record.total_personal_score = total_pers
        
        # 如果有修改字段，在审核意见后添加修改说明
        if modified_fields:
            modification_note = f"\n\n[系统记录] 审核人修改了以下字段：{', '.join(modified_fields)}"
            target_form.review_comment = (review_comment or '无') + modification_note

        db.session.add(target_form)
        delete_review_form_draft(session['user_id'], form_id)
        db.session.commit()

        return jsonify({
            'success': True, 
            'message': f'表单审核完成，状态已更新为"{new_status}"',
            'new_status': new_status,
            'modified_fields': modified_fields
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

# ==================== 数据清空功能 ====================

@admin_bp.route('/api/clear_table', methods=['POST'])
@role_required('超级管理员')
def clear_table():
    """清空指定数据表"""
    try:
        table_name = request.form.get('table_name')
        password = request.form.get('password')
        
        # 验证参数
        if not table_name or not password:
            return jsonify({'success': False, 'message': '参数不完整'})
        
        # 验证密码（验证当前用户密码）
        current_user = User.query.get(session['user_id'])
        if not current_user or not check_password_hash(current_user.password_hash, password):
            return jsonify({'success': False, 'message': '密码错误，无法执行清空操作'})
        
        # 定义允许清空的表及其对应的模型
        allowed_tables = {
            'courses': Course,
            'teachers': Teacher,
            'venues': Venue,
            'lecture_forms': LectureForm,
            'listening_bans': ListeningBan,
            'users': User,
            'departments': Department,
            'groups': Group
        }
        
        if table_name not in allowed_tables:
            return jsonify({'success': False, 'message': f'不允许清空表：{table_name}'})
        
        # 获取对应的模型类
        model_class = allowed_tables[table_name]
        
        # 统计删除前的记录数
        count_before = model_class.query.count()
        
        # 执行清空操作
        if table_name == 'users':
            # 用户表特殊处理：保留超级管理员账户
            deleted_count = model_class.query.filter(model_class.role != '超级管理员').delete()
        else:
            # 其他表直接清空
            deleted_count = model_class.query.delete()
        
        # 提交更改
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'成功清空 {table_name} 表',
            'deleted_count': deleted_count,
            'table_name': table_name
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'清空失败：{str(e)}'})


# ==================== 课程管理相关路由 ====================

@admin_bp.route('/course_management')
@role_required('超级管理员')
def course_management():
    """课程管理页面"""
    return render_template(
        'admin/course_feedback_management.html',
        page_mode='management',
        page_title='课程管理'
    )


@admin_bp.route('/course_feedback_management')
@role_required('超级管理员')
def course_feedback_management():
    """兼容旧课程反馈管理入口"""
    return redirect(url_for('admin.course_management'))

@admin_bp.route('/registration_statistics')
@role_required('超级管理员')
def registration_statistics():
    return render_template('admin/registration_statistics.html')

@admin_bp.route('/api/courses', methods=['GET'])
@login_required
def get_courses():
    """获取课程列表API - 按课程号和选课课号分组显示"""
    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 20, type=int)
        search = request.args.get('search', '')
        semester = request.args.get('semester', '')
        weekday = request.args.get('weekday', '')
        time_slot = request.args.get('time_slot', '')
        
        # 构建查询
        query = Course.query
        
        # 搜索功能
        if search:
            query = query.filter(
                db.or_(
                    Course.course_name.contains(search),
                    Course.course_code.contains(search),
                    Course.selection_code.contains(search),
                    Course.teacher_id.contains(search),
                    Course.venue_id.contains(search),
                    Course.offering_college.contains(search)
                )
            )
        
        # 学期筛选
        if semester:
            query = query.filter(Course.semester == semester)
        
        # 星期筛选
        if weekday:
            query = query.filter(Course.weekday == weekday)
        
        # 时间段筛选
        if time_slot:
            query = query.filter(Course.class_period == time_slot)
        
        # 获取所有符合条件的课程记录
        all_courses = query.all()
        
        # 按课程号和选课课号分组
        course_groups = {}
        for course in all_courses:
            group_key = f"{course.course_code}_{course.selection_code}"
            if group_key not in course_groups:
                course_groups[group_key] = {
                    'course_code': course.course_code,
                    'selection_code': course.selection_code,
                    'course_name': course.course_name,
                    'offering_college': course.offering_college,
                    'credits': course.credits,
                    'total_hours': course.total_hours,
                    'course_nature': course.course_nature,
                    'semester': course.semester,
                    'academic_year': course.academic_year,
                    'enrollment_count': course.enrollment_count,
                    'weekly_hours': course.weekly_hours,
                    'class_composition': course.class_composition,
                    'major_composition': course.major_composition,
                    'records': [],  # 存储该课程的所有记录
                    'record_count': 0,  # 记录数量
                    'expanded': False  # 默认折叠，用户可以选择展开需要查看的课程组
                }
            
            # 添加具体的课程记录
            course_record = {
                'id': course.id,
                'teacher_id': course.teacher_id,
                'teacher_name': course.teacher.name if course.teacher else '',
                'venue_id': course.venue_id,
                'venue_name': course.venue.name if course.venue else '',
                'class_size': course.class_size,
                'class_time': course.class_time,
                'class_location': course.class_location,
                'weekday': course.weekday,
                'class_period': course.class_period,
                'start_week': course.start_week,
                'venue_start_week': course.venue_start_week,
                'venue_class_period': course.venue_class_period
            }
            course_groups[group_key]['records'].append(course_record)
            course_groups[group_key]['record_count'] += 1
        
        # 转换为列表并排序
        grouped_courses = list(course_groups.values())
        grouped_courses.sort(key=lambda x: (x['course_code'], x['selection_code']))
        
        # 手动分页
        total_groups = len(grouped_courses)
        start_idx = (page - 1) * per_page
        end_idx = start_idx + per_page
        paginated_courses = grouped_courses[start_idx:end_idx]
        
        # 计算分页信息
        total_pages = (total_groups + per_page - 1) // per_page
        has_prev = page > 1
        has_next = page < total_pages
        
        # 获取筛选选项
        semesters = db.session.query(Course.semester).distinct().filter(Course.semester.isnot(None)).all()
        semesters = [s[0] for s in semesters if s[0]]  # 提取学期值并过滤空值
        
        time_slots = db.session.query(Course.class_period).distinct().filter(Course.class_period.isnot(None)).all()
        time_slots = [t[0] for t in time_slots if t[0]]  # 提取时间段值并过滤空值
        
        return jsonify({
            'success': True,
            'courses': paginated_courses,
            'pagination': {
                'page': page,
                'pages': total_pages,
                'per_page': per_page,
                'total': total_groups,
                'has_prev': has_prev,
                'has_next': has_next
            },
            'filter_options': {
                'semesters': sorted(semesters),
                'time_slots': sorted(time_slots)
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取课程列表失败：{str(e)}'})

@admin_bp.route('/api/registration_statistics', methods=['GET'])
@role_required('超级管理员')
def get_registration_statistics():
    try:
        registration_rows = db.session.query(
            CourseRegistration.course_code,
            CourseRegistration.selection_code,
            func.count(CourseRegistration.id).label('listen_count')
        ).group_by(
            CourseRegistration.course_code,
            CourseRegistration.selection_code
        ).all()

        course_rows = Course.query.with_entities(
            Course.id,
            Course.course_code,
            Course.selection_code,
            Course.course_name,
            Course.teacher_id
        ).all()

        course_group_map = {}
        for row in course_rows:
            key = f"{row.course_code}_{row.selection_code}"
            if key not in course_group_map:
                course_group_map[key] = {
                    'course_code': row.course_code,
                    'selection_code': row.selection_code,
                    'course_name': row.course_name or '未命名课程',
                    'teacher_id': row.teacher_id,
                    'course_ids': []
                }
            course_group_map[key]['course_ids'].append(row.id)

        teacher_ids = [item['teacher_id'] for item in course_group_map.values() if item['teacher_id']]
        teacher_map = {t.teacher_id: t for t in Teacher.query.filter(Teacher.teacher_id.in_(teacher_ids)).all()} if teacher_ids else {}

        teacher_stats_map = {}
        for row in registration_rows:
            key = f"{row.course_code}_{row.selection_code}"
            course_group = course_group_map.get(key)
            if not course_group:
                continue

            teacher_id = course_group['teacher_id'] or 'UNKNOWN'
            teacher_obj = teacher_map.get(course_group['teacher_id']) if course_group['teacher_id'] else None
            teacher_name = teacher_obj.name if teacher_obj else '未匹配教师'
            teacher_college = teacher_obj.college if teacher_obj else ''

            if teacher_id not in teacher_stats_map:
                teacher_stats_map[teacher_id] = {
                    'teacher_id': teacher_id if teacher_id != 'UNKNOWN' else '',
                    'teacher_name': teacher_name,
                    'teacher_college': teacher_college,
                    'total_listen_count': 0,
                    'course_ids': [],
                    'courses': []
                }

            teacher_stats_map[teacher_id]['total_listen_count'] += int(row.listen_count or 0)
            teacher_stats_map[teacher_id]['course_ids'].extend(course_group['course_ids'])
            teacher_stats_map[teacher_id]['courses'].append({
                'course_code': course_group['course_code'],
                'selection_code': course_group['selection_code'],
                'course_name': course_group['course_name'],
                'listen_count': int(row.listen_count or 0),
                'course_ids': course_group['course_ids']
            })

        teacher_stats = list(teacher_stats_map.values())
        for teacher in teacher_stats:
            teacher['course_ids'] = sorted(list(set(teacher['course_ids'])))
            teacher['courses'].sort(key=lambda x: (-x['listen_count'], x['course_name']))
            teacher['course_count'] = len(teacher['courses'])

        teacher_stats.sort(key=lambda x: (-x['total_listen_count'], x['teacher_name']))
        return jsonify({'success': True, 'teachers': teacher_stats})
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取统计失败：{str(e)}'}), 500

@admin_bp.route('/api/teachers/<teacher_id>', methods=['GET'])
@role_required('超级管理员')
def get_teacher_detail(teacher_id):
    """获取教师详细信息"""
    try:
        teacher = Teacher.query.get_or_404(teacher_id)
        
        # 获取该教师的课程
        courses = Course.query.filter_by(teacher_id=teacher_id).all()
        
        teacher_data = {
            'teacher_id': teacher.teacher_id,
            'name': teacher.name,
            'gender': teacher.gender,
            'title': teacher.title,
            'college': teacher.college,
            'department': teacher.college,  # 使用college作为department
            'phone': teacher.phone,
            'courses': [
                {
                    'course_id': course.id,
                    'course_name': course.course_name,
                    'class_time': course.class_time,
                    'class_location': course.class_location,
                    'semester': course.semester,
                    'academic_year': course.academic_year
                } for course in courses
            ]
        }
        
        return jsonify({'success': True, 'teacher': teacher_data})
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取教师信息失败：{str(e)}'})

@admin_bp.route('/api/venues/<venue_id>', methods=['GET'])
@role_required('超级管理员')
def get_venue_detail(venue_id):
    """获取场地详细信息"""
    try:
        venue = Venue.query.get_or_404(venue_id)
        
        # 获取该场地的课程
        courses = Course.query.filter_by(venue_id=venue_id).all()
        
        venue_data = {
            'venue_id': venue.venue_id,
            'name': venue.name,
            'category': venue.category,
            'type': venue.category,  # 使用category作为type
            'campus': venue.campus,
            'floor': venue.floor,
            'building': venue.building,
            'capacity': venue.capacity,
            'courses': [
                {
                    'course_id': course.id,
                    'course_name': course.course_name,
                    'class_time': course.class_time,
                    'teacher_name': course.teacher.name if course.teacher else '',
                    'semester': course.semester,
                    'academic_year': course.academic_year
                } for course in courses
            ]
        }
        
        return jsonify({'success': True, 'venue': venue_data})
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取场地信息失败：{str(e)}'})

@admin_bp.route('/api/courses/<int:course_id>/ban_users', methods=['GET'])
@role_required('超级管理员')
def get_course_ban_users(course_id):
    """获取课程的禁听用户列表"""
    try:
        # 获取该课程的所有禁听记录
        bans = ListeningBan.query.filter_by(course_id=course_id).all()
        
        ban_users = []
        for ban in bans:
            ban_users.append({
                'id': ban.id,
                'user_id': ban.user_id,
                'user_name': ban.user.name,
                'user_number': ban.user.number,
                'department': ban.user.department,
                'college': ban.user.college,
                'created_at': ban.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                'creator_name': ban.creator.name if ban.creator else ''
            })
        
        return jsonify({'success': True, 'ban_users': ban_users})
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取禁听用户列表失败：{str(e)}'})

@admin_bp.route('/api/courses/<int:course_id>/ban_users', methods=['POST'])
@role_required('超级管理员')
def add_course_ban_users(course_id):
    """为课程添加禁听用户"""
    try:
        data = request.get_json()
        user_ids = data.get('user_ids', [])
        
        if not user_ids:
            return jsonify({'success': False, 'message': '请选择要禁听的用户'})
        
        # 检查课程是否存在
        course = Course.query.get_or_404(course_id)
        
        added_count = 0
        skipped_count = 0
        
        for user_id in user_ids:
            # 检查用户是否存在
            user = User.query.get(user_id)
            if not user or not is_user_active(user):
                continue
            
            # 检查是否已经存在禁听记录
            existing_ban = ListeningBan.query.filter_by(
                course_id=course_id, user_id=user_id
            ).first()
            
            if existing_ban:
                skipped_count += 1
                continue
            
            # 创建新的禁听记录
            ban = ListeningBan(
                course_id=course_id,
                user_id=user_id,
                created_by=session.get('user_id')
            )
            db.session.add(ban)
            added_count += 1
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'成功添加 {added_count} 个禁听用户，跳过 {skipped_count} 个已存在的记录'
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'添加禁听用户失败：{str(e)}'})

@admin_bp.route('/api/courses/<int:course_id>/ban_users/<int:ban_id>', methods=['DELETE'])
@role_required('超级管理员')
def remove_course_ban_user(course_id, ban_id):
    """移除课程的禁听用户"""
    try:
        ban = ListeningBan.query.filter_by(
            id=ban_id, course_id=course_id
        ).first_or_404()
        
        db.session.delete(ban)
        db.session.commit()
        
        return jsonify({'success': True, 'message': '成功移除禁听用户'})
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'移除禁听用户失败：{str(e)}'})

@admin_bp.route('/api/courses/batch_ban_users', methods=['POST'])
@role_required('超级管理员')
def batch_ban_users():
    """批量为多个课程添加禁听用户"""
    try:
        data = request.get_json()
        course_ids = data.get('course_ids', [])
        user_ids = data.get('user_ids', [])
        
        if not course_ids:
            return jsonify({'success': False, 'message': '请选择要禁听的课程'})
        
        if not user_ids:
            return jsonify({'success': False, 'message': '请选择要禁听的用户'})
        
        # 检查课程是否存在
        courses = Course.query.filter(Course.id.in_(course_ids)).all()
        if len(courses) != len(course_ids):
            return jsonify({'success': False, 'message': '部分课程不存在'})
        
        # 检查用户是否存在
        users = _active_user_query().filter(User.id.in_(user_ids)).all()
        if len(users) != len(user_ids):
            return jsonify({'success': False, 'message': '部分用户不存在'})
        
        added_count = 0
        skipped_count = 0
        
        for course_id in course_ids:
            for user_id in user_ids:
                # 检查是否已经存在禁听记录
                existing_ban = ListeningBan.query.filter_by(
                    course_id=course_id, user_id=user_id
                ).first()
                
                if existing_ban:
                    skipped_count += 1
                    continue
                
                # 创建新的禁听记录
                ban = ListeningBan(
                    course_id=course_id,
                    user_id=user_id,
                    created_by=session.get('user_id')
                )
                db.session.add(ban)
                added_count += 1
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'成功添加 {added_count} 条禁听记录，跳过 {skipped_count} 条已存在的记录',
            'stats': {
                'added': added_count,
                'skipped': skipped_count,
                'total_courses': len(course_ids),
                'total_users': len(user_ids)
            }
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'批量禁听失败：{str(e)}'})

def _get_default_ban_target_users():
    return _active_user_query().filter(User.role == '信息员').all()


def _read_banned_teacher_import_names(file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError('请上传Excel文件')
    if not allowed_file(file_storage.filename):
        raise ValueError('请上传Excel格式文件（.xlsx 或 .xls）')

    try:
        dataframe = pd.read_excel(file_storage, header=None)
    except Exception as exc:
        raise ValueError(f'读取Excel失败：{str(exc)}')

    if dataframe.empty or dataframe.shape[1] == 0:
        return []

    names = []
    seen = set()
    ignored_headers = {'教师姓名', '禁听教师名单', '禁听教师'}
    for value in dataframe.iloc[:, 0].tolist():
        if pd.isna(value):
            continue
        name = str(value).strip()
        if not name or name.lower() == 'nan' or name in ignored_headers:
            continue
        if name not in seen:
            seen.add(name)
            names.append(name)
    return names


def _build_teacher_ban_candidate(teacher, teacher_course_map, teacher_ban_count_map, target_user_count):
    course_rows = teacher_course_map.get(teacher.teacher_id, [])
    course_ids = [row['id'] for row in course_rows]
    course_names = sorted({row['course_name'] for row in course_rows if row.get('course_name')})
    course_count = len(course_ids)
    expected_ban_count = course_count * target_user_count
    existing_ban_count = int(teacher_ban_count_map.get(teacher.teacher_id, 0) or 0)

    if course_count == 0:
        status = 'no_courses'
        status_label = '暂无课程'
    elif existing_ban_count <= 0:
        status = 'new'
        status_label = '新增禁听教师'
    elif existing_ban_count >= expected_ban_count:
        status = 'already'
        status_label = '已完整禁听'
    else:
        status = 'partial'
        status_label = '部分已禁听'

    return {
        'teacher_id': teacher.teacher_id,
        'teacher_name': teacher.name,
        'teacher_college': teacher.college or '',
        'teacher_title': teacher.title or '',
        'course_count': course_count,
        'course_ids': course_ids,
        'course_names': course_names,
        'expected_ban_count': expected_ban_count,
        'existing_ban_count': existing_ban_count,
        'missing_ban_count': max(expected_ban_count - existing_ban_count, 0),
        'status': status,
        'status_label': status_label,
        'is_new_ban': status == 'new',
        'is_already_banned': status == 'already',
        'needs_apply': status in {'new', 'partial'},
    }


def _build_banned_teacher_import_preview(import_names):
    normalized_names = []
    seen_names = set()
    for name in import_names or []:
        normalized = str(name).strip()
        if not normalized or normalized in seen_names:
            continue
        seen_names.add(normalized)
        normalized_names.append(normalized)

    target_users = _get_default_ban_target_users()
    target_user_ids = [user.id for user in target_users]
    target_user_count = len(target_user_ids)

    teachers = Teacher.query.filter(Teacher.name.in_(normalized_names)).order_by(Teacher.name.asc(), Teacher.teacher_id.asc()).all() if normalized_names else []
    teachers_by_name = defaultdict(list)
    teacher_ids = []
    for teacher in teachers:
        teachers_by_name[teacher.name].append(teacher)
        teacher_ids.append(teacher.teacher_id)

    course_rows = Course.query.with_entities(Course.id, Course.teacher_id, Course.course_name).filter(
        Course.teacher_id.in_(teacher_ids)
    ).all() if teacher_ids else []
    teacher_course_map = defaultdict(list)
    all_course_ids = []
    for row in course_rows:
        teacher_course_map[row.teacher_id].append({
            'id': row.id,
            'course_name': row.course_name or ''
        })
        all_course_ids.append(row.id)

    teacher_ban_count_map = {}
    if all_course_ids and target_user_ids:
        ban_count_rows = db.session.query(
            Course.teacher_id,
            func.count(ListeningBan.id)
        ).join(
            ListeningBan, ListeningBan.course_id == Course.id
        ).filter(
            Course.teacher_id.in_(teacher_ids),
            ListeningBan.user_id.in_(target_user_ids)
        ).group_by(Course.teacher_id).all()
        teacher_ban_count_map = {teacher_id: count for teacher_id, count in ban_count_rows}

    matched_entries = []
    ambiguous_entries = []
    unmatched_names = []
    candidate_lookup = {}

    for import_name in normalized_names:
        matched_teachers = teachers_by_name.get(import_name, [])
        if not matched_teachers:
            unmatched_names.append(import_name)
            continue

        candidates = []
        for teacher in matched_teachers:
            candidate = _build_teacher_ban_candidate(
                teacher,
                teacher_course_map,
                teacher_ban_count_map,
                target_user_count
            )
            candidate['import_name'] = import_name
            candidates.append(candidate)
            candidate_lookup[teacher.teacher_id] = candidate

        if len(candidates) == 1:
            matched_entries.append(candidates[0])
        else:
            ambiguous_entries.append({
                'import_name': import_name,
                'candidates': candidates
            })

    selected_candidates = list(matched_entries)
    summary = {
        'import_name_count': len(normalized_names),
        'matched_name_count': len(matched_entries),
        'ambiguous_name_count': len(ambiguous_entries),
        'unmatched_name_count': len(unmatched_names),
        'target_user_count': target_user_count,
        'new_teacher_count': len([item for item in selected_candidates if item['status'] == 'new']),
        'partial_teacher_count': len([item for item in selected_candidates if item['status'] == 'partial']),
        'already_banned_teacher_count': len([item for item in selected_candidates if item['status'] == 'already']),
        'no_course_teacher_count': len([item for item in selected_candidates if item['status'] == 'no_courses']),
    }

    return {
        'summary': summary,
        'import_names': normalized_names,
        'matched_teachers': matched_entries,
        'ambiguous_entries': ambiguous_entries,
        'unmatched_names': unmatched_names,
        'candidate_lookup': candidate_lookup,
    }


def _resolve_banned_teacher_import_selection(preview_data, selection_map):
    selection_map = selection_map or {}
    resolved_teachers = []
    unresolved_names = []
    selected_teacher_ids = set()

    for teacher_data in preview_data.get('matched_teachers', []):
        teacher_id = teacher_data['teacher_id']
        if teacher_id in selected_teacher_ids:
            continue
        selected_teacher_ids.add(teacher_id)
        resolved_teachers.append(teacher_data)

    candidate_lookup = preview_data.get('candidate_lookup', {})
    for entry in preview_data.get('ambiguous_entries', []):
        import_name = entry['import_name']
        chosen_ids = selection_map.get(import_name, [])
        valid_ids = []
        allowed_ids = {candidate['teacher_id'] for candidate in entry.get('candidates', [])}
        for teacher_id in chosen_ids:
            normalized_teacher_id = str(teacher_id).strip()
            if normalized_teacher_id in allowed_ids and normalized_teacher_id not in selected_teacher_ids:
                valid_ids.append(normalized_teacher_id)

        if not valid_ids:
            unresolved_names.append(import_name)
            continue

        for teacher_id in valid_ids:
            selected_teacher_ids.add(teacher_id)
            resolved_teachers.append(candidate_lookup[teacher_id])

    return resolved_teachers, unresolved_names


@admin_bp.route('/api/banned_teachers/import/template', methods=['GET'])
@role_required('超级管理员')
def download_banned_teacher_import_template():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = '禁听教师名单'
    worksheet['A1'] = '教师姓名'
    worksheet.column_dimensions['A'].width = 24
    worksheet['A1'].font = Font(bold=True)
    worksheet['A1'].alignment = Alignment(horizontal='center')

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name='禁听教师导入模板.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@admin_bp.route('/api/banned_teachers/import/preview', methods=['POST'])
@role_required('超级管理员')
def preview_banned_teacher_import():
    try:
        file = request.files.get('file')
        teacher_names = _read_banned_teacher_import_names(file)
        if not teacher_names:
            return jsonify({'success': False, 'message': 'Excel中未读取到教师姓名，请按模板填写后重试'})

        preview_data = _build_banned_teacher_import_preview(teacher_names)
        return jsonify({'success': True, 'preview': preview_data})
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    except Exception as exc:
        return jsonify({'success': False, 'message': f'预览禁听教师导入失败：{str(exc)}'}), 500


@admin_bp.route('/api/banned_teachers/import/apply', methods=['POST'])
@role_required('超级管理员')
def apply_banned_teacher_import():
    try:
        file = request.files.get('file')
        teacher_names = _read_banned_teacher_import_names(file)
        if not teacher_names:
            return jsonify({'success': False, 'message': 'Excel中未读取到教师姓名，请按模板填写后重试'})

        preview_data = _build_banned_teacher_import_preview(teacher_names)
        selection_map_raw = request.form.get('selected_teacher_ids_json', '').strip()
        try:
            selection_map = json.loads(selection_map_raw) if selection_map_raw else {}
        except json.JSONDecodeError:
            return jsonify({'success': False, 'message': '同名教师选择数据格式无效，请重新预览后再试'}), 400

        resolved_teachers, unresolved_names = _resolve_banned_teacher_import_selection(preview_data, selection_map)
        if unresolved_names:
            unresolved_display = '、'.join(unresolved_names[:10])
            suffix = ' 等' if len(unresolved_names) > 10 else ''
            return jsonify({
                'success': False,
                'message': f'以下同名教师尚未完成选择：{unresolved_display}{suffix}'
            }), 400
        if not resolved_teachers:
            return jsonify({'success': False, 'message': '未匹配到可导入的教师，请检查名单后重试'}), 400

        target_users = _get_default_ban_target_users()
        target_user_ids = [user.id for user in target_users]
        if not target_user_ids:
            return jsonify({'success': False, 'message': '当前没有可执行禁听的用户，无法导入'}), 400

        teachers_to_apply = [teacher for teacher in resolved_teachers if teacher.get('needs_apply')]
        selected_teacher_ids = [teacher['teacher_id'] for teacher in teachers_to_apply]
        all_selected_teacher_ids = [teacher['teacher_id'] for teacher in resolved_teachers]

        course_ids = []
        for teacher in teachers_to_apply:
            course_ids.extend(teacher.get('course_ids', []))
        course_ids = sorted(set(course_ids))

        existing_pairs = set()
        if course_ids:
            existing_rows = ListeningBan.query.with_entities(
                ListeningBan.course_id,
                ListeningBan.user_id
            ).filter(
                ListeningBan.course_id.in_(course_ids),
                ListeningBan.user_id.in_(target_user_ids)
            ).all()
            existing_pairs = {(course_id, user_id) for course_id, user_id in existing_rows}

        added_count = 0
        skipped_count = 0
        new_bans = []
        created_by = session.get('user_id')
        for teacher in teachers_to_apply:
            for course_id in teacher.get('course_ids', []):
                for user_id in target_user_ids:
                    pair = (course_id, user_id)
                    if pair in existing_pairs:
                        skipped_count += 1
                        continue
                    existing_pairs.add(pair)
                    new_bans.append(ListeningBan(
                        course_id=course_id,
                        user_id=user_id,
                        created_by=created_by
                    ))
                    added_count += 1

        if new_bans:
            db.session.bulk_save_objects(new_bans)
        db.session.commit()

        return jsonify({
            'success': True,
            'message': f'导入完成，新增 {added_count} 条禁听记录，跳过 {skipped_count} 条已存在记录',
            'result': {
                'selected_teacher_count': len(all_selected_teacher_ids),
                'applied_teacher_count': len(selected_teacher_ids),
                'new_teacher_names': [teacher['teacher_name'] for teacher in resolved_teachers if teacher['status'] == 'new'],
                'partial_teacher_names': [teacher['teacher_name'] for teacher in resolved_teachers if teacher['status'] == 'partial'],
                'already_banned_teacher_names': [teacher['teacher_name'] for teacher in resolved_teachers if teacher['status'] == 'already'],
                'no_course_teacher_names': [teacher['teacher_name'] for teacher in resolved_teachers if teacher['status'] == 'no_courses'],
                'unmatched_names': preview_data.get('unmatched_names', []),
            }
        })
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    except Exception as exc:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'导入禁听教师失败：{str(exc)}'}), 500


@admin_bp.route('/api/users/for_ban', methods=['GET'])
@role_required('超级管理员')
def get_users_for_ban():
    """获取可用于禁听的用户列表（按部门-小组-用户层级结构）"""
    try:
        search = request.args.get('search', '')
        
        query = _active_user_query().filter(User.role == '信息员')
        
        if search:
            query = query.filter(
                db.or_(
                    User.name.contains(search),
                    User.number.contains(search),
                    User.student_id.contains(search),
                    User.department.contains(search),
                    User.college.contains(search)
                )
            )
        
        users = query.all()
        
        # 按部门-小组-用户的层级结构组织数据
        departments = {}
        
        for user in users:
            dept_name = user.department or UNASSIGNED_DEPARTMENT_NAME
            group_name = user.group or UNASSIGNED_GROUP_NAME
            
            # 初始化部门
            if dept_name not in departments:
                departments[dept_name] = {
                    'name': dept_name,
                    'groups': {},
                    'user_count': 0
                }
            
            # 初始化小组
            if group_name not in departments[dept_name]['groups']:
                departments[dept_name]['groups'][group_name] = {
                    'name': group_name,
                    'users': [],
                    'user_count': 0
                }
            
            # 添加用户
            user_data = {
                'id': user.id,
                'name': user.name,
                'number': user.number,
                'student_id': user.student_id,
                'department': user.department,
                'group': user.group,
                'college': user.college,
                'grade': user.grade,
                'major': user.major
            }
            
            departments[dept_name]['groups'][group_name]['users'].append(user_data)
            departments[dept_name]['groups'][group_name]['user_count'] += 1
            departments[dept_name]['user_count'] += 1
        
        # 转换为列表格式
        dept_list = []
        for dept_name, dept_data in departments.items():
            group_list = []
            for group_name, group_data in dept_data['groups'].items():
                group_list.append({
                    'name': group_name,
                    'users': group_data['users'],
                    'user_count': group_data['user_count']
                })
            
            dept_list.append({
                'name': dept_name,
                'groups': group_list,
                'user_count': dept_data['user_count']
            })
        
        # 按部门名称排序
        dept_list.sort(key=lambda x: x['name'])
        for dept in dept_list:
            dept['groups'].sort(key=lambda x: x['name'])
        
        return jsonify({
            'success': True, 
            'departments': dept_list,
            'total_users': len(users)
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取用户列表失败：{str(e)}'})

@admin_bp.route('/api/courses/columns', methods=['GET'])
@login_required
def get_course_columns():
    """获取课程表格可显示的列信息"""
    # 默认列配置
    default_columns = [
        {'key': 'id', 'label': 'ID', 'default': False},
        {'key': 'course_code', 'label': '课程号', 'default': True},
        {'key': 'selection_code', 'label': '选课编号', 'default': True},
        {'key': 'course_name', 'label': '课程名称', 'default': True},
        {'key': 'teacher_id', 'label': '教工号', 'default': True},
        {'key': 'teacher_name', 'label': '教师姓名', 'default': True},
        {'key': 'venue_id', 'label': '场地编号', 'default': True},
        {'key': 'venue_name', 'label': '场地名称', 'default': False},
        {'key': 'offering_college', 'label': '开课学院', 'default': True},
        {'key': 'class_size', 'label': '教学班人数', 'default': False},
        {'key': 'credits', 'label': '学分', 'default': False},
        {'key': 'total_hours', 'label': '总学时', 'default': False},
        {'key': 'class_time', 'label': '上课时间', 'default': True},
        {'key': 'class_location', 'label': '上课地点', 'default': False},
        {'key': 'course_nature', 'label': '课程性质', 'default': False},
        {'key': 'semester', 'label': '学期', 'default': True},
        {'key': 'academic_year', 'label': '学年', 'default': False},
        {'key': 'weekday', 'label': '星期几', 'default': False},
        {'key': 'class_period', 'label': '上课节次', 'default': False},
        {'key': 'start_week', 'label': '起始周', 'default': False},
        {'key': 'enrollment_count', 'label': '选课人数', 'default': False},
        {'key': 'weekly_hours', 'label': '周学时', 'default': False},
        {'key': 'class_composition', 'label': '教学班组成', 'default': False},
        {'key': 'major_composition', 'label': '专业组成', 'default': False},
        {'key': 'venue_start_week', 'label': '场地上课起始周', 'default': False},
        {'key': 'venue_class_period', 'label': '场地上课节次', 'default': False}
    ]
    
    # 尝试从系统设置获取默认显示的列（由超级管理员设置）
    try:
        saved_setting = SystemSetting.query.filter_by(key='course_feedback_default_columns').first()
        if saved_setting and saved_setting.value:
            saved_visible_keys = json.loads(saved_setting.value)
            # 更新默认可见性
            for col in default_columns:
                col['default'] = col['key'] in saved_visible_keys
    except Exception as e:
        current_app.logger.error(f"获取默认列设置失败: {e}")
    
    return jsonify({'success': True, 'columns': default_columns})

@admin_bp.route('/api/courses/columns', methods=['POST'])
@role_required('超级管理员')
def save_course_columns():
    """保存课程表格默认显示的列信息（仅超级管理员）"""
    try:
        data = request.get_json()
        visible_columns = data.get('visible_columns', [])
        
        if not isinstance(visible_columns, list):
            return jsonify({'success': False, 'message': '数据格式错误'}), 400
            
        # 保存到系统设置
        setting = SystemSetting.query.filter_by(key='course_feedback_default_columns').first()
        if not setting:
            setting = SystemSetting(key='course_feedback_default_columns', value=json.dumps(visible_columns))
            db.session.add(setting)
        else:
            setting.value = json.dumps(visible_columns)
            
        db.session.commit()
        return jsonify({'success': True, 'message': '默认列设置已保存'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'保存失败: {str(e)}'}), 500

# 数据导出相关路由（框架性实现）
@admin_bp.route('/export_forms')
@role_required('超级管理员')
def export_forms():
    """听课记录导出（框架性实现）"""
    try:
        # TODO: 实现真正的听课记录导出功能
        # 这里暂时返回一个提示信息，避免路由错误
        flash('听课记录导出功能正在开发中，敬请期待！', 'info')
        return redirect(url_for('admin.system_management'))
    except Exception as e:
        flash(f'导出功能暂不可用：{str(e)}', 'error')
        return redirect(url_for('admin.system_management'))

@admin_bp.route('/api/export/schedule')
@role_required('超级管理员')
def export_schedule():
    """全校课表导出API（框架性实现）"""
    try:
        # TODO: 实现真正的全校课表导出功能
        return jsonify({
            'success': False,
            'message': '全校课表导出功能正在开发中，敬请期待！'
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'导出功能暂不可用：{str(e)}'
        })

@admin_bp.route('/api/export/contacts')
@role_required('超级管理员')
def export_contacts():
    """通讯录导出API
    支持参数：department（可选）、ignore_role_errors（可选）
    按主任、部长、信息员分段导出，包含文档元数据，提供进度查询。
    """
    try:
        from flask import current_app
        department = request.args.get('department')
        ignore_role_errors = (request.args.get('ignore_role_errors') or '').lower() in ('1', 'true', 'yes', 'on')
        role_error_message = '自动分配主任、部长信息时出现错误，请检查“主任”和“部长”的群组名称'

        # 生成任务ID并初始化进度
        job_id = secrets.token_hex(8)
        progress = current_app.config.setdefault('EXPORT_PROGRESS', {})
        progress[job_id] = {
            'status': 'running',
            'percent': 0,
            'message': '准备导出...'
        }

        # 查询用户
        query = _active_user_query()
        if department:
            query = query.filter_by(department=department)

        users = query.all()
        total = len(users)
        progress[job_id]['percent'] = 5
        progress[job_id]['message'] = f'查询到 {total} 条记录'

        def contact_number_sort_key(user):
            number = str(user.number or '').strip()
            if number.isdigit():
                return (0, int(number), number)
            return (1, number)

        group_ids = {u.group_id for u in users if u.group_id}
        groups_by_id = {
            group.id: group
            for group in Group.query.filter(Group.id.in_(group_ids)).all()
        } if group_ids else {}

        def current_group_name(user):
            if user.group_id and user.group_id in groups_by_id:
                return (groups_by_id[user.group_id].name or '').strip()
            return (user.group or '').strip()

        directors = [u for u in users if (u.department or '').strip() == '主任']
        ministers = [
            u for u in users
            if (u.department or '').strip() != '主任' and current_group_name(u) == '部长'
        ]
        informants = [
            u for u in users
            if (u.department or '').strip() != '主任' and current_group_name(u) != '部长'
        ]

        departments_in_scope = {
            (u.department or '').strip()
            for u in users
            if (
                (u.department or '').strip()
                and (u.department or '').strip() not in ('主任', UNASSIGNED_DEPARTMENT_NAME)
            )
        }
        minister_departments = {(u.department or '').strip() for u in ministers}
        missing_minister_departments = departments_in_scope - minister_departments
        role_error_details = []
        if not directors:
            role_error_details.append('不能检测到主任这一群组')
        if missing_minister_departments:
            missing_departments = '、'.join(sorted(missing_minister_departments))
            role_error_details.append(f'不能检测到以下部门的部长：{missing_departments}')
        if role_error_details and not ignore_role_errors:
            detailed_role_error_message = f'{role_error_message}：{"；".join(role_error_details)}'
            progress[job_id]['status'] = 'failed'
            progress[job_id]['message'] = detailed_role_error_message
            return jsonify({'success': False, 'message': detailed_role_error_message}), 400

        directors.sort(key=contact_number_sort_key)
        ministers.sort(key=contact_number_sort_key)
        informants.sort(key=contact_number_sort_key)

        # 加载模板工作簿，保留模板列宽等基础设置，再重建导出内容。
        template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
        import openpyxl
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active

        for merged_range in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merged_range))

        # 清空除模板表头外的内容，并把模板表头下移到第二行。
        max_row = ws.max_row
        if max_row > 1:
            ws.delete_rows(2, max_row - 1)
        ws.insert_rows(1)

        columns = ['编号','部门/组别','姓名','性别','年级','学院','专业','宿舍','手机号码','QQ号码','学号']
        center_alignment = Alignment(horizontal='center', vertical='center')
        title_font = Font(name='黑体', size=16)
        body_font = Font(name='微软雅黑', size=11)
        section_font = Font(name='微软雅黑', size=11, bold=True)
        grey_fill = PatternFill(fill_type='solid', start_color='FFD8D8D8', end_color='FFD8D8D8')
        thin_black_side = Side(style='thin', color='FF000000')
        thin_black_border = Border(
            left=thin_black_side,
            right=thin_black_side,
            top=thin_black_side,
            bottom=thin_black_side
        )

        def apply_content_border(cell):
            cell.border = thin_black_border

        def display_width(value):
            text = str(value or '')
            if text.isdigit():
                return len(text) * 1.25
            return sum(2 if ord(char) > 127 else 1 for char in text)

        def adjust_contact_column_widths(last_row):
            min_widths = {
                9: 16,   # 手机号码
                10: 14,  # QQ号码
                11: 18,  # 学号
            }
            for col_idx in range(1, 12):
                max_width = 0
                for row_idx in range(2, last_row + 1):
                    value = ws.cell(row=row_idx, column=col_idx).value
                    if value is not None:
                        max_width = max(max_width, display_width(value))
                column_letter = openpyxl.utils.get_column_letter(col_idx)
                min_width = min_widths.get(col_idx, 8)
                ws.column_dimensions[column_letter].width = max(min_width, min(max_width + 3, 60))
                ws.column_dimensions[column_letter].bestFit = True

        def apply_title_row(row_idx):
            for col_idx in range(1, 12):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = title_font
                cell.alignment = center_alignment
                apply_content_border(cell)
            ws.cell(row=row_idx, column=1, value='西南大学学生教学信息中心通讯录')
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
            ws.row_dimensions[row_idx].height = 28

        def apply_header_row(row_idx):
            for col_idx, header in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=header)
                cell.font = body_font
                cell.alignment = center_alignment
                cell.number_format = '@'
                apply_content_border(cell)

        def write_section_row(row_idx, title):
            for col_idx in range(1, 12):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = section_font
                cell.alignment = center_alignment
                cell.fill = grey_fill
                cell.number_format = '@'
                apply_content_border(cell)
            ws.cell(row=row_idx, column=1, value=title)
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
            ws.row_dimensions[row_idx].height = 22

        def write_user_row(row_idx, user):
            values = [
                user.number,
                user.department,
                user.name,
                user.gender,
                user.grade,
                user.college,
                user.major,
                user.dormitory,
                user.phone,
                user.qq,
                user.student_id
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else '')
                cell.font = body_font
                cell.alignment = center_alignment
                cell.number_format = '@'
                apply_content_border(cell)
                if col_idx == 1:
                    cell.fill = grey_fill

        apply_title_row(1)
        apply_header_row(2)

        row_idx = 3
        written_count = 0
        if ignore_role_errors:
            for user_item in sorted(users, key=contact_number_sort_key):
                write_user_row(row_idx, user_item)
                row_idx += 1
                written_count += 1
                if total and written_count % max(1, total // 20) == 0:
                    progress[job_id]['percent'] = 5 + int(written_count / total * 90)
                    progress[job_id]['message'] = f'已写入 {written_count}/{total}'
        else:
            sections = [
                ('中心主任', directors),
                ('中心部长', ministers),
                ('中心信息员', informants),
            ]
            for section_title, section_users in sections:
                write_section_row(row_idx, section_title)
                row_idx += 1
                for user_item in section_users:
                    write_user_row(row_idx, user_item)
                    row_idx += 1
                    written_count += 1
                    if total and written_count % max(1, total // 20) == 0:
                        progress[job_id]['percent'] = 5 + int(written_count / total * 90)
                        progress[job_id]['message'] = f'已写入 {written_count}/{total}'

        last_data_row = max(row_idx - 1, 2)
        ws.auto_filter.ref = f'A2:K{last_data_row}'
        ws.freeze_panes = 'A3'
        adjust_contact_column_widths(last_data_row)

        # 设置文档元数据
        from openpyxl.packaging.core import DocumentProperties
        wb.properties = DocumentProperties(
            creator='系统管理员',
            lastModifiedBy='系统管理员',
            title='部门通讯录',
            subject='部门通讯录导出',
            keywords='通讯录,部门,信息员',
            category='导出文件'
        )

        # 保存到导出目录
        export_dir = env_path('EXPORT_DIR', DEFAULT_EXPORT_DIR)
        os.makedirs(export_dir, exist_ok=True)
        fname_base = f"contacts_{(department or '全部')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        export_filename = f"{fname_base}.xlsx"
        export_path = os.path.join(export_dir, export_filename)
        wb.save(export_path)

        download_url = url_for('admin.download_passwords', filename=export_filename).replace('download_passwords', 'download_passwords')
        progress[job_id]['percent'] = 100
        progress[job_id]['status'] = 'completed'
        progress[job_id]['message'] = '导出完成'
        progress[job_id]['download_url'] = url_for('admin.download_passwords', filename=export_filename)

        current_app.logger.info(f'通讯录导出完成 记录数{total} 部门{department or "全部"} 文件{export_filename}')
        return jsonify({'success': True, 'job_id': job_id, 'download_url': progress[job_id]['download_url'], 'total': total})
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'导出失败：{str(e)}'
        })

@admin_bp.route('/api/export/progress/<job_id>')
@role_required('超级管理员')
def export_progress(job_id):
    """查询导出进度"""
    from flask import current_app
    prog = current_app.config.get('EXPORT_PROGRESS', {})
    if job_id not in prog:
        return jsonify({'success': False, 'message': '任务不存在'}), 404
    return jsonify({'success': True, 'progress': prog[job_id]})

# ==================== 系统管理功能 ====================

@admin_bp.route('/system_management')
@role_required('超级管理员')
def system_management():
    """统一系统设置中心。"""
    active_tab = request.args.get('tab', 'teaching')
    if active_tab not in {'teaching', 'assessment', 'automation', 'imports'}:
        active_tab = 'teaching'
    user = User.query.get(session['user_id'])
    return render_template(
        'admin/system_management.html',
        active_tab=active_tab,
        status=AutoReviewEngine().files_status(),
        available_departments=list(_get_accessible_department_users(user.id).keys()),
        account_username=user.student_id,
        is_super_admin=True,
    )

@admin_bp.route('/api/schedule/validate', methods=['POST'])
@role_required('超级管理员')
def validate_schedule_format():
    """验证全校课表文件格式"""
    try:
        if 'file' not in request.files:
            return jsonify({'success': False, 'message': '未选择文件'})
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'success': False, 'message': '未选择文件'})
        
        if not allowed_file(file.filename):
            return jsonify({'success': False, 'message': '文件格式不支持，请上传.xls或.xlsx文件'})
        
        # 读取模板文件获取标准列名（改为绝对路径，避免依赖CWD）
        template_path = env_path('SCHEDULE_TEMPLATE_PATH', DEFAULT_SCHEDULE_TEMPLATE_PATH)
        if not os.path.exists(template_path):
            return jsonify({'success': False, 'message': '模板文件不存在'})
        
        template_df = pd.read_excel(template_path)
        expected_columns = list(template_df.columns)
        
        # 读取上传的文件
        uploaded_df = pd.read_excel(file)
        uploaded_columns = list(uploaded_df.columns)
        
        # 检查列名是否匹配
        if uploaded_columns != expected_columns:
            missing_columns = set(expected_columns) - set(uploaded_columns)
            extra_columns = set(uploaded_columns) - set(expected_columns)
            
            error_msg = "文件格式不匹配：\n"
            if missing_columns:
                error_msg += f"缺少列：{', '.join(missing_columns)}\n"
            if extra_columns:
                error_msg += f"多余列：{', '.join(extra_columns)}\n"
            
            return jsonify({
                'success': False, 
                'message': error_msg,
                'expected_columns': expected_columns,
                'uploaded_columns': uploaded_columns
            })
        
        # 检查数据类型
        type_errors = []
        for col in expected_columns:
            if col in uploaded_df.columns:
                template_type = template_df[col].dtype
                uploaded_type = uploaded_df[col].dtype
                
                # 对于数值类型，检查是否兼容
                if pd.api.types.is_numeric_dtype(template_type):
                    if not pd.api.types.is_numeric_dtype(uploaded_type):
                        # 尝试转换为数值类型
                        try:
                            pd.to_numeric(uploaded_df[col], errors='coerce')
                        except:
                            type_errors.append(f"列 '{col}' 应为数值类型")
        
        if type_errors:
            return jsonify({
                'success': False,
                'message': "数据类型错误：\n" + "\n".join(type_errors)
            })
        
        return jsonify({
            'success': True,
            'message': '文件格式验证通过',
            'row_count': len(uploaded_df),
            'column_count': len(uploaded_df.columns)
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'验证失败：{str(e)}'})

@admin_bp.route('/api/schedule/import', methods=['POST'])
@role_required('超级管理员')
def import_schedule_data():
    """导入全校课表数据"""
    try:
        if 'file' not in request.files:
            return jsonify({'success': False, 'message': '未选择文件'})
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'success': False, 'message': '未选择文件'})
        
        # 先验证格式
        validation_result = validate_schedule_format()
        if not validation_result.get_json().get('success'):
            return validation_result
        
        # 读取文件
        df = pd.read_excel(file)
        
        # 统计信息
        stats = {
            'total_rows': len(df),
            'teachers_added': 0,
            'teachers_updated': 0,
            'venues_added': 0,
            'venues_updated': 0,
            'courses_added': 0,
            'errors': []
        }
        
        # 处理教师数据
        teacher_data = df[['教工号', '姓名', '性别', '职称名称', '教师所属学院', '教师联系电话']].drop_duplicates(subset=['教工号'])
        
        for _, row in teacher_data.iterrows():
            try:
                teacher_id = str(row['教工号']).strip()
                if not teacher_id or teacher_id == 'nan':
                    continue
                
                teacher = Teacher.query.get(teacher_id)
                if teacher:
                    # 更新现有教师的多值字段
                    teacher.add_value('title', row['职称名称'])
                    teacher.add_value('college', row['教师所属学院'])
                    teacher.add_value('phone', row['教师联系电话'])
                    stats['teachers_updated'] += 1
                else:
                    # 创建新教师
                    teacher = Teacher(
                        teacher_id=teacher_id,
                        name=str(row['姓名']).strip() if pd.notna(row['姓名']) else '',
                        gender=str(row['性别']).strip() if pd.notna(row['性别']) else None,
                        title=str(row['职称名称']).strip() if pd.notna(row['职称名称']) else None,
                        college=str(row['教师所属学院']).strip() if pd.notna(row['教师所属学院']) else None,
                        phone=str(row['教师联系电话']).strip() if pd.notna(row['教师联系电话']) else None
                    )
                    db.session.add(teacher)
                    stats['teachers_added'] += 1
                    
            except Exception as e:
                stats['errors'].append(f"处理教师数据时出错（教工号：{row['教工号']}）：{str(e)}")
        
        # 处理场地数据
        venue_data = df[['场地编号', '场地名称', '场地类别名称', '校区', '楼层号', '教学楼', '座位数']].dropna(subset=['场地编号']).drop_duplicates(subset=['场地编号'])
        
        for _, row in venue_data.iterrows():
            try:
                venue_id = str(row['场地编号']).strip()
                if not venue_id or venue_id == 'nan':
                    continue
                
                venue = Venue.query.get(venue_id)
                if venue:
                    stats['venues_updated'] += 1
                else:
                    venue = Venue(
                        venue_id=venue_id,
                        name=str(row['场地名称']).strip() if pd.notna(row['场地名称']) else None,
                        category=str(row['场地类别名称']).strip() if pd.notna(row['场地类别名称']) else None,
                        campus=str(row['校区']).strip() if pd.notna(row['校区']) else None,
                        floor=float(row['楼层号']) if pd.notna(row['楼层号']) else None,
                        building=str(row['教学楼']).strip() if pd.notna(row['教学楼']) else None,
                        capacity=float(row['座位数']) if pd.notna(row['座位数']) else None
                    )
                    db.session.add(venue)
                    stats['venues_added'] += 1
                    
            except Exception as e:
                stats['errors'].append(f"处理场地数据时出错（场地编号：{row['场地编号']}）：{str(e)}")
        
        # 处理课程数据 - 保留所有记录，不再合并相同课程号和选课课号的记录
        for _, row in df.iterrows():
            try:
                course_code = str(row['课程号']).strip()
                selection_code = str(row['选课课号']).strip() if pd.notna(row['选课课号']) else ''
                
                # 课程号和选课编号都必须存在
                if not course_code or course_code == 'nan' or not selection_code or selection_code == 'nan':
                    stats['errors'].append(f"课程号或选课编号为空（行：{row.name + 2}）")
                    continue
                
                # 直接创建新课程记录，不再查找现有记录
                course = Course(
                    course_code=course_code,
                    selection_code=selection_code,
                    start_week=str(row['起始周']).strip() if pd.notna(row['起始周']) else None,
                    weekday=int(row['星期几']) if pd.notna(row['星期几']) else None,
                    class_period=str(row['上课节次']).strip() if pd.notna(row['上课节次']) else None,
                    course_name=str(row['课程名称']).strip() if pd.notna(row['课程名称']) else '',
                    venue_start_week=str(row['场地上课起始周']).strip() if pd.notna(row['场地上课起始周']) else None,
                    venue_class_period=str(row['场地上课节次']).strip() if pd.notna(row['场地上课节次']) else None,
                    class_size=int(row['教学班人数']) if pd.notna(row['教学班人数']) else None,
                    class_composition=str(row['教学班组成']).strip() if pd.notna(row['教学班组成']) else None,
                    credits=float(row['学分']) if pd.notna(row['学分']) else None,
                    total_hours=float(row['总学时']) if pd.notna(row['总学时']) else None,
                    offering_college=str(row['开课学院']).strip() if pd.notna(row['开课学院']) else None,
                    enrollment_count=int(row['选课人数']) if pd.notna(row['选课人数']) else None,
                    weekly_hours=str(row['周学时']).strip() if pd.notna(row['周学时']) else None,
                    class_time=str(row['上课时间']).strip() if pd.notna(row['上课时间']) else None,
                    class_location=str(row['上课地点']).strip() if pd.notna(row['上课地点']) else None,
                    course_nature=str(row['课程性质']).strip() if pd.notna(row['课程性质']) else None,
                    major_composition=str(row['专业组成']).strip() if pd.notna(row['专业组成']) else None,
                    teacher_id=str(row['教工号']).strip() if pd.notna(row['教工号']) else None,
                    venue_id=str(row['场地编号']).strip() if pd.notna(row['场地编号']) else None,
                    semester=str(row['学期']) if pd.notna(row['学期']) else None,
                    academic_year=str(row['学年']).strip() if pd.notna(row['学年']) else None
                )
                db.session.add(course)
                stats['courses_added'] += 1
                    
            except Exception as e:
                stats['errors'].append(f"处理课程数据时出错（课程号：{course_code}-{selection_code}）：{str(e)}")
        
        # 提交数据库更改
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '数据导入成功',
            'stats': stats
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'导入失败：{str(e)}'})

# ==================== 部门管理 API 路由 ====================

@admin_bp.route('/api/departments', methods=['GET'])
@role_required('管理员')
def get_departments():
    """获取所有部门及其小组和用户信息"""
    user = User.query.get(session['user_id'])
    manage_permission = get_user_manage_permission(user.id)
    
    if not manage_permission:
        return jsonify({'success': False, 'message': '无权访问'}), 403
        
    try:
        # 确定要查询的部门范围
        if manage_permission == '超级管理员':
            departments = Department.query.all()
        else:
            departments = Department.query.filter_by(name=user.department).all()
        has_real_unassigned_department = any(dept.name == UNASSIGNED_DEPARTMENT_NAME for dept in departments)
            
        result = []
        
        for dept in departments:
            groups = Group.query.filter_by(department=dept.name).all()
            
            # 如果是"管理部门小组"权限，只显示自己所在的小组
            if manage_permission == '管理部门小组':
                # 假设user.group存储的是小组名称
                if user.group:
                    groups = [g for g in groups if g.name == user.group]
                else:
                    groups = []
            
            # 计算部门总用户数
            total_users = _active_user_query().filter_by(department=dept.name).count()
            if dept.name == UNASSIGNED_DEPARTMENT_NAME and total_users == 0:
                continue
            
            # 计算部门评价信息 (仅超级管理员和部门管理员可见，小组管理员不可见)
            total_deduction = 0.0
            member_deductions = []
            
            if manage_permission in ['超级管理员', '管理部门']:
                # 1. 获取部门内所有用户
                dept_users = _active_user_query().filter_by(department=dept.name).all()
                dept_user_numbers = [u.number for u in dept_users]
                
                if dept_user_numbers:
                    # 2. 获取这些用户的所有评分记录
                    forms = LectureForm.query.filter(LectureForm.listener_number.in_(dept_user_numbers)).with_entities(LectureForm.id).all()
                    form_ids = [f.id for f in forms]
                    
                    if form_ids:
                        # 部门累计扣分 (假设为 total_department_score 之和)
                        total_deduction = db.session.query(func.sum(ScoreRecord.total_department_score))\
                            .filter(ScoreRecord.form_id.in_(form_ids)).scalar() or 0.0
                        
                        # 部员扣分排行
                        member_stats = db.session.query(
                            LectureForm.listener_number,
                            func.sum(ScoreRecord.total_department_score).label('score')
                        ).join(ScoreRecord, LectureForm.id == ScoreRecord.form_id)\
                         .filter(LectureForm.listener_number.in_(dept_user_numbers))\
                         .group_by(LectureForm.listener_number)\
                         .order_by(func.sum(ScoreRecord.total_department_score).desc())\
                         .all()
                         
                        user_map = {u.number: u.name for u in dept_users}
                        
                        for ms in member_stats:
                            if ms.score > 0:
                                member_deductions.append({
                                    'name': user_map.get(ms.listener_number, ms.listener_number),
                                    'score': ms.score
                                })

            dept_data = {
                'id': dept.id,
                'name': dept.name,
                'description': dept.description,
                'head': dept.head,
                'manager_id': dept.manager_id,
                'total_users': total_users,
                'total_deduction': total_deduction,
                'member_deductions': member_deductions,
                'groups': [],
                'is_virtual': False,
            }
            
            for group in groups:
                users = _active_user_query().filter_by(group_id=group.id).all()
                if group.name == UNASSIGNED_GROUP_NAME and not users:
                    continue
                dept_data['groups'].append(_build_group_payload(group, users))
            
            # 添加没有小组的用户（仅当权限允许时）
            if manage_permission != '管理部门小组':
                no_group_users = _active_user_query().filter_by(department=dept.name, group_id=None).all()
                if no_group_users:
                    no_group_data = _build_group_payload(
                        None,
                        no_group_users,
                        name=UNASSIGNED_GROUP_NAME,
                        department=dept.name,
                        description='未分配到小组的用户',
                        max_members='无限制',
                        is_virtual=True
                    )
                    dept_data['groups'].append(no_group_data)
            
            result.append(dept_data)

        if manage_permission == '超级管理员' and not has_real_unassigned_department:
            unassigned_users = _active_user_query().filter_by(department=UNASSIGNED_DEPARTMENT_NAME).all()
            if unassigned_users:
                result.append({
                    'id': None,
                    'name': UNASSIGNED_DEPARTMENT_NAME,
                    'description': '未分配到部门的用户',
                    'head': None,
                    'manager_id': None,
                    'total_users': len(unassigned_users),
                    'total_deduction': 0.0,
                    'member_deductions': [],
                    'groups': [
                        _build_group_payload(
                            None,
                            unassigned_users,
                            name=UNASSIGNED_GROUP_NAME,
                            department=UNASSIGNED_DEPARTMENT_NAME,
                            description='未分配到小组的用户',
                            max_members='无限制',
                            is_virtual=True
                        )
                    ],
                    'is_virtual': True,
                })
        
        return jsonify({'success': True, 'data': result})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/departments', methods=['POST'])
@role_required('超级管理员')
def add_department():
    """添加新部门"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if manage_permission != '超级管理员':
             return jsonify({'success': False, 'message': '只有超级管理员可以添加部门'}), 403
             
        data = request.get_json()
        name = data.get('name', '').strip()
        description = data.get('description', '').strip()
        manager_id_raw = data.get('manager_id')
        head = data.get('head', '').strip()
        
        if not name:
            return jsonify({'success': False, 'message': '部门名称不能为空'}), 400
        
        # 检查部门名称是否已存在
        existing_dept = Department.query.filter_by(name=name).first()
        if existing_dept:
            return jsonify({'success': False, 'message': '部门名称已存在'}), 400
        
        # 处理负责人ID（优先使用manager_id）
        manager_id = None
        head_name = head if head else None
        manager_user = None
        if manager_id_raw is not None and str(manager_id_raw).strip() != '':
            try:
                manager_id = int(manager_id_raw)
            except ValueError:
                return jsonify({'success': False, 'message': '负责人ID格式不正确'}), 400
            manager_user = User.query.get(manager_id)
            if not manager_user or not is_user_active(manager_user):
                return jsonify({'success': False, 'message': '指定的负责人不存在'}), 400
            head_name = manager_user.name
        
        new_dept = Department(name=name, description=description, manager_id=manager_id, head=head_name)
        db.session.add(new_dept)
        db.session.commit()
        
        # 如果指定了负责人，更新负责人的department
        if manager_user:
            if manager_user.department != name:
                manager_user.department = name
                # 清除旧的小组关联（因为部门变了）
                manager_user.group_id = None
                manager_user.group = UNASSIGNED_GROUP_NAME
                db.session.commit()
        
        return jsonify({'success': True, 'message': '部门添加成功', 'data': {
            'id': new_dept.id,
            'name': new_dept.name,
            'description': new_dept.description,
            'head': new_dept.head,
            'manager_id': new_dept.manager_id
        }})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/departments/<int:dept_id>', methods=['GET'])
@role_required('管理员')
def get_department(dept_id):
    """获取单个部门信息"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权查看部门信息'}), 403
             
        dept = Department.query.get_or_404(dept_id)
        
        # 权限检查
        if manage_permission != '超级管理员':
            if dept.name != user.department:
                 return jsonify({'success': False, 'message': '无权查看其他部门信息'}), 403
        
        return jsonify({'success': True, 'data': {
            'id': dept.id,
            'name': dept.name,
            'description': dept.description,
            'head': dept.head,
            'manager_id': dept.manager_id
        }})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/departments/<int:dept_id>', methods=['PUT'])
@role_required('管理员')
def update_department(dept_id):
    """更新部门信息"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权修改部门信息'}), 403
             
        dept = Department.query.get_or_404(dept_id)
        
        # 权限检查
        if manage_permission != '超级管理员':
            if dept.name != user.department:
                 return jsonify({'success': False, 'message': '无权修改其他部门信息'}), 403
        
        data = request.get_json()
        
        old_name = dept.name
        new_name = data.get('name', '').strip()
        description = data.get('description', '').strip()
        manager_id_raw = data.get('manager_id')
        head = data.get('head', '').strip()
        
        if not new_name:
            return jsonify({'success': False, 'message': '部门名称不能为空'}), 400
        
        # 检查是否尝试修改部门名称
        if new_name != old_name:
            if manage_permission != '超级管理员':
                return jsonify({'success': False, 'message': '无权修改部门名称'}), 403
                
            existing_dept = Department.query.filter_by(name=new_name).first()
            if existing_dept:
                return jsonify({'success': False, 'message': '部门名称已存在'}), 400
            
            # 更新相关用户和小组的部门信息
            User.query.filter_by(department=old_name).update({'department': new_name})
            Group.query.filter_by(department=old_name).update({'department': new_name})
        
        dept.name = new_name
        dept.description = description
        
        # 更新负责人（优先使用manager_id）
        if manager_id_raw is not None:
            if str(manager_id_raw).strip() == '':
                dept.manager_id = None
                dept.head = None
            else:
                try:
                    manager_id = int(manager_id_raw)
                except ValueError:
                    return jsonify({'success': False, 'message': '负责人ID格式不正确'}), 400
                manager_user = User.query.get(manager_id)
                if not manager_user or not is_user_active(manager_user):
                    return jsonify({'success': False, 'message': '指定的负责人不存在'}), 400
                
                # 检查负责人是否属于该部门
                if manager_user.department != new_name: # new_name 是更新后的部门名
                     # 自动将负责人分配至该部门
                     manager_user.department = new_name
                     # 如果之前有小组且小组不属于新部门，则清除小组（或保持不变，视业务逻辑而定）
                     # 简单起见，如果部门变了，小组关联也应该重置，除非小组也迁移了
                     if manager_user.group_id:
                         group = Group.query.get(manager_user.group_id)
                         if group and group.department != new_name:
                             manager_user.group_id = None
                             manager_user.group = UNASSIGNED_GROUP_NAME

                dept.manager_id = manager_id
                dept.head = manager_user.name
        else:
            # 兼容旧字段
            dept.head = head
        
        db.session.commit()
        
        return jsonify({'success': True, 'message': '部门更新成功'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/departments/<int:dept_id>/disband', methods=['POST'])
@admin_bp.route('/api/departments/<int:dept_id>', methods=['DELETE'])
@role_required('超级管理员')
def disband_department(dept_id):
    """解散部门：删除组织记录，成员转入未分配部门。"""
    try:
        user = User.query.get(session['user_id'])
        password = request.get_json().get('password') if request.get_json() else None
        
        # 验证密码
        if not password or not check_password_hash(user.password_hash, password):
            return jsonify({'success': False, 'message': '密码验证失败，无法解散部门'}), 403
            
        dept = Department.query.get_or_404(dept_id)
        if dept.name == UNASSIGNED_DEPARTMENT_NAME:
            return jsonify({'success': False, 'message': '未分配部门不能解散'}), 400
        
        dept_name = dept.name
        dept_users = User.query.filter_by(department=dept_name).all()
        moved_count = 0
        for target_user in dept_users:
            before_snapshot = _snapshot_user_for_movement(target_user)
            _set_user_unassigned(target_user)
            after_snapshot = _snapshot_user_for_movement(target_user)
            moved_count += 1
            _create_personnel_movement_record(
                user.id,
                {
                    **after_snapshot,
                    'department_before': before_snapshot.get('department'),
                    'department_after': after_snapshot.get('department'),
                    'group_before': before_snapshot.get('group'),
                    'group_after': after_snapshot.get('group'),
                },
                '部门解散',
                f"{after_snapshot.get('name')}（{after_snapshot.get('number')}）因部门解散转入未分配部门",
                [{
                    'field': 'department',
                    'label': '部门',
                    'before': before_snapshot.get('department') or '未设置',
                    'after': UNASSIGNED_DEPARTMENT_NAME
                }, {
                    'field': 'group',
                    'label': '小组',
                    'before': before_snapshot.get('group') or UNASSIGNED_GROUP_NAME,
                    'after': UNASSIGNED_GROUP_NAME
                }]
            )

        # 删除部门下的小组记录，用户已转入未分配部门，不再保留原小组。
        groups = Group.query.filter_by(department=dept.name).all()
        for group in groups:
            db.session.delete(group)
        
        db.session.delete(dept)
        db.session.commit()
        
        return jsonify({'success': True, 'message': f'部门已解散，{moved_count} 名用户已转入未分配部门'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'解散失败: {str(e)}'}), 500

# ==================== 小组管理 API 路由 ====================

@admin_bp.route('/api/groups', methods=['GET'])
@role_required('管理员')
def get_groups_by_department():
    """按部门获取小组"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权查看小组'}), 403
             
        department = request.args.get('department')
        
        # 权限过滤
        if manage_permission == '超级管理员':
            if department:
                groups = Group.query.filter_by(department=department).all()
            else:
                groups = Group.query.all()
        elif manage_permission == '管理部门':
            # 只能查看本部门的小组
            # 即使前端没传department参数，也只返回本部门的
            # 如果前端传了其他部门，则返回空或报错
            if department and department != user.department:
                return jsonify({'success': False, 'message': '无权查看其他部门的小组'}), 403
            groups = Group.query.filter_by(department=user.department).all()
        elif manage_permission == '管理部门小组':
            # 只能查看本小组
            if user.group:
                # 假设user.group存储的是小组名称
                groups = Group.query.filter_by(department=user.department, name=user.group).all()
            else:
                groups = []
        
        groups_data = []
        for group in groups:
            groups_data.append({
                'id': group.id,
                'name': group.name,
                'department': group.department,
                'leader': group.leader,
                'description': group.description,
                'max_members': group.max_members
            })
        
        return jsonify({'success': True, 'data': groups_data})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/groups', methods=['POST'])
@role_required('管理员')
def add_group():
    """添加新小组"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
            return jsonify({'success': False, 'message': '无权添加小组'}), 403
            
        data = request.get_json()
        name = data.get('name', '').strip()
        department = data.get('department', '').strip()
        leader = data.get('leader', '').strip()
        description = data.get('description', '').strip()
        max_members = data.get('max_members', 10)
        
        if not name:
            return jsonify({'success': False, 'message': '小组名称不能为空'}), 400
        
        if not department:
            return jsonify({'success': False, 'message': '所属部门不能为空'}), 400
            
        # 权限检查
        if manage_permission != '超级管理员':
            if department != user.department:
                return jsonify({'success': False, 'message': '只能在自己的部门添加小组'}), 403
        
        # 检查部门是否存在
        dept = Department.query.filter_by(name=department).first()
        if not dept and department != UNASSIGNED_DEPARTMENT_NAME:
            return jsonify({'success': False, 'message': '指定的部门不存在'}), 400
        
        # 检查小组名称在该部门内是否已存在
        existing_group = Group.query.filter_by(name=name, department=department).first()
        if existing_group:
            return jsonify({'success': False, 'message': '该部门内已存在同名小组'}), 400
        
        # 处理组长ID
        leader_id = data.get('leader_id')
        leader_name = leader
        
        if leader_id:
            leader_user = User.query.get(leader_id)
            if leader_user and is_user_active(leader_user):
                leader_name = leader_user.name
                # 稍后在new_group创建后，我们需要更新leader_user的group_id
                # 但由于new_group还没ID，我们只能在commit之后更新，或者先add再commit再更新
            else:
                return jsonify({'success': False, 'message': '指定的小组长不存在或已离任'}), 400
        elif leader:
            # 尝试根据名字查找
            leader_user = _active_user_query().filter_by(name=leader, department=department).first()
            if leader_user:
                leader_id = leader_user.id
        
        new_group = Group(
            name=name,
            department=department,
            # department_id=dept.id, # Group表无此字段
            leader=leader_name,
            leader_id=leader_id,
            description=description,
            max_members=max_members
        )
        db.session.add(new_group)
        db.session.commit()
        
        # 如果指定了组长，更新组长的group_id
        if leader_id:
            leader_user = User.query.get(leader_id)
            if leader_user and is_user_active(leader_user):
                leader_user.group_id = new_group.id
                leader_user.group = new_group.name
                if leader_user.department != department:
                    leader_user.department = department
                db.session.commit()
        
        return jsonify({'success': True, 'message': '小组添加成功', 'data': {
            'id': new_group.id,
            'name': new_group.name,
            'department': new_group.department,
            'leader': new_group.leader,
            'leader_id': new_group.leader_id,
            'description': new_group.description,
            'max_members': new_group.max_members
        }})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/groups/<int:group_id>', methods=['GET'])
@role_required('管理员')
def get_group(group_id):
    """获取单个小组信息"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权查看小组信息'}), 403
             
        group = Group.query.get_or_404(group_id)
        
        # 权限检查
        if manage_permission != '超级管理员':
            if group.department != user.department:
                 return jsonify({'success': False, 'message': '无权查看其他部门的小组'}), 403
            
            if manage_permission == '管理部门小组':
                # 假设user.group存储的是小组名称
                if group.name != user.group:
                     return jsonify({'success': False, 'message': '无权查看其他小组'}), 403
        
        return jsonify({'success': True, 'data': {
            'id': group.id,
            'name': group.name,
            'department': group.department,
            'leader': group.leader,
            'leader_id': group.leader_id,
            'description': group.description,
            'max_members': group.max_members
        }})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/groups/<int:group_id>', methods=['PUT'])
@role_required('管理员')
def update_group(group_id):
    """更新小组信息"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权修改小组'}), 403
             
        group = Group.query.get_or_404(group_id)
        
        # 权限检查
        if manage_permission != '超级管理员':
            if manage_permission == '管理部门':
                if group.department != user.department:
                     return jsonify({'success': False, 'message': '无权修改其他部门的小组'}), 403
            elif manage_permission == '管理部门小组':
                # 假设user.group存储的是小组名称
                if group.name != user.group or group.department != user.department:
                     return jsonify({'success': False, 'message': '无权修改其他小组'}), 403
        
        data = request.get_json()
        
        name = data.get('name', '').strip()
        department = data.get('department', '').strip()
        leader = data.get('leader', '').strip()
        description = data.get('description', '').strip()
        max_members = data.get('max_members', 10)
        
        if not name:
            return jsonify({'success': False, 'message': '小组名称不能为空'}), 400
        
        if not department:
            return jsonify({'success': False, 'message': '所属部门不能为空'}), 400

        old_group_name = group.name
        old_department = group.department
            
        # 检查部门变更权限
        if manage_permission != '超级管理员':
            if department != group.department:
                 return jsonify({'success': False, 'message': '无权更改小组所属部门'}), 403
        
        # 检查部门是否存在
        dept = Department.query.filter_by(name=department).first()
        if not dept and department != UNASSIGNED_DEPARTMENT_NAME:
            return jsonify({'success': False, 'message': '指定的部门不存在'}), 400
        
        # 如果名称或部门发生变化，检查新组合是否已存在
        if name != old_group_name or department != old_department:
            existing_group = Group.query.filter_by(name=name, department=department).first()
            if existing_group and existing_group.id != group_id:
                return jsonify({'success': False, 'message': '该部门内已存在同名小组'}), 400
        
        # 处理组长ID
        leader_id = data.get('leader_id')
        leader_name = leader
        
        if leader_id:
            try:
                # 确保leader_id是整数
                leader_id = int(leader_id) if leader_id else None
            except ValueError:
                leader_id = None
                
            if leader_id:
                leader_user = User.query.get(leader_id)
                if leader_user and is_user_active(leader_user):
                    leader_name = leader_user.name
                    
                    # 自动将组长分配至该小组
                    if leader_user.group_id != group.id:
                        leader_user.group_id = group.id
                        leader_user.group = name # 更新组名
                        # 同时也需确保部门一致
                        if leader_user.department != department:
                            leader_user.department = department
                else:
                    return jsonify({'success': False, 'message': '指定的小组长不存在或已离任'}), 400
        elif leader:
            # 尝试根据名字查找
            leader_user = _active_user_query().filter_by(name=leader, department=department).first()
            if leader_user:
                leader_id = leader_user.id
        
        group.name = name
        group.department = department
        # group.department_id = dept.id # Group表无此字段
        group.leader = leader_name
        group.leader_id = leader_id
        group.description = description
        group.max_members = max_members
        
        # 用户表仍保留 group 文本字段做兼容，改名/改部门时同步冗余字段。
        member_updates = {}
        if name != old_group_name:
            member_updates['group'] = name
        if department != old_department:
            member_updates['department'] = department
        if member_updates:
            User.query.filter_by(group_id=group.id).update(member_updates)
        
        db.session.commit()
        
        return jsonify({'success': True, 'message': '小组更新成功'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/groups/<int:group_id>/disband', methods=['POST'])
@admin_bp.route('/api/groups/<int:group_id>', methods=['DELETE'])
@role_required('管理员')
def disband_group(group_id):
    """解散小组：删除小组记录，成员转入未分配小组。"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权解散小组'}), 403
              
        group = Group.query.get_or_404(group_id)
        
        # 权限检查
        if manage_permission != '超级管理员':
            if manage_permission == '管理部门':
                if group.department != user.department:
                     return jsonify({'success': False, 'message': '无权解散其他部门的小组'}), 403
            else:
                 return jsonify({'success': False, 'message': '无权解散小组'}), 403
         
        # 验证密码
        data = request.get_json() or {}
        password = data.get('password')
        if not password or not check_password_hash(user.password_hash, password):
            return jsonify({'success': False, 'message': '密码验证失败，无法解散小组'}), 403

        group_users = User.query.filter(
            db.or_(
                User.group_id == group.id,
                db.and_(User.department == group.department, User.group == group.name)
            )
        ).all()
        moved_count = 0
        for target_user in group_users:
            before_snapshot = _snapshot_user_for_movement(target_user)
            target_user.group_id = None
            target_user.group = UNASSIGNED_GROUP_NAME
            after_snapshot = _snapshot_user_for_movement(target_user)
            moved_count += 1
            _create_personnel_movement_record(
                user.id,
                {
                    **after_snapshot,
                    'department_before': before_snapshot.get('department'),
                    'department_after': after_snapshot.get('department'),
                    'group_before': before_snapshot.get('group'),
                    'group_after': after_snapshot.get('group'),
                },
                '小组解散',
                f"{after_snapshot.get('name')}（{after_snapshot.get('number')}）因小组解散转入未分配小组",
                [{
                    'field': 'group',
                    'label': '小组',
                    'before': before_snapshot.get('group') or UNASSIGNED_GROUP_NAME,
                    'after': UNASSIGNED_GROUP_NAME
                }]
            )
         
        db.session.delete(group)
        db.session.commit()
         
        return jsonify({'success': True, 'message': f'小组已解散，{moved_count} 名用户已转入未分配小组'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'解散失败: {str(e)}'}), 500

@admin_bp.route('/api/groups/<int:group_id>/move_members', methods=['POST'])
@role_required('管理员')
def move_members_to_group(group_id):
    """批量移动成员到小组"""
    try:
        user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(user.id)
        
        if not manage_permission:
             return jsonify({'success': False, 'message': '无权操作'}), 403
             
        group = Group.query.get_or_404(group_id)
        
        # 权限检查
        if manage_permission != '超级管理员':
            if manage_permission == '管理部门':
                if group.department != user.department:
                     return jsonify({'success': False, 'message': '无权操作其他部门的小组'}), 403
            elif manage_permission == '管理部门小组':
                # 假设user.group存储的是小组名称
                if group.name != user.group or group.department != user.department:
                     return jsonify({'success': False, 'message': '无权操作其他小组'}), 403
        
        data = request.get_json()
        user_ids = data.get('user_ids', [])
        
        if not user_ids:
            return jsonify({'success': False, 'message': '请选择要移动的用户'}), 400
            
        # 验证用户是否属于该部门（只能在同部门内移动到小组）
        # 除非是超级管理员，可能允许跨部门（但业务逻辑上小组属于部门，跨部门移动到小组意味着换部门）
        # 这里简化逻辑：移动到小组 = 更新 group_id 和 group
        # 同时也检查/更新 department
        
        users = _active_user_query().filter(User.id.in_(user_ids)).all()
        count = 0
        
        for u in users:
            # 权限检查：能否操作该用户
            if manage_permission != '超级管理员':
                if u.department != user.department:
                    continue # 跳过非本部门用户
            
            u.group_id = group.id
            u.group = group.name
            
            # 如果部门不一致，更新部门（注意：这可能需要更高权限，这里假设移动到小组就隐含了部门变更）
            if u.department != group.department:
                if manage_permission == '超级管理员':
                    u.department = group.department
                else:
                    # 普通管理员不能跨部门拉人
                    continue 
            
            count += 1
            
        db.session.commit()
        
        return jsonify({'success': True, 'message': f'成功移动 {count} 名用户到小组'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/departments/<int:dept_id>/move_members', methods=['POST'])
@role_required('超级管理员')
def move_members_to_department(dept_id):
    """批量移动成员到部门"""
    try:
        dept = Department.query.get_or_404(dept_id)
        
        data = request.get_json()
        user_ids = data.get('user_ids', [])
        
        if not user_ids:
            return jsonify({'success': False, 'message': '请选择要移动的用户'}), 400
            
        users = _active_user_query().filter(User.id.in_(user_ids)).all()
        
        for u in users:
            u.department = dept.name
            u.group_id = None
            u.group = UNASSIGNED_GROUP_NAME
            
        db.session.commit()
        
        return jsonify({'success': True, 'message': f'成功移动 {len(users)} 名用户到部门'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

# ==================== 用户管理 API 路由 ====================

@admin_bp.route('/api/users', methods=['GET'])
@role_required('管理员')
def get_users():
    """获取用户（支持按部门筛选）"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        if not manage_permission:
            return jsonify({'success': False, 'message': '无权查看用户'}), 403
            
        department = request.args.get('department')
        
        query = _active_user_query()
        
        # 权限过滤
        if manage_permission == '超级管理员':
            if department:
                query = query.filter_by(department=department)
        elif manage_permission == '管理部门':
            # 只能查看本部门
            query = query.filter_by(department=current_user.department)
        elif manage_permission == '管理部门小组':
            # 只能查看本小组
            if current_user.group:
                # 假设user.group存储的是小组名称
                query = query.filter_by(department=current_user.department, group=current_user.group)
            else:
                # 未分配小组的用户无法看到任何用户，或者只能看到自己？这里暂且返回空
                return jsonify({'success': True, 'data': []})
        
        users = query.all()
        
        users_data = [_serialize_user_basic(user) for user in users]
        
        return jsonify({'success': True, 'data': users_data})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users', methods=['POST'])
@role_required('管理员')
def add_user():
    """添加新用户"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        # 只有超级管理员和具有"管理部门"权限的管理员可以添加用户
        if manage_permission != '超级管理员' and manage_permission != '管理部门':
             return jsonify({'success': False, 'message': '无权添加用户'}), 403
        
        data = request.get_json()
        
        # 必填字段验证
        name = data.get('name', '').strip()
        department = data.get('department', '').strip()
        gender = data.get('gender', '').strip()
        grade = data.get('grade', '').strip()
        college = data.get('college', '').strip()
        major = data.get('major', '').strip()
        dormitory = data.get('dormitory', '').strip()
        phone = data.get('phone', '').strip()
        qq = data.get('qq', '').strip()
        student_id = data.get('student_id', '').strip()
        role = data.get('role', '').strip()
        
        # 权限检查：非超级管理员只能添加本部门用户
        if manage_permission != '超级管理员':
            if department != current_user.department:
                return jsonify({'success': False, 'message': '只能添加本部门的用户'}), 403
            # 非超级管理员不能添加超级管理员
            if role == '超级管理员':
                return jsonify({'success': False, 'message': '无权创建超级管理员'}), 403
        
        # 检查所有必填字段
        required_fields = {
            'name': name,
            'department': department,
            'gender': gender,
            'grade': grade,
            'college': college,
            'major': major,
            'dormitory': dormitory,
            'phone': phone,
            'qq': qq,
            'student_id': student_id,
            'role': role
        }
        
        missing_fields = []
        for field_name, field_value in required_fields.items():
            if not field_value:
                field_labels = {
                    'name': '姓名',
                    'department': '部门',
                    'gender': '性别',
                    'grade': '年级',
                    'college': '学院',
                    'major': '专业',
                    'dormitory': '宿舍',
                    'phone': '手机号码',
                    'qq': 'QQ号码',
                    'student_id': '学号',
                    'role': '角色'
                }
                missing_fields.append(field_labels.get(field_name, field_name))
        
        if missing_fields:
            return jsonify({'success': False, 'message': f'以下字段为必填项：{", ".join(missing_fields)}'}), 400
        
        # 检查学号是否存在
        if User.query.filter_by(student_id=student_id).first():
            return jsonify({'success': False, 'message': '该学号已存在'}), 400
            
        # 自动生成编号：获取当前最大编号并加1
        max_user = User.query.order_by(User.number.desc()).first()
        if max_user and max_user.number.isdigit():
            number = str(int(max_user.number) + 1).zfill(4)  # 4位数字，不足补0
        else:
            number = "0001"  # 如果没有用户或编号不是数字，从0001开始
        
        # 检查部门是否存在
        dept = Department.query.filter_by(name=department).first()
        if not dept and department != UNASSIGNED_DEPARTMENT_NAME:
            return jsonify({'success': False, 'message': '指定的部门不存在'}), 400
        
        # 获取可选字段
        password_field = data.get('password', '').strip()  # 用户可能设置自定义密码
        group_id = data.get('group_id')
        group_name = None
        
        # 如果指定了小组，检查小组是否存在且属于指定部门
        if group_id:
            try:
                group_id = int(group_id)
            except ValueError:
                return jsonify({'success': False, 'message': '小组ID格式不正确'}), 400
            
            if group_id:
                group = Group.query.get(group_id)
                if not group:
                    return jsonify({'success': False, 'message': '指定的小组不存在'}), 400
                if group.department != department:
                    return jsonify({'success': False, 'message': '小组不属于指定的部门'}), 400
                group_name = group.name
        else:
            # 如果没有指定小组，设置为未分配小组
            group_id = None
            group_name = UNASSIGNED_GROUP_NAME
        
        # 生成密码：如果用户提供了密码就使用用户的密码，否则生成随机密码
        if password_field:
            password = password_field
        else:
            password = generate_random_password()
        password_hash = generate_password_hash(password)
        
        new_user = User(
            number=number,
            name=name,
            gender=gender,
            grade=grade,
            college=college,
            major=major,
            dormitory=dormitory,
            phone=phone,
            qq=qq,
            student_id=student_id,
            password_hash=password_hash,
            role=role,
            department=department,
            group_id=group_id,
            group=group_name # 设置 group 字段
        )
        
        db.session.add(new_user)
        db.session.flush()
        record_password_audit(
            actor_user_id=session.get('user_id'),
            target_user_id=new_user.id,
            action='admin_create_user_password_init',
            details={'custom_password': bool(password_field)}
        )
        db.session.commit()
        
        return jsonify({'success': True, 'message': '用户添加成功', 'data': {
            'id': new_user.id,
            'number': new_user.number,
            'name': new_user.name,
            'password': password  # 返回明文密码供管理员记录
        }})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users/<int:user_id>', methods=['GET'])
@role_required('管理员')
def get_user(user_id):
    """获取单个用户信息"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        # 检查权限
        if not manage_permission:
            return jsonify({'success': False, 'message': '无权查看用户'}), 403
            
        user = User.query.get_or_404(user_id)
        if not is_user_active(user):
            return jsonify({'success': False, 'message': '用户已离任'}), 404
        
        # 权限范围检查
        if manage_permission != '超级管理员':
            if user.department != current_user.department:
                return jsonify({'success': False, 'message': '无权查看其他部门的用户'}), 403
            if manage_permission == '管理部门小组':
                # 假设user.group存储的是小组名称
                if user.group != current_user.group:
                    return jsonify({'success': False, 'message': '无权查看其他小组的用户'}), 403
        
        return jsonify({'success': True, 'data': _serialize_user_basic(user)})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/form/<int:form_id>', methods=['DELETE'])
@login_required
def delete_form(form_id):
    """删除单个表单版本"""
    try:
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限'}), 403

        form_to_delete = LectureForm.query.get_or_404(form_id)

        reviewable_user_ids = get_reviewable_users(session['user_id'])
        form_user = _active_user_query().filter_by(number=form_to_delete.listener_number).first()
        if not form_user or form_user.id not in reviewable_user_ids:
            return jsonify({'success': False, 'message': '您无权删除该表单'}), 403

        registration_id = form_to_delete.registration_id
        score_record = ScoreRecord.query.filter_by(form_id=form_id).first()
        if score_record:
            db.session.delete(score_record)
        db.session.delete(form_to_delete)
        db.session.flush()

        if registration_id:
            remains = LectureForm.query.filter_by(registration_id=registration_id).count()
            if remains == 0:
                registration = CourseRegistration.query.get(registration_id)
                if registration:
                    db.session.delete(registration)

        db.session.commit()
        return jsonify({'success': True, 'message': '表单版本删除成功'})

    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/review/group/<int:group_id>', methods=['DELETE'])
@login_required
def delete_form_group(group_id):
    """删除整个表单组及关联数据"""
    try:
        permission = get_user_review_permission(session['user_id'])
        if not permission:
            return jsonify({'success': False, 'message': '您不具有审表权限'}), 403

        group_forms = LectureForm.query.filter(
            db.or_(LectureForm.unique_id == group_id, LectureForm.id == group_id)
        ).all()
        if not group_forms:
            return jsonify({'success': False, 'message': '表单组不存在'}), 404

        reviewable_user_ids = get_reviewable_users(session['user_id'])
        for form in group_forms:
            form_user = _active_user_query().filter_by(number=form.listener_number).first()
            if not form_user or form_user.id not in reviewable_user_ids:
                return jsonify({'success': False, 'message': '您无权删除该表单组'}), 403

        form_ids = [form.id for form in group_forms]
        registration_ids = list({form.registration_id for form in group_forms if form.registration_id})

        score_records = ScoreRecord.query.filter(ScoreRecord.form_id.in_(form_ids)).all()
        for score_record in score_records:
            db.session.delete(score_record)

        for form in group_forms:
            db.session.delete(form)
        db.session.flush()

        deleted_registration_count = 0
        for registration_id in registration_ids:
            remains = LectureForm.query.filter_by(registration_id=registration_id).count()
            if remains == 0:
                registration = CourseRegistration.query.get(registration_id)
                if registration:
                    db.session.delete(registration)
                    deleted_registration_count += 1

        db.session.commit()
        return jsonify({
            'success': True,
            'message': f'已删除表单组，共删除 {len(form_ids)} 个版本，清理 {deleted_registration_count} 条课程登记记录'
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users/<int:user_id>', methods=['PUT'])
@role_required('管理员')
def update_user(user_id):
    """更新用户信息"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        # 只有超级管理员和具有"管理部门"权限的管理员可以修改用户信息
        # "管理部门小组"权限不能修改用户信息
        if manage_permission != '超级管理员' and manage_permission != '管理部门':
             return jsonify({'success': False, 'message': '无权修改用户信息'}), 403
             
        user = User.query.get_or_404(user_id)
        if not is_user_active(user):
            return jsonify({'success': False, 'message': '用户已离任'}), 404
        
        # 如果不是超级管理员，只能修改本部门用户
        if manage_permission != '超级管理员' and user.department != current_user.department:
             return jsonify({'success': False, 'message': '只能修改本部门用户'}), 403
             
        # "管理部门"权限不能修改超级管理员的信息
        if manage_permission != '超级管理员' and user.role == '超级管理员':
             return jsonify({'success': False, 'message': '无法修改超级管理员的信息'}), 403
             
        data = request.get_json()
        before_snapshot = _snapshot_user_for_movement(user)
        password_changed = False
        if 'new_password' not in data and 'password' in data:
            data['new_password'] = data.get('password')
        
        # "管理部门"权限的限制：只能修改小组和密码，不能修改角色、部门等敏感信息
        if manage_permission == '管理部门':
            # 允许修改的字段：小组
            if 'group_id' in data:
                group_id = data['group_id']
                if group_id:
                    group = Group.query.get(group_id)
                    if not group:
                        return jsonify({'success': False, 'message': '指定的小组不存在'}), 400
                    if group.department != user.department:
                        return jsonify({'success': False, 'message': '小组不属于用户所在部门'}), 400
                    user.group_id = group_id
                    user.group = group.name # 同步更新名称
                else:
                    user.group_id = None
                    user.group = UNASSIGNED_GROUP_NAME # 同步更新名称

            if 'new_password' in data and data['new_password']:
                user.password_hash = generate_password_hash(data['new_password'])
                password_changed = True
            
            # 检查是否尝试修改受限字段
            restricted_fields = ['name', 'department', 'role', 'student_id']
            for field in restricted_fields:
                if field in data and data[field]:
                    # 如果尝试修改且值确实改变了
                    if field == 'department' and data[field] != user.department:
                         return jsonify({'success': False, 'message': '无权修改用户部门'}), 403
                    elif field == 'role' and data[field] != user.role:
                         return jsonify({'success': False, 'message': '无权修改用户角色'}), 403
                    # 其他字段暂不严格限制报错，只是忽略
        else:
            # 超级管理员可以修改所有字段
            
            # 基本信息更新
            if 'name' in data: user.name = data['name'].strip()
            if 'gender' in data: user.gender = data['gender'].strip()
            if 'grade' in data: user.grade = data['grade'].strip()
            if 'college' in data: user.college = data['college'].strip()
            if 'major' in data: user.major = data['major'].strip()
            if 'dormitory' in data: user.dormitory = data['dormitory'].strip()
            if 'phone' in data: user.phone = data['phone'].strip()
            if 'qq' in data: user.qq = data['qq'].strip()
            if 'student_id' in data: user.student_id = data['student_id'].strip()
            if 'role' in data: user.role = data['role'].strip()
            
            # 部门更新
            if 'department' in data:
                new_department = data['department'].strip()
                if new_department != user.department:
                    # 检查新部门是否存在
                    dept = Department.query.filter_by(name=new_department).first()
                    if not dept and new_department != UNASSIGNED_DEPARTMENT_NAME:
                        return jsonify({'success': False, 'message': '指定的部门不存在'}), 400
                    user.department = new_department
                    # 如果更换部门，清除小组关联
                    user.group_id = None
                    user.group = UNASSIGNED_GROUP_NAME
            
            # 小组更新
            if 'group_id' in data:
                group_id = data['group_id']
                if group_id:
                    group = Group.query.get(group_id)
                    if not group:
                        return jsonify({'success': False, 'message': '指定的小组不存在'}), 400
                    if group.department != user.department:
                        return jsonify({'success': False, 'message': '小组不属于用户所在部门'}), 400
                    user.group_id = group_id
                    user.group = group.name
                else:
                    user.group_id = None
                    user.group = UNASSIGNED_GROUP_NAME
            
            # 密码更新
            if 'new_password' in data and data['new_password']:
                user.password_hash = generate_password_hash(data['new_password'])
                password_changed = True
        
        after_snapshot = _snapshot_user_for_movement(user)
        movement_changes = []
        if before_snapshot.get('department') != after_snapshot.get('department'):
            movement_changes.append({
                'field': 'department',
                'label': '部门',
                'before': before_snapshot.get('department') or '未设置',
                'after': after_snapshot.get('department') or '未设置'
            })
        if before_snapshot.get('group') != after_snapshot.get('group'):
            movement_changes.append({
                'field': 'group',
                'label': '小组',
                'before': before_snapshot.get('group') or UNASSIGNED_GROUP_NAME,
                'after': after_snapshot.get('group') or UNASSIGNED_GROUP_NAME
            })

        profile_field_labels = {
            'name': '姓名',
            'gender': '性别',
            'grade': '年级',
            'college': '学院',
            'major': '专业',
            'dormitory': '宿舍',
            'phone': '电话',
            'qq': 'QQ',
            'student_id': '学号',
            'role': '角色'
        }
        profile_changes = []
        for field, label in profile_field_labels.items():
            if before_snapshot.get(field) != after_snapshot.get(field):
                profile_changes.append({
                    'field': field,
                    'label': label,
                    'before': before_snapshot.get(field) or '',
                    'after': after_snapshot.get(field) or ''
                })

        if movement_changes:
            movement_snapshot = dict(after_snapshot)
            movement_snapshot['department_before'] = before_snapshot.get('department')
            movement_snapshot['department_after'] = after_snapshot.get('department')
            movement_snapshot['group_before'] = before_snapshot.get('group')
            movement_snapshot['group_after'] = after_snapshot.get('group')
            _create_personnel_movement_record(
                current_user.id,
                movement_snapshot,
                '部门/小组变动',
                f"{after_snapshot.get('name')}（{after_snapshot.get('number')}）发生部门/小组变动",
                movement_changes
            )

        if profile_changes:
            _create_personnel_movement_record(
                current_user.id,
                after_snapshot,
                '个人资料修改',
                f"{after_snapshot.get('name')}（{after_snapshot.get('number')}）修改了个人资料",
                profile_changes
            )

        if password_changed:
            record_password_audit(
                actor_user_id=current_user.id,
                target_user_id=user.id,
                action='admin_reset_user_password',
                details={'manage_permission': manage_permission}
            )
        
        db.session.commit()
        
        return jsonify({'success': True, 'message': '用户信息更新成功'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users/<int:user_id>/depart', methods=['POST'])
@role_required('管理员')
def depart_user(user_id):
    """用户离任：软删除账号，保留历史记录。"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)

        if manage_permission != '超级管理员' and manage_permission != '管理部门':
            return jsonify({'success': False, 'message': '无权办理用户离任'}), 403

        user_to_depart = User.query.get_or_404(user_id)
        if user_to_depart.id == current_user.id:
            return jsonify({'success': False, 'message': '不能将当前登录账号设为离任'}), 400
        if user_to_depart.role == '超级管理员':
            return jsonify({'success': False, 'message': '超级管理员不能办理离任，请先调整角色或使用超级管理员删除'}), 400
        if manage_permission != '超级管理员' and user_to_depart.department != current_user.department:
            return jsonify({'success': False, 'message': '只能办理本部门用户离任'}), 403

        data = request.get_json() or {}
        password = data.get('password')
        if not password or not check_password_hash(current_user.password_hash, password):
            return jsonify({'success': False, 'message': '密码验证失败，无法办理离任'}), 403

        if not is_user_active(user_to_depart):
            return jsonify({'success': True, 'message': '该用户已处于离任状态'})

        before_snapshot = _snapshot_user_for_movement(user_to_depart)
        user_to_depart.is_active = False
        user_to_depart.department = '离任'
        user_to_depart.group_id = None
        user_to_depart.group = '离任'
        RolePermission.query.filter_by(role=f'特殊角色_{user_to_depart.id}').delete()
        after_snapshot = _snapshot_user_for_movement(user_to_depart)

        _create_personnel_movement_record(
            current_user.id,
            {
                **after_snapshot,
                'department_before': before_snapshot.get('department'),
                'department_after': after_snapshot.get('department'),
                'group_before': before_snapshot.get('group'),
                'group_after': after_snapshot.get('group'),
            },
            '用户离任',
            f"{before_snapshot.get('name')}（{before_snapshot.get('number')}）已离任",
            [{
                'field': 'is_active',
                'label': '状态',
                'before': '在任',
                'after': '离任'
            }, {
                'field': 'department',
                'label': '部门',
                'before': before_snapshot.get('department') or '未设置',
                'after': '离任'
            }, {
                'field': 'group',
                'label': '小组',
                'before': before_snapshot.get('group') or UNASSIGNED_GROUP_NAME,
                'after': '离任'
            }]
        )

        db.session.commit()
        return jsonify({'success': True, 'message': '用户已离任，后续不再参与活动或页面展示'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users/<int:user_id>', methods=['DELETE'])
@role_required('管理员')
def delete_user(user_id):
    """删除用户（及相关权限）"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        # 只有超级管理员保留物理删除用户的能力
        if manage_permission != '超级管理员':
             return jsonify({'success': False, 'message': '无权删除用户，请使用离任操作'}), 403
             
        user_to_delete = User.query.get_or_404(user_id)
        
        if user_to_delete.id == current_user.id:
             return jsonify({'success': False, 'message': '不能删除当前登录账号'}), 400
        
        # 验证密码
        data = request.get_json() or {}
        password = data.get('password')
        if not password or not check_password_hash(current_user.password_hash, password):
            return jsonify({'success': False, 'message': '密码验证失败，无法删除用户'}), 403
        
        # 检查是否有相关的听课表单
        forms_count = LectureForm.query.filter_by(listener_number=user_to_delete.number).count()
        if forms_count > 0:
            return jsonify({'success': False, 'message': f'无法删除用户，该用户有 {forms_count} 条听课记录'}), 400
        
        delete_snapshot = _snapshot_user_for_movement(user_to_delete)

        # 删除相关的特殊角色权限
        RolePermission.query.filter_by(role=f'特殊角色_{user_to_delete.id}').delete()

        _create_personnel_movement_record(
            current_user.id,
            delete_snapshot,
            '删除用户',
            f"删除用户 {delete_snapshot.get('name')}（{delete_snapshot.get('number')}）",
            [{
                'field': 'delete',
                'label': '删除',
                'before': f"{delete_snapshot.get('department') or '未设置'} / {delete_snapshot.get('group') or UNASSIGNED_GROUP_NAME}",
                'after': '用户已删除'
            }]
        )
        
        db.session.delete(user_to_delete)
        db.session.commit()
        
        return jsonify({'success': True, 'message': '用户及关联权限已删除'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500


@admin_bp.route('/api/personnel-movement-records', methods=['GET'])
@role_required('管理员')
def list_personnel_movement_records():
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        if not manage_permission:
            return jsonify({'success': False, 'message': '无权查看人员流动记录'}), 403

        query = PersonnelMovementRecord.query
        if manage_permission != '超级管理员':
            query = query.filter(PersonnelMovementRecord.department_name == current_user.department)

        limit = request.args.get('limit', 100)
        try:
            limit = max(1, min(int(limit), 300))
        except Exception:
            limit = 100

        records = query.order_by(
            PersonnelMovementRecord.created_at.desc(),
            PersonnelMovementRecord.id.desc()
        ).limit(limit).all()

        operator_ids = {record.operator_user_id for record in records if record.operator_user_id}
        operator_map = {
            user.id: user
            for user in User.query.filter(User.id.in_(list(operator_ids))).all()
        } if operator_ids else {}

        result = []
        for record in records:
            try:
                details = json.loads(record.details_json) if record.details_json else {}
            except Exception:
                details = {}
            operator = operator_map.get(record.operator_user_id)
            result.append({
                'id': record.id,
                'target_user_id': record.target_user_id,
                'target_user_number': record.target_user_number,
                'target_user_name': record.target_user_name,
                'department_name': record.department_name,
                'group_name': record.group_name,
                'action_type': record.action_type,
                'summary': record.summary,
                'details': details.get('changes', []),
                'operator_name': operator.name if operator else '',
                'operator_number': operator.number if operator else '',
                'created_at': record.created_at.strftime('%Y-%m-%d %H:%M:%S') if record.created_at else ''
            })
        return jsonify({'success': True, 'records': result})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

# ==================== 权限管理 API 路由 ====================

@admin_bp.route('/api/permissions', methods=['GET'])
@role_required('超级管理员')
def get_permissions():
    """获取所有权限列表"""
    try:
        permissions = Permission.query.all()
        result = []
        for perm in permissions:
            result.append({
                'id': perm.id,
                'name': perm.name,
                'description': perm.description
            })
        return jsonify({'success': True, 'data': result})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users/<int:user_id>/permissions', methods=['GET'])
@role_required('管理员')
def get_user_permissions(user_id):
    """获取用户的权限列表"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        # 只有超级管理员和具有"管理部门"权限的管理员可以查看权限
        # "管理部门小组"权限不能管理用户权限
        if manage_permission != '超级管理员' and manage_permission != '管理部门':
             return jsonify({'success': False, 'message': '无权查看权限'}), 403
             
        user = User.query.get_or_404(user_id)
        if not is_user_active(user):
            return jsonify({'success': False, 'message': '用户已离任'}), 404
        
        # 如果不是超级管理员，只能查看本部门用户的权限
        if manage_permission != '超级管理员' and user.department != current_user.department:
             return jsonify({'success': False, 'message': '只能查看本部门用户的权限'}), 403
        
        # 获取所有权限
        all_permissions = Permission.query.all()
        permissions = []
        for perm in all_permissions:
            permissions.append({
                'id': perm.id,
                'name': perm.name,
                'description': perm.description
            })
        
        # 获取用户当前拥有的权限ID列表
        user_permissions = []
        # 首先检查是否有自定义权限（使用特殊角色标识符存储）
        custom_role_permissions = db.session.query(Permission).join(RolePermission).filter(
            RolePermission.role == f'特殊角色_{user.id}'
        ).all()
        
        if custom_role_permissions:
            # 有自定义权限，使用自定义权限
            user_permissions = [perm.id for perm in custom_role_permissions]
        else:
            # 没有自定义权限，查询角色默认权限
            role_permissions = db.session.query(Permission).join(RolePermission).filter(
                RolePermission.role == user.role
            ).all()
            user_permissions = [perm.id for perm in role_permissions]
        
        return jsonify({
            'success': True, 
            'permissions': permissions,
            'user_permissions': user_permissions
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/users/<int:user_id>/permissions', methods=['PUT'])
@role_required('管理员')
def update_user_permissions(user_id):
    """更新用户权限（自动设置为管理员）"""
    try:
        current_user = User.query.get(session['user_id'])
        manage_permission = get_user_manage_permission(current_user.id)
        
        # 只有超级管理员和具有"管理部门"权限的管理员可以修改权限
        if manage_permission != '超级管理员' and manage_permission != '管理部门':
             return jsonify({'success': False, 'message': '无权修改权限'}), 403
             
        user = User.query.get_or_404(user_id)
        if not is_user_active(user):
            return jsonify({'success': False, 'message': '用户已离任'}), 404
        
        # 如果不是超级管理员，只能修改本部门用户的权限
        if manage_permission != '超级管理员' and user.department != current_user.department:
             return jsonify({'success': False, 'message': '只能修改本部门用户的权限'}), 403
             
        # "管理部门"权限不能修改超级管理员的权限
        if manage_permission != '超级管理员' and user.role == '超级管理员':
             return jsonify({'success': False, 'message': '无法修改超级管理员的权限'}), 403
        
        data = request.get_json()
        permission_ids = data.get('permission_ids', [])
        
        # 删除用户现有的自定义权限（如果有特殊角色权限）
        RolePermission.query.filter_by(role=f'特殊角色_{user.id}').delete()
        
        # 添加新的权限
        for perm_id in permission_ids:
            permission = Permission.query.get(perm_id)
            if permission:
                # "管理部门"权限不能赋予"管理部门"或"超级管理员"权限，防止权限提升
                if manage_permission != '超级管理员' and (permission.name == '管理部门' or permission.name == '超级管理员'):
                    continue
                    
                role_perm = RolePermission(role=f'特殊角色_{user.id}', permission_id=perm_id)
                db.session.add(role_perm)
        
        # 更新用户角色为管理员
        if user.role != '超级管理员':
            user.role = '管理员'
        
        db.session.commit()
        
        return jsonify({'success': True, 'message': '用户权限更新成功，角色已设置为管理员'})
    except Exception as e:
        print(f"DEBUG: 发生错误: {str(e)}")
        print(f"DEBUG: 错误类型: {type(e)}")
        import traceback
        print(f"DEBUG: 错误堆栈: {traceback.format_exc()}")
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500

STATISTICS_APPROVED_STATUSES = {'已审核', '部门已审核', '中心已审核'}
STATISTICS_DIMENSIONS = {'department', 'user', 'teacher', 'college', 'date'}
STATISTICS_TEACHER_KEY_SEPARATOR = '\u241f'


def _statistics_query(current_user, args):
    """Build the shared, permission-scoped query used by statistics views and exports."""
    query = db.session.query(LectureForm, User).join(
        User,
        LectureForm.listener_number == User.number,
    ).filter(active_user_filter())

    if current_user and current_user.role != '超级管理员':
        query = query.filter(User.department == current_user.department)

    start_date = (args.get('start_date') or '').strip()
    end_date = (args.get('end_date') or '').strip()
    if start_date:
        try:
            datetime.strptime(start_date, '%Y-%m-%d')
        except ValueError:
            start_date = ''
    if end_date:
        try:
            datetime.strptime(end_date, '%Y-%m-%d')
        except ValueError:
            end_date = ''
    if start_date:
        query = query.filter(LectureForm.lecture_date >= start_date)
    if end_date:
        query = query.filter(LectureForm.lecture_date <= end_date)

    department = (args.get('department') or '').strip()
    if department and current_user and current_user.role == '超级管理员':
        query = query.filter(User.department == department)

    status = (args.get('status') or '').strip()
    if status == '已审核':
        query = query.filter(LectureForm.status.in_(STATISTICS_APPROVED_STATUSES))
    elif status in {'待审核', '部门已审核', '中心已审核', '已驳回'}:
        query = query.filter(LectureForm.status == status)

    return query


def _statistics_dimension(args):
    dimension = (args.get('dimension') or 'department').strip()
    return dimension if dimension in STATISTICS_DIMENSIONS else 'department'


def _build_statistics_rows(records, dimension):
    groups = {}
    for form, listener in records:
        if dimension == 'user':
            key = ('user', listener.number or '')
            label = listener.name or form.listener_name or '未知用户'
            metadata = {
                'department': listener.department or '未分配部门',
                'detail_value': listener.number or '',
            }
        elif dimension == 'teacher':
            teacher_name = form.teacher_name or '未知教师'
            teacher_college = form.teacher_college or '未知学院'
            key = ('teacher', teacher_name, teacher_college)
            label = teacher_name
            metadata = {
                'college': teacher_college,
                'detail_value': STATISTICS_TEACHER_KEY_SEPARATOR.join((teacher_name, teacher_college)),
            }
        elif dimension == 'college':
            label = form.teacher_college or '未知学院'
            key = ('college', label)
            metadata = {'detail_value': label}
        elif dimension == 'date':
            label = str(form.lecture_date or '未知日期')
            key = ('date', label)
            metadata = {'detail_value': label}
        else:
            label = listener.department or '未分配部门'
            key = ('department', label)
            metadata = {'detail_value': label}

        item = groups.setdefault(key, {
            'name': label,
            'total': 0,
            'pending': 0,
            'approved': 0,
            'rejected': 0,
            **metadata,
        })
        item['total'] += 1
        if form.status == '待审核':
            item['pending'] += 1
        elif form.status in STATISTICS_APPROVED_STATUSES:
            item['approved'] += 1
        elif form.status == '已驳回':
            item['rejected'] += 1

    return sorted(
        groups.values(),
        key=lambda item: (str(item['name']).casefold(), str(item.get('detail_value', '')).casefold()),
    )


def _apply_statistics_detail_filter(query, dimension, detail_value):
    detail_value = (detail_value or '').strip()
    if not detail_value:
        return query
    if dimension == 'user':
        return query.filter(User.number == detail_value)
    if dimension == 'teacher':
        teacher_name, separator, teacher_college = detail_value.partition(STATISTICS_TEACHER_KEY_SEPARATOR)
        query = query.filter(LectureForm.teacher_name == teacher_name)
        return query.filter(LectureForm.teacher_college == teacher_college) if separator else query
    if dimension == 'college':
        return query.filter(LectureForm.teacher_college == detail_value)
    if dimension == 'date':
        return query.filter(LectureForm.lecture_date == detail_value)
    return query.filter(User.department == detail_value)


def _safe_statistics_excel_text(value):
    """Prevent aggregate labels from being interpreted as spreadsheet formulas."""
    text = str(value or '')
    if text.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + text
    return text


@admin_bp.route('/statistics')
@role_required('管理员')
def statistics():
    """统计分析页面"""
    current_user = User.query.get(session['user_id'])
    manage_permission = get_user_manage_permission(session['user_id'])
    can_access_extended_stats = manage_permission in ['超级管理员', '管理部门']
    can_manage_department_leave = manage_permission == '超级管理员'
    available_departments = list(_get_accessible_department_users(session['user_id']).keys()) if can_access_extended_stats else []

    form_query = _statistics_query(current_user, request.args)
    records = form_query.order_by(LectureForm.created_at.asc(), LectureForm.id.asc()).all()
    forms = [form for form, _listener in records]
    total_forms = len(forms)
    logical_form_keys = {
        ('unique_id', str(unique_id).strip())
        if unique_id is not None and str(unique_id).strip()
        else ('id', str(form_id))
        for unique_id, form_id in ((form.unique_id, form.id) for form in forms)
    }
    logical_form_count = len(logical_form_keys)
    pending_forms = sum(form.status == '待审核' for form in forms)
    approved_forms = sum(form.status in STATISTICS_APPROVED_STATUSES for form in forms)
    rejected_forms = sum(form.status == '已驳回' for form in forms)

    trend_counts = {}
    for created_at in (form.created_at for form in forms):
        if created_at:
            label = created_at.strftime('%Y-%m-%d')
            trend_counts[label] = trend_counts.get(label, 0) + 1
    trend_labels = sorted(trend_counts)
    trend_data = [trend_counts[label] for label in trend_labels]
    if not trend_labels:
        visualization_state='NO_DATA'
    elif len(trend_labels) == 1:
        visualization_state='INSUFFICIENT_DATA'
    else:
        visualization_state='DATA_READY'

    statistics_dimension = _statistics_dimension(request.args)
    return render_template(
        'admin/statistics.html',
        can_access_extended_stats=can_access_extended_stats,
        can_manage_department_leave=can_manage_department_leave,
        total_forms=total_forms,
        logical_form_count=logical_form_count,
        pending_forms=pending_forms,
        approved_forms=approved_forms,
        rejected_forms=rejected_forms,
        departments=available_departments,
        statistics_data=_build_statistics_rows(records, statistics_dimension),
        statistics_dimension=statistics_dimension,
        visualization_state=visualization_state,
        trend_point_count=trend_data[0] if len(trend_data) == 1 else None,
        trend_labels=json.dumps(trend_labels, ensure_ascii=False),
        trend_data=json.dumps(trend_data, ensure_ascii=False),
    )


@admin_bp.route('/export_statistics')
@role_required('管理员')
def export_statistics():
    """Export the currently filtered aggregate statistics as an XLSX workbook."""
    current_user = User.query.get(session['user_id'])
    dimension = _statistics_dimension(request.args)
    records = _statistics_query(current_user, request.args).order_by(
        LectureForm.created_at.asc(),
        LectureForm.id.asc(),
    ).all()
    rows = _build_statistics_rows(records, dimension)

    dimension_headers = {
        'department': ['部门'],
        'user': ['用户', '部门'],
        'teacher': ['教师', '学院'],
        'college': ['学院'],
        'date': ['日期'],
    }
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = '统计汇总'
    headers = dimension_headers[dimension] + ['总数', '待审核', '已通过', '已驳回', '通过率']
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for item in rows:
        if dimension == 'user':
            identity = [item['name'], item.get('department', '')]
        elif dimension == 'teacher':
            identity = [item['name'], item.get('college', '')]
        else:
            identity = [item['name']]
        identity = [_safe_statistics_excel_text(value) for value in identity]
        pass_rate = item['approved'] / item['total'] if item['total'] else 0
        worksheet.append(identity + [
            item['total'],
            item['pending'],
            item['approved'],
            item['rejected'],
            pass_rate,
        ])
        worksheet.cell(row=worksheet.max_row, column=len(headers)).number_format = '0.0%'

    worksheet.freeze_panes = 'A2'
    for column_index, header in enumerate(headers, start=1):
        values = [str(worksheet.cell(row=row, column=column_index).value or '') for row in range(1, worksheet.max_row + 1)]
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(column_index)].width = min(
            max(len(header) + 2, max((len(value) for value in values), default=0) + 2),
            32,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f'统计分析_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )

def _has_assessment_stats_access(user_id):
    manage_permission = get_user_manage_permission(user_id)
    return manage_permission in ['超级管理员', '管理部门']


SNAPSHOT_TYPE_SUBMISSION_REWARD = 'submission_reward'
SNAPSHOT_TYPE_DEPARTMENT_MONTHLY = 'department_monthly_assessment'


def _snapshot_base_query(snapshot_type, current_user_id):
    query = StatisticsSnapshot.query.filter_by(snapshot_type=snapshot_type)
    if get_user_manage_permission(current_user_id) != '超级管理员':
        query = query.filter_by(created_by=current_user_id)
    return query.order_by(StatisticsSnapshot.created_at.desc(), StatisticsSnapshot.id.desc())


def _build_snapshot_list_items(snapshot_type, current_user_id):
    records = _snapshot_base_query(snapshot_type, current_user_id).all()
    creator_ids = {record.created_by for record in records if record.created_by}
    creators = {
        user.id: user
        for user in User.query.filter(User.id.in_(list(creator_ids))).all()
    } if creator_ids else {}
    items = []
    for record in records:
        try:
            filters_data = json.loads(record.filters_json or '{}')
        except Exception:
            filters_data = {}
        creator = creators.get(record.created_by)
        items.append({
            'id': record.id,
            'title': record.title,
            'created_at': record.created_at.strftime('%Y-%m-%d %H:%M:%S') if record.created_at else '',
            'creator_name': creator.name if creator else '',
            'range_start': filters_data.get('start_date', ''),
            'range_end': filters_data.get('end_date', ''),
            'period_label': filters_data.get('period_label', ''),
            'scope_label': filters_data.get('scope_label', ''),
        })
    return items


def _get_snapshot_record_or_404(snapshot_id, snapshot_type, current_user_id):
    record = StatisticsSnapshot.query.filter_by(id=snapshot_id, snapshot_type=snapshot_type).first()
    if not record:
        return None, jsonify({'success': False, 'message': '历史记录不存在'}), 404
    if get_user_manage_permission(current_user_id) != '超级管理员' and record.created_by != current_user_id:
        return None, jsonify({'success': False, 'message': '无权访问该历史记录'}), 403
    return record, None, None


def _create_statistics_snapshot(snapshot_type, title, filters_data, payload_data, current_user_id):
    record = StatisticsSnapshot(
        snapshot_type=snapshot_type,
        title=title,
        filters_json=json.dumps(filters_data, ensure_ascii=False),
        payload_json=json.dumps(payload_data, ensure_ascii=False),
        created_by=current_user_id
    )
    db.session.add(record)
    db.session.commit()
    return record


def _load_snapshot_payload(record):
    try:
        filters_data = json.loads(record.filters_json or '{}')
    except Exception:
        filters_data = {}
    try:
        payload_data = json.loads(record.payload_json or '{}')
    except Exception:
        payload_data = {}
    return filters_data, payload_data

@admin_bp.route('/review-assessment-stats')
@role_required('管理员')
def review_assessment_stats():
    if not _has_assessment_stats_access(session['user_id']):
        flash_forbidden('审表考评统计')
        return redirect(url_for('main.index'))
    manage_permission = get_user_manage_permission(session['user_id'])
    can_manual_assessment_import = manage_permission == '超级管理员'
    return render_template('admin/review_assessment_stats.html', can_manual_assessment_import=can_manual_assessment_import)

@admin_bp.route('/submission-count-stats')
@role_required('管理员')
def submission_count_stats():
    if not _has_assessment_stats_access(session['user_id']):
        flash_forbidden('交表数量统计')
        return redirect(url_for('main.index'))
    return render_template('admin/submission_count_stats.html')


@admin_bp.route('/department-monthly-assessment-stats')
@role_required('管理员')
def department_monthly_assessment_stats():
    if not _has_assessment_stats_access(session['user_id']):
        flash_forbidden('部门月度考评')
        return redirect(url_for('main.index'))
    available_departments = list(_get_accessible_department_users(session['user_id']).keys())
    return render_template('admin/department_monthly_assessment_stats.html', available_departments=available_departments)

def _parse_assessment_range(start_date_str, end_date_str):
    if not start_date_str or not end_date_str:
        return None, None, '请先选择开始和结束时间'
    try:
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d') + timedelta(days=1)
        return start_date, end_date, None
    except Exception:
        return None, None, '时间格式错误，请使用YYYY-MM-DD'

def _resolve_assessment_users(current_user_id, user_ids):
    current_user = User.query.get(current_user_id)
    if not current_user:
        return []
    manage_permission = get_user_manage_permission(current_user_id)
    if manage_permission == '管理部门':
        scoped_users = _active_user_query().filter(
            User.role.in_(['信息员', '管理员']),
            User.department == current_user.department,
            User.id != current_user_id
        ).all()
        reviewable_ids = set(u.id for u in scoped_users)
    else:
        reviewable_ids = set(get_reviewable_users(current_user_id))
    if user_ids:
        selected_ids = set()
        for uid in user_ids:
            try:
                selected_ids.add(int(uid))
            except Exception:
                continue
        target_ids = reviewable_ids.intersection(selected_ids)
    else:
        target_ids = reviewable_ids
    if not target_ids:
        return []
    users = _active_user_query().filter(User.id.in_(list(target_ids))).all()
    return users


def _get_accessible_department_users(current_user_id):
    users = _resolve_assessment_users(current_user_id, [])
    department_map = {}
    for user in users:
        department_name = (user.department or '').strip()
        if not department_name:
            continue
        department_map.setdefault(department_name, []).append(user)
    return dict(sorted(department_map.items(), key=lambda item: item[0]))


def _resolve_selected_departments(current_user_id, department_names):
    department_map = _get_accessible_department_users(current_user_id)
    if department_names:
        normalized = []
        for name in department_names:
            value = (name or '').strip()
            if value and value in department_map and value not in normalized:
                normalized.append(value)
        selected_names = normalized
    else:
        selected_names = list(department_map.keys())
    return selected_names, {name: department_map[name] for name in selected_names if name in department_map}


def _resolve_department_leave_users(current_user_id, user_ids=None):
    current_user = User.query.get(current_user_id)
    if not current_user or get_user_manage_permission(current_user_id) != '超级管理员':
        return []

    query = _active_user_query().filter(User.role == '信息员')
    if user_ids:
        selected_ids = set()
        for uid in user_ids:
            try:
                selected_ids.add(int(uid))
            except (TypeError, ValueError):
                continue
        if not selected_ids:
            return []
        query = query.filter(User.id.in_(list(selected_ids)))

    return query.order_by(User.department.asc(), User.group.asc(), User.number.asc(), User.name.asc(), User.id.asc()).all()


@admin_bp.route('/api/leave-management/status', methods=['GET'])
@login_required
def get_leave_management_status():
    if get_user_manage_permission(session['user_id']) != '超级管理员':
        return forbidden_json('成员请假设置')

    settings, settings_err = get_leave_teaching_settings()
    current_week = None
    if not settings_err:
        current_week, week_err = get_leave_current_teaching_week(settings)
        if week_err:
            settings_err = week_err

    users = _resolve_department_leave_users(session['user_id'])
    user_map = {user.id: user for user in users}
    user_ids = list(user_map.keys())
    leave_records = AssessmentOverride.query.filter(
        AssessmentOverride.user_id.in_(user_ids),
        AssessmentOverride.override_type == LEAVE_OVERRIDE_TYPE,
    ).order_by(AssessmentOverride.start_week.asc(), AssessmentOverride.end_week.asc()).all() if user_ids else []

    current_leave_user_ids = set()
    active_items = []
    pending_items = []
    completed_items = []
    if current_week and settings:
        for record in leave_records:
            user = user_map.get(record.user_id)
            payload = build_leave_status_payload(record, current_week, user=user, settings=settings)
            if not payload:
                continue
            if payload['status'] == 'active':
                active_items.append(payload)
                current_leave_user_ids.add(record.user_id)
            elif payload['status'] == 'pending':
                pending_items.append(payload)
            elif payload['status'] == 'completed':
                completed_items.append(payload)

    candidates = []
    for user in sorted(users, key=lambda item: (item.number or '', item.name or '')):
        candidates.append({
            'user_id': user.id,
            'name': user.name,
            'number': user.number,
            'department': user.department,
            'group': user.group,
            'role': user.role,
            'has_current_leave': user.id in current_leave_user_ids,
        })

    return jsonify({
        'success': True,
        'current_week': current_week,
        'settings_error': settings_err,
        'required_submission': int((settings or {}).get('required_submission') or 0),
        'summary': {
            'active_count': len(active_items),
            'pending_makeup_count': len(pending_items),
            'completed_makeup_count': len(completed_items),
        },
        'candidates': candidates,
        'active': active_items,
        'pending': pending_items,
        'completed': completed_items,
    })


@admin_bp.route('/api/leave-management/current-week', methods=['POST'])
@login_required
def create_current_week_leave():
    if get_user_manage_permission(session['user_id']) != '超级管理员':
        return forbidden_json('成员请假设置')

    settings, settings_err = get_leave_teaching_settings()
    if settings_err:
        return jsonify({'success': False, 'message': settings_err}), 400
    current_week, week_err = get_leave_current_teaching_week(settings)
    if week_err or not current_week:
        return jsonify({'success': False, 'message': week_err or '当前教学周不可用'}), 400

    data = request.get_json() or {}
    user_id = data.get('user_id')
    target_users = _resolve_department_leave_users(session['user_id'], [str(user_id)])
    if not target_users:
        return jsonify({'success': False, 'message': '无权限操作该成员'}), 403
    target_user = target_users[0]

    existing = AssessmentOverride.query.filter(
        AssessmentOverride.user_id == target_user.id,
        AssessmentOverride.override_type == LEAVE_OVERRIDE_TYPE,
        AssessmentOverride.start_week <= current_week,
        AssessmentOverride.end_week >= current_week,
    ).first()
    if existing:
        return jsonify({'success': False, 'message': f'{target_user.name} 当前教学周已处于请假中'}), 400

    reason = (data.get('reason') or '').strip() or '超级管理员发起当前教学周请假'
    record = AssessmentOverride(
        user_id=target_user.id,
        start_week=current_week,
        end_week=current_week,
        override_type=LEAVE_OVERRIDE_TYPE,
        reason=reason,
        created_by=session['user_id'],
    )
    db.session.add(record)
    db.session.commit()
    return jsonify({
        'success': True,
        'message': f'已为 {target_user.name} 设置第{current_week}周请假',
        'leave': build_leave_status_payload(record, current_week, user=target_user, settings=settings),
    })


def _get_super_admin_leave_record(override_id):
    if get_user_manage_permission(session['user_id']) != '超级管理员':
        return None, jsonify(build_forbidden_payload('成员请假设置')), 403
    record = AssessmentOverride.query.filter_by(
        id=override_id,
        override_type=LEAVE_OVERRIDE_TYPE,
    ).first()
    if not record:
        return None, jsonify({'success': False, 'message': '请假记录不存在'}), 404
    return record, None, None


def _serialize_leave_form_option(group_data, settings, selected_unique_ids):
    latest_form = group_data.get('latest_form')
    if not latest_form:
        return None
    unique_id = latest_form.unique_id or latest_form.id
    try:
        normalized_unique_id = int(unique_id)
    except (TypeError, ValueError):
        normalized_unique_id = unique_id
    effective_week = get_leave_form_effective_week_no(latest_form, settings) if settings else None
    return {
        'form_id': latest_form.id,
        'unique_id': normalized_unique_id,
        'course_title': latest_form.course_title or '',
        'teacher_name': latest_form.teacher_name or '',
        'lecture_date': latest_form.lecture_date or '',
        'status': latest_form.status or '',
        'effective_week': effective_week,
        'created_at': latest_form.created_at.strftime('%Y-%m-%d %H:%M:%S') if latest_form.created_at else '',
        'selected': normalized_unique_id in selected_unique_ids,
    }


@admin_bp.route('/api/leave-management/<int:override_id>/forms', methods=['GET'])
@login_required
def get_leave_makeup_form_options(override_id):
    record, error_response, status_code = _get_super_admin_leave_record(override_id)
    if error_response:
        return error_response, status_code
    user = User.query.get(record.user_id)
    if not user:
        return jsonify({'success': False, 'message': '请假人员不存在'}), 404

    settings, settings_err = get_leave_teaching_settings()
    selected_unique_ids = set()
    for item in get_leave_makeup_forms(record):
        value = item.get('unique_id') or item.get('form_id')
        try:
            value = int(value)
        except (TypeError, ValueError):
            pass
        selected_unique_ids.add(value)

    form_options = []
    for group_data in _latest_form_groups_for_users([user.number]):
        item = _serialize_leave_form_option(group_data, settings, selected_unique_ids)
        if item:
            form_options.append(item)
    form_options.sort(key=lambda item: item.get('created_at') or '', reverse=True)

    return jsonify({
        'success': True,
        'settings_error': settings_err,
        'leave': {
            'override_id': record.id,
            'user_id': record.user_id,
            'name': user.name,
            'number': user.number,
            'department': user.department,
            'group': user.group,
            'start_week': record.start_week,
            'end_week': record.end_week,
            'reason': record.reason or '',
        },
        'forms': form_options,
    })


@admin_bp.route('/api/leave-management/<int:override_id>/makeup-forms', methods=['PUT'])
@login_required
def update_leave_makeup_forms(override_id):
    record, error_response, status_code = _get_super_admin_leave_record(override_id)
    if error_response:
        return error_response, status_code
    user = User.query.get(record.user_id)
    if not user:
        return jsonify({'success': False, 'message': '请假人员不存在'}), 404

    data = request.get_json() or {}
    selected_unique_ids = set()
    for raw_value in data.get('unique_ids') or []:
        try:
            selected_unique_ids.add(int(raw_value))
        except (TypeError, ValueError):
            continue

    selected_forms = []
    for group_data in _latest_form_groups_for_users([user.number]):
        latest_form = group_data.get('latest_form')
        if not latest_form:
            continue
        unique_id = latest_form.unique_id or latest_form.id
        try:
            normalized_unique_id = int(unique_id)
        except (TypeError, ValueError):
            continue
        if normalized_unique_id in selected_unique_ids:
            selected_forms.append(latest_form)

    set_leave_makeup_forms(
        record,
        selected_forms,
        source='manual',
        operator_user_id=session['user_id'],
    )
    db.session.commit()

    settings, settings_err = get_leave_teaching_settings()
    current_week = None
    if not settings_err:
        current_week, _ = get_leave_current_teaching_week(settings)
    return jsonify({
        'success': True,
        'message': '补交表单已更新',
        'leave': build_leave_status_payload(record, current_week, user=user, settings=settings) if current_week else None,
    })


def _latest_forms_in_range(listener_numbers, start_date, end_date):
    if not listener_numbers:
        return []
    forms = LectureForm.query.filter(
        LectureForm.listener_number.in_(listener_numbers),
        LectureForm.created_at >= start_date,
        LectureForm.created_at < end_date
    ).order_by(LectureForm.created_at.asc(), LectureForm.id.asc()).all()
    latest_map = {}
    for form in forms:
        uid = form.unique_id or form.id
        latest_map[uid] = form
    return list(latest_map.values())

def _latest_form_groups_for_users(listener_numbers):
    if not listener_numbers:
        return []
    forms = LectureForm.query.filter(
        LectureForm.listener_number.in_(listener_numbers)
    ).order_by(LectureForm.unique_id.asc(), LectureForm.created_at.asc(), LectureForm.id.asc()).all()
    group_map = {}
    for form in forms:
        uid = form.unique_id or form.id
        if uid not in group_map:
            group_map[uid] = []
        group_map[uid].append(form)
    groups = []
    for uid, form_list in group_map.items():
        sorted_forms = sorted(form_list, key=lambda f: ((f.created_at or datetime.min), f.id), reverse=True)
        latest_form = sorted_forms[0]
        groups.append({
            'unique_id': uid,
            'latest_form': latest_form,
            'forms': sorted_forms
        })
    return groups

def _assessment_form_latest_timestamp(form):
    if not form:
        return None
    return form.updated_at or form.created_at

def _assessment_latest_form_groups_for_users(listener_numbers):
    if not listener_numbers:
        return []
    forms = LectureForm.query.filter(
        LectureForm.listener_number.in_(listener_numbers)
    ).order_by(LectureForm.unique_id.asc(), LectureForm.id.asc()).all()
    group_map = {}
    for form in forms:
        uid = form.unique_id or form.id
        group_map.setdefault(uid, []).append(form)

    groups = []
    for uid, form_list in group_map.items():
        sorted_forms = sorted(
            form_list,
            key=lambda f: (_assessment_form_latest_timestamp(f) or datetime.min, f.id),
            reverse=True
        )
        latest_form = sorted_forms[0]
        groups.append({
            'unique_id': uid,
            'latest_form': latest_form,
            'latest_timestamp': _assessment_form_latest_timestamp(latest_form),
            'forms': sorted_forms
        })

    groups.sort(
        key=lambda group: (
            group['latest_timestamp'] or datetime.min,
            group['latest_form'].id if group['latest_form'] else 0
        ),
        reverse=True
    )
    return groups

def _assessment_latest_forms_in_range(listener_numbers, start_date, end_date):
    if not listener_numbers:
        return []
    latest_forms = []
    for group_data in _assessment_latest_form_groups_for_users(listener_numbers):
        latest_form = group_data['latest_form']
        latest_timestamp = group_data['latest_timestamp']
        if not latest_form or not latest_timestamp:
            continue
        if start_date <= latest_timestamp < end_date:
            latest_forms.append(latest_form)
    return latest_forms

def _get_teaching_reward_settings():
    first_week_raw = SystemSetting.get('teaching_first_week_monday')
    if not first_week_raw:
        return None, '请先在制度设置中配置第一周起始日期'
    try:
        first_week_date = datetime.strptime(first_week_raw, '%Y-%m-%d').date()
    except Exception:
        return None, '制度设置中的第一周起始日期格式错误'
    try:
        week_start_day = int(SystemSetting.get('teaching_week_start_day', '0') or 0)
    except Exception:
        week_start_day = 0
    if week_start_day < 0 or week_start_day > 6:
        week_start_day = 0
    try:
        required_submission = int(SystemSetting.get('teaching_required_submission', '1') or 1)
    except Exception:
        required_submission = 1
    if required_submission < 0:
        required_submission = 0
    return {
        'first_week_date': first_week_date,
        'week_start_day': week_start_day,
        'required_submission': required_submission
    }, None


def _parse_lecture_date_value(raw_value):
    if not raw_value:
        return None
    if isinstance(raw_value, datetime):
        return raw_value.date()
    text = str(raw_value).strip()
    if not text:
        return None
    normalized = text.replace('年', '-').replace('月', '-').replace('日', '')
    for fmt in ['%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d']:
        try:
            return datetime.strptime(normalized, fmt).date()
        except Exception:
            continue
    match = re.search(r'(\d{4})\D+(\d{1,2})\D+(\d{1,2})', text)
    if not match:
        return None
    try:
        return datetime(int(match.group(1)), int(match.group(2)), int(match.group(3))).date()
    except Exception:
        return None


def _get_teaching_week_no(date_obj, first_week_date, week_start_day):
    if not date_obj or not first_week_date:
        return None
    teaching_start = first_week_date - timedelta(days=(first_week_date.weekday() - week_start_day) % 7)
    diff = (date_obj - teaching_start).days
    if diff < 0:
        return None
    return (diff // 7) + 1

def _compute_teaching_week_window(start_date, end_date, first_week_date, week_start_day):
    range_start = start_date.date()
    range_end = (end_date - timedelta(days=1)).date()
    if range_end < range_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    first_full_start = range_start + timedelta(days=(week_start_day - range_start.weekday()) % 7)
    last_candidate_start = range_end - timedelta(days=6)
    if last_candidate_start < first_full_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    last_full_start = last_candidate_start - timedelta(days=(last_candidate_start.weekday() - week_start_day) % 7)
    teaching_start = first_week_date - timedelta(days=(first_week_date.weekday() - week_start_day) % 7)
    if first_full_start < teaching_start:
        first_full_start = teaching_start
    if last_full_start < first_full_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    start_week = ((first_full_start - teaching_start).days // 7) + 1
    end_week = ((last_full_start - teaching_start).days // 7) + 1
    if end_week < 1:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    if start_week < 1:
        start_week = 1
    week_count = ((last_full_start - first_full_start).days // 7) + 1
    label = str(start_week) if start_week == end_week else f'{start_week}-{end_week}'
    window_start = datetime.combine(first_full_start, datetime.min.time())
    window_end = datetime.combine(last_full_start + timedelta(days=7), datetime.min.time())
    return {
        'has_full_weeks': True,
        'start_week': start_week,
        'end_week': end_week,
        'week_count': week_count,
        'label': label,
        'window_start': window_start,
        'window_end': window_end
    }


def _get_teaching_term_start(first_week_date, week_start_day):
    return first_week_date - timedelta(days=(first_week_date.weekday() - week_start_day) % 7)


def _get_form_effective_week_no(form, reward_settings):
    if not form:
        return None
    parsed_tag = parse_audit_tag(form.audit_tag)
    if parsed_tag.get('week_correction_week_no') is not None:
        return parsed_tag['week_correction_week_no']
    lecture_date = _parse_lecture_date_value(form.lecture_date)
    week_no = _get_teaching_week_no(
        lecture_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )
    if week_no is None:
        return None
    if parsed_tag.get('legacy_late_tag') == LATE_TAG_LATE:
        week_no -= 1
    if week_no < 1:
        return None
    return week_no


def _build_teaching_month_templates(reward_window, reward_settings):
    if not reward_window.get('has_full_weeks'):
        return []
    start_week = reward_window.get('start_week')
    end_week = reward_window.get('end_week')
    if start_week is None or end_week is None:
        return []
    teaching_start = _get_teaching_term_start(
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )

    # 尝试加载自定义教学月设置
    custom_months = _load_custom_month_definitions()

    if custom_months:
        # 使用自定义教学月设置
        month_templates = []
        for idx, m_def in enumerate(custom_months):
            m_start = m_def['start_week']
            m_end = m_def['end_week']
            # 仅包含完整落在选择范围内的教学月
            if m_start < start_week or m_end > end_week:
                continue
            month_no = idx + 1
            month_start_date = teaching_start + timedelta(days=(m_start - 1) * 7)
            month_end_date = teaching_start + timedelta(days=(m_end - 1) * 7 + 6)
            weeks = []
            for week_no in range(m_start, m_end + 1):
                week_start_date = teaching_start + timedelta(days=(week_no - 1) * 7)
                week_end_date = week_start_date + timedelta(days=6)
                weeks.append({
                    'week_no': week_no,
                    'week_label': f'第{week_no}周',
                    'start_date': week_start_date.strftime('%Y-%m-%d'),
                    'end_date': week_end_date.strftime('%Y-%m-%d'),
                })
            month_label = m_def.get('label') or f'第{month_no}教学月'
            month_templates.append({
                'month_no': month_no,
                'month_label': month_label,
                'week_range_label': f'第{m_start}-{m_end}周',
                'start_week': m_start,
                'end_week': m_end,
                'start_date': month_start_date.strftime('%Y-%m-%d'),
                'end_date': month_end_date.strftime('%Y-%m-%d'),
                'weeks': weeks,
            })
        return month_templates

    # 默认逻辑：每4周为一个教学月
    month_templates = []
    start_month = ((start_week - 1) // 4) + 1
    end_month = ((end_week - 1) // 4) + 1
    for month_no in range(start_month, end_month + 1):
        month_start_week = (month_no - 1) * 4 + 1
        month_end_week = month_start_week + 3
        if month_start_week < start_week or month_end_week > end_week:
            continue
        month_start_date = teaching_start + timedelta(days=(month_start_week - 1) * 7)
        month_end_date = month_start_date + timedelta(days=27)
        weeks = []
        for week_no in range(month_start_week, month_end_week + 1):
            week_start_date = teaching_start + timedelta(days=(week_no - 1) * 7)
            week_end_date = week_start_date + timedelta(days=6)
            weeks.append({
                'week_no': week_no,
                'week_label': f'第{week_no}周',
                'start_date': week_start_date.strftime('%Y-%m-%d'),
                'end_date': week_end_date.strftime('%Y-%m-%d'),
            })
        month_templates.append({
            'month_no': month_no,
            'month_label': f'第{month_no}教学月',
            'week_range_label': f'第{month_start_week}-{month_end_week}周',
            'start_week': month_start_week,
            'end_week': month_end_week,
            'start_date': month_start_date.strftime('%Y-%m-%d'),
            'end_date': month_end_date.strftime('%Y-%m-%d'),
            'weeks': weeks,
        })
    return month_templates


def _load_custom_month_definitions():
    """加载自定义教学月设置，返回列表或 None"""
    raw = SystemSetting.get('teaching_month_definitions')
    if not raw:
        return None
    try:
        data = json.loads(raw)
        if not isinstance(data, list) or len(data) == 0:
            return None
        result = []
        for item in data:
            s = int(item.get('start_week'))
            e = int(item.get('end_week'))
            if s < 1 or e < 1 or s > e:
                continue
            result.append({
                'start_week': s,
                'end_week': e,
                'label': item.get('label', '').strip() or None
            })
        return result if result else None
    except (json.JSONDecodeError, TypeError, ValueError, KeyError):
        return None

def _resolve_error_unique_ids(candidate_groups):
    form_to_unique = {}
    for group_data in candidate_groups:
        unique_id = group_data['unique_id']
        for form in group_data['forms']:
            form_to_unique[form.id] = unique_id
    if not form_to_unique:
        return set()
    error_form_ids = db.session.query(ScoreRecord.form_id).join(
        ScoreItem, ScoreItem.score_record_id == ScoreRecord.id
    ).filter(
        ScoreRecord.form_id.in_(list(form_to_unique.keys())),
        (ScoreItem.department_score > 0) | (ScoreItem.personal_score > 0)
    ).distinct().all()
    error_unique_ids = set()
    for row in error_form_ids:
        unique_id = form_to_unique.get(row.form_id)
        if unique_id is not None:
            error_unique_ids.add(unique_id)
    return error_unique_ids

def _form_group_created_at(group_data):
    if not group_data:
        return None
    created_candidates = [
        form.created_at for form in (group_data.get('forms') or [])
        if getattr(form, 'created_at', None)
    ]
    if created_candidates:
        return min(created_candidates)
    latest_form = group_data.get('latest_form')
    return latest_form.created_at if latest_form and latest_form.created_at else None

def _sort_submission_count_rows(rows, sort_by):
    if sort_by == 'reward':
        rows.sort(key=lambda x: (-x['reward_form_count'], -x['submission_group_count'], x['number'] or '', x['name'] or ''))
    else:
        rows.sort(key=lambda x: (-x['submission_group_count'], -x['reward_form_count'], x['number'] or '', x['name'] or ''))
    for idx, row in enumerate(rows, start=1):
        row['rank'] = idx
    return rows

def _build_review_assessment_rows(users, start_date, end_date):
    by_number = {u.number: u for u in users if u.number}
    user_stats = {
        u.id: {
            'user_id': u.id,
            'name': u.name,
            'number': u.number,
            'department': u.department,
            'group': u.group,
            'linked_department_score': 0.0,
            'linked_personal_score': 0.0
        } for u in users
    }

    latest_forms = _assessment_latest_forms_in_range(list(by_number.keys()), start_date, end_date)
    for form in latest_forms:
        user = by_number.get(form.listener_number)
        if not user or not form.score_record:
            continue
        user_stats[user.id]['linked_department_score'] += float(form.score_record.total_department_score or 0.0)
        user_stats[user.id]['linked_personal_score'] += float(form.score_record.total_personal_score or 0.0)

    rows = []
    for row in user_stats.values():
        row['total_department_score'] = row['linked_department_score']
        row['total_personal_score'] = row['linked_personal_score']
        row['total_score'] = row['total_personal_score']
        rows.append(row)

    rows.sort(key=lambda x: x['total_personal_score'], reverse=True)
    for idx, row in enumerate(rows, start=1):
        row['rank'] = idx
    return rows

def _build_export_filename(filename_title, default_title):
    title = (filename_title or '').strip() or default_title
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        title = title.replace(char, '_')
    title = title.strip().strip('.')
    if not title:
        title = default_title
    if not title.lower().endswith('.xlsx'):
        title = f'{title}.xlsx'
    return title

def _autosize_worksheet(ws, min_width=12, max_width=40):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = '' if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        ws.column_dimensions[column_letter].width = max(min_width, min(max_length + 2, max_width))

def _workbook_response(wb, filename):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

def _build_reward_week_templates(reward_window, required_submission):
    templates = []
    if not reward_window.get('has_full_weeks') or not reward_window.get('window_start'):
        return templates
    window_start = reward_window['window_start']
    start_week = reward_window['start_week'] or 1
    for index in range(reward_window.get('week_count') or 0):
        week_start = window_start + timedelta(days=7 * index)
        week_end = week_start + timedelta(days=6)
        week_no = start_week + index
        templates.append({
            'week_index': index,
            'week_no': week_no,
            'week_label': f'第{week_no}周',
            'start_date': week_start.strftime('%Y-%m-%d'),
            'end_date': week_end.strftime('%Y-%m-%d'),
            'required_submission': int(required_submission or 0),
            'effective_count': 0,
            'error_count': 0,
            'reward_form_count': 0
        })
    return templates

def _get_reward_week_index(group_data, reward_window, reward_settings):
    if not reward_window.get('has_full_weeks'):
        return None
    latest_form = group_data.get('latest_form') if group_data else None
    if not latest_form:
        return None
    week_no = _get_form_effective_week_no(latest_form, reward_settings)
    if week_no is None:
        return None
    start_week = reward_window.get('start_week')
    end_week = reward_window.get('end_week')
    if start_week is None or end_week is None or week_no < start_week or week_no > end_week:
        return None
    return week_no - start_week

def _build_submission_count_rows(users, start_date, end_date, reward_settings, reward_window, time_filter_type='created'):
    by_number = {u.number: u for u in users if u.number}
    week_templates = _build_reward_week_templates(reward_window, reward_settings['required_submission'])
    user_stats = {
        u.id: {
            'user_id': u.id,
            'name': u.name,
            'number': u.number,
            'department': u.department,
            'group': u.group,
            'submission_group_count': 0,
            'reward_effective_count': 0,
            'reward_error_count': 0,
            'reward_form_count': 0,
            'reward_week_details': [dict(template) for template in week_templates]
        } for u in users
    }

    groups = _latest_form_groups_for_users(list(by_number.keys()))
    reward_candidates = []
    for group_data in groups:
        latest_form = group_data['latest_form']
        if not latest_form or not latest_form.created_at:
            continue
        user = by_number.get(latest_form.listener_number)
        if not user:
            continue
        filter_dt = _build_review_form_filter_datetime(group_data, time_filter_type)
        if filter_dt and start_date <= filter_dt < end_date:
            user_stats[user.id]['submission_group_count'] += 1
        week_index = _get_reward_week_index(group_data, reward_window, reward_settings)
        if week_index is None or latest_form.status != '中心已审核':
            continue
        reward_candidates.append({
            'group_data': group_data,
            'user_id': user.id,
            'week_index': week_index
        })

    error_unique_ids = _resolve_error_unique_ids([item['group_data'] for item in reward_candidates])
    for item in reward_candidates:
        row = user_stats[item['user_id']]
        week_detail = row['reward_week_details'][item['week_index']]
        row['reward_effective_count'] += 1
        week_detail['effective_count'] += 1
        if item['group_data']['unique_id'] in error_unique_ids:
            row['reward_error_count'] += 1
            week_detail['error_count'] += 1

    for row in user_stats.values():
        total_reward_forms = 0
        for week_detail in row['reward_week_details']:
            week_reward = max(
                int(week_detail['effective_count']) - int(week_detail['error_count']) - int(week_detail['required_submission']),
                0
            )
            week_detail['reward_form_count'] = week_reward
            total_reward_forms += week_reward
        row['reward_form_count'] = total_reward_forms

    return list(user_stats.values())


def _format_score_value_label(value):
    try:
        numeric = float(value or 0)
    except Exception:
        numeric = 0.0
    if numeric.is_integer():
        return str(int(numeric))
    return f'{numeric:.2f}'.rstrip('0').rstrip('.')


def _build_department_monthly_assessment_payload(current_user_id, start_date, end_date, department_names):
    reward_settings, reward_err = _get_teaching_reward_settings()
    if reward_err:
        return None, reward_err

    reward_window = _compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )
    month_templates = _build_teaching_month_templates(reward_window, reward_settings)
    selected_departments, department_user_map = _resolve_selected_departments(current_user_id, department_names)
    payload = {
        'departments': [],
        'meta': {
            'range_start': start_date.strftime('%Y-%m-%d'),
            'range_end': (end_date - timedelta(days=1)).strftime('%Y-%m-%d'),
            'week_label': reward_window.get('label'),
            'has_full_weeks': reward_window.get('has_full_weeks', False),
            'has_full_months': len(month_templates) > 0,
            'month_count': len(month_templates),
            'month_labels': [item['month_label'] for item in month_templates],
            'months': month_templates,
            'required_submission_per_member': reward_settings['required_submission'],
            'scope_label': '、'.join(selected_departments) if selected_departments else '未选择部门',
            'period_label': '、'.join(item['month_label'] for item in month_templates) if month_templates else '无完整教学月',
            'disclaimer': '考评情况为系统根据当前考评规则和数据自动给出，不代表真实考评情况，请认真核实，建议核实无误后保存或导出数据。'
        }
    }

    if not selected_departments:
        return payload, None
    if not month_templates:
        for department_name in selected_departments:
            payload['departments'].append({
                'department': department_name,
                'member_count': len(department_user_map.get(department_name, [])),
                'months': []
            })
        return payload, None

    selected_week_nos = {
        week['week_no']
        for month in month_templates
        for week in month.get('weeks', [])
    }

    # 查询所有相关用户的考核豁免规则
    all_user_ids_for_exemption = []
    for dept_users in department_user_map.values():
        all_user_ids_for_exemption.extend([u.id for u in dept_users])
    exempt_map = {}  # {user_id: [(start_week, end_week), ...]}
    if all_user_ids_for_exemption:
        exemption_records = AssessmentOverride.query.filter(
            AssessmentOverride.user_id.in_(all_user_ids_for_exemption),
            AssessmentOverride.override_type.in_(ASSESSMENT_EXEMPT_OVERRIDE_TYPES)
        ).all()
        for ov in exemption_records:
            exempt_map.setdefault(ov.user_id, []).append((ov.start_week, ov.end_week))

    def _is_user_exempted(user_id, week_no):
        for s, e in exempt_map.get(user_id, []):
            if s <= week_no <= e:
                return True
        return False

    user_by_number = {}
    department_rows = {}
    required_submission_per_member = int(reward_settings['required_submission'] or 0)
    for department_name in selected_departments:
        users = department_user_map.get(department_name, [])
        user_by_number.update({user.number: user for user in users if user.number})
        month_rows = []
        week_map = {}
        member_count = len(users)
        for month in month_templates:
            weeks = []
            for week in month.get('weeks', []):
                wno = week['week_no']
                active_count = sum(1 for u in users if not _is_user_exempted(u.id, wno))
                week_required = active_count * required_submission_per_member
                week_row = {
                    'week_no': wno,
                    'week_label': week['week_label'],
                    'start_date': week['start_date'],
                    'end_date': week['end_date'],
                    'required_submission': week_required,
                    'effective_count': 0,
                    'error_count': 0,
                    'extra_count': 0,
                    'shortage_count': 0,
                    'missing_count': 0,
                    'missing_penalty': 0,
                    'assessment_total_score': 0.0,
                    'assessment_actual_score': 0.0,
                    'raw_score': 0.0,
                    'final_score': 0.0,
                    'assessment_tiers': [],
                    '_assessment_tier_map': {},
                    '_active_member_count': active_count
                }
                weeks.append(week_row)
                week_map[wno] = week_row
            month_summary_required = sum(w['required_submission'] for w in weeks)
            month_rows.append({
                'month_no': month['month_no'],
                'month_label': month['month_label'],
                'week_range_label': month['week_range_label'],
                'start_date': month['start_date'],
                'end_date': month['end_date'],
                'weeks': weeks,
                'summary': {
                    'required_submission': month_summary_required,
                    'effective_count': 0,
                    'error_count': 0,
                    'extra_count': 0,
                    'shortage_count': 0,
                    'missing_count': 0,
                    'missing_penalty': 0,
                    'assessment_total_score': 0.0,
                    'assessment_actual_score': 0.0,
                    'raw_score': 0.0,
                    'final_score': 0.0
                }
            })
        department_rows[department_name] = {
            'department': department_name,
            'member_count': member_count,
            'required_submission_per_member': required_submission_per_member,
            'required_submission_per_week': member_count * required_submission_per_member,
            'users': users,
            'months': month_rows,
            '_week_map': week_map
        }

    groups = _latest_form_groups_for_users(list(user_by_number.keys()))
    effective_form_candidates = []
    user_week_effective = {}
    for group_data in groups:
        latest_form = group_data.get('latest_form')
        if not latest_form:
            continue
        user = user_by_number.get(latest_form.listener_number)
        if not user:
            continue
        department_name = user.department
        department_row = department_rows.get(department_name)
        if not department_row:
            continue
        week_no = _get_form_effective_week_no(latest_form, reward_settings)
        if week_no not in selected_week_nos:
            continue
        week_row = department_row['_week_map'].get(week_no)
        if not week_row:
            continue

        # 跳过被豁免用户的表单
        if _is_user_exempted(user.id, week_no):
            continue

        if latest_form.status == '中心已审核':
            week_row['effective_count'] += 1
            user_week_effective[(user.id, week_no)] = user_week_effective.get((user.id, week_no), 0) + 1
            effective_form_candidates.append({
                'group_data': group_data,
                'department': department_name,
                'week_no': week_no,
            })

        if latest_form.score_record:
            for item in latest_form.score_record.items:
                score_value = float(item.personal_score or 0.0)
                if score_value <= 0:
                    continue
                week_row['assessment_total_score'] += score_value
                score_key = _format_score_value_label(score_value)
                tier_entry = week_row['_assessment_tier_map'].setdefault(score_key, {
                    'score_value': score_value,
                    'score_label': score_key,
                    'item_count': 0,
                    'total_score': 0.0
                })
                tier_entry['item_count'] += 1
                tier_entry['total_score'] += score_value

    error_unique_ids = _resolve_error_unique_ids([item['group_data'] for item in effective_form_candidates])
    for item in effective_form_candidates:
        if item['group_data']['unique_id'] in error_unique_ids:
            department_rows[item['department']]['_week_map'][item['week_no']]['error_count'] += 1

    for department_name, department_row in department_rows.items():
        users = department_row['users']
        for week_no, week_row in department_row['_week_map'].items():
            for user in users:
                # 跳过被豁免的用户
                if _is_user_exempted(user.id, week_no):
                    continue
                effective_count = user_week_effective.get((user.id, week_no), 0)
                diff = effective_count - department_row['required_submission_per_member']
                if diff > 0:
                    week_row['extra_count'] += diff
                elif diff < 0:
                    week_row['shortage_count'] += abs(diff)

            week_row['missing_count'] = max(
                int(week_row['required_submission']) - int(week_row['effective_count']),
                0
            )
            week_row['missing_penalty'] = week_row['missing_count'] * 3
            week_row['raw_score'] = float(week_row['extra_count']) - float(week_row['shortage_count']) * 2
            week_row['assessment_actual_score'] = max(float(week_row['assessment_total_score']) - 5, 0.0)
            week_row['final_score'] = (
                float(week_row['raw_score'])
                - float(week_row['missing_penalty'])
                - float(week_row['assessment_actual_score'])
            )
            week_row['assessment_tiers'] = sorted(
                week_row['_assessment_tier_map'].values(),
                key=lambda item: (float(item['score_value']), item['score_label'])
            )
            week_row.pop('_assessment_tier_map', None)

        for month_row in department_row['months']:
            for week_row in month_row['weeks']:
                for key in ['required_submission', 'effective_count', 'error_count', 'extra_count', 'shortage_count', 'missing_count', 'missing_penalty', 'assessment_total_score', 'assessment_actual_score', 'raw_score', 'final_score']:
                    month_row['summary'][key] += float(week_row[key])
            for int_key in ['required_submission', 'effective_count', 'error_count', 'extra_count', 'shortage_count', 'missing_count', 'missing_penalty']:
                month_row['summary'][int_key] = int(month_row['summary'][int_key])

        department_row.pop('users', None)
        department_row.pop('_week_map', None)
        payload['departments'].append({
            'department': department_row['department'],
            'member_count': department_row['member_count'],
            'required_submission_per_member': department_row['required_submission_per_member'],
            'required_submission_per_week': department_row['required_submission_per_week'],
            'months': department_row['months'],
        })

    return payload, None


def _build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type='created'):
    reward_settings, reward_err = _get_teaching_reward_settings()
    if reward_err:
        return None, reward_err
    reward_window = _compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )
    rows = _build_submission_count_rows(
        users,
        start_date,
        end_date,
        reward_settings,
        reward_window,
        time_filter_type=time_filter_type
    )
    _sort_submission_count_rows(rows, sort_by)
    return {
        'users': rows,
        'sort_by': sort_by,
        'time_filter_type': time_filter_type,
        'reward_meta': {
            'range_start': start_date.strftime('%Y-%m-%d'),
            'range_end': (end_date - timedelta(days=1)).strftime('%Y-%m-%d'),
            'week_label': reward_window['label'],
            'start_week': reward_window['start_week'],
            'end_week': reward_window['end_week'],
            'week_count': reward_window['week_count'],
            'has_full_weeks': reward_window['has_full_weeks'],
            'required_submission_per_week': reward_settings['required_submission'],
            'disclaimer': '奖励表情况为系统根据当前奖励表规则和数据自动给出，不代表真实考评情况，请认真核实，建议核实无误后保存或导出数据。'
        }
    }, None


def _build_submission_snapshot_workbook(payload):
    rows = payload.get('users', [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '交表统计'
    headers = ['排名', '姓名', '编号', '部门', '小组', '交表数量', '奖励表数量', '有效表单数', '错误表单数']
    ws.append(headers)
    for row in rows:
        ws.append([
            row.get('rank'),
            row.get('name'),
            row.get('number'),
            row.get('department'),
            row.get('group'),
            row.get('submission_group_count'),
            row.get('reward_form_count'),
            row.get('reward_effective_count'),
            row.get('reward_error_count')
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    detail_ws = wb.create_sheet('奖励表详情')
    detail_headers = ['姓名', '编号', '周次', '开始日期', '结束日期', '有效表单数', '错误表单数', '每周需交表数', '奖励表数']
    detail_ws.append(detail_headers)
    for row in rows:
        for week in row.get('reward_week_details', []):
            detail_ws.append([
                row.get('name'),
                row.get('number'),
                week.get('week_label'),
                week.get('start_date'),
                week.get('end_date'),
                week.get('effective_count'),
                week.get('error_count'),
                week.get('required_submission'),
                week.get('reward_form_count')
            ])
    for cell in detail_ws[1]:
        cell.font = Font(bold=True)

    _autosize_worksheet(ws)
    _autosize_worksheet(detail_ws)
    return wb


def _build_department_monthly_workbook(payload):
    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = '月度考评汇总'
    summary_headers = ['部门', '教学月', '周次范围', '成员数', '月多交表', '月少交表', '月错误表单扣分（原始分）', '月错误表单扣分（实际扣分）', '月总考评分']
    summary_ws.append(summary_headers)
    for department_row in payload.get('departments', []):
        for month_row in department_row.get('months', []):
            summary = month_row.get('summary', {})
            summary_ws.append([
                department_row.get('department'),
                month_row.get('month_label'),
                month_row.get('week_range_label'),
                department_row.get('member_count'),
                summary.get('extra_count'),
                summary.get('shortage_count'),
                summary.get('assessment_total_score'),
                summary.get('assessment_actual_score'),
                summary.get('final_score'),
            ])
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)

    detail_ws = wb.create_sheet('周度考评明细')
    detail_headers = ['部门', '教学月', '周次', '时间范围', '需交表数', '有效表单数', '错误表单数', '多交表（+1分/张）', '少交表（-2分/张）', '实际交表少于应交表数（-3/张）', '错误表单扣分（原始分）', '错误表单扣分（实际扣分）', '周考评分', '扣分档位汇总']
    detail_ws.append(detail_headers)
    for department_row in payload.get('departments', []):
        for month_row in department_row.get('months', []):
            for week_row in month_row.get('weeks', []):
                tier_text = '；'.join(
                    f"{tier.get('score_label')}分 x {tier.get('item_count')}项 = {tier.get('total_score')}"
                    for tier in week_row.get('assessment_tiers', [])
                )
                detail_ws.append([
                    department_row.get('department'),
                    month_row.get('month_label'),
                    week_row.get('week_label'),
                    f"{week_row.get('start_date')} ~ {week_row.get('end_date')}",
                    week_row.get('required_submission'),
                    week_row.get('effective_count'),
                    week_row.get('error_count'),
                    week_row.get('extra_count'),
                    week_row.get('shortage_count'),
                    week_row.get('missing_count'),
                    week_row.get('assessment_total_score'),
                    week_row.get('assessment_actual_score'),
                    week_row.get('final_score'),
                    tier_text
                ])
    for cell in detail_ws[1]:
        cell.font = Font(bold=True)

    _autosize_worksheet(summary_ws)
    _autosize_worksheet(detail_ws, max_width=60)
    return wb

@admin_bp.route('/api/review/submission-count/stats', methods=['GET'])
@login_required
def get_submission_count_stats():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = _resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': True, 'users': []})
    reward_settings, reward_err = _get_teaching_reward_settings()
    if reward_err:
        return jsonify({'success': False, 'message': reward_err}), 400
    reward_window = _compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )
    sort_by = request.args.get('sort_by', 'submission')
    if sort_by not in ['submission', 'reward']:
        sort_by = 'submission'
    time_filter_type = _normalize_review_form_time_filter(request.args.get('time_filter_type'))
    payload, payload_err = _build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type)
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    rows = []
    for row in payload.get('users', []):
        copied_row = dict(row)
        copied_row.pop('reward_week_details', None)
        rows.append(copied_row)
    _sort_submission_count_rows(rows, sort_by)
    return jsonify({
        'success': True,
        'users': rows,
        'sort_by': sort_by,
        'time_filter_type': time_filter_type,
        'reward_meta': payload.get('reward_meta', {})
    })

@admin_bp.route('/api/review/submission-count/export', methods=['GET'])
@login_required
def export_submission_count_stats():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = _resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': False, 'message': '无可导出用户'}), 400
    sort_by = request.args.get('sort_by', 'submission')
    if sort_by not in ['submission', 'reward']:
        sort_by = 'submission'
    time_filter_type = _normalize_review_form_time_filter(request.args.get('time_filter_type'))
    payload, payload_err = _build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type)
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    wb = _build_submission_snapshot_workbook(payload)

    default_title = f'{start_date_str}至{end_date_str}交表统计'
    filename = _build_export_filename(request.args.get('filename_title'), default_title)
    return _workbook_response(wb, filename)

@admin_bp.route('/api/review/submission-count/reward-detail/<int:user_id>', methods=['GET'])
@login_required
def get_submission_reward_detail(user_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    users = _resolve_assessment_users(session['user_id'], [str(user_id)])
    if not users:
        return jsonify({'success': False, 'message': '无权限查看该用户'}), 403
    reward_settings, reward_err = _get_teaching_reward_settings()
    if reward_err:
        return jsonify({'success': False, 'message': reward_err}), 400
    reward_window = _compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )
    row = _build_submission_count_rows(users, start_date, end_date, reward_settings, reward_window)[0]
    return jsonify({
        'success': True,
        'user': {
            'id': users[0].id,
            'name': users[0].name,
            'number': users[0].number,
            'department': users[0].department,
            'group': users[0].group
        },
        'summary': {
            'submission_group_count': row['submission_group_count'],
            'reward_effective_count': row['reward_effective_count'],
            'reward_error_count': row['reward_error_count'],
            'reward_form_count': row['reward_form_count']
        },
        'weeks': row['reward_week_details'],
        'reward_meta': {
            'week_label': reward_window['label'],
            'start_week': reward_window['start_week'],
            'end_week': reward_window['end_week'],
            'week_count': reward_window['week_count'],
            'has_full_weeks': reward_window['has_full_weeks'],
            'required_submission_per_week': reward_settings['required_submission']
        }
    })

@admin_bp.route('/api/review/submission-count/detail/<int:user_id>', methods=['GET'])
@login_required
def get_submission_count_detail(user_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    users = _resolve_assessment_users(session['user_id'], [str(user_id)])
    if not users:
        return jsonify({'success': False, 'message': '无权限查看该用户'}), 403
    time_filter_type = _normalize_review_form_time_filter(request.args.get('time_filter_type'))
    user = users[0]
    if not user.number:
        return jsonify({'success': True, 'user': {'id': user.id, 'name': user.name, 'number': user.number, 'department': user.department, 'group': user.group}, 'groups': []})

    reviewer_display_mode = get_reviewer_display_mode()
    groups = _latest_form_groups_for_users([user.number])
    filtered_groups = []
    reviewer_ids = set()
    for group_data in groups:
        latest_form = group_data['latest_form']
        if not latest_form:
            continue
        filter_dt = _build_review_form_filter_datetime(group_data, time_filter_type)
        if not filter_dt or filter_dt < start_date or filter_dt >= end_date:
            continue
        for form in group_data['forms']:
            if form.reviewer_id:
                reviewer_ids.add(form.reviewer_id)
    reviewer_map = {}
    if reviewer_ids:
        reviewers = User.query.filter(User.id.in_(list(reviewer_ids))).all()
        reviewer_map = {r.id: r for r in reviewers}

    for group_data in groups:
        latest_form = group_data['latest_form']
        if not latest_form or not latest_form.created_at:
            continue
        if latest_form.created_at < start_date or latest_form.created_at >= end_date:
            continue
        forms_data = []
        for form in group_data['forms']:
            reviewer = reviewer_map.get(form.reviewer_id) if form.reviewer_id else None
            if reviewer_display_mode == 'number':
                reviewer_display = reviewer.number if reviewer else '-'
            else:
                reviewer_display = reviewer.name if reviewer else '-'
            forms_data.append({
                'id': form.id,
                'status': form.status,
                'created_at': form.created_at.strftime('%Y-%m-%d %H:%M:%S') if form.created_at else '-',
                'updated_at': form.updated_at.strftime('%Y-%m-%d %H:%M:%S') if form.updated_at else '-',
                'review_comment': form.review_comment or '',
                'reviewer_display': reviewer_display,
                'teacher_name': form.teacher_name or '',
                'course_title': form.course_title or '',
                'lecture_date': form.lecture_date or ''
            })
        filtered_groups.append({
            'unique_id': group_data['unique_id'],
            'latest_created_at': latest_form.created_at.strftime('%Y-%m-%d %H:%M:%S') if latest_form.created_at else '-',
            'latest_updated_at': latest_form.updated_at.strftime('%Y-%m-%d %H:%M:%S') if latest_form.updated_at else '-',
            'filter_datetime': filter_dt.strftime('%Y-%m-%d %H:%M:%S') if filter_dt else '-',
            'teacher_name': latest_form.teacher_name or '',
            'course_title': latest_form.course_title or '',
            'lecture_date': latest_form.lecture_date or '',
            'forms': forms_data
        })

    filtered_groups.sort(key=lambda x: x['latest_created_at'], reverse=True)
    return jsonify({
        'success': True,
        'user': {
            'id': user.id,
            'name': user.name,
            'number': user.number,
            'department': user.department,
            'group': user.group
        },
        'groups': filtered_groups,
        'time_filter_type': time_filter_type
    })


@admin_bp.route('/api/review/submission-count/snapshots', methods=['GET'])
@login_required
def list_submission_count_snapshots():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    return jsonify({
        'success': True,
        'snapshots': _build_snapshot_list_items(SNAPSHOT_TYPE_SUBMISSION_REWARD, session['user_id'])
    })


@admin_bp.route('/api/review/submission-count/snapshots', methods=['POST'])
@login_required
def save_submission_count_snapshot():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    data = request.get_json() or {}
    start_date, end_date, err = _parse_assessment_range(data.get('start_date'), data.get('end_date'))
    if err:
        return jsonify({'success': False, 'message': err}), 400
    users = _resolve_assessment_users(session['user_id'], data.get('user_ids') or [])
    if not users:
        return jsonify({'success': False, 'message': '无可保存用户'}), 400
    sort_by = data.get('sort_by', 'submission')
    if sort_by not in ['submission', 'reward']:
        sort_by = 'submission'
    time_filter_type = _normalize_review_form_time_filter(data.get('time_filter_type'))
    payload, payload_err = _build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type)
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    if not payload.get('reward_meta', {}).get('has_full_weeks'):
        return jsonify({'success': False, 'message': '当前筛选范围内无完整教学周，无法保存奖励表历史'}), 400

    filters_data = {
        'start_date': data.get('start_date'),
        'end_date': data.get('end_date'),
        'sort_by': sort_by,
        'time_filter_type': time_filter_type,
        'user_ids': [user.id for user in users],
        'scope_label': f'共 {len(users)} 人',
        'period_label': payload.get('reward_meta', {}).get('week_label', '')
    }
    title = (data.get('title') or '').strip() or f"{data.get('start_date')}至{data.get('end_date')}交表统计"
    record = _create_statistics_snapshot(
        SNAPSHOT_TYPE_SUBMISSION_REWARD,
        title,
        filters_data,
        payload,
        session['user_id']
    )
    return jsonify({
        'success': True,
        'message': '奖励表历史已保存',
        'snapshot': {
            'id': record.id,
            'title': record.title
        }
    })


@admin_bp.route('/api/review/submission-count/snapshots/<int:snapshot_id>', methods=['GET'])
@login_required
def get_submission_count_snapshot(snapshot_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_SUBMISSION_REWARD, session['user_id'])
    if error_response:
        return error_response, status_code
    filters_data, payload_data = _load_snapshot_payload(record)
    return jsonify({
        'success': True,
        'snapshot': {
            'id': record.id,
            'title': record.title,
            'created_at': record.created_at.strftime('%Y-%m-%d %H:%M:%S') if record.created_at else '',
            'filters': filters_data,
            'payload': payload_data
        }
    })


@admin_bp.route('/api/review/submission-count/snapshots/<int:snapshot_id>/export', methods=['GET'])
@login_required
def export_submission_count_snapshot(snapshot_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_SUBMISSION_REWARD, session['user_id'])
    if error_response:
        return error_response, status_code
    _, payload_data = _load_snapshot_payload(record)
    wb = _build_submission_snapshot_workbook(payload_data)
    filename = _build_export_filename(request.args.get('filename_title'), record.title)
    return _workbook_response(wb, filename)


@admin_bp.route('/api/review/department-monthly-assessment/stats', methods=['GET'])
@login_required
def get_department_monthly_assessment_stats():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    payload, payload_err = _build_department_monthly_assessment_payload(
        session['user_id'],
        start_date,
        end_date,
        request.args.getlist('departments')
    )
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    return jsonify({'success': True, **payload})


@admin_bp.route('/api/review/department-monthly-assessment/export', methods=['GET'])
@login_required
def export_department_monthly_assessment_stats():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    payload, payload_err = _build_department_monthly_assessment_payload(
        session['user_id'],
        start_date,
        end_date,
        request.args.getlist('departments')
    )
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    if not payload.get('meta', {}).get('has_full_months'):
        return jsonify({'success': False, 'message': '当前筛选范围内无完整教学月，无法导出月度考评'}), 400
    wb = _build_department_monthly_workbook(payload)
    default_title = f'{start_date_str}至{end_date_str}部门月度考评统计'
    filename = _build_export_filename(request.args.get('filename_title'), default_title)
    return _workbook_response(wb, filename)


@admin_bp.route('/api/review/department-monthly-assessment/snapshots', methods=['GET'])
@login_required
def list_department_monthly_assessment_snapshots():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    return jsonify({
        'success': True,
        'snapshots': _build_snapshot_list_items(SNAPSHOT_TYPE_DEPARTMENT_MONTHLY, session['user_id'])
    })


@admin_bp.route('/api/review/department-monthly-assessment/snapshots', methods=['POST'])
@login_required
def save_department_monthly_assessment_snapshot():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    data = request.get_json() or {}
    start_date, end_date, err = _parse_assessment_range(data.get('start_date'), data.get('end_date'))
    if err:
        return jsonify({'success': False, 'message': err}), 400
    departments = data.get('departments') or []
    payload, payload_err = _build_department_monthly_assessment_payload(
        session['user_id'],
        start_date,
        end_date,
        departments
    )
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    if not payload.get('meta', {}).get('has_full_months'):
        return jsonify({'success': False, 'message': '当前筛选范围内无完整教学月，无法保存月度考评历史'}), 400
    filters_data = {
        'start_date': data.get('start_date'),
        'end_date': data.get('end_date'),
        'departments': departments,
        'scope_label': payload.get('meta', {}).get('scope_label', ''),
        'period_label': payload.get('meta', {}).get('period_label', '')
    }
    title = (data.get('title') or '').strip() or f"{data.get('start_date')}至{data.get('end_date')}部门月度考评统计"
    record = _create_statistics_snapshot(
        SNAPSHOT_TYPE_DEPARTMENT_MONTHLY,
        title,
        filters_data,
        payload,
        session['user_id']
    )
    return jsonify({
        'success': True,
        'message': '部门月度考评历史已保存',
        'snapshot': {
            'id': record.id,
            'title': record.title
        }
    })


@admin_bp.route('/api/review/department-monthly-assessment/snapshots/<int:snapshot_id>', methods=['GET'])
@login_required
def get_department_monthly_assessment_snapshot(snapshot_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_DEPARTMENT_MONTHLY, session['user_id'])
    if error_response:
        return error_response, status_code
    filters_data, payload_data = _load_snapshot_payload(record)
    return jsonify({
        'success': True,
        'snapshot': {
            'id': record.id,
            'title': record.title,
            'created_at': record.created_at.strftime('%Y-%m-%d %H:%M:%S') if record.created_at else '',
            'filters': filters_data,
            'payload': payload_data
        }
    })


@admin_bp.route('/api/review/department-monthly-assessment/snapshots/<int:snapshot_id>/export', methods=['GET'])
@login_required
def export_department_monthly_assessment_snapshot(snapshot_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_DEPARTMENT_MONTHLY, session['user_id'])
    if error_response:
        return error_response, status_code
    _, payload_data = _load_snapshot_payload(record)
    wb = _build_department_monthly_workbook(payload_data)
    filename = _build_export_filename(request.args.get('filename_title'), record.title)
    return _workbook_response(wb, filename)

@admin_bp.route('/api/review/assessment/stats', methods=['GET'])
@login_required
def get_review_assessment_stats():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('审表考评统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = _resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': True, 'users': []})

    rows = _build_review_assessment_rows(users, start_date, end_date)
    return jsonify({'success': True, 'users': rows})

@admin_bp.route('/api/review/assessment/export', methods=['GET'])
@login_required
def export_review_assessment_stats():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('审表考评统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = _resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': False, 'message': '无可导出用户'}), 400

    rows = _build_review_assessment_rows(users, start_date, end_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '考评统计'
    headers = ['排名', '姓名', '编号', '部门', '小组', '部门扣分', '个人扣分', '总个人扣分']
    ws.append(headers)
    for row in rows:
        ws.append([
            row['rank'],
            row['name'],
            row['number'],
            row['department'],
            row['group'],
            float(row['linked_department_score'] or 0.0),
            float(row['linked_personal_score'] or 0.0),
            float(row['total_personal_score'] or 0.0)
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    _autosize_worksheet(ws)

    default_title = f'{start_date_str}至{end_date_str}考评统计'
    filename = _build_export_filename(request.args.get('filename_title'), default_title)
    return _workbook_response(wb, filename)

@admin_bp.route('/api/review/assessment/detail/<int:user_id>', methods=['GET'])
@login_required
def get_review_assessment_detail(user_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('审表考评统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = _parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    users = _resolve_assessment_users(session['user_id'], [str(user_id)])
    if not users:
        return jsonify({'success': False, 'message': '无权限查看该用户'}), 403
    user = users[0]

    entries = []
    groups = []
    linked_department_total = 0.0
    linked_personal_total = 0.0

    latest_forms = _assessment_latest_forms_in_range([user.number], start_date, end_date)
    for form in latest_forms:
        if not form.score_record:
            continue
        display_time = _assessment_form_latest_timestamp(form)
        group_items = []
        group_department_total = 0.0
        group_personal_total = 0.0
        for item in form.score_record.items:
            dept_score = float(item.department_score or 0.0)
            pers_score = float(item.personal_score or 0.0)
            linked_department_total += dept_score
            linked_personal_total += pers_score
            group_department_total += dept_score
            group_personal_total += pers_score
            item_data = {
                'issue_detail': item.reason or '',
                'department_score': dept_score,
                'personal_score': pers_score
            }
            group_items.append(item_data)
            entries.append({
                'source': '反馈表单',
                'date': display_time.strftime('%Y-%m-%d %H:%M') if display_time else '-',
                'teacher_name': form.teacher_name or '',
                'course_title': form.course_title or '',
                **item_data,
                'sort_time': display_time.isoformat() if display_time else ''
            })
        groups.append({
            'form_id': form.id,
            'unique_id': form.unique_id or '',
            'date': display_time.strftime('%Y-%m-%d %H:%M') if display_time else '-',
            'teacher_name': form.teacher_name or '',
            'course_title': form.course_title or '',
            'total_department_score': group_department_total,
            'total_personal_score': group_personal_total,
            'item_count': len(group_items),
            'items': group_items,
            'sort_time': display_time.isoformat() if display_time else ''
        })

    groups.sort(key=lambda x: x['sort_time'], reverse=True)
    entries.sort(key=lambda x: x['sort_time'], reverse=True)
    for group in groups:
        group.pop('sort_time', None)
    for entry in entries:
        entry.pop('sort_time', None)

    return jsonify({
        'success': True,
        'user': {
            'id': user.id,
            'name': user.name,
            'number': user.number,
            'department': user.department,
            'group': user.group
        },
        'summary': {
            'linked_department_score': linked_department_total,
            'linked_personal_score': linked_personal_total,
            'total_department_score': linked_department_total,
            'total_personal_score': linked_personal_total,
            'total_score': linked_personal_total
        },
        'groups': groups,
        'entries': entries
    })

def _read_manual_assessment_file(file_storage):
    if not file_storage or not file_storage.filename:
        return None, '请先选择导入文件'
    if not allowed_file(file_storage.filename):
        return None, '仅支持xls或xlsx文件'
    try:
        return pd.read_excel(file_storage), None
    except Exception as e:
        return None, f'读取文件失败：{str(e)}'

def _normalize_manual_assessment_items(items):
    normalized_items = []
    for item in items or []:
        if isinstance(item, dict):
            reason = item.get('reason')
            department_score = item.get('department_score')
            personal_score = item.get('personal_score')
        else:
            reason = getattr(item, 'reason', '')
            department_score = getattr(item, 'department_score', 0.0)
            personal_score = getattr(item, 'personal_score', 0.0)
        normalized_items.append({
            'reason': str(reason or '').strip() or '手动导入考评项',
            'department_score': float(department_score or 0.0),
            'personal_score': float(personal_score or 0.0)
        })
    normalized_items.sort(key=lambda item: (item['reason'], item['department_score'], item['personal_score']))
    return normalized_items

def _build_score_snapshot(items, score_record=None):
    normalized_items = _normalize_manual_assessment_items(items)
    if score_record:
        total_department_score = float(score_record.total_department_score or 0.0)
        total_personal_score = float(score_record.total_personal_score or 0.0)
    else:
        total_department_score = sum(item['department_score'] for item in normalized_items)
        total_personal_score = sum(item['personal_score'] for item in normalized_items)
    return {
        'total_department_score': total_department_score,
        'total_personal_score': total_personal_score,
        'items': normalized_items
    }

def _build_existing_score_snapshot(form):
    if not form or not form.score_record:
        return _build_score_snapshot([])
    return _build_score_snapshot(form.score_record.items, score_record=form.score_record)

def _manual_assessment_snapshots_equal(before_snapshot, after_snapshot):
    return (
        float(before_snapshot.get('total_department_score') or 0.0) == float(after_snapshot.get('total_department_score') or 0.0)
        and float(before_snapshot.get('total_personal_score') or 0.0) == float(after_snapshot.get('total_personal_score') or 0.0)
        and before_snapshot.get('items', []) == after_snapshot.get('items', [])
    )

def _parse_manual_assessment_rows(df):
    required_columns = ['表单ID', '考评项', '部门扣分', '个人扣分']
    for col in required_columns:
        if col not in df.columns:
            return None, f'缺少字段：{col}'

    valid_row_count = 0
    skipped_count = 0
    skipped_details = []
    grouped_items = {}
    forms_by_id = {}

    for idx, row in df.iterrows():
        form_id_value = row.get('表单ID')
        reason_value = row.get('考评项')
        dept_value = row.get('部门扣分')
        pers_value = row.get('个人扣分')

        if pd.isna(form_id_value) and pd.isna(reason_value) and pd.isna(dept_value) and pd.isna(pers_value):
            continue

        form_id = _to_int_or_none(form_id_value)
        if not form_id:
            skipped_count += 1
            if len(skipped_details) < 15:
                skipped_details.append(f'第{idx + 2}行：表单ID无效')
            continue

        form = LectureForm.query.get(form_id)
        if not form:
            skipped_count += 1
            if len(skipped_details) < 15:
                skipped_details.append(f'第{idx + 2}行：表单ID {form_id} 不存在')
            continue

        reason = str(reason_value).strip() if not pd.isna(reason_value) else ''
        try:
            department_score = float(dept_value) if not pd.isna(dept_value) else 0.0
            personal_score = float(pers_value) if not pd.isna(pers_value) else 0.0
        except Exception:
            skipped_count += 1
            if len(skipped_details) < 15:
                skipped_details.append(f'第{idx + 2}行：扣分字段格式错误')
            continue

        if not reason and department_score == 0 and personal_score == 0:
            skipped_count += 1
            if len(skipped_details) < 15:
                skipped_details.append(f'第{idx + 2}行：考评项和扣分不能同时为空/0')
            continue

        normalized_item = {
            'reason': reason or '手动导入考评项',
            'department_score': department_score,
            'personal_score': personal_score
        }
        valid_row_count += 1
        forms_by_id[form_id] = form
        grouped_items.setdefault(form_id, []).append(normalized_item)

    return {
        'valid_row_count': valid_row_count,
        'skipped_count': skipped_count,
        'skipped_details': skipped_details,
        'grouped_items': grouped_items,
        'forms_by_id': forms_by_id
    }, None

def _build_manual_assessment_preview(parsed_result):
    changed_forms = []
    changed_grouped_items = {}
    unchanged_form_ids = []
    changed_row_count = 0

    for form_id, items in parsed_result['grouped_items'].items():
        form = parsed_result['forms_by_id'][form_id]
        before_snapshot = _build_existing_score_snapshot(form)
        after_snapshot = _build_score_snapshot(items)
        if _manual_assessment_snapshots_equal(before_snapshot, after_snapshot):
            unchanged_form_ids.append(form_id)
            continue

        changed_grouped_items[form_id] = _normalize_manual_assessment_items(items)
        changed_row_count += len(items)
        changed_forms.append({
            'form_id': form.id,
            'listener_number': form.listener_number or '',
            'teacher_name': form.teacher_name or '',
            'course_title': form.course_title or '',
            'change_type': 'replace' if form.score_record else 'create',
            'before': before_snapshot,
            'after': after_snapshot
        })

    changed_forms.sort(key=lambda item: item['form_id'])
    unchanged_form_ids.sort()
    return {
        'summary': {
            'valid_row_count': parsed_result['valid_row_count'],
            'changed_form_count': len(changed_forms),
            'changed_row_count': changed_row_count,
            'unchanged_form_count': len(unchanged_form_ids),
            'skipped_count': parsed_result['skipped_count']
        },
        'changed_forms': changed_forms,
        'changed_grouped_items': changed_grouped_items,
        'unchanged_form_ids': unchanged_form_ids,
        'skipped_details': parsed_result['skipped_details']
    }

def _apply_manual_assessment_changes(changed_grouped_items, operator_id, import_time):
    imported_row_count = 0
    imported_form_count = 0
    for form_id, items in changed_grouped_items.items():
        old_record = ScoreRecord.query.filter_by(form_id=form_id).first()
        if old_record:
            db.session.delete(old_record)
            db.session.flush()

        score_snapshot = _build_score_snapshot(items)
        score_record = ScoreRecord(
            form_id=form_id,
            reviewer_id=operator_id,
            total_department_score=score_snapshot['total_department_score'],
            total_personal_score=score_snapshot['total_personal_score'],
            created_at=import_time,
            updated_at=import_time
        )
        db.session.add(score_record)
        db.session.flush()

        for item in score_snapshot['items']:
            db.session.add(ScoreItem(
                score_record_id=score_record.id,
                reason=item['reason'],
                department_score=item['department_score'],
                personal_score=item['personal_score'],
                is_auto_generated=False
            ))

        form = LectureForm.query.get(form_id)
        if form:
            form.reviewer_id = operator_id
            form.review_time = import_time
            db.session.add(form)

        imported_row_count += len(score_snapshot['items'])
        imported_form_count += 1

    return imported_row_count, imported_form_count

@admin_bp.route('/api/review/assessment/import/template', methods=['GET'])
@role_required('超级管理员')
def download_manual_assessment_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '手动导入考评模板'
    headers = ['表单ID', '考评项', '部门扣分', '个人扣分']
    ws.append(headers)
    ws.append([1001, '课堂秩序管理欠佳', 1, 2])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name='手动导入考评模板.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@admin_bp.route('/api/review/assessment/import/preview', methods=['POST'])
@role_required('超级管理员')
def preview_manual_assessment_import():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400

    df, read_error = _read_manual_assessment_file(request.files['file'])
    if read_error:
        return jsonify({'success': False, 'message': read_error}), 400

    parsed_result, parse_error = _parse_manual_assessment_rows(df)
    if parse_error:
        return jsonify({'success': False, 'message': parse_error}), 400

    preview_result = _build_manual_assessment_preview(parsed_result)
    return jsonify({
        'success': True,
        'summary': preview_result['summary'],
        'changed_forms': preview_result['changed_forms'],
        'unchanged_form_ids': preview_result['unchanged_form_ids'],
        'skipped_details': preview_result['skipped_details']
    })

@admin_bp.route('/api/review/assessment/import', methods=['POST'])
@role_required('超级管理员')
def import_manual_assessment():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400

    df, read_error = _read_manual_assessment_file(request.files['file'])
    if read_error:
        return jsonify({'success': False, 'message': read_error}), 400

    parsed_result, parse_error = _parse_manual_assessment_rows(df)
    if parse_error:
        return jsonify({'success': False, 'message': parse_error}), 400

    preview_result = _build_manual_assessment_preview(parsed_result)
    operator_id = session.get('user_id')
    import_time = datetime.now()
    imported_count, imported_form_count = _apply_manual_assessment_changes(
        preview_result['changed_grouped_items'],
        operator_id,
        import_time
    )

    db.session.commit()
    return jsonify({
        'success': True,
        'imported_count': imported_count,
        'imported_form_count': imported_form_count,
        'valid_row_count': parsed_result['valid_row_count'],
        'unchanged_form_count': len(preview_result['unchanged_form_ids']),
        'unchanged_form_ids': preview_result['unchanged_form_ids'],
        'skipped_count': parsed_result['skipped_count'],
        'skipped_details': parsed_result['skipped_details']
    })

@admin_bp.route('/api/statistics_data')
@role_required('管理员')
def get_statistics_data():
    """获取统计数据API"""
    user = User.query.get(session['user_id'])
    
    # 获取筛选参数
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    department = request.args.get('department')
    status = request.args.get('status')
    
    # 构建基础查询
    if user.role == '超级管理员':
        query = LectureForm.query.join(User, LectureForm.listener_number == User.number).filter(active_user_filter())
    else:
        # 管理员只能查看本部门数据
        query = LectureForm.query.join(User, LectureForm.listener_number == User.number)\
                                 .filter(User.department == user.department, active_user_filter())
    
    # 应用筛选条件（先不按状态过滤，后续按“表单组最新状态”过滤）
    if start_date:
        # 注意：lecture_date 是字符串字段，此处比较仅作为简单示例，可能需要根据实际存储格式调整
        query = query.filter(LectureForm.lecture_date >= start_date)
    if end_date:
        query = query.filter(LectureForm.lecture_date <= end_date)
    if department and user.role == '超级管理员':
        query = query.filter(User.department == department, active_user_filter())
    
    forms = query.order_by(LectureForm.updated_at.desc(), LectureForm.id.desc()).all()

    latest_forms_by_group = {}
    for form in forms:
        uid = form.unique_id or form.id
        if uid not in latest_forms_by_group:
            latest_forms_by_group[uid] = form

    latest_forms = list(latest_forms_by_group.values())
    
    review_permission = get_user_review_permission(user.id)
    if review_permission == '审表_中心':
        pending_statuses = {'待审核', '部门已审核'}
    else:
        pending_statuses = {'待审核'}

    if status:
        latest_forms = [f for f in latest_forms if f.status == status]

    pending_forms = len([f for f in latest_forms if f.status in pending_statuses])
    
    # 按部门统计
    dept_stats = {}
    for form in latest_forms:
        listener = _active_user_query().filter_by(number=form.listener_number).first()
        if listener:
            dept = listener.department
            if dept not in dept_stats:
                dept_stats[dept] = {'pending': 0}
            
            # 使用集合去重 unique_id
            if 'pending_ids' not in dept_stats[dept]:
                dept_stats[dept]['pending_ids'] = set()
                
            if form.status in pending_statuses:
                dept_stats[dept]['pending_ids'].add(form.unique_id or form.id)
                
    # 整理部门统计数据
    for dept in dept_stats:
        dept_stats[dept]['pending'] = len(dept_stats[dept].pop('pending_ids'))
    
    # 按月份统计
    month_stats = {}
    for form in latest_forms:
        if form.lecture_date:
            # 假设 lecture_date 格式为 YYYY-MM-DD
            try:
                date_str = str(form.lecture_date)
                # 提取年份和月份
                if '-' in date_str:
                    parts = date_str.split('-')
                    if len(parts) >= 2:
                        month_key = f"{parts[0]}-{parts[1]}"
                    else:
                        month_key = '未知'
                else:
                    month_key = '未知'
            except:
                month_key = '未知'

            if month_key not in month_stats:
                month_stats[month_key] = {'pending': 0, 'pending_ids': set()}
            
            if form.status in pending_statuses:
                month_stats[month_key]['pending_ids'].add(form.unique_id or form.id)
                
    # 整理月份统计数据
    for month in month_stats:
        month_stats[month]['pending'] = len(month_stats[month].pop('pending_ids'))
    
    # 按听课类型统计 (模型无此字段，暂留空)
    type_stats = {}
    
    # 计算平均评分 (模型无此字段，暂为0)
    avg_score = 0
    
    return jsonify({
        'overview': {
            'pending_forms': pending_forms,
            'avg_score': avg_score
        },
        'department_stats': dept_stats,
        'month_stats': month_stats,
        'type_stats': type_stats
    })

@admin_bp.route('/api/statistics_detail')
@role_required('管理员')
def get_statistics_detail():
    """Return detail rows from the same permission and filter scope as the overview."""
    current_user = User.query.get(session['user_id'])
    page = request.args.get('page', 1, type=int)
    per_page = 10
    dimension = _statistics_dimension({'dimension': request.args.get('detail_dimension')})
    query = _statistics_query(current_user, request.args)
    query = _apply_statistics_detail_filter(
        query,
        dimension,
        request.args.get('detail_value'),
    ).order_by(LectureForm.created_at.desc(), LectureForm.id.desc())
    pagination = query.paginate(page=max(page, 1), per_page=per_page, error_out=False)

    form_data = []
    for form, listener in pagination.items:
        form_data.append({
            'id': form.id,
            'listener_name': form.listener_name,
            'listener_department': listener.department or '未知',
            'teacher_name': form.teacher_name,
            'course_name': form.course_title,
            'listen_date': str(form.lecture_date or ''),
            'listen_type': '普通听课',
            'overall_rating': form.overall_effect,
            'status': form.status,
            'created_at': form.created_at.strftime('%Y-%m-%d %H:%M') if form.created_at else ''
        })
    
    return jsonify({
        'forms': form_data,
        'pagination': {
            'page': pagination.page,
            'pages': pagination.pages,
            'per_page': pagination.per_page,
            'total': pagination.total,
            'has_prev': pagination.has_prev,
            'has_next': pagination.has_next
        }
    })

@admin_bp.route('/api/review/auto_check', methods=['POST'])
@login_required
def auto_check_form():
    """对表单数据进行自动审核"""
    try:
        # 检查权限
        permission = get_user_review_permission(session['user_id'])
        if not permission:
             return jsonify({'success': False, 'message': '您不具有审表权限'}), 403

        form_data = request.get_json()
        
        # 构造类似 LectureForm 的对象供 auto_review 使用
        class FormLike:
            pass
        
        form = FormLike()
        # 映射字段
        setattr(form, 'id', form_data.get('id'))
        setattr(form, 'listener_name', form_data.get('listener_name'))
        setattr(form, 'listener_number', form_data.get('listener_number'))
        # reviewer_id 留空，避免与数据库ID混淆，auto_review现在优先使用listener_number
        # setattr(form, 'reviewer_id', form_data.get('listener_number')) 
        setattr(form, 'lecture_date', form_data.get('lecture_date'))
        setattr(form, 'class_period', form_data.get('class_period'))
        setattr(form, 'lecture_location', form_data.get('lecture_location'))
        setattr(form, 'teacher_name', form_data.get('teacher_name'))
        setattr(form, 'teacher_college', form_data.get('teacher_college'))
        setattr(form, 'course_title', form_data.get('course_title'))
        setattr(form, 'class_composition', form_data.get('student_grade_class')) # 映射到 student_grade_class
        setattr(form, 'teaching_method', form_data.get('teaching_method'))
        setattr(form, 'classroom_discipline', form_data.get('classroom_discipline'))
        setattr(form, 'classroom_atmosphere', form_data.get('classroom_atmosphere'))
        setattr(form, 'courseware_quality', form_data.get('courseware_quality'))
        setattr(form, 'overall_effect', form_data.get('overall_effect'))
        setattr(form, 'quality_case', form_data.get('quality_case'))
        setattr(form, 'course_feedback', form_data.get('course_feedback'))
        setattr(form, 'suggestions', form_data.get('suggestions'))
        setattr(form, 'phone', form_data.get('contact_phone1'))

        engine = AutoReviewEngine()
        result = engine.review_any(form)
        
        return jsonify({'success': True, 'result': result})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

# ==================== 导出功能实现 ====================

def _parse_export_status_list(status_list=None):
    if not status_list:
        return []
    if isinstance(status_list, str):
        status_list = status_list.split(',')
    return [str(s).strip() for s in status_list if str(s).strip()]


def _parse_export_form_ids(raw_form_ids=None):
    if not raw_form_ids:
        return []
    if isinstance(raw_form_ids, str):
        raw_items = raw_form_ids.split(',')
    else:
        raw_items = raw_form_ids
    form_ids = []
    for item in raw_items:
        try:
            form_id = int(item)
        except (TypeError, ValueError):
            continue
        if form_id > 0:
            form_ids.append(form_id)
    return list(dict.fromkeys(form_ids))


def get_export_query(scope, status_list=None, form_ids=None):
    """构建导出查询对象"""
    query = LectureForm.query
    scope = scope if scope in ['latest', 'selected', 'all'] else 'latest'

    # 1. 范围筛选
    if scope == 'latest':
        last_id = int(SystemSetting.get('last_exported_form_id', 0) or 0)
        query = query.filter(LectureForm.id > last_id)
    elif scope == 'selected':
        form_ids = _parse_export_form_ids(form_ids)
        query = query.filter(LectureForm.id.in_(form_ids))
    # scope == 'all' 不做额外ID/时间筛选

    # 2. 状态筛选
    status_list = _parse_export_status_list(status_list)
    if status_list:
        query = query.filter(LectureForm.status.in_(status_list))

    return query


def _format_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return value


def _sort_export_forms(forms, sort_by):
    if sort_by == 'course':
        forms.sort(key=lambda x: (x.teacher_college or '', x.teacher_name or '', x.course_title or ''))
    else:
        forms.sort(key=lambda x: x.lecture_date or '', reverse=True)


def _build_standard_forms_workbook(forms):
    # 生成特定规范 Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "听课反馈表"

    # 创建第二个空Sheet
    wb.create_sheet("其他信息")

    # === 样式定义 ===
    font_title = Font(name='宋体', size=20, bold=True)
    font_header = Font(name='宋体', size=12, bold=True)
    font_data = Font(name='宋体', size=12)

    thin_border = Side(style='thin')
    thick_border = Side(style='medium')
    border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    ws.merge_cells('A1:L1')
    ws['A1'] = "西南大学学生教学质量监控信息（意见）处理笺"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_center
    ws.row_dimensions[1].height = 48

    headers = [
        "序号", "听课人", "听课时间", "第几节", "听课地点",
        "授课教师", "教师所属学院", "课程", "专业年级",
        "异常情况反映", "课程反馈", "不足及建议"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_all

    ws.row_dimensions[2].height = 25

    max_id = 0
    for row_idx, form in enumerate(forms, 3):
        if form.id > max_id:
            max_id = form.id

        row_data = [
            row_idx - 2,
            form.listener_number,
            form.lecture_date,
            form.class_period,
            form.lecture_location,
            form.teacher_name,
            form.teacher_college,
            form.course_title,
            form.student_grade_class,
            form.abnormal_situation,
            form.course_feedback,
            form.suggestions
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else '')
            cell.font = font_data
            cell.alignment = align_left
            cell.border = border_all

        ws.row_dimensions[row_idx].height = 15

    for col_idx in range(1, len(headers) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 15

    max_row = len(forms) + 2

    for col in range(1, 13):
        cell = ws.cell(row=1, column=col)
        current_border = cell.border
        cell.border = Border(
            top=thick_border,
            bottom=current_border.bottom,
            left=current_border.left,
            right=current_border.right
        )

    for col in range(1, 13):
        cell = ws.cell(row=max_row, column=col)
        current_border = cell.border
        cell.border = Border(
            top=current_border.top,
            bottom=thick_border,
            left=current_border.left,
            right=current_border.right
        )

    for row in range(1, max_row + 1):
        cell = ws.cell(row=row, column=1)
        current_border = cell.border
        cell.border = Border(
            top=current_border.top,
            bottom=current_border.bottom,
            left=thick_border,
            right=current_border.right
        )

    for row in range(1, max_row + 1):
        cell = ws.cell(row=row, column=12)
        current_border = cell.border
        cell.border = Border(
            top=current_border.top,
            bottom=current_border.bottom,
            left=current_border.left,
            right=thick_border
        )

    ws['A1'].border = Border(top=thick_border, left=thick_border, right=thick_border, bottom=thin_border)
    ws.cell(row=1, column=12).border = Border(top=thick_border, right=thick_border, bottom=thin_border)

    return wb, max_id


def _build_all_data_forms_workbook(forms):
    # 导出 lecture_forms 表中所有字段
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表单数据库数据"

    columns = [
        ('id', 'ID'),
        ('listener_name', '听课人姓名+学院'),
        ('listener_number', '听课人编号'),
        ('course_changes', '课程信息变化'),
        ('lecture_date', '听课时间'),
        ('class_period', '第几节'),
        ('lecture_location', '听课地点'),
        ('teacher_name', '授课教师'),
        ('teacher_college', '教师所属学院'),
        ('course_title', '课程总标题'),
        ('student_grade_class', '专业年级'),
        ('abnormal_situation', '异常情况反映'),
        ('teaching_method', '主要教学方法'),
        ('classroom_discipline', '管理课堂纪律'),
        ('classroom_atmosphere', '调动课堂气氛'),
        ('courseware_quality', '课件制作质量'),
        ('overall_effect', '整体教学效果'),
        ('quality_case', '优质案例推荐'),
        ('course_feedback', '课程反馈'),
        ('suggestions', '不足及建议'),
        ('student_signature1', '听课班级同学签名1'),
        ('contact_phone1', '联系电话1'),
        ('student_signature2', '听课班级同学签名2'),
        ('contact_phone2', '联系电话2'),
        ('status', '审核状态'),
        ('reviewer_id', '审核人ID'),
        ('review_time', '审核时间'),
        ('review_comment', '审核意见'),
        ('registration_id', '关联登记记录ID'),
        ('audit_tag', '审核标签'),
        ('unique_id', '唯一标志ID'),
        ('created_at', '创建时间'),
        ('updated_at', '更新时间'),
    ]

    header_font = Font(name='微软雅黑', size=11, bold=True)
    data_font = Font(name='微软雅黑', size=11)
    header_fill = PatternFill(fill_type='solid', fgColor='D8D8D8')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, (_, header) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    max_id = 0
    for row_idx, form in enumerate(forms, 2):
        if form.id > max_id:
            max_id = form.id
        for col_idx, (attr, _) in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_format_excel_value(getattr(form, attr, None)))
            cell.font = data_font
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.border = border

    return wb, max_id

@admin_bp.route('/api/forms/export/check', methods=['POST'])
@role_required('超级管理员')
def check_export_forms():
    """检查符合导出条件的表单数量"""
    try:
        data = request.get_json() or {}
        scope = data.get('scope', 'latest')
        status_list = data.get('status_list', []) # 默认为空列表，或者前端应该传所有选中的状态
        form_ids = data.get('form_ids', [])
        
        # 兼容旧代码：如果传了 status_approved
        if 'status_approved' in data and not status_list:
            if data['status_approved']:
                status_list = ['已审核', '部门已审核', '中心已审核']
            else:
                status_list = [] # 不筛选，即所有状态
        
        # 获取符合条件的查询
        query = get_export_query(scope, status_list, form_ids)
        total_count = query.count()
        
        # 计算被排除的数量
        excluded_count = 0
        if status_list:
            # 构建一个不带状态筛选的查询
            all_status_query = get_export_query(scope, None, form_ids)
            full_count = all_status_query.count()
            excluded_count = full_count - total_count
            
        return jsonify({
            'success': True,
            'total_count': total_count,
            'excluded_count': excluded_count
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@admin_bp.route('/api/forms/export/download', methods=['GET'])
@role_required('超级管理员')
def download_export_forms():
    """执行导出并下载Excel"""
    try:
        # 获取参数
        scope = request.args.get('scope', 'latest')
        status_list = request.args.get('status_list') # 逗号分隔的字符串
        form_ids = request.args.get('form_ids')
        export_mode = request.args.get('export_mode', 'standard')
        
        # 兼容旧代码
        status_approved = request.args.get('status_approved')
        if status_approved and not status_list:
             if str(status_approved).lower() == 'true':
                 status_list = '已审核,部门已审核,中心已审核'
        
        sort_by = request.args.get('sort_by', 'time')
        filename_title = (request.args.get('filename_title') or '').strip()
            
        # 1. 获取数据
        query = get_export_query(scope, status_list, form_ids)
        forms = query.all()
        
        if not forms:
            flash('没有符合条件的记录可导出', 'warning')
            return redirect(url_for('admin.view_forms'))
            
        # 2. 排序
        _sort_export_forms(forms, sort_by)
            
        # 3. 生成Excel
        if export_mode == 'all_data':
            wb, max_id = _build_all_data_forms_workbook(forms)
        else:
            export_mode = 'standard'
            wb, max_id = _build_standard_forms_workbook(forms)
        
        # 4. 更新SystemSetting (如果是latest模式)
        if scope == 'latest' and max_id > 0:
            SystemSetting.set('last_exported_form_id', str(max_id))
            
        # 5. 返回文件
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        date_str = datetime.now().strftime('%Y%m%d')
        if export_mode == 'all_data':
            filename = f"听课反馈表-{date_str}.xlsx"
        elif filename_title:
            filename = f"学生教学信息中心听课反馈表汇总-{filename_title} {date_str}.xlsx"
        else:
            filename = f"学生教学信息中心听课反馈表汇总-未命名 {date_str}.xlsx"
        
        # 处理中文文件名
        try:
            filename_encoded = filename.encode('latin-1').decode('latin-1')
        except UnicodeEncodeError:
            # 如果包含非latin-1字符（如中文），quote它
            from urllib.parse import quote
            filename_encoded = quote(filename)
            
        rv = send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        
        # 设置Content-Disposition header以支持中文文件名
        rv.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{filename_encoded}"
        
        return rv
        
    except Exception as e:
        flash(f'导出失败: {str(e)}', 'error')
        return redirect(url_for('admin.view_forms'))

@admin_bp.route('/api/review/reference_data', methods=['POST'])
@login_required
def get_reference_data():
    """获取参考数据（课表和通讯录）"""
    try:
        # 检查权限
        permission = get_user_review_permission(session['user_id'])
        if not permission:
             return jsonify({'success': False, 'message': '您不具有审表权限'}), 403

        form_data = request.get_json()
        
        engine = AutoReviewEngine()
        result = engine.search_reference_data(form_data)
        
        return jsonify({'success': True, 'result': result})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@admin_bp.route('/assessment-exemption-settings')
@role_required('管理员')
def assessment_exemption_settings():
    user = User.query.get(session['user_id'])
    if user.role == '超级管理员':
        return redirect(url_for('admin.system_management', tab='assessment'))
    if not _has_assessment_stats_access(session['user_id']):
        flash_forbidden('考评规则设置')
        return redirect(url_for('main.index'))
    available_departments = list(_get_accessible_department_users(session['user_id']).keys())
    is_super_admin = get_user_manage_permission(session['user_id']) == '超级管理员'
    return render_template(
        'admin/assessment_exemption_settings.html',
        available_departments=available_departments,
        is_super_admin=is_super_admin,
    )


@admin_bp.route('/api/assessment-overrides/list', methods=['GET'])
@login_required
def list_assessment_overrides():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('考评规则设置')
    department_names = request.args.getlist('departments')
    selected_departments, department_user_map = _resolve_selected_departments(
        session['user_id'],
        department_names,
    )
    if not selected_departments:
        return jsonify({'success': True, 'departments': []})

    all_user_ids = []
    for users in department_user_map.values():
        all_user_ids.extend([u.id for u in users])

    overrides_by_user = {}
    if all_user_ids:
        overrides = AssessmentOverride.query.filter(
            AssessmentOverride.user_id.in_(all_user_ids)
        ).order_by(AssessmentOverride.start_week.asc(), AssessmentOverride.end_week.asc()).all()
        for ov in overrides:
            overrides_by_user.setdefault(ov.user_id, []).append({
                'id': ov.id,
                'override_type': ov.override_type,
                'override_value': ov.override_value,
                'start_week': ov.start_week,
                'end_week': ov.end_week,
                'reason': ov.reason,
                'created_at': ov.created_at.strftime('%Y-%m-%d %H:%M:%S') if ov.created_at else '',
            })

    result_departments = []
    for dept_name in selected_departments:
        users = department_user_map.get(dept_name, [])
        group_map = {}
        for user in users:
            group_name = user.group or UNASSIGNED_GROUP_NAME
            group_map.setdefault(group_name, []).append(user)
        groups = []
        for group_name in sorted(group_map.keys()):
            members = []
            for user in sorted(group_map[group_name], key=lambda u: (u.number or '', u.name or '')):
                members.append({
                    'user_id': user.id,
                    'name': user.name,
                    'number': user.number,
                    'overrides': overrides_by_user.get(user.id, []),
                })
            groups.append({
                'group_name': group_name,
                'members': members,
            })
        result_departments.append({
            'department': dept_name,
            'groups': groups,
        })

    return jsonify({'success': True, 'departments': result_departments})


@admin_bp.route('/api/assessment-overrides', methods=['POST'])
@login_required
def create_assessment_override():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('考评规则设置')
    data = request.get_json() or {}
    user_ids = data.get('user_ids') or []
    if not user_ids:
        return jsonify({'success': False, 'message': '请选择至少一个成员'}), 400
    try:
        start_week = int(data.get('start_week'))
        end_week = int(data.get('end_week'))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'message': '起止周次必须为整数'}), 400
    if start_week < 1 or end_week < 1:
        return jsonify({'success': False, 'message': '周次必须大于0'}), 400
    if start_week > end_week:
        return jsonify({'success': False, 'message': '起始周不能大于结束周'}), 400
    override_type = (data.get('override_type') or 'exempt').strip()
    if override_type not in ('exempt', 'custom_requirement', 'partial_exempt', LEAVE_OVERRIDE_TYPE):
        return jsonify({'success': False, 'message': '不支持的规则类型'}), 400
    if override_type == LEAVE_OVERRIDE_TYPE and get_user_manage_permission(session['user_id']) != '超级管理员':
        return jsonify({'success': False, 'message': '请假规则请在统计分析页为当前教学周发起'}), 403
    reason = (data.get('reason') or '').strip()
    if not reason:
        return jsonify({'success': False, 'message': '请填写理由'}), 400
    override_value = data.get('override_value')

    accessible_users = _resolve_assessment_users(session['user_id'], [str(uid) for uid in user_ids])
    accessible_ids = {u.id for u in accessible_users}
    if not accessible_ids:
        return jsonify({'success': False, 'message': '无权限操作所选成员'}), 403

    created_count = 0
    skipped_count = 0
    for uid in user_ids:
        try:
            uid_int = int(uid)
        except (TypeError, ValueError):
            continue
        if uid_int not in accessible_ids:
            continue
        existing = AssessmentOverride.query.filter_by(
            user_id=uid_int,
            start_week=start_week,
            end_week=end_week,
            override_type=override_type,
        ).first()
        if existing:
            skipped_count += 1
            continue
        record = AssessmentOverride(
            user_id=uid_int,
            start_week=start_week,
            end_week=end_week,
            override_type=override_type,
            override_value=override_value,
            reason=reason,
            created_by=session['user_id'],
        )
        db.session.add(record)
        created_count += 1

    db.session.commit()
    message = f'成功添加 {created_count} 条规则'
    if skipped_count > 0:
        message += f'，跳过 {skipped_count} 条已存在的规则'
    return jsonify({
        'success': True,
        'message': message,
        'created_count': created_count,
        'skipped_count': skipped_count,
    })


@admin_bp.route('/api/assessment-overrides/<int:override_id>', methods=['DELETE'])
@login_required
def delete_assessment_override(override_id):
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('考评规则设置')
    record = db.session.get(AssessmentOverride, override_id)
    if not record:
        return jsonify({'success': False, 'message': '规则不存在'}), 404
    manage_permission = get_user_manage_permission(session['user_id'])
    if manage_permission != '超级管理员':
        accessible_users = _resolve_assessment_users(session['user_id'], [str(record.user_id)])
        if not accessible_users:
            return jsonify({'success': False, 'message': '无权限删除该规则'}), 403
    db.session.delete(record)
    db.session.commit()
    return jsonify({'success': True, 'message': '规则已删除'})


@admin_bp.route('/api/teaching-month-definitions', methods=['GET'])
@login_required
def get_teaching_month_definitions():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('教学月设置')
    custom = _load_custom_month_definitions()
    if custom:
        return jsonify({'success': True, 'is_custom': True, 'months': custom})
    defaults = []
    for i in range(5):
        start_week = i * 4 + 1
        end_week = start_week + 3
        defaults.append({
            'start_week': start_week,
            'end_week': end_week,
            'label': f'第{i + 1}教学月',
        })
    return jsonify({'success': True, 'is_custom': False, 'months': defaults})


@admin_bp.route('/api/teaching-month-definitions', methods=['POST'])
@login_required
def save_teaching_month_definitions():
    if not _has_assessment_stats_access(session['user_id']):
        return forbidden_json('教学月设置')
    data = request.get_json() or {}
    months = data.get('months')
    if not months:
        SystemSetting.set('teaching_month_definitions', '')
        return jsonify({'success': True, 'message': '已恢复为默认教学月设置（每4周一个教学月）'})
    if not isinstance(months, list):
        return jsonify({'success': False, 'message': 'months 必须是数组'}), 400

    validated = []
    for idx, item in enumerate(months):
        try:
            start_week = int(item.get('start_week'))
            end_week = int(item.get('end_week'))
        except (TypeError, ValueError):
            return jsonify({'success': False, 'message': f'第{idx + 1}项的周次必须为整数'}), 400
        if start_week < 1 or end_week < 1:
            return jsonify({'success': False, 'message': f'第{idx + 1}项的周次必须大于0'}), 400
        if start_week > end_week:
            return jsonify({
                'success': False,
                'message': f'第{idx + 1}项的起始周({start_week})不能大于结束周({end_week})',
            }), 400
        label = str(item.get('label') or '').strip()
        validated.append({'start_week': start_week, 'end_week': end_week, 'label': label})

    SystemSetting.set('teaching_month_definitions', json.dumps(validated, ensure_ascii=False))
    return jsonify({'success': True, 'message': f'已保存自定义教学月设置（{len(validated)}个教学月）'})
