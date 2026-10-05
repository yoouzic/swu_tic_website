# -*- coding: utf-8 -*-
# Phase 1 mechanical split from app/blueprints/admin.py
# Module: contacts

from flask import request, redirect, url_for, flash, session, jsonify, send_file, current_app
from app.models import User, Department, Group, db
from datetime import datetime
from app.security import role_required
from app.utils.password_audit import record_password_audit
from app.utils.user_status import UNASSIGNED_DEPARTMENT_NAME, UNASSIGNED_GROUP_NAME
from app.utils.env_config import env_path
from app.services.import_state import (
    get_committed_import_result,
    import_preview_exists,
    is_password_artifact_published,
)
import pandas as pd
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import os
import re
from werkzeug.security import generate_password_hash
from . import admin_bp
from .shared import _active_user_query, allowed_file, generate_random_password


DEFAULT_CONTACT_TEMPLATE_PATH = os.path.join('data', 'storage', 'templates', 'contacts', '通讯录.xlsx')


DEFAULT_EXPORT_DIR = os.path.join('data', 'storage', 'exports', 'contacts')


@admin_bp.route('/download_template')
@role_required('超级管理员')
def download_template():
    """下载通讯录模板文件"""
    template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
    return send_file(template_path, as_attachment=True)


@admin_bp.route('/view_sample')
@role_required('超级管理员')
def view_sample():
    """在线查看模板示例"""
    template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
    return send_file(template_path, as_attachment=False)


def _safe_export_filename(filename):
    """Validate a generated export filename without breaking Chinese contact names."""
    if not isinstance(filename, str) or not filename:
        return None
    # 拒绝任何路径分隔符、父目录、绝对/驱动器路径和 Windows ADS 冒号。
    if '/' in filename or '\\' in filename:
        return None
    if filename != os.path.basename(filename):
        return None
    if '..' in filename or ':' in filename:
        return None
    if any(ord(ch) < 32 for ch in filename):
        return None
    if not filename.lower().endswith('.xlsx'):
        return None
    return filename


def _export_file_within_root(export_dir, filename):
    """Return the realpath if candidate is safely inside export_dir, else None."""
    export_real = os.path.realpath(export_dir)
    candidate_real = os.path.realpath(os.path.join(export_real, filename))
    try:
        inside = os.path.commonpath([export_real, candidate_real]) == export_real
    except ValueError:
        return None
    return candidate_real if inside else None


# 新式 R2 密码工件：文件名内嵌 16-hex import_id + 32-hex claim_token，发布以
# durable COMMITTED marker 为界（fail closed：row 缺失 / AVAILABLE / CLAIMED /
# token 不匹配一律拒绝，row 缺失不再视为已发布）。
_TOKENIZED_PASSWORD_ARTIFACT_RE = re.compile(
    r'^passwords_([0-9a-f]{16})_([0-9a-f]{32})\.xlsx$'
)
# P1/P1-R1 legacy 密码工件：仅内嵌 import_id（无 claim token）。历史文件没有
# commit marker，无法可靠区分成功工件与孤儿工件，保留 P1-R1 兼容行为
# （LEGACY_PASSWORD_ARTIFACT_PROVENANCE_AMBIGUITY）。
_PASSWORD_ARTIFACT_RE = re.compile(r'^passwords_([0-9a-f]{16})\.xlsx$')


def _password_artifact_filename(claim):
    """Single filename builder for confirm-import password artifacts.

    Claim-token-scoped physical identity: claim A → file A, claim B → file B,
    so a stale worker's artifact path can never collide with the winner's.
    """
    return f'passwords_{claim.import_id}_{claim.claim_token}.xlsx'


def _build_import_success_response(imported_count, updated_count, skipped_count,
                                   total_users, password_artifact):
    """Single success-JSON builder for first success and COMMITTED replay (R3).

    Both paths must render through this helper so the first response and the
    idempotent replay cannot drift apart.  ``message`` is rebuilt from counts
    (never persisted), and the URL is regenerated from the artifact filename.
    """
    return {
        'success': True,
        'message': f'导入完成：新增 {imported_count}，更新 {updated_count}，跳过 {skipped_count}',
        'imported_count': imported_count,
        'updated_count': updated_count,
        'skipped_count': skipped_count,
        'total_users': total_users,
        'password_file_url': (url_for('admin.download_passwords', filename=password_artifact) if password_artifact else None),
    }


@admin_bp.route('/download_passwords/<filename>')
@role_required('超级管理员')
def download_passwords(filename):
    """下载密码文件"""
    safe_filename = _safe_export_filename(filename)
    if safe_filename is None:
        flash('文件不存在', 'error')
        return redirect(url_for('admin.super_admin_dashboard'))
    # Publish gate（新式 token 工件）：只有 durable COMMITTED marker 精确匹配
    # (import_id, claim_token, filename) 才允许下载；其余一律拒绝。
    token_match = _TOKENIZED_PASSWORD_ARTIFACT_RE.match(safe_filename)
    if token_match:
        if not is_password_artifact_published(
            token_match.group(1), token_match.group(2), safe_filename,
        ):
            flash('文件不存在', 'error')
            return redirect(url_for('admin.super_admin_dashboard'))
    else:
        # Legacy P1/P1-R1 工件（仅 import_id）：preview row 仍存在（AVAILABLE
        # 或 CLAIMED）说明业务 commit 尚未发生，拒绝下载；row 缺失视为已发布
        # （历史兼容行为，provenance 歧义已知）。contacts 导出与 legacy 时间戳
        # 文件不受影响。
        artifact_match = _PASSWORD_ARTIFACT_RE.match(safe_filename)
        if artifact_match and import_preview_exists(artifact_match.group(1)):
            flash('文件不存在', 'error')
            return redirect(url_for('admin.super_admin_dashboard'))
    export_dir = env_path('EXPORT_DIR', DEFAULT_EXPORT_DIR)
    file_path = _export_file_within_root(export_dir, safe_filename)
    if file_path and os.path.isfile(file_path):
        return send_file(file_path, as_attachment=True)
    flash('文件不存在', 'error')
    return redirect(url_for('admin.super_admin_dashboard'))


class ContactIdentityConflict(ValueError):
    """A personnel number belongs to a different student identity."""


def _contact_identity_conflict(number, student_id):
    """Use the same identity rule for preview and the execution preflight."""
    by_student = User.query.filter_by(student_id=student_id).first()
    by_number = User.query.filter_by(number=number).first()
    if by_number is not None and (by_student is None or by_number.id != by_student.id):
        return f'编号{number}已属于其他学号，不能覆盖，请核对编号和学号'
    return None


@admin_bp.route('/preview_import', methods=['POST'])
@role_required('超级管理员')
def preview_import():
    """预览通讯录并进行格式校验（支持 xls/xlsx，严格对标模板列）"""
    try:
        from flask import current_app
        current_app.logger.info('通讯录预览开始')
        if 'file' not in request.files:
            return jsonify({'success': False, 'message': '未收到文件'}), 400
        file = request.files['file']
        if file.filename == '':
            return jsonify({'success': False, 'message': '请选择文件'}), 400
        if not allowed_file(file.filename):
            return jsonify({'success': False, 'message': '文件格式不支持，仅支持xls/xlsx'}), 400

        ext = file.filename.rsplit('.', 1)[1].lower()
        engine = 'openpyxl' if ext == 'xlsx' else 'xlrd'
        # 以字符串读入，保留学号/手机号等前导零
        df = pd.read_excel(file, dtype=str, engine=engine, keep_default_na=False)

        # 加载模板列
        template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
        template_df = pd.read_excel(template_path, dtype=str, engine='openpyxl', keep_default_na=False)
        expected_cols = [c.strip() for c in template_df.columns.tolist()]
        incoming_cols = [c.strip() for c in df.columns.tolist()]
        if incoming_cols != expected_cols:
            return jsonify({'success': False, 'message': 'Excel列不匹配，请使用模板列名与顺序'}), 400

        errors = []
        preview_rows = []
        valid_rows = 0
        seen_numbers = set()
        seen_student_ids = set()

        for i, row in df.iterrows():
            rec = {col: str(row[col]).strip() for col in expected_cols}
            dep_group = rec['部门/组别']
            if '/' in dep_group:
                department = dep_group.split('/', 1)[0]
                group_name = dep_group.split('/', 1)[1]
            else:
                department = dep_group
                group_name = UNASSIGNED_GROUP_NAME

            row_errors = []
            # 基本校验
            if not rec['编号'] or not rec['编号'].isdigit():
                row_errors.append('编号必须为纯数字')
            if not rec['姓名']:
                row_errors.append('姓名不能为空')
            if rec['性别'] not in ('男', '女'):
                row_errors.append('性别必须为男/女')
            grade = rec['年级']
            if not grade or not (grade.isdigit() or (grade.endswith('级') and grade[:-1].isdigit())):
                row_errors.append('年级格式不正确')
            phone = rec['手机号码']
            if phone and not (phone.isdigit() and len(phone) >= 6):
                row_errors.append('手机号码必须为数字')
            qq = rec['QQ号码']
            if qq and not (qq.isdigit() and 5 <= len(qq) <= 12):
                row_errors.append('QQ号码格式不正确')
            student_id = rec['学号']
            if not student_id or not student_id.isdigit():
                row_errors.append('学号必须为纯数字')

            # 重复检查（以学号判定）
            existing_user_by_sid = User.query.filter_by(student_id=student_id).first()
            is_duplicate = existing_user_by_sid is not None
            conflict = _contact_identity_conflict(rec['编号'], student_id)
            if conflict:
                row_errors.append(conflict)
            if rec['编号'] in seen_numbers:
                row_errors.append('文件内编号重复')
            if student_id in seen_student_ids:
                row_errors.append('文件内学号重复')
            seen_numbers.add(rec['编号'])
            seen_student_ids.add(student_id)

            preview_rows.append({
                'row_number': i + 1,
                'number': rec['编号'],
                'department': department,
                'name': rec['姓名'],
                'gender': rec['性别'],
                'grade': rec['年级'],
                'college': rec['学院'],
                'major': rec['专业'],
                'dormitory': rec['宿舍'],
                'phone': phone,
                'qq': qq,
                'student_id': student_id,
                'role': '信息员',
                'group': group_name,
                'has_error': bool(row_errors),
                'errors': row_errors,
                'is_duplicate': is_duplicate
            })
            if not row_errors:
                valid_rows += 1
            errors.extend([f"第{i + 1}行: {e}" for e in row_errors])

        from app.services.import_state import save_import_preview
        import_id = save_import_preview(preview_rows, expected_cols)

        return jsonify({
            'success': True,
            'import_id': import_id,
            'columns': expected_cols,
            'total_rows': len(preview_rows),
            'valid_rows': valid_rows,
            'error_rows': len(preview_rows) - valid_rows,
            'errors': errors,
            'preview': preview_rows[:50]
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'预览失败：{str(e)}'}), 500


def _resolve_group_id_for_import(user):
    """Resolve/create the department scoped Group and fill user.group_id."""
    if not user.department or not user.group or user.group in ('待分配', UNASSIGNED_GROUP_NAME):
        return None
    group = Group.query.filter_by(name=user.group, department=user.department).first()
    if group is None:
        group = Group(name=user.group, department=user.department)
        db.session.add(group)
        db.session.flush()
    return group.id


def _cleanup_generated_password_file(path):
    """Best-effort removal of a generated password workbook; never masks the original error."""
    if not path:
        return
    try:
        os.remove(path)
    except FileNotFoundError:
        return
    except OSError as exc:
        current_app.logger.warning('密码文件清理失败 %s: %s', path, exc)


@admin_bp.route('/confirm_import', methods=['POST'])
@role_required('超级管理员')
def confirm_import():
    """确认导入到数据库，生成随机密码并保存哈希，返回总用户数"""
    current_app.logger.info('通讯录导入开始')
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({'success': False, 'message': '请求数据必须为JSON对象'}), 400
    import_id = data.get('import_id')
    overwrite = data.get('overwrite', False)
    if not isinstance(overwrite, bool):
        return jsonify({'success': False, 'message': 'overwrite必须为布尔值'}), 400
    if not isinstance(import_id, str) or not import_id:
        return jsonify({'success': False, 'message': '导入会话已失效，请重新预览'}), 400

    # 所有权先行：同一 import_id 只有一个 confirm winner（durable claim，
    # CAS UPDATE 所有权，崩溃后 lease 到期可恢复）。
    from app.services.import_state import (
        claim_import_preview,
        finalize_import_preview_claim,
        release_import_preview_claim,
    )
    claim = claim_import_preview(import_id)
    if claim is None:
        # R3 committed-aware 分支：claim 失败有两种原因——
        # 1) 该 import 已 COMMITTED（v3 含 replay metadata）→ 幂等重放原成功
        #    结果（RESULT REPLAY，绝不 RETRY EXECUTION：不重新生成密码/不重建
        #    用户/不写审计/不重写工件，COMMITTED 仍是 terminal）；
        # 2) 其余（missing/过期/active CLAIMED/COMMITTED v2/malformed）→
        #    维持既有 400。
        committed = get_committed_import_result(import_id)
        if committed is not None:
            response = _build_import_success_response(
                committed.imported_count,
                committed.updated_count,
                committed.skipped_count,
                committed.total_users,
                committed.password_artifact,
            )
            current_app.logger.info(f'通讯录导入结果重放 import_id={import_id}')
            return jsonify(response)
        return jsonify({'success': False, 'message': '导入会话已失效，请重新预览'}), 400

    export_path = None
    try:
        rows = claim.payload['rows']
        actor_user_id = session.get('user_id')
        imported_count = 0
        updated_count = 0
        skipped_count = 0
        password_list = []

        # Preview is advisory: recheck the complete valid batch before any writes.
        seen_numbers = set()
        seen_student_ids = set()
        for r in rows:
            if r['has_error']:
                continue
            conflict = _contact_identity_conflict(r['number'], r['student_id'])
            if conflict:
                raise ContactIdentityConflict(conflict)
            if r['number'] in seen_numbers or r['student_id'] in seen_student_ids:
                raise ContactIdentityConflict('文件内编号或学号重复，请重新预览')
            seen_numbers.add(r['number'])
            seen_student_ids.add(r['student_id'])

        for r in rows:
            if r['has_error']:
                skipped_count += 1
                continue

            # 以学号判定唯一
            user = User.query.filter_by(student_id=r['student_id']).first()
            if user:
                if not overwrite:
                    skipped_count += 1
                    continue
                # 通讯录不能降级管理员或重置管理员凭据。
                preserve_admin_credentials = user.role in ('管理员', '超级管理员')
                user.number = r['number']
                user.department = r['department']
                user.name = r['name']
                user.gender = r['gender']
                user.grade = r['grade']
                user.college = r['college']
                user.major = r['major']
                user.dormitory = r['dormitory']
                user.phone = r['phone']
                user.qq = r['qq']
                if not preserve_admin_credentials:
                    user.role = r['role']
                user.group = r['group']
                # 保障部门/小组存在，并同步真实 group_id
                if user.department and not Department.query.filter_by(name=user.department).first():
                    db.session.add(Department(name=user.department))
                if user.group and user.department:
                    if not Group.query.filter_by(name=user.group, department=user.department).first() and user.group not in ['待分配', UNASSIGNED_GROUP_NAME]:
                        db.session.add(Group(name=user.group, department=user.department))
                user.group_id = _resolve_group_id_for_import(user)
                updated_count += 1
                if not preserve_admin_credentials:
                    password = generate_random_password()
                    user.password_hash = generate_password_hash(password)
                    record_password_audit(
                        actor_user_id=actor_user_id,
                        target_user_id=user.id,
                        action='import_overwrite_reset_password',
                        details={'match_type': 'student_id', 'overwrite': True}
                    )
                    password_list.append({
                        'number': user.number,
                        'name': user.name,
                        'student_id': user.student_id,
                        'password': password
                    })
            else:
                # 身份仅按学号匹配，编号不能覆盖另一人。
                existed_by_number = User.query.filter_by(number=r['number']).first()
                if existed_by_number:
                    raise ContactIdentityConflict(f"编号{r['number']}已属于其他学号，不能覆盖")
                else:
                    # 新增用户
                    password = generate_random_password()
                    user = User(
                        number=r['number'],
                        department=r['department'],
                        name=r['name'],
                        gender=r['gender'],
                        grade=r['grade'],
                        college=r['college'],
                        major=r['major'],
                        dormitory=r['dormitory'],
                        phone=r['phone'],
                        qq=r['qq'],
                        student_id=r['student_id'],
                        password_hash=generate_password_hash(password),
                        role=r['role'],
                        group=r['group']
                    )
                    # 保障部门/小组存在
                    if user.department and not Department.query.filter_by(name=user.department).first():
                        db.session.add(Department(name=user.department))
                    if user.group and user.department:
                        if not Group.query.filter_by(name=user.group, department=user.department).first() and user.group not in ['待分配', UNASSIGNED_GROUP_NAME]:
                            db.session.add(Group(name=user.group, department=user.department))
                    db.session.add(user)
                    db.session.flush()
                    user.group_id = _resolve_group_id_for_import(user)
                    record_password_audit(
                        actor_user_id=actor_user_id,
                        target_user_id=user.id,
                        action='import_create_user_password_init',
                        details={'overwrite': overwrite}
                    )
                    imported_count += 1
                    password_list.append({
                        'number': user.number,
                        'name': user.name,
                        'student_id': user.student_id,
                        'password': password
                    })

        # 密码工件必须在业务 commit 前生成完成：文件名内嵌 claim token，物理上
        # 隔离不同 claim 的工件（stale worker 的路径绝不与 winner 冲突）；
        # 工件失败则不进入 commit（§14/§15）。
        export_filename = None
        if password_list:
            export_filename = _password_artifact_filename(claim)
            password_df = pd.DataFrame(password_list)
            export_dir = env_path('EXPORT_DIR', DEFAULT_EXPORT_DIR)
            os.makedirs(export_dir, exist_ok=True)
            export_path = os.path.join(export_dir, export_filename)
            password_df.to_excel(export_path, index=False)

        total_users = _active_user_query().count()
        # 响应在 commit 前构造完成：commit 之后只剩低风险的日志与返回（§21）。
        # 首次成功与 replay 共用同一 builder，避免字段漂移（R3 §12）。
        response = _build_import_success_response(
            imported_count, updated_count, skipped_count, total_users, export_filename,
        )
        # COMMITTED 发布证明 + replay result 与业务变更同一个数据库事务：
        # CAS CLAIMED → COMMITTED（v3 记录 artifact filename 与结果计数快照），
        # CAS 失败（ownership 丢失）即抛出并整体回滚，绝不提交用户修改（§10）。
        finalize_import_preview_claim(
            claim, db.session,
            password_artifact=export_filename,
            result={
                'imported_count': imported_count,
                'updated_count': updated_count,
                'skipped_count': skipped_count,
                'total_users': total_users,
            },
        )
        db.session.commit()
    except ContactIdentityConflict as e:
        db.session.rollback()
        _cleanup_generated_password_file(export_path)
        release_import_preview_claim(claim)
        return jsonify({'success': False, 'message': f'导入冲突：{str(e)}'}), 400
    except Exception as e:
        db.session.rollback()
        _cleanup_generated_password_file(export_path)
        # 业务未 commit：CAS 释放 claim 回 AVAILABLE，管理员可用同一 import_id
        # 重试；cleanup 只作用于当前 claim token 的工件路径，绝不影响其他
        # owner 的文件（§14）。ownership 已被 lease 恢复取走时 release 返回
        # False，不覆盖新 owner（§13）。
        release_import_preview_claim(claim)
        return jsonify({'success': False, 'message': f'导入失败：{str(e)}'}), 500

    current_app.logger.info(f'通讯录导入完成 新增{imported_count} 更新{updated_count} 跳过{skipped_count}')
    return jsonify(response)


@admin_bp.route('/api/export/contacts')
@role_required('超级管理员')
def export_contacts():
    """通讯录导出API
    支持参数：department（可选）、ignore_role_errors（可选）
    按主任、部长、信息员分段导出，包含文档元数据，提供进度查询。
    """
    try:
        from flask import current_app
        department = request.args.get('department')
        ignore_role_errors = (request.args.get('ignore_role_errors') or '').lower() in ('1', 'true', 'yes', 'on')
        role_error_message = '自动分配主任、部长信息时出现错误，请检查“主任”和“部长”的群组名称'

        # 生成任务ID并初始化进度
        from app.services.import_state import get_export_job, start_export_job, update_export_job
        job_id = start_export_job()

        # 查询用户
        query = _active_user_query()
        if department:
            query = query.filter_by(department=department)

        users = query.all()
        total = len(users)
        update_export_job(job_id, percent=5, message=f'查询到 {total} 条记录')

        def contact_number_sort_key(user):
            number = str(user.number or '').strip()
            if number.isdigit():
                return (0, int(number), number)
            return (1, number)

        group_ids = {u.group_id for u in users if u.group_id}
        groups_by_id = {
            group.id: group
            for group in Group.query.filter(Group.id.in_(group_ids)).all()
        } if group_ids else {}

        def current_group_name(user):
            if user.group_id and user.group_id in groups_by_id:
                return (groups_by_id[user.group_id].name or '').strip()
            return (user.group or '').strip()

        directors = [u for u in users if (u.department or '').strip() == '主任']
        ministers = [
            u for u in users
            if (u.department or '').strip() != '主任' and current_group_name(u) == '部长'
        ]
        informants = [
            u for u in users
            if (u.department or '').strip() != '主任' and current_group_name(u) != '部长'
        ]

        departments_in_scope = {
            (u.department or '').strip()
            for u in users
            if (
                (u.department or '').strip()
                and (u.department or '').strip() not in ('主任', UNASSIGNED_DEPARTMENT_NAME)
            )
        }
        minister_departments = {(u.department or '').strip() for u in ministers}
        missing_minister_departments = departments_in_scope - minister_departments
        role_error_details = []
        if not directors:
            role_error_details.append('不能检测到主任这一群组')
        if missing_minister_departments:
            missing_departments = '、'.join(sorted(missing_minister_departments))
            role_error_details.append(f'不能检测到以下部门的部长：{missing_departments}')
        if role_error_details and not ignore_role_errors:
            detailed_role_error_message = f'{role_error_message}：{"；".join(role_error_details)}'
            update_export_job(job_id, status='failed', message=detailed_role_error_message)
            return jsonify({'success': False, 'message': detailed_role_error_message}), 400

        directors.sort(key=contact_number_sort_key)
        ministers.sort(key=contact_number_sort_key)
        informants.sort(key=contact_number_sort_key)

        # 加载模板工作簿，保留模板列宽等基础设置，再重建导出内容。
        template_path = env_path('CONTACT_TEMPLATE_PATH', DEFAULT_CONTACT_TEMPLATE_PATH)
        import openpyxl
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active

        for merged_range in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merged_range))

        # 清空除模板表头外的内容，并把模板表头下移到第二行。
        max_row = ws.max_row
        if max_row > 1:
            ws.delete_rows(2, max_row - 1)
        ws.insert_rows(1)

        columns = ['编号','部门/组别','姓名','性别','年级','学院','专业','宿舍','手机号码','QQ号码','学号']
        center_alignment = Alignment(horizontal='center', vertical='center')
        title_font = Font(name='黑体', size=16)
        body_font = Font(name='微软雅黑', size=11)
        section_font = Font(name='微软雅黑', size=11, bold=True)
        grey_fill = PatternFill(fill_type='solid', start_color='FFD8D8D8', end_color='FFD8D8D8')
        thin_black_side = Side(style='thin', color='FF000000')
        thin_black_border = Border(
            left=thin_black_side,
            right=thin_black_side,
            top=thin_black_side,
            bottom=thin_black_side
        )

        def apply_content_border(cell):
            cell.border = thin_black_border

        def display_width(value):
            text = str(value or '')
            if text.isdigit():
                return len(text) * 1.25
            return sum(2 if ord(char) > 127 else 1 for char in text)

        def adjust_contact_column_widths(last_row):
            min_widths = {
                9: 16,   # 手机号码
                10: 14,  # QQ号码
                11: 18,  # 学号
            }
            for col_idx in range(1, 12):
                max_width = 0
                for row_idx in range(2, last_row + 1):
                    value = ws.cell(row=row_idx, column=col_idx).value
                    if value is not None:
                        max_width = max(max_width, display_width(value))
                column_letter = openpyxl.utils.get_column_letter(col_idx)
                min_width = min_widths.get(col_idx, 8)
                ws.column_dimensions[column_letter].width = max(min_width, min(max_width + 3, 60))
                ws.column_dimensions[column_letter].bestFit = True

        def apply_title_row(row_idx):
            for col_idx in range(1, 12):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = title_font
                cell.alignment = center_alignment
                apply_content_border(cell)
            ws.cell(row=row_idx, column=1, value='西南大学学生教学信息中心通讯录')
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
            ws.row_dimensions[row_idx].height = 28

        def apply_header_row(row_idx):
            for col_idx, header in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=header)
                cell.font = body_font
                cell.alignment = center_alignment
                cell.number_format = '@'
                apply_content_border(cell)

        def write_section_row(row_idx, title):
            for col_idx in range(1, 12):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = section_font
                cell.alignment = center_alignment
                cell.fill = grey_fill
                cell.number_format = '@'
                apply_content_border(cell)
            ws.cell(row=row_idx, column=1, value=title)
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
            ws.row_dimensions[row_idx].height = 22

        def write_user_row(row_idx, user):
            values = [
                user.number,
                user.department,
                user.name,
                user.gender,
                user.grade,
                user.college,
                user.major,
                user.dormitory,
                user.phone,
                user.qq,
                user.student_id
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else '')
                cell.font = body_font
                cell.alignment = center_alignment
                cell.number_format = '@'
                apply_content_border(cell)
                if col_idx == 1:
                    cell.fill = grey_fill

        apply_title_row(1)
        apply_header_row(2)

        row_idx = 3
        written_count = 0
        if ignore_role_errors:
            for user_item in sorted(users, key=contact_number_sort_key):
                write_user_row(row_idx, user_item)
                row_idx += 1
                written_count += 1
                if total and written_count % max(1, total // 20) == 0:
                    update_export_job(job_id, percent=5 + int(written_count / total * 90), message=f'已写入 {written_count}/{total}')
        else:
            sections = [
                ('中心主任', directors),
                ('中心部长', ministers),
                ('中心信息员', informants),
            ]
            for section_title, section_users in sections:
                write_section_row(row_idx, section_title)
                row_idx += 1
                for user_item in section_users:
                    write_user_row(row_idx, user_item)
                    row_idx += 1
                    written_count += 1
                    if total and written_count % max(1, total // 20) == 0:
                        update_export_job(job_id, percent=5 + int(written_count / total * 90), message=f'已写入 {written_count}/{total}')

        last_data_row = max(row_idx - 1, 2)
        ws.auto_filter.ref = f'A2:K{last_data_row}'
        ws.freeze_panes = 'A3'
        adjust_contact_column_widths(last_data_row)

        # 设置文档元数据
        from openpyxl.packaging.core import DocumentProperties
        wb.properties = DocumentProperties(
            creator='系统管理员',
            lastModifiedBy='系统管理员',
            title='部门通讯录',
            subject='部门通讯录导出',
            keywords='通讯录,部门,信息员',
            category='导出文件'
        )

        # 保存到导出目录
        export_dir = env_path('EXPORT_DIR', DEFAULT_EXPORT_DIR)
        os.makedirs(export_dir, exist_ok=True)
        fname_base = f"contacts_{(department or '全部')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        export_filename = f"{fname_base}.xlsx"
        export_path = os.path.join(export_dir, export_filename)
        wb.save(export_path)

        download_url = url_for('admin.download_passwords', filename=export_filename)
        update_export_job(
            job_id,
            percent=100,
            status='completed',
            message='导出完成',
            download_url=download_url,
        )

        current_app.logger.info(f'通讯录导出完成 记录数{total} 部门{department or "全部"} 文件{export_filename}')
        return jsonify({'success': True, 'job_id': job_id, 'download_url': download_url, 'total': total})
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'导出失败：{str(e)}'
        })


@admin_bp.route('/api/export/progress/<job_id>')
@role_required('超级管理员')
def export_progress(job_id):
    """查询导出进度"""
    from app.services.import_state import get_export_job
    progress = get_export_job(job_id)
    if progress is None:
        return jsonify({'success': False, 'message': '任务不存在'}), 404
    return jsonify({'success': True, 'progress': progress})


