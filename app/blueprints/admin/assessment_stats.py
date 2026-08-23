# -*- coding: utf-8 -*-
# Phase 1 mechanical split from app/blueprints/admin.py
# Module: assessment_stats

from flask import render_template, request, redirect, url_for, session, jsonify
from app.models import User, LectureForm, db, SystemSetting, AssessmentOverride
from datetime import datetime
from app.security import login_required, role_required
from app.services.assessment_calc import (
    assessment_form_latest_timestamp,
    assessment_latest_forms_in_range,
    build_department_monthly_assessment_payload,
    build_department_monthly_workbook,
    build_review_assessment_rows,
    build_submission_count_rows,
    build_submission_snapshot_payload,
    build_submission_snapshot_workbook,
    compute_teaching_week_window,
    get_form_effective_week_no,
    get_teaching_reward_settings,
    get_teaching_term_start,
    get_teaching_week_no,
    load_custom_month_definitions,
    parse_assessment_range,
    sort_submission_count_rows,
)
from app.services.assessment_scope import (
    get_accessible_department_users,
    has_assessment_stats_access,
    resolve_assessment_users,
    resolve_selected_departments,
)
from app.services.manual_assessment import (
    apply_manual_assessment_changes,
    build_manual_assessment_preview,
    parse_manual_assessment_rows,
    read_manual_assessment_file,
)
from app.services.review_form_queries import (
    build_review_form_filter_datetime,
    latest_form_groups_for_users,
    normalize_review_form_time_filter,
)
from app.services.stat_snapshots import (
    build_snapshot_list_items,
    can_access_snapshot,
    create_statistics_snapshot,
    find_snapshot_record,
    load_snapshot_payload,
)
from app.services.workbook import autosize_worksheet, build_export_filename, workbook_response
from app.utils.permission_feedback import build_forbidden_message, forbidden_json, flash_forbidden
from app.utils.manage_permissions import get_user_manage_permission
from app.utils.leave_management import LEAVE_OVERRIDE_TYPE
from app.utils.user_status import UNASSIGNED_GROUP_NAME
import openpyxl
from openpyxl.styles import Font
import json
from . import admin_bp
from .shared import get_reviewer_display_mode


# Compatibility aliases kept for tests/private callers that historically
# imported these teaching-week calculation helpers from this admin module.
_get_teaching_week_no = get_teaching_week_no
_get_teaching_term_start = get_teaching_term_start
_get_form_effective_week_no = get_form_effective_week_no
_compute_teaching_week_window = compute_teaching_week_window
_get_teaching_reward_settings = get_teaching_reward_settings


SNAPSHOT_TYPE_DEPARTMENT_MONTHLY = 'department_monthly_assessment'


SNAPSHOT_TYPE_SUBMISSION_REWARD = 'submission_reward'


def _get_snapshot_record_or_404(snapshot_id, snapshot_type, current_user_id):
    record = find_snapshot_record(snapshot_id, snapshot_type)
    if not record:
        return None, jsonify({'success': False, 'message': '历史记录不存在'}), 404
    if not can_access_snapshot(record, current_user_id):
        return None, jsonify({
            'success': False,
            'message': build_forbidden_message(
                '统计历史记录',
                '当前账号没有访问该历史记录的权限。',
                action='查看',
            ),
        }), 403
    return record, None, None


@admin_bp.route('/review-assessment-stats')
@role_required('管理员')
def review_assessment_stats():
    if not has_assessment_stats_access(session['user_id']):
        flash_forbidden('审表考评统计')
        return redirect(url_for('main.index'))
    manage_permission = get_user_manage_permission(session['user_id'])
    can_manual_assessment_import = manage_permission == '超级管理员'
    return render_template('admin/review_assessment_stats.html', can_manual_assessment_import=can_manual_assessment_import)


@admin_bp.route('/submission-count-stats')
@role_required('管理员')
def submission_count_stats():
    if not has_assessment_stats_access(session['user_id']):
        flash_forbidden('交表数量统计')
        return redirect(url_for('main.index'))
    return render_template('admin/submission_count_stats.html')


@admin_bp.route('/department-monthly-assessment-stats')
@role_required('管理员')
def department_monthly_assessment_stats():
    if not has_assessment_stats_access(session['user_id']):
        flash_forbidden('部门月度考评')
        return redirect(url_for('main.index'))
    available_departments = list(get_accessible_department_users(session['user_id']).keys())
    return render_template('admin/department_monthly_assessment_stats.html', available_departments=available_departments)


@admin_bp.route('/api/review/submission-count/stats', methods=['GET'])
@login_required
def get_submission_count_stats():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': True, 'users': []})
    reward_settings, reward_err = get_teaching_reward_settings()
    if reward_err:
        return jsonify({'success': False, 'message': reward_err}), 400
    reward_window = compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day'],
        reward_settings.get('total_weeks'),
    )
    sort_by = request.args.get('sort_by', 'submission')
    if sort_by not in ['submission', 'reward']:
        sort_by = 'submission'
    time_filter_type = normalize_review_form_time_filter(request.args.get('time_filter_type'))
    payload, payload_err = build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type)
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    rows = []
    for row in payload.get('users', []):
        copied_row = dict(row)
        copied_row.pop('reward_week_details', None)
        rows.append(copied_row)
    sort_submission_count_rows(rows, sort_by)
    return jsonify({
        'success': True,
        'users': rows,
        'sort_by': sort_by,
        'time_filter_type': time_filter_type,
        'reward_meta': payload.get('reward_meta', {})
    })


@admin_bp.route('/api/review/submission-count/export', methods=['GET'])
@login_required
def export_submission_count_stats():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': False, 'message': '无可导出用户'}), 400
    sort_by = request.args.get('sort_by', 'submission')
    if sort_by not in ['submission', 'reward']:
        sort_by = 'submission'
    time_filter_type = normalize_review_form_time_filter(request.args.get('time_filter_type'))
    payload, payload_err = build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type)
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    wb = build_submission_snapshot_workbook(payload)

    default_title = f'{start_date_str}至{end_date_str}交表统计'
    filename = build_export_filename(request.args.get('filename_title'), default_title)
    return workbook_response(wb, filename)


@admin_bp.route('/api/review/submission-count/reward-detail/<int:user_id>', methods=['GET'])
@login_required
def get_submission_reward_detail(user_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    users = resolve_assessment_users(session['user_id'], [str(user_id)])
    if not users:
        return jsonify({
            'success': False,
            'message': build_forbidden_message(
                '审表考评统计',
                '当前账号没有查看统计成员明细的权限。',
                action='查看',
            ),
        }), 403
    reward_settings, reward_err = get_teaching_reward_settings()
    if reward_err:
        return jsonify({'success': False, 'message': reward_err}), 400
    reward_window = compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day'],
        reward_settings.get('total_weeks'),
    )
    row = build_submission_count_rows(users, start_date, end_date, reward_settings, reward_window)[0]
    return jsonify({
        'success': True,
        'user': {
            'id': users[0].id,
            'name': users[0].name,
            'number': users[0].number,
            'department': users[0].department,
            'group': users[0].group
        },
        'summary': {
            'submission_group_count': row['submission_group_count'],
            'reward_effective_count': row['reward_effective_count'],
            'reward_error_count': row['reward_error_count'],
            'reward_form_count': row['reward_form_count']
        },
        'weeks': row['reward_week_details'],
        'reward_meta': {
            'week_label': reward_window['label'],
            'start_week': reward_window['start_week'],
            'end_week': reward_window['end_week'],
            'week_count': reward_window['week_count'],
            'has_full_weeks': reward_window['has_full_weeks'],
            'required_submission_per_week': reward_settings['required_submission']
        }
    })


@admin_bp.route('/api/review/submission-count/detail/<int:user_id>', methods=['GET'])
@login_required
def get_submission_count_detail(user_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表数量统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    users = resolve_assessment_users(session['user_id'], [str(user_id)])
    if not users:
        return jsonify({
            'success': False,
            'message': build_forbidden_message(
                '交表数量统计',
                '当前账号没有查看统计成员明细的权限。',
                action='查看',
            ),
        }), 403
    time_filter_type = normalize_review_form_time_filter(request.args.get('time_filter_type'))
    user = users[0]
    if not user.number:
        return jsonify({'success': True, 'user': {'id': user.id, 'name': user.name, 'number': user.number, 'department': user.department, 'group': user.group}, 'groups': []})

    reviewer_display_mode = get_reviewer_display_mode()
    groups = latest_form_groups_for_users([user.number])
    filtered_groups = []
    reviewer_ids = set()
    for group_data in groups:
        latest_form = group_data['latest_form']
        if not latest_form:
            continue
        filter_dt = build_review_form_filter_datetime(group_data, time_filter_type)
        if not filter_dt or filter_dt < start_date or filter_dt >= end_date:
            continue
        for form in group_data['forms']:
            if form.reviewer_id:
                reviewer_ids.add(form.reviewer_id)
    reviewer_map = {}
    if reviewer_ids:
        reviewers = User.query.filter(User.id.in_(list(reviewer_ids))).all()
        reviewer_map = {r.id: r for r in reviewers}

    for group_data in groups:
        latest_form = group_data['latest_form']
        if not latest_form or not latest_form.created_at:
            continue
        if latest_form.created_at < start_date or latest_form.created_at >= end_date:
            continue
        forms_data = []
        for form in group_data['forms']:
            reviewer = reviewer_map.get(form.reviewer_id) if form.reviewer_id else None
            if reviewer_display_mode == 'number':
                reviewer_display = reviewer.number if reviewer else '-'
            else:
                reviewer_display = reviewer.name if reviewer else '-'
            forms_data.append({
                'id': form.id,
                'status': form.status,
                'created_at': form.created_at.strftime('%Y-%m-%d %H:%M:%S') if form.created_at else '-',
                'updated_at': form.updated_at.strftime('%Y-%m-%d %H:%M:%S') if form.updated_at else '-',
                'review_comment': form.review_comment or '',
                'reviewer_display': reviewer_display,
                'teacher_name': form.teacher_name or '',
                'course_title': form.course_title or '',
                'lecture_date': form.lecture_date or ''
            })
        filtered_groups.append({
            'unique_id': group_data['unique_id'],
            'latest_created_at': latest_form.created_at.strftime('%Y-%m-%d %H:%M:%S') if latest_form.created_at else '-',
            'latest_updated_at': latest_form.updated_at.strftime('%Y-%m-%d %H:%M:%S') if latest_form.updated_at else '-',
            'filter_datetime': filter_dt.strftime('%Y-%m-%d %H:%M:%S') if filter_dt else '-',
            'teacher_name': latest_form.teacher_name or '',
            'course_title': latest_form.course_title or '',
            'lecture_date': latest_form.lecture_date or '',
            'forms': forms_data
        })

    filtered_groups.sort(key=lambda x: x['latest_created_at'], reverse=True)
    return jsonify({
        'success': True,
        'user': {
            'id': user.id,
            'name': user.name,
            'number': user.number,
            'department': user.department,
            'group': user.group
        },
        'groups': filtered_groups,
        'time_filter_type': time_filter_type
    })


@admin_bp.route('/api/review/submission-count/snapshots', methods=['GET'])
@login_required
def list_submission_count_snapshots():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    return jsonify({
        'success': True,
        'snapshots': build_snapshot_list_items(SNAPSHOT_TYPE_SUBMISSION_REWARD, session['user_id'])
    })


@admin_bp.route('/api/review/submission-count/snapshots', methods=['POST'])
@login_required
def save_submission_count_snapshot():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    data = request.get_json() or {}
    start_date, end_date, err = parse_assessment_range(data.get('start_date'), data.get('end_date'))
    if err:
        return jsonify({'success': False, 'message': err}), 400
    users = resolve_assessment_users(session['user_id'], data.get('user_ids') or [])
    if not users:
        return jsonify({'success': False, 'message': '无可保存用户'}), 400
    sort_by = data.get('sort_by', 'submission')
    if sort_by not in ['submission', 'reward']:
        sort_by = 'submission'
    time_filter_type = normalize_review_form_time_filter(data.get('time_filter_type'))
    payload, payload_err = build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type)
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    if not payload.get('reward_meta', {}).get('has_full_weeks'):
        return jsonify({'success': False, 'message': '当前筛选范围内无完整教学周，无法保存奖励表历史'}), 400

    filters_data = {
        'start_date': data.get('start_date'),
        'end_date': data.get('end_date'),
        'sort_by': sort_by,
        'time_filter_type': time_filter_type,
        'user_ids': [user.id for user in users],
        'scope_label': f'共 {len(users)} 人',
        'period_label': payload.get('reward_meta', {}).get('week_label', '')
    }
    title = (data.get('title') or '').strip() or f"{data.get('start_date')}至{data.get('end_date')}交表统计"
    record = create_statistics_snapshot(
        SNAPSHOT_TYPE_SUBMISSION_REWARD,
        title,
        filters_data,
        payload,
        session['user_id']
    )
    return jsonify({
        'success': True,
        'message': '奖励表历史已保存',
        'snapshot': {
            'id': record.id,
            'title': record.title
        }
    })


@admin_bp.route('/api/review/submission-count/snapshots/<int:snapshot_id>', methods=['GET'])
@login_required
def get_submission_count_snapshot(snapshot_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_SUBMISSION_REWARD, session['user_id'])
    if error_response:
        return error_response, status_code
    filters_data, payload_data = load_snapshot_payload(record)
    return jsonify({
        'success': True,
        'snapshot': {
            'id': record.id,
            'title': record.title,
            'created_at': record.created_at.strftime('%Y-%m-%d %H:%M:%S') if record.created_at else '',
            'filters': filters_data,
            'payload': payload_data
        }
    })


@admin_bp.route('/api/review/submission-count/snapshots/<int:snapshot_id>/export', methods=['GET'])
@login_required
def export_submission_count_snapshot(snapshot_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('交表统计快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_SUBMISSION_REWARD, session['user_id'])
    if error_response:
        return error_response, status_code
    _, payload_data = load_snapshot_payload(record)
    wb = build_submission_snapshot_workbook(payload_data)
    filename = build_export_filename(request.args.get('filename_title'), record.title)
    return workbook_response(wb, filename)


@admin_bp.route('/api/review/department-monthly-assessment/stats', methods=['GET'])
@login_required
def get_department_monthly_assessment_stats():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    payload, payload_err = build_department_monthly_assessment_payload(
        session['user_id'],
        start_date,
        end_date,
        request.args.getlist('departments')
    )
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    return jsonify({'success': True, **payload})


@admin_bp.route('/api/review/department-monthly-assessment/export', methods=['GET'])
@login_required
def export_department_monthly_assessment_stats():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    payload, payload_err = build_department_monthly_assessment_payload(
        session['user_id'],
        start_date,
        end_date,
        request.args.getlist('departments')
    )
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    if not payload.get('meta', {}).get('has_full_months'):
        return jsonify({'success': False, 'message': '当前筛选范围内无完整教学月，无法导出月度考评'}), 400
    wb = build_department_monthly_workbook(payload)
    default_title = f'{start_date_str}至{end_date_str}部门月度考评统计'
    filename = build_export_filename(request.args.get('filename_title'), default_title)
    return workbook_response(wb, filename)


@admin_bp.route('/api/review/department-monthly-assessment/snapshots', methods=['GET'])
@login_required
def list_department_monthly_assessment_snapshots():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    return jsonify({
        'success': True,
        'snapshots': build_snapshot_list_items(SNAPSHOT_TYPE_DEPARTMENT_MONTHLY, session['user_id'])
    })


@admin_bp.route('/api/review/department-monthly-assessment/snapshots', methods=['POST'])
@login_required
def save_department_monthly_assessment_snapshot():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    data = request.get_json() or {}
    start_date, end_date, err = parse_assessment_range(data.get('start_date'), data.get('end_date'))
    if err:
        return jsonify({'success': False, 'message': err}), 400
    departments = data.get('departments') or []
    payload, payload_err = build_department_monthly_assessment_payload(
        session['user_id'],
        start_date,
        end_date,
        departments
    )
    if payload_err:
        return jsonify({'success': False, 'message': payload_err}), 400
    if not payload.get('meta', {}).get('has_full_months'):
        return jsonify({'success': False, 'message': '当前筛选范围内无完整教学月，无法保存月度考评历史'}), 400
    filters_data = {
        'start_date': data.get('start_date'),
        'end_date': data.get('end_date'),
        'departments': departments,
        'scope_label': payload.get('meta', {}).get('scope_label', ''),
        'period_label': payload.get('meta', {}).get('period_label', '')
    }
    title = (data.get('title') or '').strip() or f"{data.get('start_date')}至{data.get('end_date')}部门月度考评统计"
    record = create_statistics_snapshot(
        SNAPSHOT_TYPE_DEPARTMENT_MONTHLY,
        title,
        filters_data,
        payload,
        session['user_id']
    )
    return jsonify({
        'success': True,
        'message': '部门月度考评历史已保存',
        'snapshot': {
            'id': record.id,
            'title': record.title
        }
    })


@admin_bp.route('/api/review/department-monthly-assessment/snapshots/<int:snapshot_id>', methods=['GET'])
@login_required
def get_department_monthly_assessment_snapshot(snapshot_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_DEPARTMENT_MONTHLY, session['user_id'])
    if error_response:
        return error_response, status_code
    filters_data, payload_data = load_snapshot_payload(record)
    return jsonify({
        'success': True,
        'snapshot': {
            'id': record.id,
            'title': record.title,
            'created_at': record.created_at.strftime('%Y-%m-%d %H:%M:%S') if record.created_at else '',
            'filters': filters_data,
            'payload': payload_data
        }
    })


@admin_bp.route('/api/review/department-monthly-assessment/snapshots/<int:snapshot_id>/export', methods=['GET'])
@login_required
def export_department_monthly_assessment_snapshot(snapshot_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('部门月度考评快照')
    record, error_response, status_code = _get_snapshot_record_or_404(snapshot_id, SNAPSHOT_TYPE_DEPARTMENT_MONTHLY, session['user_id'])
    if error_response:
        return error_response, status_code
    _, payload_data = load_snapshot_payload(record)
    wb = build_department_monthly_workbook(payload_data)
    filename = build_export_filename(request.args.get('filename_title'), record.title)
    return workbook_response(wb, filename)


@admin_bp.route('/api/review/assessment/stats', methods=['GET'])
@login_required
def get_review_assessment_stats():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('审表考评统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': True, 'users': []})

    rows = build_review_assessment_rows(users, start_date, end_date)
    return jsonify({'success': True, 'users': rows})


@admin_bp.route('/api/review/assessment/export', methods=['GET'])
@login_required
def export_review_assessment_stats():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('审表考评统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    user_ids = request.args.getlist('user_ids')
    users = resolve_assessment_users(session['user_id'], user_ids)
    if not users:
        return jsonify({'success': False, 'message': '无可导出用户'}), 400

    rows = build_review_assessment_rows(users, start_date, end_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '考评统计'
    headers = ['排名', '姓名', '编号', '部门', '小组', '部门扣分', '个人扣分', '总个人扣分']
    ws.append(headers)
    for row in rows:
        ws.append([
            row['rank'],
            row['name'],
            row['number'],
            row['department'],
            row['group'],
            float(row['linked_department_score'] or 0.0),
            float(row['linked_personal_score'] or 0.0),
            float(row['total_personal_score'] or 0.0)
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    autosize_worksheet(ws)

    default_title = f'{start_date_str}至{end_date_str}考评统计'
    filename = build_export_filename(request.args.get('filename_title'), default_title)
    return workbook_response(wb, filename)


@admin_bp.route('/api/review/assessment/detail/<int:user_id>', methods=['GET'])
@login_required
def get_review_assessment_detail(user_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('审表考评统计')
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    start_date, end_date, err = parse_assessment_range(start_date_str, end_date_str)
    if err:
        return jsonify({'success': False, 'message': err}), 400

    users = resolve_assessment_users(session['user_id'], [str(user_id)])
    if not users:
        return jsonify({
            'success': False,
            'message': build_forbidden_message(
                '审表考评统计',
                '当前账号没有查看统计成员明细的权限。',
                action='查看',
            ),
        }), 403
    user = users[0]

    entries = []
    groups = []
    linked_department_total = 0.0
    linked_personal_total = 0.0

    latest_forms = assessment_latest_forms_in_range([user.number], start_date, end_date)
    for form in latest_forms:
        if not form.score_record:
            continue
        display_time = assessment_form_latest_timestamp(form)
        group_items = []
        group_department_total = 0.0
        group_personal_total = 0.0
        for item in form.score_record.items:
            dept_score = float(item.department_score or 0.0)
            pers_score = float(item.personal_score or 0.0)
            linked_department_total += dept_score
            linked_personal_total += pers_score
            group_department_total += dept_score
            group_personal_total += pers_score
            item_data = {
                'issue_detail': item.reason or '',
                'department_score': dept_score,
                'personal_score': pers_score
            }
            group_items.append(item_data)
            entries.append({
                'source': '反馈表单',
                'date': display_time.strftime('%Y-%m-%d %H:%M') if display_time else '-',
                'teacher_name': form.teacher_name or '',
                'course_title': form.course_title or '',
                **item_data,
                'sort_time': display_time.isoformat() if display_time else ''
            })
        groups.append({
            'form_id': form.id,
            'unique_id': form.unique_id or '',
            'date': display_time.strftime('%Y-%m-%d %H:%M') if display_time else '-',
            'teacher_name': form.teacher_name or '',
            'course_title': form.course_title or '',
            'total_department_score': group_department_total,
            'total_personal_score': group_personal_total,
            'item_count': len(group_items),
            'items': group_items,
            'sort_time': display_time.isoformat() if display_time else ''
        })

    groups.sort(key=lambda x: x['sort_time'], reverse=True)
    entries.sort(key=lambda x: x['sort_time'], reverse=True)
    for group in groups:
        group.pop('sort_time', None)
    for entry in entries:
        entry.pop('sort_time', None)

    return jsonify({
        'success': True,
        'user': {
            'id': user.id,
            'name': user.name,
            'number': user.number,
            'department': user.department,
            'group': user.group
        },
        'summary': {
            'linked_department_score': linked_department_total,
            'linked_personal_score': linked_personal_total,
            'total_department_score': linked_department_total,
            'total_personal_score': linked_personal_total,
            'total_score': linked_personal_total
        },
        'groups': groups,
        'entries': entries
    })


@admin_bp.route('/api/review/assessment/import/template', methods=['GET'])
@role_required('超级管理员')
def download_manual_assessment_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '手动导入考评模板'
    headers = ['表单ID', '考评项', '部门扣分', '个人扣分']
    ws.append(headers)
    ws.append([1001, '课堂秩序管理欠佳', 1, 2])
    return workbook_response(wb, '手动导入考评模板.xlsx')


@admin_bp.route('/api/review/assessment/import/preview', methods=['POST'])
@role_required('超级管理员')
def preview_manual_assessment_import():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400

    df, read_error = read_manual_assessment_file(request.files['file'])
    if read_error:
        return jsonify({'success': False, 'message': read_error}), 400

    parsed_result, parse_error = parse_manual_assessment_rows(df)
    if parse_error:
        return jsonify({'success': False, 'message': parse_error}), 400

    preview_result = build_manual_assessment_preview(parsed_result)
    return jsonify({
        'success': True,
        'summary': preview_result['summary'],
        'changed_forms': preview_result['changed_forms'],
        'unchanged_form_ids': preview_result['unchanged_form_ids'],
        'skipped_details': preview_result['skipped_details']
    })


@admin_bp.route('/api/review/assessment/import', methods=['POST'])
@role_required('超级管理员')
def import_manual_assessment():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400

    df, read_error = read_manual_assessment_file(request.files['file'])
    if read_error:
        return jsonify({'success': False, 'message': read_error}), 400

    parsed_result, parse_error = parse_manual_assessment_rows(df)
    if parse_error:
        return jsonify({'success': False, 'message': parse_error}), 400

    preview_result = build_manual_assessment_preview(parsed_result)
    operator_id = session.get('user_id')
    import_time = datetime.now()
    imported_count, imported_form_count = apply_manual_assessment_changes(
        preview_result['changed_grouped_items'],
        operator_id,
        import_time
    )

    db.session.commit()
    return jsonify({
        'success': True,
        'imported_count': imported_count,
        'imported_form_count': imported_form_count,
        'valid_row_count': parsed_result['valid_row_count'],
        'unchanged_form_count': len(preview_result['unchanged_form_ids']),
        'unchanged_form_ids': preview_result['unchanged_form_ids'],
        'skipped_count': parsed_result['skipped_count'],
        'skipped_details': parsed_result['skipped_details']
    })


@admin_bp.route('/assessment-exemption-settings')
@role_required('管理员')
def assessment_exemption_settings():
    user = User.query.get(session['user_id'])
    if user.role == '超级管理员':
        return redirect(url_for('admin.system_management', tab='assessment'))
    if not has_assessment_stats_access(session['user_id']):
        flash_forbidden('考评规则设置')
        return redirect(url_for('main.index'))
    available_departments = list(get_accessible_department_users(session['user_id']).keys())
    is_super_admin = get_user_manage_permission(session['user_id']) == '超级管理员'
    return render_template(
        'admin/assessment_exemption_settings.html',
        available_departments=available_departments,
        is_super_admin=is_super_admin,
    )


@admin_bp.route('/api/assessment-overrides/list', methods=['GET'])
@login_required
def list_assessment_overrides():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('考评规则设置')
    department_names = request.args.getlist('departments')
    selected_departments, department_user_map = resolve_selected_departments(
        session['user_id'],
        department_names,
    )
    if not selected_departments:
        return jsonify({'success': True, 'departments': []})

    all_user_ids = []
    for users in department_user_map.values():
        all_user_ids.extend([u.id for u in users])

    overrides_by_user = {}
    if all_user_ids:
        overrides = AssessmentOverride.query.filter(
            AssessmentOverride.user_id.in_(all_user_ids)
        ).order_by(AssessmentOverride.start_week.asc(), AssessmentOverride.end_week.asc()).all()
        for ov in overrides:
            overrides_by_user.setdefault(ov.user_id, []).append({
                'id': ov.id,
                'override_type': ov.override_type,
                'override_value': ov.override_value,
                'start_week': ov.start_week,
                'end_week': ov.end_week,
                'reason': ov.reason,
                'created_at': ov.created_at.strftime('%Y-%m-%d %H:%M:%S') if ov.created_at else '',
            })

    result_departments = []
    for dept_name in selected_departments:
        users = department_user_map.get(dept_name, [])
        group_map = {}
        for user in users:
            group_name = user.group or UNASSIGNED_GROUP_NAME
            group_map.setdefault(group_name, []).append(user)
        groups = []
        for group_name in sorted(group_map.keys()):
            members = []
            for user in sorted(group_map[group_name], key=lambda u: (u.number or '', u.name or '')):
                members.append({
                    'user_id': user.id,
                    'name': user.name,
                    'number': user.number,
                    'overrides': overrides_by_user.get(user.id, []),
                })
            groups.append({
                'group_name': group_name,
                'members': members,
            })
        result_departments.append({
            'department': dept_name,
            'groups': groups,
        })

    return jsonify({'success': True, 'departments': result_departments})


@admin_bp.route('/api/assessment-overrides', methods=['POST'])
@login_required
def create_assessment_override():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('考评规则设置')
    data = request.get_json() or {}
    user_ids = data.get('user_ids') or []
    if not user_ids:
        return jsonify({'success': False, 'message': '请选择至少一个成员'}), 400
    try:
        start_week = int(data.get('start_week'))
        end_week = int(data.get('end_week'))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'message': '起止周次必须为整数'}), 400
    if start_week < 1 or end_week < 1:
        return jsonify({'success': False, 'message': '周次必须大于0'}), 400
    if start_week > end_week:
        return jsonify({'success': False, 'message': '起始周不能大于结束周'}), 400
    override_type = (data.get('override_type') or 'exempt').strip()
    if override_type not in ('exempt', 'custom_requirement', 'partial_exempt', LEAVE_OVERRIDE_TYPE):
        return jsonify({'success': False, 'message': '不支持的规则类型'}), 400
    if override_type == LEAVE_OVERRIDE_TYPE and get_user_manage_permission(session['user_id']) != '超级管理员':
        return jsonify({'success': False, 'message': '请假规则请在统计分析页为当前教学周发起'}), 403
    reason = (data.get('reason') or '').strip()
    if not reason:
        return jsonify({'success': False, 'message': '请填写理由'}), 400
    override_value = data.get('override_value')

    accessible_users = resolve_assessment_users(session['user_id'], [str(uid) for uid in user_ids])
    accessible_ids = {u.id for u in accessible_users}
    if not accessible_ids:
        return jsonify({
            'success': False,
            'message': build_forbidden_message(
                '考评规则设置',
                '当前账号没有操作所选成员规则的权限。',
                action='操作',
            ),
        }), 403

    created_count = 0
    skipped_count = 0
    for uid in user_ids:
        try:
            uid_int = int(uid)
        except (TypeError, ValueError):
            continue
        if uid_int not in accessible_ids:
            continue
        existing = AssessmentOverride.query.filter_by(
            user_id=uid_int,
            start_week=start_week,
            end_week=end_week,
            override_type=override_type,
        ).first()
        if existing:
            skipped_count += 1
            continue
        record = AssessmentOverride(
            user_id=uid_int,
            start_week=start_week,
            end_week=end_week,
            override_type=override_type,
            override_value=override_value,
            reason=reason,
            created_by=session['user_id'],
        )
        db.session.add(record)
        created_count += 1

    db.session.commit()
    message = f'成功添加 {created_count} 条规则'
    if skipped_count > 0:
        message += f'，跳过 {skipped_count} 条已存在的规则'
    return jsonify({
        'success': True,
        'message': message,
        'created_count': created_count,
        'skipped_count': skipped_count,
    })


@admin_bp.route('/api/assessment-overrides/<int:override_id>', methods=['DELETE'])
@login_required
def delete_assessment_override(override_id):
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('考评规则设置')
    record = db.session.get(AssessmentOverride, override_id)
    if not record:
        return jsonify({'success': False, 'message': '规则不存在'}), 404
    manage_permission = get_user_manage_permission(session['user_id'])
    if manage_permission != '超级管理员':
        accessible_users = resolve_assessment_users(session['user_id'], [str(record.user_id)])
        if not accessible_users:
            return jsonify({
                'success': False,
                'message': build_forbidden_message(
                    '考评规则设置',
                    '当前账号没有删除该规则的权限。',
                    action='删除',
                ),
            }), 403
    db.session.delete(record)
    db.session.commit()
    return jsonify({'success': True, 'message': '规则已删除'})


@admin_bp.route('/api/teaching-month-definitions', methods=['GET'])
@login_required
def get_teaching_month_definitions():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('教学月设置')
    custom = load_custom_month_definitions()
    if custom:
        return jsonify({'success': True, 'is_custom': True, 'months': custom})
    defaults = []
    for i in range(5):
        start_week = i * 4 + 1
        end_week = start_week + 3
        defaults.append({
            'start_week': start_week,
            'end_week': end_week,
            'label': f'第{i + 1}教学月',
        })
    return jsonify({'success': True, 'is_custom': False, 'months': defaults})


@admin_bp.route('/api/teaching-month-definitions', methods=['POST'])
@login_required
def save_teaching_month_definitions():
    if not has_assessment_stats_access(session['user_id']):
        return forbidden_json('教学月设置')
    data = request.get_json() or {}
    months = data.get('months')
    if not months:
        SystemSetting.set('teaching_month_definitions', '')
        return jsonify({'success': True, 'message': '已恢复为默认教学月设置（每4周一个教学月）'})
    if not isinstance(months, list):
        return jsonify({'success': False, 'message': 'months 必须是数组'}), 400

    validated = []
    for idx, item in enumerate(months):
        try:
            start_week = int(item.get('start_week'))
            end_week = int(item.get('end_week'))
        except (TypeError, ValueError):
            return jsonify({'success': False, 'message': f'第{idx + 1}项的周次必须为整数'}), 400
        if start_week < 1 or end_week < 1:
            return jsonify({'success': False, 'message': f'第{idx + 1}项的周次必须大于0'}), 400
        if start_week > end_week:
            return jsonify({
                'success': False,
                'message': f'第{idx + 1}项的起始周({start_week})不能大于结束周({end_week})',
            }), 400
        label = str(item.get('label') or '').strip()
        validated.append({'start_week': start_week, 'end_week': end_week, 'label': label})

    SystemSetting.set('teaching_month_definitions', json.dumps(validated, ensure_ascii=False))
    return jsonify({'success': True, 'message': f'已保存自定义教学月设置（{len(validated)}个教学月）'})


