# -*- coding: utf-8 -*-
# Phase 1 mechanical split from app/blueprints/admin.py
# Module: statistics

from flask import Blueprint, render_template, request, redirect, url_for, flash, session, jsonify, send_file, current_app
from app.models import User, Department, Group, LectureForm, Permission, RolePermission, Teacher, Venue, Course, ListeningBan, CourseRegistration, db, SystemSetting, ScoreRecord, ScoreItem, StatisticsSnapshot, PersonnelMovementRecord, AssessmentOverride
from sqlalchemy import func
from datetime import datetime, timedelta
from app.utils.auto_review import AutoReviewEngine, SETTING_KEY_SEMESTER_MONDAY, SETTING_KEY_SCHEDULE_PATH, SETTING_KEY_CONTACTS_PATH, SETTING_KEY_FEEDBACK_PATH
from app.blueprints.auth import login_required, role_required
from app.utils.review_permissions import (
    get_user_review_permission, 
    can_review_status, 
    get_user_structure_for_review, 
    get_reviewable_users,
    get_reviewable_status_list,
    get_next_status_after_review,
    get_review_permission_presentation,
)
from app.utils.permission_feedback import (
    build_forbidden_message,
    build_forbidden_payload,
    forbidden_json,
    flash_forbidden,
)
from app.utils.manage_permissions import (
    get_user_manage_permission,
    check_manage_permission
)
from app.utils.password_audit import record_password_audit
from app.utils.audit_tags import (
    REVIEW_TAG_OPTIONS,
    LATE_TAG_OPTIONS,
    LATE_TAG_LATE,
    REVIEW_TAG_REQUIRED,
    build_audit_tag,
    parse_audit_tag,
    validate_audit_tag,
    is_auto_review_allowed,
)
from app.utils.leave_management import (
    ASSESSMENT_EXEMPT_OVERRIDE_TYPES,
    LEAVE_OVERRIDE_TYPE,
    build_leave_status_payload,
    get_current_teaching_week as get_leave_current_teaching_week,
    get_form_effective_week_no as get_leave_form_effective_week_no,
    get_leave_makeup_forms,
    get_teaching_settings as get_leave_teaching_settings,
    parse_lecture_date_value,
    set_leave_makeup_forms,
)
from app.utils.user_status import (
    UNASSIGNED_DEPARTMENT_NAME,
    UNASSIGNED_GROUP_NAME,
    active_user_filter,
    is_user_active,
)
from app.utils.profile_settings import (
    PROFILE_EDITABLE_FIELD_OPTIONS,
    SETTING_KEY_PROFILE_EDITABLE_FIELDS,
    get_profile_editable_fields,
    normalize_profile_editable_fields,
)
from app.utils.course_registration_limits import (
    SETTING_KEY_COURSE_WEEKLY_LIMIT_COUNT,
    SETTING_KEY_COURSE_WEEKLY_LIMIT_ENABLED,
    get_course_weekly_limit_settings,
    normalize_course_weekly_limit_count,
)
from app.utils.review_drafts import (
    delete_review_form_draft,
    load_review_form_draft,
    normalize_review_draft_payload,
    parse_review_form_draft,
    save_review_form_draft,
)
from app.utils.env_config import env_path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
import os
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
import secrets
import string
from datetime import datetime
import json
import numpy as np
import re
from collections import defaultdict

from . import admin_bp  # noqa: F401
from .shared import _active_user_query, _get_accessible_department_users  # noqa: F401

STATISTICS_APPROVED_STATUSES = {'已审核', '部门已审核', '中心已审核'}


STATISTICS_DIMENSIONS = {'department', 'user', 'teacher', 'college', 'date'}


STATISTICS_TEACHER_KEY_SEPARATOR = '\u241f'


def _statistics_query(current_user, args):
    """Build the shared, permission-scoped query used by statistics views and exports."""
    query = db.session.query(LectureForm, User).join(
        User,
        LectureForm.listener_number == User.number,
    ).filter(active_user_filter())

    if current_user and current_user.role != '超级管理员':
        query = query.filter(User.department == current_user.department)

    start_date = (args.get('start_date') or '').strip()
    end_date = (args.get('end_date') or '').strip()
    if start_date:
        try:
            datetime.strptime(start_date, '%Y-%m-%d')
        except ValueError:
            start_date = ''
    if end_date:
        try:
            datetime.strptime(end_date, '%Y-%m-%d')
        except ValueError:
            end_date = ''
    if start_date:
        query = query.filter(LectureForm.lecture_date >= start_date)
    if end_date:
        query = query.filter(LectureForm.lecture_date <= end_date)

    department = (args.get('department') or '').strip()
    if department and current_user and current_user.role == '超级管理员':
        query = query.filter(User.department == department)

    status = (args.get('status') or '').strip()
    if status == '已审核':
        query = query.filter(LectureForm.status.in_(STATISTICS_APPROVED_STATUSES))
    elif status in {'待审核', '部门已审核', '中心已审核', '已驳回'}:
        query = query.filter(LectureForm.status == status)

    return query


def _statistics_dimension(args):
    dimension = (args.get('dimension') or 'department').strip()
    return dimension if dimension in STATISTICS_DIMENSIONS else 'department'


def _build_statistics_rows(records, dimension):
    groups = {}
    for form, listener in records:
        if dimension == 'user':
            key = ('user', listener.number or '')
            label = listener.name or form.listener_name or '未知用户'
            metadata = {
                'department': listener.department or '未分配部门',
                'detail_value': listener.number or '',
            }
        elif dimension == 'teacher':
            teacher_name = form.teacher_name or '未知教师'
            teacher_college = form.teacher_college or '未知学院'
            key = ('teacher', teacher_name, teacher_college)
            label = teacher_name
            metadata = {
                'college': teacher_college,
                'detail_value': STATISTICS_TEACHER_KEY_SEPARATOR.join((teacher_name, teacher_college)),
            }
        elif dimension == 'college':
            label = form.teacher_college or '未知学院'
            key = ('college', label)
            metadata = {'detail_value': label}
        elif dimension == 'date':
            label = str(form.lecture_date or '未知日期')
            key = ('date', label)
            metadata = {'detail_value': label}
        else:
            label = listener.department or '未分配部门'
            key = ('department', label)
            metadata = {'detail_value': label}

        item = groups.setdefault(key, {
            'name': label,
            'total': 0,
            'pending': 0,
            'approved': 0,
            'rejected': 0,
            **metadata,
        })
        item['total'] += 1
        if form.status == '待审核':
            item['pending'] += 1
        elif form.status in STATISTICS_APPROVED_STATUSES:
            item['approved'] += 1
        elif form.status == '已驳回':
            item['rejected'] += 1

    return sorted(
        groups.values(),
        key=lambda item: (str(item['name']).casefold(), str(item.get('detail_value', '')).casefold()),
    )


def _apply_statistics_detail_filter(query, dimension, detail_value):
    detail_value = (detail_value or '').strip()
    if not detail_value:
        return query
    if dimension == 'user':
        return query.filter(User.number == detail_value)
    if dimension == 'teacher':
        teacher_name, separator, teacher_college = detail_value.partition(STATISTICS_TEACHER_KEY_SEPARATOR)
        query = query.filter(LectureForm.teacher_name == teacher_name)
        return query.filter(LectureForm.teacher_college == teacher_college) if separator else query
    if dimension == 'college':
        return query.filter(LectureForm.teacher_college == detail_value)
    if dimension == 'date':
        return query.filter(LectureForm.lecture_date == detail_value)
    return query.filter(User.department == detail_value)


def _safe_statistics_excel_text(value):
    """Prevent aggregate labels from being interpreted as spreadsheet formulas."""
    text = str(value or '')
    if text.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + text
    return text


@admin_bp.route('/statistics')
@role_required('管理员')
def statistics():
    """统计分析页面"""
    current_user = User.query.get(session['user_id'])
    manage_permission = get_user_manage_permission(session['user_id'])
    can_access_extended_stats = manage_permission in ['超级管理员', '管理部门']
    can_manage_department_leave = manage_permission == '超级管理员'
    available_departments = list(_get_accessible_department_users(session['user_id']).keys()) if can_access_extended_stats else []

    form_query = _statistics_query(current_user, request.args)
    records = form_query.order_by(LectureForm.created_at.asc(), LectureForm.id.asc()).all()
    forms = [form for form, _listener in records]
    total_forms = len(forms)
    logical_form_keys = {
        ('unique_id', str(unique_id).strip())
        if unique_id is not None and str(unique_id).strip()
        else ('id', str(form_id))
        for unique_id, form_id in ((form.unique_id, form.id) for form in forms)
    }
    logical_form_count = len(logical_form_keys)
    pending_forms = sum(form.status == '待审核' for form in forms)
    approved_forms = sum(form.status in STATISTICS_APPROVED_STATUSES for form in forms)
    rejected_forms = sum(form.status == '已驳回' for form in forms)

    trend_counts = {}
    for created_at in (form.created_at for form in forms):
        if created_at:
            label = created_at.strftime('%Y-%m-%d')
            trend_counts[label] = trend_counts.get(label, 0) + 1
    trend_labels = sorted(trend_counts)
    trend_data = [trend_counts[label] for label in trend_labels]
    if not trend_labels:
        visualization_state='NO_DATA'
    elif len(trend_labels) == 1:
        visualization_state='INSUFFICIENT_DATA'
    else:
        visualization_state='DATA_READY'

    statistics_dimension = _statistics_dimension(request.args)
    return render_template(
        'admin/statistics.html',
        can_access_extended_stats=can_access_extended_stats,
        can_manage_department_leave=can_manage_department_leave,
        total_forms=total_forms,
        logical_form_count=logical_form_count,
        pending_forms=pending_forms,
        approved_forms=approved_forms,
        rejected_forms=rejected_forms,
        departments=available_departments,
        statistics_data=_build_statistics_rows(records, statistics_dimension),
        statistics_dimension=statistics_dimension,
        visualization_state=visualization_state,
        trend_point_count=trend_data[0] if len(trend_data) == 1 else None,
        trend_labels=json.dumps(trend_labels, ensure_ascii=False),
        trend_data=json.dumps(trend_data, ensure_ascii=False),
    )


@admin_bp.route('/export_statistics')
@role_required('管理员')
def export_statistics():
    """Export the currently filtered aggregate statistics as an XLSX workbook."""
    current_user = User.query.get(session['user_id'])
    dimension = _statistics_dimension(request.args)
    records = _statistics_query(current_user, request.args).order_by(
        LectureForm.created_at.asc(),
        LectureForm.id.asc(),
    ).all()
    rows = _build_statistics_rows(records, dimension)

    dimension_headers = {
        'department': ['部门'],
        'user': ['用户', '部门'],
        'teacher': ['教师', '学院'],
        'college': ['学院'],
        'date': ['日期'],
    }
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = '统计汇总'
    headers = dimension_headers[dimension] + ['总数', '待审核', '已通过', '已驳回', '通过率']
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for item in rows:
        if dimension == 'user':
            identity = [item['name'], item.get('department', '')]
        elif dimension == 'teacher':
            identity = [item['name'], item.get('college', '')]
        else:
            identity = [item['name']]
        identity = [_safe_statistics_excel_text(value) for value in identity]
        pass_rate = item['approved'] / item['total'] if item['total'] else 0
        worksheet.append(identity + [
            item['total'],
            item['pending'],
            item['approved'],
            item['rejected'],
            pass_rate,
        ])
        worksheet.cell(row=worksheet.max_row, column=len(headers)).number_format = '0.0%'

    worksheet.freeze_panes = 'A2'
    for column_index, header in enumerate(headers, start=1):
        values = [str(worksheet.cell(row=row, column=column_index).value or '') for row in range(1, worksheet.max_row + 1)]
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(column_index)].width = min(
            max(len(header) + 2, max((len(value) for value in values), default=0) + 2),
            32,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f'统计分析_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@admin_bp.route('/api/statistics_data')
@role_required('管理员')
def get_statistics_data():
    """获取统计数据API"""
    user = User.query.get(session['user_id'])
    
    # 获取筛选参数
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    department = request.args.get('department')
    status = request.args.get('status')
    
    # 构建基础查询
    if user.role == '超级管理员':
        query = LectureForm.query.join(User, LectureForm.listener_number == User.number).filter(active_user_filter())
    else:
        # 管理员只能查看本部门数据
        query = LectureForm.query.join(User, LectureForm.listener_number == User.number)\
                                 .filter(User.department == user.department, active_user_filter())
    
    # 应用筛选条件（先不按状态过滤，后续按“表单组最新状态”过滤）
    if start_date:
        # 注意：lecture_date 是字符串字段，此处比较仅作为简单示例，可能需要根据实际存储格式调整
        query = query.filter(LectureForm.lecture_date >= start_date)
    if end_date:
        query = query.filter(LectureForm.lecture_date <= end_date)
    if department and user.role == '超级管理员':
        query = query.filter(User.department == department, active_user_filter())
    
    forms = query.order_by(LectureForm.updated_at.desc(), LectureForm.id.desc()).all()

    latest_forms_by_group = {}
    for form in forms:
        uid = form.unique_id or form.id
        if uid not in latest_forms_by_group:
            latest_forms_by_group[uid] = form

    latest_forms = list(latest_forms_by_group.values())
    
    review_permission = get_user_review_permission(user.id)
    if review_permission == '审表_中心':
        pending_statuses = {'待审核', '部门已审核'}
    else:
        pending_statuses = {'待审核'}

    if status:
        latest_forms = [f for f in latest_forms if f.status == status]

    pending_forms = len([f for f in latest_forms if f.status in pending_statuses])
    
    # 按部门统计
    dept_stats = {}
    for form in latest_forms:
        listener = _active_user_query().filter_by(number=form.listener_number).first()
        if listener:
            dept = listener.department
            if dept not in dept_stats:
                dept_stats[dept] = {'pending': 0}
            
            # 使用集合去重 unique_id
            if 'pending_ids' not in dept_stats[dept]:
                dept_stats[dept]['pending_ids'] = set()
                
            if form.status in pending_statuses:
                dept_stats[dept]['pending_ids'].add(form.unique_id or form.id)
                
    # 整理部门统计数据
    for dept in dept_stats:
        dept_stats[dept]['pending'] = len(dept_stats[dept].pop('pending_ids'))
    
    # 按月份统计
    month_stats = {}
    for form in latest_forms:
        if form.lecture_date:
            # 假设 lecture_date 格式为 YYYY-MM-DD
            try:
                date_str = str(form.lecture_date)
                # 提取年份和月份
                if '-' in date_str:
                    parts = date_str.split('-')
                    if len(parts) >= 2:
                        month_key = f"{parts[0]}-{parts[1]}"
                    else:
                        month_key = '未知'
                else:
                    month_key = '未知'
            except:
                month_key = '未知'

            if month_key not in month_stats:
                month_stats[month_key] = {'pending': 0, 'pending_ids': set()}
            
            if form.status in pending_statuses:
                month_stats[month_key]['pending_ids'].add(form.unique_id or form.id)
                
    # 整理月份统计数据
    for month in month_stats:
        month_stats[month]['pending'] = len(month_stats[month].pop('pending_ids'))
    
    # 按听课类型统计 (模型无此字段，暂留空)
    type_stats = {}
    
    # 计算平均评分 (模型无此字段，暂为0)
    avg_score = 0
    
    return jsonify({
        'overview': {
            'pending_forms': pending_forms,
            'avg_score': avg_score
        },
        'department_stats': dept_stats,
        'month_stats': month_stats,
        'type_stats': type_stats
    })


@admin_bp.route('/api/statistics_detail')
@role_required('管理员')
def get_statistics_detail():
    """Return detail rows from the same permission and filter scope as the overview."""
    current_user = User.query.get(session['user_id'])
    page = request.args.get('page', 1, type=int)
    per_page = 10
    dimension = _statistics_dimension({'dimension': request.args.get('detail_dimension')})
    query = _statistics_query(current_user, request.args)
    query = _apply_statistics_detail_filter(
        query,
        dimension,
        request.args.get('detail_value'),
    ).order_by(LectureForm.created_at.desc(), LectureForm.id.desc())
    pagination = query.paginate(page=max(page, 1), per_page=per_page, error_out=False)

    form_data = []
    for form, listener in pagination.items:
        form_data.append({
            'id': form.id,
            'listener_name': form.listener_name,
            'listener_department': listener.department or '未知',
            'teacher_name': form.teacher_name,
            'course_name': form.course_title,
            'listen_date': str(form.lecture_date or ''),
            'listen_type': '普通听课',
            'overall_rating': form.overall_effect,
            'status': form.status,
            'created_at': form.created_at.strftime('%Y-%m-%d %H:%M') if form.created_at else ''
        })
    
    return jsonify({
        'forms': form_data,
        'pagination': {
            'page': pagination.page,
            'pages': pagination.pages,
            'per_page': pagination.per_page,
            'total': pagination.total,
            'has_prev': pagination.has_prev,
            'has_next': pagination.has_next
        }
    })


