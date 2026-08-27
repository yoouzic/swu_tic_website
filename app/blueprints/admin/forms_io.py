# -*- coding: utf-8 -*-
# Phase 1 mechanical split from app/blueprints/admin.py
# Module: forms_io

from flask import render_template, request, redirect, url_for, flash, session, jsonify, send_file
from app.models import User, LectureForm, db, SystemSetting
from datetime import datetime
from app.security import login_required, role_required
from app.services.workbook import workbook_response
from app.utils.audit_tags import REVIEW_TAG_REQUIRED, validate_audit_tag
from app.utils.user_status import active_user_filter
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
from . import admin_bp
from .shared import _active_user_query, _excel_cell_to_text, allowed_file, get_reviewer_display_mode


FORM_IMPORT_COLUMN_ALIASES = {
    'listener_number': ['听课人编号', '听课人(填写编号)', '编号', '信息员编号', '学号'],
    'listener_name': ['听课人姓名+学院', '听课人', '听课人姓名'],
    'course_changes': ['课程信息变化', '课程信息变化(一般三个方面:老师、教室、时间,没有变化填“无”)'],
    'lecture_date': ['听课时间', '听课时间(如:2023/10/19星期四)', '日期', '上课时间'],
    'class_period': ['第几节', '第几节(如:第1-3节)', '节次', '上课节次'],
    'lecture_location': ['听课地点', '听课地点(如:32-302)', '上课地点', '教室'],
    'teacher_name': ['授课教师', '授课教师(谨防错别字)', '教师', '教师姓名'],
    'teacher_college': ['教师所属学院', '教师所属学院(对照全校课表填写)', '教师学院', '学院'],
    'course_title': ['课程名称', '课程(总标题)', '课程', '课程名'],
    'student_grade_class': ['专业年级', '专业年级(如:2018级植物生产类05、06班)', '班级信息'],
    'abnormal_situation': ['异常情况反映', '异常情况反映(根据事实,没有则填"无")'],
    'teaching_method': ['主要教学方法'],
    'classroom_discipline': ['管理课堂纪律', '课堂纪律'],
    'classroom_atmosphere': ['调动课堂气氛', '课堂气氛'],
    'courseware_quality': ['课件制作质量', '课件质量'],
    'overall_effect': ['整体教学效果'],
    'quality_case': ['优质案例推荐', '是否推荐优质案例'],
    'course_feedback': ['课程反馈', '课程反馈(优点,五十字以上,评价的内容实在且有针对性,结尾不需要句号)'],
    'suggestions': ['不足及建议', '不足及建议(根据事实,没有则填"无")', '建议'],
    'student_signature1': ['听课班级同学签名1', '听课班级同学签名'],
    'contact_phone1': ['联系电话1', '联系电话'],
    'student_signature2': ['听课班级同学签名2'],
    'contact_phone2': ['联系电话2'],
    'status': ['status', '状态'],
    'audit_tag': ['audit_tag', '审核标签'],
    'review_comment': ['review_comment', '审核意见'],
    'created_at': ['created_at', '创建时间'],
    'updated_at': ['updated_at', '更新时间']
}


FORM_IMPORT_TEMPLATE_HEADERS = [
    '听课人编号', '听课人姓名+学院', '课程信息变化', '听课时间', '第几节', '听课地点',
    '授课教师', '教师所属学院', '课程名称', '专业年级', '异常情况反映', '主要教学方法',
    '管理课堂纪律', '调动课堂气氛', '课件制作质量', '整体教学效果', '优质案例推荐',
    '课程反馈', '不足及建议', '听课班级同学签名1', '联系电话1', '听课班级同学签名2', '联系电话2',
    '状态', '审核标签', '审核意见', '创建时间', '更新时间'
]


LEGACY_IMPORT_INDEX_MAP = {
    'listener_name': 1,
    'listener_number': 2,
    'course_changes': 3,
    'lecture_date': 4,
    'class_period': 5,
    'lecture_location': 6,
    'teacher_name': 7,
    'teacher_college': 8,
    'course_title': 9,
    'student_grade_class': 10,
    'abnormal_situation': 11,
    'teaching_method': 12,
    'classroom_discipline': 13,
    'classroom_atmosphere': 14,
    'courseware_quality': 15,
    'overall_effect': 16,
    'quality_case': 17,
    'course_feedback': 18,
    'suggestions': 19,
    'student_signature1': 21,
    'contact_phone1': 22,
    'student_signature2': 23,
    'contact_phone2': 24
}


def _normalize_listener_number(value):
    text = _excel_cell_to_text(value).replace(' ', '')
    if text.endswith('.0'):
        candidate = text[:-2]
        if candidate.isdigit():
            return candidate
    return text


def _extract_named_form_rows(df):
    columns = [str(c).strip() for c in df.columns]
    alias_to_column = {}
    for key, aliases in FORM_IMPORT_COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in columns:
                alias_to_column[key] = alias
                break
    if not all(k in alias_to_column for k in ['listener_number', 'teacher_name', 'course_title']):
        return []
    rows = []
    for _, row in df.iterrows():
        if row.isna().all():
            continue
        item = {}
        for key, col_name in alias_to_column.items():
            item[key] = _excel_cell_to_text(row.get(col_name))
        if item.get('listener_number') or item.get('teacher_name') or item.get('course_title'):
            rows.append(item)
    return rows


def _extract_legacy_form_rows(df):
    rows = []
    for _, row in df.iterrows():
        item = {}
        has_content = False
        for key, idx in LEGACY_IMPORT_INDEX_MAP.items():
            value = _excel_cell_to_text(row[idx]) if idx < len(row) else ''
            item[key] = value
            if key in ['listener_number', 'teacher_name', 'course_title'] and value:
                has_content = True
        if has_content:
            rows.append(item)
    return rows


def _parse_form_import_rows(file_storage):
    try:
        df_named = pd.read_excel(file_storage, sheet_name=0)
    except Exception:
        return [], 'Excel读取失败，请检查文件格式'
    rows = _extract_named_form_rows(df_named)
    if rows:
        return rows, None
    try:
        file_storage.stream.seek(0)
        df_legacy = pd.read_excel(file_storage, sheet_name=0, header=None)
    except Exception:
        return [], 'Excel读取失败，请检查文件格式'
    legacy_rows = _extract_legacy_form_rows(df_legacy)
    if legacy_rows:
        return legacy_rows, None
    return [], '未识别到可导入的数据，请使用系统模板或检查表头'


def _parse_optional_datetime(value, fallback):
    if value is None:
        return fallback
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    text = _excel_cell_to_text(value)
    if not text:
        return fallback
    normalized = text.replace('T', ' ').replace('/', '-').strip()
    formats = [
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%Y-%m-%d'
    ]
    for fmt in formats:
        try:
            return datetime.strptime(normalized, fmt)
        except Exception:
            continue
    try:
        return datetime.fromisoformat(normalized)
    except Exception:
        return fallback


def _prepare_import_form_row(row, row_no, importer_id, import_time):
    errors = []
    listener_number = _normalize_listener_number(row.get('listener_number', ''))
    if not listener_number:
        errors.append('听课人编号为空')
        return {'ok': False, 'errors': errors, 'row_no': row_no}
    listener_user = _active_user_query().filter_by(number=listener_number).first()
    if not listener_user:
        errors.append(f'听课人编号 {listener_number} 不存在')
        return {'ok': False, 'errors': errors, 'row_no': row_no, 'listener_number': listener_number}

    lecture_date = _excel_cell_to_text(row.get('lecture_date'))
    class_period = _excel_cell_to_text(row.get('class_period'))
    lecture_location = _excel_cell_to_text(row.get('lecture_location'))
    teacher_name = _excel_cell_to_text(row.get('teacher_name'))
    teacher_college = _excel_cell_to_text(row.get('teacher_college'))
    course_title = _excel_cell_to_text(row.get('course_title'))
    student_grade_class = _excel_cell_to_text(row.get('student_grade_class'))
    teaching_method = _excel_cell_to_text(row.get('teaching_method'))
    classroom_discipline = _excel_cell_to_text(row.get('classroom_discipline'))
    classroom_atmosphere = _excel_cell_to_text(row.get('classroom_atmosphere'))
    overall_effect = _excel_cell_to_text(row.get('overall_effect'))
    quality_case = _excel_cell_to_text(row.get('quality_case'))
    course_feedback = _excel_cell_to_text(row.get('course_feedback'))
    student_signature1 = _excel_cell_to_text(row.get('student_signature1'))
    contact_phone1 = _excel_cell_to_text(row.get('contact_phone1'))

    required_pairs = [
        ('听课时间', lecture_date),
        ('第几节', class_period),
        ('听课地点', lecture_location),
        ('授课教师', teacher_name),
        ('教师所属学院', teacher_college),
        ('课程名称', course_title),
        ('专业年级', student_grade_class),
        ('主要教学方法', teaching_method),
        ('管理课堂纪律', classroom_discipline),
        ('调动课堂气氛', classroom_atmosphere),
        ('整体教学效果', overall_effect),
        ('优质案例推荐', quality_case),
        ('课程反馈', course_feedback),
        ('听课班级同学签名1', student_signature1),
        ('联系电话1', contact_phone1)
    ]
    missing_fields = [name for name, value in required_pairs if not value]
    if missing_fields:
        errors.append(f'缺少必填字段 {",".join(missing_fields[:5])}')

    status = _excel_cell_to_text(row.get('status')) or '部门已审核'
    allowed_status = {'待审核', '部门已审核', '中心已审核', '已驳回'}
    if status not in allowed_status:
        errors.append('status 不合法')
    audit_tag = _excel_cell_to_text(row.get('audit_tag')) or REVIEW_TAG_REQUIRED
    errors.extend(validate_audit_tag(audit_tag))

    courseware_quality = _excel_cell_to_text(row.get('courseware_quality')) or '无'
    if 'PPT演示法' not in teaching_method and not _excel_cell_to_text(row.get('courseware_quality')):
        courseware_quality = '无'

    listener_name = _excel_cell_to_text(row.get('listener_name')) or f'{listener_user.name}（{listener_user.college}）'
    created_at = _parse_optional_datetime(row.get('created_at'), import_time)
    updated_at = _parse_optional_datetime(row.get('updated_at'), import_time)
    payload = {
        'listener_name': listener_name,
        'listener_number': listener_number,
        'course_changes': _excel_cell_to_text(row.get('course_changes')) or '无',
        'lecture_date': lecture_date,
        'class_period': class_period,
        'lecture_location': lecture_location,
        'teacher_name': teacher_name,
        'teacher_college': teacher_college,
        'course_title': course_title,
        'student_grade_class': student_grade_class,
        'abnormal_situation': _excel_cell_to_text(row.get('abnormal_situation')) or '无',
        'teaching_method': teaching_method,
        'classroom_discipline': classroom_discipline,
        'classroom_atmosphere': classroom_atmosphere,
        'courseware_quality': courseware_quality,
        'overall_effect': overall_effect,
        'quality_case': quality_case,
        'course_feedback': course_feedback,
        'suggestions': _excel_cell_to_text(row.get('suggestions')) or '无',
        'student_signature1': student_signature1,
        'contact_phone1': contact_phone1,
        'student_signature2': _excel_cell_to_text(row.get('student_signature2')),
        'contact_phone2': _excel_cell_to_text(row.get('contact_phone2')),
        'registration_id': None,
        'status': status,
        'audit_tag': audit_tag,
        'reviewer_id': importer_id,
        'review_time': import_time,
        'review_comment': _excel_cell_to_text(row.get('review_comment')),
        'created_at': created_at,
        'updated_at': updated_at
    }
    if errors:
        return {
            'ok': False,
            'errors': errors,
            'row_no': row_no,
            'listener_number': listener_number,
            'teacher_name': teacher_name,
            'course_title': course_title,
            'status': status
        }
    return {
        'ok': True,
        'payload': payload,
        'row_no': row_no,
        'listener_number': listener_number,
        'teacher_name': teacher_name,
        'course_title': course_title,
        'status': status
    }


@admin_bp.route('/view_forms')
@login_required
def view_forms():
    """查看听课表单 - 按unique_id分组显示"""
    user = User.query.get(session['user_id'])
    page = request.args.get('page', 1, type=int)
    per_page = 20
    
    # 搜索参数
    search = request.args.get('search', '')
    date_filter = request.args.get('date', '')
    status_filter = request.args.get('status', '')
    
    # 基础查询
    if user.role == '超级管理员':
        # 超级管理员可以查看所有表单
        base_query = LectureForm.query.join(User, LectureForm.listener_number == User.number).filter(active_user_filter())
    elif user.role == '管理员':
        # 管理员只能查看本部门的表单
        base_query = LectureForm.query.join(User, LectureForm.listener_number == User.number)\
                                     .filter(User.department == user.department, active_user_filter())
    else:
        # 信息员只能查看自己的表单
        base_query = LectureForm.query.filter_by(listener_number=user.number)
    
    # 应用搜索过滤
    if search:
        base_query = base_query.filter(
            db.or_(
                LectureForm.listener_name.contains(search),
                LectureForm.teacher_name.contains(search),
                LectureForm.course_title.contains(search)
            )
        )
    
    if date_filter:
        base_query = base_query.filter(LectureForm.lecture_date.contains(date_filter))
    
    if status_filter:
        if status_filter == '已审核':
            base_query = base_query.filter(LectureForm.status.in_(['已审核', '部门已审核', '中心已审核']))
        else:
            base_query = base_query.filter(LectureForm.status == status_filter)
    
    # 获取所有符合条件的表单，按unique_id分组
    all_forms = base_query.order_by(LectureForm.unique_id, LectureForm.updated_at.desc()).all()
    
    # 按unique_id分组，每组只取最新的表单作为代表
    form_groups = {}
    for form in all_forms:
        unique_id = form.unique_id or f"single_{form.id}"  # 处理没有unique_id的旧表单
        if unique_id not in form_groups:
            form_groups[unique_id] = {
                'latest_form': form,
                'versions': []
            }
        form_groups[unique_id]['versions'].append(form)
    
    # 转换为列表并按创建时间排序（使用组内第一个版本的创建时间）
    grouped_forms = list(form_groups.values())
    
    # 获取每个组的最早版本时间作为排序依据
    for group in grouped_forms:
        # 在 form_groups 构建时，versions 是按 updated_at desc 排列的
        # 所以 versions[-1] 是最早的版本（如果按时间倒序）
        # 或者更稳妥地，重新排序 versions 找到最早的 created_at
        versions = group['versions']
        if versions:
            # 找到最早的创建时间
            earliest_time = min(v.created_at for v in versions if v.created_at)
            group['sort_time'] = earliest_time
        else:
            group['sort_time'] = datetime.now() # Fallback

    # 按最早创建时间倒序排序
    grouped_forms.sort(key=lambda x: x['sort_time'], reverse=True)
    
    # 手动实现分页
    total = len(grouped_forms)
    start = (page - 1) * per_page
    end = start + per_page
    page_forms = grouped_forms[start:end]
    
    # 创建分页对象
    class Pagination:
        def __init__(self, page, per_page, total, items):
            self.page = page
            self.per_page = per_page
            self.total = total
            self.items = items
            self.pages = (total + per_page - 1) // per_page
            self.has_prev = page > 1
            self.has_next = page < self.pages
            self.prev_num = page - 1 if self.has_prev else None
            self.next_num = page + 1 if self.has_next else None
        
        def iter_pages(self, left_edge=2, left_current=2, right_current=3, right_edge=2):
            last = self.pages
            for num in range(1, last + 1):
                if num <= left_edge or \
                   (self.page - left_current - 1 < num < self.page + right_current) or \
                   num > last - right_edge:
                    yield num
    
    forms = Pagination(page, per_page, total, page_forms)

    def page_url(target_page):
        args = request.args.to_dict(flat=True)
        args['page'] = target_page
        return url_for('admin.view_forms', **args)

    reviewer_display_mode = get_reviewer_display_mode()
    return render_template(
        'admin/view_forms.html',
        forms=forms,
        user=user,
        reviewer_display_mode=reviewer_display_mode,
        page_url=page_url,
    )


@admin_bp.route('/import_forms_excel')
@role_required('超级管理员')
def import_forms_excel_page():
    return render_template('admin/import_forms_excel.html')


@admin_bp.route('/api/forms/import/template', methods=['GET'])
@role_required('超级管理员')
def download_forms_import_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表单导入模板'
    ws.append(FORM_IMPORT_TEMPLATE_HEADERS)
    ws.append([
        '20230001', '张三（计算机与信息科学学院）', '无', '2026/03/15星期三', '第3-4节', '32-302',
        '李老师', '计算机与信息科学学院', '数据结构', '2023级计算机1班', '无', 'PPT演示法；案例教学法',
        '秩序良好', '互动积极', '图文清晰', '整体较好', '可推荐', '课堂目标明确，内容组织清晰，学生参与度高。',
        '建议增加课堂练习时间', '王同学', '13800000000', '', '',
        '部门已审核', '需要人工审核', '', '2026-03-16 10:00:00', '2026-03-16 10:05:00'
    ])
    for col_idx in range(1, len(FORM_IMPORT_TEMPLATE_HEADERS) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
    return workbook_response(wb, '听课表单导入模板.xlsx')


@admin_bp.route('/api/forms/import', methods=['POST'])
@role_required('超级管理员')
def import_forms_from_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'success': False, 'message': '请先选择导入文件'}), 400
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'message': '仅支持xls或xlsx文件'}), 400

    rows, err = _parse_form_import_rows(file)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    if not rows:
        return jsonify({'success': False, 'message': '未读取到可导入的数据'}), 400

    operator_id = session.get('user_id')
    import_time = datetime.now()
    imported_count = 0
    skipped_count = 0
    skipped_details = []
    max_detail = 15
    imported_form_ids = []

    for idx, row in enumerate(rows, start=1):
        prepared = _prepare_import_form_row(row, idx, operator_id, import_time)
        if not prepared['ok']:
            skipped_count += 1
            if len(skipped_details) < max_detail:
                skipped_details.append(f"第{idx}行：{'；'.join(prepared['errors'][:2])}")
            continue
        payload = prepared['payload']
        form = LectureForm(**payload)
        try:
            db.session.add(form)
            db.session.flush()
            imported_updated_at = payload.get('updated_at') or import_time
            LectureForm.query.filter_by(id=form.id).update({
                'unique_id': form.id,
                'updated_at': imported_updated_at
            }, synchronize_session=False)
            db.session.commit()
            imported_count += 1
            imported_form_ids.append(form.id)
        except Exception as e:
            db.session.rollback()
            skipped_count += 1
            if len(skipped_details) < max_detail:
                skipped_details.append(f'第{idx}行：写入失败 {str(e)}')

    export_url = None
    if imported_form_ids:
        export_url = url_for('admin.export_imported_forms_excel', form_ids=','.join(str(fid) for fid in imported_form_ids))

    return jsonify({
        'success': True,
        'imported_count': imported_count,
        'skipped_count': skipped_count,
        'skipped_details': skipped_details,
        'imported_form_ids': imported_form_ids,
        'export_url': export_url
    })


@admin_bp.route('/api/forms/import/preview', methods=['POST'])
@role_required('超级管理员')
def preview_forms_from_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请先选择预览文件'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'success': False, 'message': '请先选择预览文件'}), 400
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'message': '仅支持xls或xlsx文件'}), 400

    rows, err = _parse_form_import_rows(file)
    if err:
        return jsonify({'success': False, 'message': err}), 400
    if not rows:
        return jsonify({'success': False, 'message': '未读取到可预览的数据'}), 400

    operator_id = session.get('user_id')
    import_time = datetime.now()
    preview_rows = []
    valid_rows = 0
    invalid_rows = 0
    for idx, row in enumerate(rows, start=1):
        prepared = _prepare_import_form_row(row, idx, operator_id, import_time)
        if prepared['ok']:
            valid_rows += 1
            preview_rows.append({
                'row_no': idx,
                'valid': True,
                'listener_number': prepared.get('listener_number', ''),
                'teacher_name': prepared.get('teacher_name', ''),
                'course_title': prepared.get('course_title', ''),
                'status': prepared.get('status', '部门已审核'),
                'errors': []
            })
        else:
            invalid_rows += 1
            preview_rows.append({
                'row_no': idx,
                'valid': False,
                'listener_number': prepared.get('listener_number', ''),
                'teacher_name': prepared.get('teacher_name', ''),
                'course_title': prepared.get('course_title', ''),
                'status': prepared.get('status', '部门已审核'),
                'errors': prepared.get('errors', [])
            })
    return jsonify({
        'success': True,
        'summary': {
            'total_rows': len(rows),
            'valid_rows': valid_rows,
            'invalid_rows': invalid_rows
        },
        'rows': preview_rows[:50]
    })


@admin_bp.route('/api/forms/import/export', methods=['GET'])
@role_required('超级管理员')
def export_imported_forms_excel():
    form_ids_raw = request.args.get('form_ids', '').strip()
    if not form_ids_raw:
        return jsonify({'success': False, 'message': '缺少form_ids参数'}), 400
    form_ids = []
    for part in form_ids_raw.split(','):
        part = part.strip()
        if not part:
            continue
        try:
            form_ids.append(int(part))
        except Exception:
            continue
    if not form_ids:
        return jsonify({'success': False, 'message': 'form_ids参数无效'}), 400
    forms = LectureForm.query.filter(LectureForm.id.in_(form_ids)).order_by(LectureForm.id.asc()).all()
    if not forms:
        return jsonify({'success': False, 'message': '未找到可导出的表单'}), 404
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '导入结果明细'
    headers = [
        '表单ID', '表单组ID', '听课人编号', '听课人姓名+学院', '课程信息变化', '听课时间', '第几节', '听课地点',
        '授课教师', '教师所属学院', '课程名称', '专业年级', '异常情况反映', '主要教学方法', '管理课堂纪律', '调动课堂气氛',
        '课件制作质量', '整体教学效果', '优质案例推荐', '课程反馈', '不足及建议', '听课班级同学签名1', '联系电话1',
        '听课班级同学签名2', '联系电话2', '状态', '审核标签', '审核人ID', '审核时间', '审核意见',
        '登记ID', '创建时间', '更新时间'
    ]
    ws.append(headers)
    for form in forms:
        ws.append([
            form.id,
            form.unique_id,
            form.listener_number or '',
            form.listener_name or '',
            form.course_changes or '',
            form.lecture_date or '',
            form.class_period or '',
            form.lecture_location or '',
            form.teacher_name or '',
            form.teacher_college or '',
            form.course_title or '',
            form.student_grade_class or '',
            form.abnormal_situation or '',
            form.teaching_method or '',
            form.classroom_discipline or '',
            form.classroom_atmosphere or '',
            form.courseware_quality or '',
            form.overall_effect or '',
            form.quality_case or '',
            form.course_feedback or '',
            form.suggestions or '',
            form.student_signature1 or '',
            form.contact_phone1 or '',
            form.student_signature2 or '',
            form.contact_phone2 or '',
            form.status or '',
            form.audit_tag or '',
            form.reviewer_id if form.reviewer_id is not None else '',
            form.review_time.strftime('%Y-%m-%d %H:%M:%S') if form.review_time else '',
            form.review_comment or '',
            form.registration_id if form.registration_id is not None else '',
            form.created_at.strftime('%Y-%m-%d %H:%M:%S') if form.created_at else '',
            form.updated_at.strftime('%Y-%m-%d %H:%M:%S') if form.updated_at else ''
        ])
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
    filename = f'导入表单明细_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return workbook_response(wb, filename)


@admin_bp.route('/export_forms')
@role_required('超级管理员')
def export_forms():
    """听课记录导出（框架性实现）"""
    try:
        # TODO: 实现真正的听课记录导出功能
        # 这里暂时返回一个提示信息，避免路由错误
        flash('听课记录导出功能正在开发中，敬请期待！', 'info')
        return redirect(url_for('admin.system_management'))
    except Exception as e:
        flash(f'导出功能暂不可用：{str(e)}', 'error')
        return redirect(url_for('admin.system_management'))


def _parse_export_status_list(status_list=None):
    if not status_list:
        return []
    if isinstance(status_list, str):
        status_list = status_list.split(',')
    return [str(s).strip() for s in status_list if str(s).strip()]


def _parse_export_form_ids(raw_form_ids=None):
    if not raw_form_ids:
        return []
    if isinstance(raw_form_ids, str):
        raw_items = raw_form_ids.split(',')
    else:
        raw_items = raw_form_ids
    form_ids = []
    for item in raw_items:
        try:
            form_id = int(item)
        except (TypeError, ValueError):
            continue
        if form_id > 0:
            form_ids.append(form_id)
    return list(dict.fromkeys(form_ids))


def get_export_query(scope, status_list=None, form_ids=None):
    """构建导出查询对象"""
    query = LectureForm.query
    scope = scope if scope in ['latest', 'selected', 'all'] else 'latest'

    # 1. 范围筛选
    if scope == 'latest':
        last_id = int(SystemSetting.get('last_exported_form_id', 0) or 0)
        query = query.filter(LectureForm.id > last_id)
    elif scope == 'selected':
        form_ids = _parse_export_form_ids(form_ids)
        query = query.filter(LectureForm.id.in_(form_ids))
    # scope == 'all' 不做额外ID/时间筛选

    # 2. 状态筛选
    status_list = _parse_export_status_list(status_list)
    if status_list:
        query = query.filter(LectureForm.status.in_(status_list))

    return query


def _format_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return value


def _sort_export_forms(forms, sort_by):
    if sort_by == 'course':
        forms.sort(key=lambda x: (x.teacher_college or '', x.teacher_name or '', x.course_title or ''))
    else:
        forms.sort(key=lambda x: x.lecture_date or '', reverse=True)


def _build_standard_forms_workbook(forms):
    # 生成特定规范 Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "听课反馈表"

    # 创建第二个空Sheet
    wb.create_sheet("其他信息")

    # === 样式定义 ===
    font_title = Font(name='宋体', size=20, bold=True)
    font_header = Font(name='宋体', size=12, bold=True)
    font_data = Font(name='宋体', size=12)

    thin_border = Side(style='thin')
    thick_border = Side(style='medium')
    border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    ws.merge_cells('A1:L1')
    ws['A1'] = "西南大学学生教学质量监控信息（意见）处理笺"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_center
    ws.row_dimensions[1].height = 48

    headers = [
        "序号", "听课人", "听课时间", "第几节", "听课地点",
        "授课教师", "教师所属学院", "课程", "专业年级",
        "异常情况反映", "课程反馈", "不足及建议"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_all

    ws.row_dimensions[2].height = 25

    max_id = 0
    for row_idx, form in enumerate(forms, 3):
        if form.id > max_id:
            max_id = form.id

        row_data = [
            row_idx - 2,
            form.listener_number,
            form.lecture_date,
            form.class_period,
            form.lecture_location,
            form.teacher_name,
            form.teacher_college,
            form.course_title,
            form.student_grade_class,
            form.abnormal_situation,
            form.course_feedback,
            form.suggestions
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else '')
            cell.font = font_data
            cell.alignment = align_left
            cell.border = border_all

        ws.row_dimensions[row_idx].height = 15

    for col_idx in range(1, len(headers) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 15

    max_row = len(forms) + 2

    for col in range(1, 13):
        cell = ws.cell(row=1, column=col)
        current_border = cell.border
        cell.border = Border(
            top=thick_border,
            bottom=current_border.bottom,
            left=current_border.left,
            right=current_border.right
        )

    for col in range(1, 13):
        cell = ws.cell(row=max_row, column=col)
        current_border = cell.border
        cell.border = Border(
            top=current_border.top,
            bottom=thick_border,
            left=current_border.left,
            right=current_border.right
        )

    for row in range(1, max_row + 1):
        cell = ws.cell(row=row, column=1)
        current_border = cell.border
        cell.border = Border(
            top=current_border.top,
            bottom=current_border.bottom,
            left=thick_border,
            right=current_border.right
        )

    for row in range(1, max_row + 1):
        cell = ws.cell(row=row, column=12)
        current_border = cell.border
        cell.border = Border(
            top=current_border.top,
            bottom=current_border.bottom,
            left=current_border.left,
            right=thick_border
        )

    ws['A1'].border = Border(top=thick_border, left=thick_border, right=thick_border, bottom=thin_border)
    ws.cell(row=1, column=12).border = Border(top=thick_border, right=thick_border, bottom=thin_border)

    return wb, max_id


def _build_all_data_forms_workbook(forms):
    # 导出 lecture_forms 表中所有字段
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表单数据库数据"

    columns = [
        ('id', 'ID'),
        ('listener_name', '听课人姓名+学院'),
        ('listener_number', '听课人编号'),
        ('course_changes', '课程信息变化'),
        ('lecture_date', '听课时间'),
        ('class_period', '第几节'),
        ('lecture_location', '听课地点'),
        ('teacher_name', '授课教师'),
        ('teacher_college', '教师所属学院'),
        ('course_title', '课程总标题'),
        ('student_grade_class', '专业年级'),
        ('abnormal_situation', '异常情况反映'),
        ('teaching_method', '主要教学方法'),
        ('classroom_discipline', '管理课堂纪律'),
        ('classroom_atmosphere', '调动课堂气氛'),
        ('courseware_quality', '课件制作质量'),
        ('overall_effect', '整体教学效果'),
        ('quality_case', '优质案例推荐'),
        ('course_feedback', '课程反馈'),
        ('suggestions', '不足及建议'),
        ('student_signature1', '听课班级同学签名1'),
        ('contact_phone1', '联系电话1'),
        ('student_signature2', '听课班级同学签名2'),
        ('contact_phone2', '联系电话2'),
        ('status', '审核状态'),
        ('reviewer_id', '审核人ID'),
        ('review_time', '审核时间'),
        ('review_comment', '审核意见'),
        ('registration_id', '关联登记记录ID'),
        ('audit_tag', '审核标签'),
        ('unique_id', '唯一标志ID'),
        ('created_at', '创建时间'),
        ('updated_at', '更新时间'),
    ]

    header_font = Font(name='微软雅黑', size=11, bold=True)
    data_font = Font(name='微软雅黑', size=11)
    header_fill = PatternFill(fill_type='solid', fgColor='D8D8D8')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, (_, header) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    max_id = 0
    for row_idx, form in enumerate(forms, 2):
        if form.id > max_id:
            max_id = form.id
        for col_idx, (attr, _) in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_format_excel_value(getattr(form, attr, None)))
            cell.font = data_font
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.border = border

    return wb, max_id


@admin_bp.route('/api/forms/export/check', methods=['POST'])
@role_required('超级管理员')
def check_export_forms():
    """检查符合导出条件的表单数量"""
    try:
        data = request.get_json() or {}
        scope = data.get('scope', 'latest')
        status_list = data.get('status_list', []) # 默认为空列表，或者前端应该传所有选中的状态
        form_ids = data.get('form_ids', [])
        
        # 兼容旧代码：如果传了 status_approved
        if 'status_approved' in data and not status_list:
            if data['status_approved']:
                status_list = ['已审核', '部门已审核', '中心已审核']
            else:
                status_list = [] # 不筛选，即所有状态
        
        # 获取符合条件的查询
        query = get_export_query(scope, status_list, form_ids)
        total_count = query.count()
        
        # 计算被排除的数量
        excluded_count = 0
        if status_list:
            # 构建一个不带状态筛选的查询
            all_status_query = get_export_query(scope, None, form_ids)
            full_count = all_status_query.count()
            excluded_count = full_count - total_count
            
        return jsonify({
            'success': True,
            'total_count': total_count,
            'excluded_count': excluded_count
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@admin_bp.route('/api/forms/export/download', methods=['GET'])
@role_required('超级管理员')
def download_export_forms():
    """执行导出并下载Excel"""
    try:
        # 获取参数
        scope = request.args.get('scope', 'latest')
        status_list = request.args.get('status_list') # 逗号分隔的字符串
        form_ids = request.args.get('form_ids')
        export_mode = request.args.get('export_mode', 'standard')
        
        # 兼容旧代码
        status_approved = request.args.get('status_approved')
        if status_approved and not status_list:
             if str(status_approved).lower() == 'true':
                 status_list = '已审核,部门已审核,中心已审核'
        
        sort_by = request.args.get('sort_by', 'time')
        filename_title = (request.args.get('filename_title') or '').strip()
            
        # 1. 获取数据
        query = get_export_query(scope, status_list, form_ids)
        forms = query.all()
        
        if not forms:
            flash('没有符合条件的记录可导出', 'warning')
            return redirect(url_for('admin.view_forms'))
            
        # 2. 排序
        _sort_export_forms(forms, sort_by)
            
        # 3. 生成Excel
        if export_mode == 'all_data':
            wb, max_id = _build_all_data_forms_workbook(forms)
        else:
            export_mode = 'standard'
            wb, max_id = _build_standard_forms_workbook(forms)
        
        # 4. 更新SystemSetting (如果是latest模式)
        if scope == 'latest' and max_id > 0:
            SystemSetting.set('last_exported_form_id', str(max_id))
            
        # 5. 返回文件
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        date_str = datetime.now().strftime('%Y%m%d')
        if export_mode == 'all_data':
            filename = f"听课反馈表-{date_str}.xlsx"
        elif filename_title:
            filename = f"学生教学信息中心听课反馈表汇总-{filename_title} {date_str}.xlsx"
        else:
            filename = f"学生教学信息中心听课反馈表汇总-未命名 {date_str}.xlsx"
        
        # 处理中文文件名
        try:
            filename_encoded = filename.encode('latin-1').decode('latin-1')
        except UnicodeEncodeError:
            # 如果包含非latin-1字符（如中文），quote它
            from urllib.parse import quote
            filename_encoded = quote(filename)
            
        rv = send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        
        # 设置Content-Disposition header以支持中文文件名
        rv.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{filename_encoded}"
        
        return rv
        
    except Exception as e:
        flash(f'导出失败: {str(e)}', 'error')
        return redirect(url_for('admin.view_forms'))


