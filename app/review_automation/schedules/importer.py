"""Safe, staged imports for synthetic and administrator-supplied schedules."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
import uuid
from pathlib import Path
from typing import Iterable, Mapping

from flask import current_app
from openpyxl import load_workbook

from app.models import db
from app.review_automation.contracts import DatasetStatus
from app.review_automation.models import (
    ListenerClassMapping,
    PersonalScheduleSlot,
    ScheduleDataset,
    ScheduleImportIssue,
    SchoolScheduleEntry,
)

from .aliases import (
    CURRENT_SCHOOL_SCHEDULE_ALIASES,
    HISTORICAL_SCHOOL_SCHEDULE_ALIASES,
    LISTENER_CLASS_MAPPING_ALIASES,
    PERSONAL_SCHEDULE_ALIASES,
)
from .normalization import (
    NormalizationError,
    normalize_class_name,
    normalize_college,
    normalize_identifier,
    normalize_name,
    normalize_weekday,
    parse_period_range,
    parse_weeks,
    resolve_columns,
)


SCHOOL_KIND = 'school'
CLASS_MAPPING_KIND = 'class_mapping'
PERSONAL_KIND = 'personal'

_KIND_ALIASES = {
    SCHOOL_KIND: SCHOOL_KIND,
    'school_schedule': SCHOOL_KIND,
    'school-schedule': SCHOOL_KIND,
    '全校课表': SCHOOL_KIND,
    CLASS_MAPPING_KIND: CLASS_MAPPING_KIND,
    'listener_class': CLASS_MAPPING_KIND,
    'listener_class_mapping': CLASS_MAPPING_KIND,
    'admin_class': CLASS_MAPPING_KIND,
    '行政班映射': CLASS_MAPPING_KIND,
    PERSONAL_KIND: PERSONAL_KIND,
    'personal_schedule': PERSONAL_KIND,
    'personal-schedule': PERSONAL_KIND,
    '个人课表': PERSONAL_KIND,
}

_SCHOOL_ALIASES = {
    field: tuple(dict.fromkeys(
        tuple(HISTORICAL_SCHOOL_SCHEDULE_ALIASES.get(field, ()))
        + tuple(CURRENT_SCHOOL_SCHEDULE_ALIASES.get(field, ()))
    ))
    for field in set(HISTORICAL_SCHOOL_SCHEDULE_ALIASES)
    | set(CURRENT_SCHOOL_SCHEDULE_ALIASES)
}
_SCHOOL_REQUIRED = {
    'course_title', 'teacher_name', 'teacher_college', 'teaching_class',
    'weeks', 'weekday', 'periods', 'location',
}
_CLASS_REQUIRED = {'admin_class'}
_PERSONAL_REQUIRED = {'course_title', 'semester', 'weeks', 'weekday', 'periods'}
_IDENTITY_FIELDS = ('listener_number', 'student_id')


class ScheduleImportError(ValueError):
    """A safe import failure with a stable error code."""

    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(message)


def normalize_dataset_kind(kind: str) -> str:
    key = str(kind or '').strip().lower()
    try:
        return _KIND_ALIASES[key]
    except KeyError as exc:
        raise ScheduleImportError('unsupported_kind', f'unsupported dataset kind: {kind}') from exc


def _read_stream(stream) -> bytes:
    if isinstance(stream, bytes):
        return stream
    if isinstance(stream, bytearray):
        return bytes(stream)
    if not hasattr(stream, 'read'):
        raise ScheduleImportError('invalid_stream', 'stream must be a readable binary object')
    data = stream.read()
    if not isinstance(data, (bytes, bytearray)):
        raise ScheduleImportError('invalid_stream', 'stream did not return binary data')
    return bytes(data)


def _safe_filename(filename: str) -> tuple[str, str]:
    original = Path(str(filename or '')).name
    suffix = Path(original).suffix.lower()
    if not original or original in {'.', '..'}:
        raise ScheduleImportError('invalid_filename', 'filename is required')
    if suffix != '.xlsx':
        raise ScheduleImportError('unsupported_workbook', 'only .xlsx workbooks are accepted')
    return original, suffix


def _storage_path(suffix: str, data: bytes) -> tuple[str, Path]:
    root = Path(current_app.config['AUTOMATION_UPLOAD_DIR']).resolve()
    root.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(data).hexdigest()
    for _ in range(5):
        stored_name = f'{uuid.uuid4().hex}{suffix}'
        candidate = root / stored_name
        if candidate.parent != root:
            raise ScheduleImportError('unsafe_storage_path', 'generated storage path escaped upload root')
        try:
            with candidate.open('xb') as handle:
                handle.write(data)
            return digest, candidate
        except FileExistsError:
            continue
    raise ScheduleImportError('storage_collision', 'could not allocate generated storage name')


def _header_positions(
    headers: Iterable[object],
    aliases: Mapping[str, Iterable[str] | str],
    required: set[str],
    identity_required: bool = False,
) -> dict[str, int]:
    headers = list(headers)
    resolved = resolve_columns(headers, aliases, required)
    positions = {field: headers.index(header) for field, header in resolved.items()}
    if identity_required:
        identity_positions = {}
        for field in _IDENTITY_FIELDS:
            try:
                header = resolve_columns(headers, aliases, {field})[field]
            except NormalizationError as exc:
                if exc.code == 'missing_required_column':
                    continue
                raise
            identity_positions[field] = headers.index(header)
        if not identity_positions:
            raise NormalizationError('columns', 'missing_identity_column', value='identity')
        positions.update(identity_positions)
    return positions


def _find_sheet_and_header(workbook, kind: str):
    if kind == SCHOOL_KIND:
        aliases, required, identity = _SCHOOL_ALIASES, _SCHOOL_REQUIRED, False
    elif kind == CLASS_MAPPING_KIND:
        aliases, required, identity = LISTENER_CLASS_MAPPING_ALIASES, _CLASS_REQUIRED, True
    else:
        aliases, required, identity = PERSONAL_SCHEDULE_ALIASES, _PERSONAL_REQUIRED, True

    candidates = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        for row_number, row in enumerate(rows, start=1):
            try:
                positions = _header_positions(row, aliases, required, identity)
            except NormalizationError:
                continue
            candidates.append((len(positions), sheet.title, row_number, positions, rows))
            break
    if not candidates:
        raise ScheduleImportError('missing_header', 'no recognizable schedule header was found')
    _, sheet_name, header_row, positions, rows = max(candidates, key=lambda item: item[0])
    return sheet_name, header_row, positions, rows


def _cell(row, positions: Mapping[str, int], field: str):
    index = positions.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _optional(value, normalizer):
    return None if _blank(value) else normalizer(value)


def _optional_identifier(value):
    return _optional(value, normalize_identifier)


def _weeks_json(weeks) -> str:
    return json.dumps(sorted(weeks), ensure_ascii=False, separators=(',', ':'))


def _dedupe_key(values: Mapping[str, object]) -> tuple:
    return tuple(values[key] for key in sorted(values))


def _redacted_excerpt(field: str, value) -> str:
    if field in {'listener_number', 'student_id', 'identity'}:
        return f'{field}=[redacted]'
    text = str(value or '')
    text = re.sub(r'\d{4,}', '***', text)
    return f'{field}={text[:80]}'


def _add_issue(dataset: ScheduleDataset, row_number: int, exc, field: str | None = None, value=None):
    if isinstance(exc, NormalizationError):
        code = exc.code
        message = str(exc)
        issue_field = exc.field
        issue_value = exc.value
    elif isinstance(exc, ScheduleImportError):
        code = exc.code
        message = str(exc)
        issue_field = field or 'row'
        issue_value = value
    else:
        code = 'invalid_row'
        message = 'row could not be normalized'
        issue_field = field or 'row'
        issue_value = value
    db.session.add(ScheduleImportIssue(
        dataset_id=dataset.id,
        row_number=row_number,
        code=code,
        message=message,
        value_excerpt=_redacted_excerpt(issue_field, issue_value),
    ))


def _school_values(row, positions):
    weeks = parse_weeks(_cell(row, positions, 'weeks'))
    start, end = parse_period_range(_cell(row, positions, 'periods'))
    return {
        'teacher_name': normalize_name(_cell(row, positions, 'teacher_name')),
        'teacher_college': normalize_college(_cell(row, positions, 'teacher_college')),
        'course_title': normalize_name(_cell(row, positions, 'course_title')),
        'teaching_class': normalize_class_name(_cell(row, positions, 'teaching_class')),
        'major': _optional(_cell(row, positions, 'major'), normalize_name),
        'weeks_json': _weeks_json(weeks),
        'weekday': normalize_weekday(_cell(row, positions, 'weekday')),
        'start_period': start,
        'end_period': end,
        'location': normalize_name(_cell(row, positions, 'location')),
    }


def _mapping_values(row, positions):
    listener_number = _optional_identifier(_cell(row, positions, 'listener_number'))
    student_id = _optional_identifier(_cell(row, positions, 'student_id'))
    if not listener_number and not student_id:
        raise NormalizationError(
            'identity', 'missing_identity', value=None, message='listener identity is required',
        )
    return {
        'listener_number': listener_number,
        'student_id': student_id,
        'admin_class': normalize_class_name(_cell(row, positions, 'admin_class')),
    }


def _semester_text(value):
    if _blank(value):
        raise NormalizationError(
            'semester', 'missing_semester', value=value, message='academic semester is required',
        )
    return ' '.join(unicodedata.normalize('NFKC', str(value)).split())


def _personal_values(row, positions, expected_semester):
    source_semester = _semester_text(_cell(row, positions, 'semester'))
    if source_semester != _semester_text(expected_semester):
        raise NormalizationError(
            'semester', 'semester_mismatch', value=source_semester,
            message='row semester does not match dataset semester',
        )
    listener_number = _optional_identifier(_cell(row, positions, 'listener_number'))
    student_id = _optional_identifier(_cell(row, positions, 'student_id'))
    if not listener_number and not student_id:
        raise NormalizationError(
            'identity', 'missing_identity', value=None, message='listener identity is required',
        )
    weeks = parse_weeks(_cell(row, positions, 'weeks'))
    start, end = parse_period_range(_cell(row, positions, 'periods'))
    return {
        'listener_number': listener_number,
        'student_id': student_id,
        'course_title': normalize_name(_cell(row, positions, 'course_title')),
        'weeks_json': _weeks_json(weeks),
        'weekday': normalize_weekday(_cell(row, positions, 'weekday')),
        'start_period': start,
        'end_period': end,
    }


def _import_rows(
    dataset: ScheduleDataset,
    kind: str,
    rows,
    header_row: int,
    positions,
    expected_semester=None,
):
    seen = set()
    valid_rows = 0
    duplicate_rows = 0
    source_rows = 0
    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        if all(_blank(_cell(row, positions, field)) for field in positions):
            continue
        source_rows += 1
        try:
            if kind == SCHOOL_KIND:
                values = _school_values(row, positions)
                key = _dedupe_key(values)
                if key in seen:
                    duplicate_rows += 1
                    continue
                seen.add(key)
                db.session.add(SchoolScheduleEntry(
                    dataset_id=dataset.id,
                    source_row=row_number,
                    **values,
                ))
            elif kind == CLASS_MAPPING_KIND:
                values = _mapping_values(row, positions)
                key = _dedupe_key(values)
                if key in seen:
                    duplicate_rows += 1
                    continue
                seen.add(key)
                db.session.add(ListenerClassMapping(
                    dataset_id=dataset.id,
                    source_row=row_number,
                    match_status='matched',
                    **values,
                ))
            else:
                values = _personal_values(row, positions, expected_semester)
                key = _dedupe_key(values)
                if key in seen:
                    duplicate_rows += 1
                    continue
                seen.add(key)
                db.session.add(PersonalScheduleSlot(
                    dataset_id=dataset.id,
                    source_row=row_number,
                    **values,
                ))
            valid_rows += 1
        except (NormalizationError, ScheduleImportError) as exc:
            _add_issue(dataset, row_number, exc)

    return valid_rows, source_rows, duplicate_rows


def preview_dataset(kind, semester, stream, filename, actor_id):
    """Persist a validated dataset as staged; never changes active data."""
    kind = normalize_dataset_kind(kind)
    original_filename, suffix = _safe_filename(filename)
    data = _read_stream(stream)
    if not data:
        raise ScheduleImportError('empty_workbook', 'workbook is empty')
    digest, stored_path = _storage_path(suffix, data)
    dataset = ScheduleDataset(
        kind=kind,
        semester=str(semester),
        sha256=digest,
        original_filename=original_filename,
        status=DatasetStatus.STAGED.value,
        created_by=actor_id,
        summary_json='{}',
    )
    db.session.add(dataset)
    db.session.flush()

    summary = {
        'kind': kind,
        'storage_filename': stored_path.name,
        'source_rows': 0,
        'duplicate_rows': 0,
        'selected_fields': [],
        'macro_execution': False,
        'external_links_loaded': False,
    }
    workbook = None
    try:
        workbook = load_workbook(
            filename=__import__('io').BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
            keep_vba=False,
        )
        sheet_name, header_row, positions, rows = _find_sheet_and_header(workbook, kind)
        valid_rows, source_rows, duplicate_rows = _import_rows(
            dataset, kind, rows, header_row, positions, expected_semester=dataset.semester,
        )
        summary.update({
            'sheet_name': sheet_name,
            'header_row': header_row,
            'selected_fields': sorted(positions),
            'selected_column_count': len(positions),
            'source_rows': source_rows,
            'duplicate_rows': duplicate_rows,
        })
        dataset.row_count = valid_rows
        dataset.error_count = ScheduleImportIssue.query.filter_by(dataset_id=dataset.id).count()
        dataset.status = DatasetStatus.STAGED.value
    except (ScheduleImportError, NormalizationError) as exc:
        dataset.status = DatasetStatus.FAILED.value
        dataset.error_count = 1
        _add_issue(dataset, 0, exc)
        summary.update({'status_error': getattr(exc, 'code', 'invalid_dataset')})
    except Exception as exc:
        dataset.status = DatasetStatus.FAILED.value
        dataset.error_count = 1
        _add_issue(dataset, 0, ScheduleImportError('unreadable_workbook', 'workbook could not be read'))
        summary.update({'status_error': 'unreadable_workbook'})
    finally:
        if workbook is not None:
            workbook.close()

    dataset.summary_json = json.dumps(summary, ensure_ascii=False, sort_keys=True)
    db.session.commit()
    return dataset


__all__ = [
    'CLASS_MAPPING_KIND',
    'PERSONAL_KIND',
    'SCHOOL_KIND',
    'ScheduleImportError',
    'normalize_dataset_kind',
    'preview_dataset',
]
