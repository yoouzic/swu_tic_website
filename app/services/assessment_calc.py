# -*- coding: utf-8 -*-
"""Assessment statistics calculation / query composition service.

This module contains the domain calculations that were originally embedded in
the admin ``assessment_stats`` blueprint.  It intentionally keeps HTTP
adaptation (request/session/jsonify/send_file) out; blueprints call these
functions and shape the responses.
"""
from datetime import datetime, timedelta
import json

import openpyxl
from openpyxl.styles import Font

from app.models import LectureForm, SystemSetting, ScoreRecord, ScoreItem, AssessmentOverride, db
from app.services.assessment_scope import resolve_selected_departments
from app.services.excel_utils import autosize_worksheet
from app.services.form_week_semantics import effective_form_week
from app.services.review_form_queries import (
    build_review_form_filter_datetime,
    latest_form_groups_for_users,
)
from app.services.teaching_calendar import TeachingCalendarConfig, teaching_term_start, teaching_week_number
from app.services.teaching_calendar_settings import (
    INVALID_FIRST_WEEK,
    MISSING_FIRST_WEEK,
    load_teaching_calendar_config,
)
from app.utils.leave_management import ASSESSMENT_EXEMPT_OVERRIDE_TYPES


def parse_assessment_range(start_date_str, end_date_str):
    if not start_date_str or not end_date_str:
        return None, None, '请先选择开始和结束时间'
    try:
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d') + timedelta(days=1)
        return start_date, end_date, None
    except Exception:
        return None, None, '时间格式错误，请使用YYYY-MM-DD'


def latest_forms_in_range(listener_numbers, start_date, end_date):
    if not listener_numbers:
        return []
    forms = LectureForm.query.filter(
        LectureForm.listener_number.in_(listener_numbers),
        LectureForm.created_at >= start_date,
        LectureForm.created_at < end_date
    ).order_by(LectureForm.created_at.asc(), LectureForm.id.asc()).all()
    latest_map = {}
    for form in forms:
        uid = form.unique_id or form.id
        latest_map[uid] = form
    return list(latest_map.values())


def assessment_form_latest_timestamp(form):
    if not form:
        return None
    return form.updated_at or form.created_at


def assessment_latest_form_groups_for_users(listener_numbers):
    if not listener_numbers:
        return []
    forms = LectureForm.query.filter(
        LectureForm.listener_number.in_(listener_numbers)
    ).order_by(LectureForm.unique_id.asc(), LectureForm.id.asc()).all()
    group_map = {}
    for form in forms:
        uid = form.unique_id or form.id
        group_map.setdefault(uid, []).append(form)

    groups = []
    for uid, form_list in group_map.items():
        sorted_forms = sorted(
            form_list,
            key=lambda f: (assessment_form_latest_timestamp(f) or datetime.min, f.id),
            reverse=True
        )
        latest_form = sorted_forms[0]
        groups.append({
            'unique_id': uid,
            'latest_form': latest_form,
            'latest_timestamp': assessment_form_latest_timestamp(latest_form),
            'forms': sorted_forms
        })

    groups.sort(
        key=lambda group: (
            group['latest_timestamp'] or datetime.min,
            group['latest_form'].id if group['latest_form'] else 0
        ),
        reverse=True
    )
    return groups


def assessment_latest_forms_in_range(listener_numbers, start_date, end_date):
    if not listener_numbers:
        return []
    latest_forms = []
    for group_data in assessment_latest_form_groups_for_users(listener_numbers):
        latest_form = group_data['latest_form']
        latest_timestamp = group_data['latest_timestamp']
        if not latest_form or not latest_timestamp:
            continue
        if start_date <= latest_timestamp < end_date:
            latest_forms.append(latest_form)
    return latest_forms


def get_teaching_reward_settings():
    config, error = load_teaching_calendar_config()
    if error == MISSING_FIRST_WEEK:
        return None, '请先在制度设置中配置第一周起始日期'
    if error == INVALID_FIRST_WEEK:
        return None, '制度设置中的第一周起始日期格式错误'
    if config is None:
        return None, '请先在制度设置中配置第一周起始日期'
    try:
        required_submission = int(SystemSetting.get('teaching_required_submission', '1') or 1)
    except Exception:
        required_submission = 1
    if required_submission < 0:
        required_submission = 0
    return {
        'first_week_date': config.first_week_date,
        'week_start_day': config.week_start_day,
        'required_submission': required_submission,
        'total_weeks': config.total_weeks
    }, None


def get_teaching_week_no(date_obj, first_week_date, week_start_day, total_weeks=None):
    if not date_obj or not first_week_date:
        return None
    config = TeachingCalendarConfig(
        first_week_date=first_week_date,
        week_start_day=week_start_day,
        total_weeks=total_weeks if total_weeks is not None else 52,
    )
    return teaching_week_number(date_obj, config)


def compute_teaching_week_window(start_date, end_date, first_week_date, week_start_day, total_weeks=None):
    range_start = start_date.date()
    range_end = (end_date - timedelta(days=1)).date()
    if range_end < range_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    teaching_start = first_week_date - timedelta(days=(first_week_date.weekday() - week_start_day) % 7)
    effective_start = max(range_start, teaching_start)
    teaching_end_exclusive = teaching_start + timedelta(days=total_weeks * 7) if total_weeks is not None else None
    if teaching_end_exclusive is not None:
        effective_end = min(range_end, teaching_end_exclusive - timedelta(days=1))
    else:
        effective_end = range_end
    if effective_end < effective_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    first_full_start = effective_start + timedelta(days=(week_start_day - effective_start.weekday()) % 7)
    last_candidate_start = effective_end - timedelta(days=6)
    if last_candidate_start < first_full_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    last_full_start = last_candidate_start - timedelta(days=(last_candidate_start.weekday() - week_start_day) % 7)
    if last_full_start < first_full_start:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    start_week = ((first_full_start - teaching_start).days // 7) + 1
    end_week = ((last_full_start - teaching_start).days // 7) + 1
    if end_week < 1:
        return {
            'has_full_weeks': False,
            'start_week': None,
            'end_week': None,
            'week_count': 0,
            'label': '无完整教学周',
            'window_start': None,
            'window_end': None
        }
    if start_week < 1:
        start_week = 1
    week_count = ((last_full_start - first_full_start).days // 7) + 1
    label = str(start_week) if start_week == end_week else f'{start_week}-{end_week}'
    window_start = datetime.combine(first_full_start, datetime.min.time())
    window_end = datetime.combine(last_full_start + timedelta(days=7), datetime.min.time())
    return {
        'has_full_weeks': True,
        'start_week': start_week,
        'end_week': end_week,
        'week_count': week_count,
        'label': label,
        'window_start': window_start,
        'window_end': window_end
    }


def get_teaching_term_start(first_week_date, week_start_day):
    return teaching_term_start(TeachingCalendarConfig(
        first_week_date=first_week_date,
        week_start_day=week_start_day,
        total_weeks=52,
    ))


def get_form_effective_week_no(form, reward_settings):
    if not form:
        return None
    config = TeachingCalendarConfig(
        first_week_date=reward_settings['first_week_date'],
        week_start_day=reward_settings['week_start_day'],
        total_weeks=reward_settings.get('total_weeks', 52),
    )
    return effective_form_week(form.lecture_date, form.audit_tag, config)


def build_teaching_month_templates(reward_window, reward_settings):
    if not reward_window.get('has_full_weeks'):
        return []
    start_week = reward_window.get('start_week')
    end_week = reward_window.get('end_week')
    if start_week is None or end_week is None:
        return []
    teaching_start = get_teaching_term_start(
        reward_settings['first_week_date'],
        reward_settings['week_start_day']
    )

    # 尝试加载自定义教学月设置
    custom_months = load_custom_month_definitions()

    if custom_months:
        # 使用自定义教学月设置
        month_templates = []
        for idx, m_def in enumerate(custom_months):
            m_start = m_def['start_week']
            m_end = m_def['end_week']
            # 仅包含完整落在选择范围内的教学月
            if m_start < start_week or m_end > end_week:
                continue
            month_no = idx + 1
            month_start_date = teaching_start + timedelta(days=(m_start - 1) * 7)
            month_end_date = teaching_start + timedelta(days=(m_end - 1) * 7 + 6)
            weeks = []
            for week_no in range(m_start, m_end + 1):
                week_start_date = teaching_start + timedelta(days=(week_no - 1) * 7)
                week_end_date = week_start_date + timedelta(days=6)
                weeks.append({
                    'week_no': week_no,
                    'week_label': f'第{week_no}周',
                    'start_date': week_start_date.strftime('%Y-%m-%d'),
                    'end_date': week_end_date.strftime('%Y-%m-%d'),
                })
            month_label = m_def.get('label') or f'第{month_no}教学月'
            month_templates.append({
                'month_no': month_no,
                'month_label': month_label,
                'week_range_label': f'第{m_start}-{m_end}周',
                'start_week': m_start,
                'end_week': m_end,
                'start_date': month_start_date.strftime('%Y-%m-%d'),
                'end_date': month_end_date.strftime('%Y-%m-%d'),
                'weeks': weeks,
            })
        return month_templates

    # 默认逻辑：每4周为一个教学月
    month_templates = []
    start_month = ((start_week - 1) // 4) + 1
    end_month = ((end_week - 1) // 4) + 1
    for month_no in range(start_month, end_month + 1):
        month_start_week = (month_no - 1) * 4 + 1
        month_end_week = month_start_week + 3
        if month_start_week < start_week or month_end_week > end_week:
            continue
        month_start_date = teaching_start + timedelta(days=(month_start_week - 1) * 7)
        month_end_date = month_start_date + timedelta(days=27)
        weeks = []
        for week_no in range(month_start_week, month_end_week + 1):
            week_start_date = teaching_start + timedelta(days=(week_no - 1) * 7)
            week_end_date = week_start_date + timedelta(days=6)
            weeks.append({
                'week_no': week_no,
                'week_label': f'第{week_no}周',
                'start_date': week_start_date.strftime('%Y-%m-%d'),
                'end_date': week_end_date.strftime('%Y-%m-%d'),
            })
        month_templates.append({
            'month_no': month_no,
            'month_label': f'第{month_no}教学月',
            'week_range_label': f'第{month_start_week}-{month_end_week}周',
            'start_week': month_start_week,
            'end_week': month_end_week,
            'start_date': month_start_date.strftime('%Y-%m-%d'),
            'end_date': month_end_date.strftime('%Y-%m-%d'),
            'weeks': weeks,
        })
    return month_templates


def load_custom_month_definitions():
    """加载自定义教学月设置，返回列表或 None"""
    raw = SystemSetting.get('teaching_month_definitions')
    if not raw:
        return None
    try:
        data = json.loads(raw)
        if not isinstance(data, list) or len(data) == 0:
            return None
        result = []
        for item in data:
            s = int(item.get('start_week'))
            e = int(item.get('end_week'))
            if s < 1 or e < 1 or s > e:
                continue
            result.append({
                'start_week': s,
                'end_week': e,
                'label': item.get('label', '').strip() or None
            })
        return result if result else None
    except (json.JSONDecodeError, TypeError, ValueError, KeyError):
        return None


def resolve_error_unique_ids(candidate_groups):
    form_to_unique = {}
    for group_data in candidate_groups:
        unique_id = group_data['unique_id']
        for form in group_data['forms']:
            form_to_unique[form.id] = unique_id
    if not form_to_unique:
        return set()
    error_form_ids = db.session.query(ScoreRecord.form_id).join(
        ScoreItem, ScoreItem.score_record_id == ScoreRecord.id
    ).filter(
        ScoreRecord.form_id.in_(list(form_to_unique.keys())),
        (ScoreItem.department_score > 0) | (ScoreItem.personal_score > 0)
    ).distinct().all()
    error_unique_ids = set()
    for row in error_form_ids:
        unique_id = form_to_unique.get(row.form_id)
        if unique_id is not None:
            error_unique_ids.add(unique_id)
    return error_unique_ids


def form_group_created_at(group_data):
    if not group_data:
        return None
    created_candidates = [
        form.created_at for form in (group_data.get('forms') or [])
        if getattr(form, 'created_at', None)
    ]
    if created_candidates:
        return min(created_candidates)
    latest_form = group_data.get('latest_form')
    return latest_form.created_at if latest_form and latest_form.created_at else None


def sort_submission_count_rows(rows, sort_by):
    if sort_by == 'reward':
        rows.sort(key=lambda x: (-x['reward_form_count'], -x['submission_group_count'], x['number'] or '', x['name'] or ''))
    else:
        rows.sort(key=lambda x: (-x['submission_group_count'], -x['reward_form_count'], x['number'] or '', x['name'] or ''))
    for idx, row in enumerate(rows, start=1):
        row['rank'] = idx
    return rows


def build_review_assessment_rows(users, start_date, end_date):
    by_number = {u.number: u for u in users if u.number}
    user_stats = {
        u.id: {
            'user_id': u.id,
            'name': u.name,
            'number': u.number,
            'department': u.department,
            'group': u.group,
            'linked_department_score': 0.0,
            'linked_personal_score': 0.0
        } for u in users
    }

    latest_forms = assessment_latest_forms_in_range(list(by_number.keys()), start_date, end_date)
    for form in latest_forms:
        user = by_number.get(form.listener_number)
        if not user or not form.score_record:
            continue
        user_stats[user.id]['linked_department_score'] += float(form.score_record.total_department_score or 0.0)
        user_stats[user.id]['linked_personal_score'] += float(form.score_record.total_personal_score or 0.0)

    rows = []
    for row in user_stats.values():
        row['total_department_score'] = row['linked_department_score']
        row['total_personal_score'] = row['linked_personal_score']
        row['total_score'] = row['total_personal_score']
        rows.append(row)

    rows.sort(key=lambda x: x['total_personal_score'], reverse=True)
    for idx, row in enumerate(rows, start=1):
        row['rank'] = idx
    return rows


def build_reward_week_templates(reward_window, required_submission):
    templates = []
    if not reward_window.get('has_full_weeks') or not reward_window.get('window_start'):
        return templates
    window_start = reward_window['window_start']
    start_week = reward_window['start_week'] or 1
    for index in range(reward_window.get('week_count') or 0):
        week_start = window_start + timedelta(days=7 * index)
        week_end = week_start + timedelta(days=6)
        week_no = start_week + index
        templates.append({
            'week_index': index,
            'week_no': week_no,
            'week_label': f'第{week_no}周',
            'start_date': week_start.strftime('%Y-%m-%d'),
            'end_date': week_end.strftime('%Y-%m-%d'),
            'required_submission': int(required_submission or 0),
            'effective_count': 0,
            'error_count': 0,
            'reward_form_count': 0
        })
    return templates


def get_reward_week_index(group_data, reward_window, reward_settings):
    if not reward_window.get('has_full_weeks'):
        return None
    latest_form = group_data.get('latest_form') if group_data else None
    if not latest_form:
        return None
    week_no = get_form_effective_week_no(latest_form, reward_settings)
    if week_no is None:
        return None
    start_week = reward_window.get('start_week')
    end_week = reward_window.get('end_week')
    if start_week is None or end_week is None or week_no < start_week or week_no > end_week:
        return None
    return week_no - start_week


def build_submission_count_rows(users, start_date, end_date, reward_settings, reward_window, time_filter_type='created'):
    by_number = {u.number: u for u in users if u.number}
    week_templates = build_reward_week_templates(reward_window, reward_settings['required_submission'])
    user_stats = {
        u.id: {
            'user_id': u.id,
            'name': u.name,
            'number': u.number,
            'department': u.department,
            'group': u.group,
            'submission_group_count': 0,
            'reward_effective_count': 0,
            'reward_error_count': 0,
            'reward_form_count': 0,
            'reward_week_details': [dict(template) for template in week_templates]
        } for u in users
    }

    groups = latest_form_groups_for_users(list(by_number.keys()))
    reward_candidates = []
    for group_data in groups:
        latest_form = group_data['latest_form']
        if not latest_form or not latest_form.created_at:
            continue
        user = by_number.get(latest_form.listener_number)
        if not user:
            continue
        filter_dt = build_review_form_filter_datetime(group_data, time_filter_type)
        if filter_dt and start_date <= filter_dt < end_date:
            user_stats[user.id]['submission_group_count'] += 1
        week_index = get_reward_week_index(group_data, reward_window, reward_settings)
        if week_index is None or latest_form.status != '中心已审核':
            continue
        reward_candidates.append({
            'group_data': group_data,
            'user_id': user.id,
            'week_index': week_index
        })

    error_unique_ids = resolve_error_unique_ids([item['group_data'] for item in reward_candidates])
    for item in reward_candidates:
        row = user_stats[item['user_id']]
        week_detail = row['reward_week_details'][item['week_index']]
        row['reward_effective_count'] += 1
        week_detail['effective_count'] += 1
        if item['group_data']['unique_id'] in error_unique_ids:
            row['reward_error_count'] += 1
            week_detail['error_count'] += 1

    for row in user_stats.values():
        total_reward_forms = 0
        for week_detail in row['reward_week_details']:
            week_reward = max(
                int(week_detail['effective_count']) - int(week_detail['error_count']) - int(week_detail['required_submission']),
                0
            )
            week_detail['reward_form_count'] = week_reward
            total_reward_forms += week_reward
        row['reward_form_count'] = total_reward_forms

    return list(user_stats.values())


def format_score_value_label(value):
    try:
        numeric = float(value or 0)
    except Exception:
        numeric = 0.0
    if numeric.is_integer():
        return str(int(numeric))
    return f'{numeric:.2f}'.rstrip('0').rstrip('.')


def build_department_monthly_assessment_payload(current_user_id, start_date, end_date, department_names):
    reward_settings, reward_err = get_teaching_reward_settings()
    if reward_err:
        return None, reward_err

    reward_window = compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day'],
        reward_settings.get('total_weeks'),
    )
    month_templates = build_teaching_month_templates(reward_window, reward_settings)
    selected_departments, department_user_map = resolve_selected_departments(current_user_id, department_names)
    payload = {
        'departments': [],
        'meta': {
            'range_start': start_date.strftime('%Y-%m-%d'),
            'range_end': (end_date - timedelta(days=1)).strftime('%Y-%m-%d'),
            'week_label': reward_window.get('label'),
            'has_full_weeks': reward_window.get('has_full_weeks', False),
            'has_full_months': len(month_templates) > 0,
            'month_count': len(month_templates),
            'month_labels': [item['month_label'] for item in month_templates],
            'months': month_templates,
            'required_submission_per_member': reward_settings['required_submission'],
            'scope_label': '、'.join(selected_departments) if selected_departments else '未选择部门',
            'period_label': '、'.join(item['month_label'] for item in month_templates) if month_templates else '无完整教学月',
            'disclaimer': '考评情况为系统根据当前考评规则和数据自动给出，不代表真实考评情况，请认真核实，建议核实无误后保存或导出数据。'
        }
    }

    if not selected_departments:
        return payload, None
    if not month_templates:
        for department_name in selected_departments:
            payload['departments'].append({
                'department': department_name,
                'member_count': len(department_user_map.get(department_name, [])),
                'months': []
            })
        return payload, None

    selected_week_nos = {
        week['week_no']
        for month in month_templates
        for week in month.get('weeks', [])
    }

    # 查询所有相关用户的考核豁免规则
    all_user_ids_for_exemption = []
    for dept_users in department_user_map.values():
        all_user_ids_for_exemption.extend([u.id for u in dept_users])
    exempt_map = {}  # {user_id: [(start_week, end_week), ...]}
    if all_user_ids_for_exemption:
        exemption_records = AssessmentOverride.query.filter(
            AssessmentOverride.user_id.in_(all_user_ids_for_exemption),
            AssessmentOverride.override_type.in_(ASSESSMENT_EXEMPT_OVERRIDE_TYPES)
        ).all()
        for ov in exemption_records:
            exempt_map.setdefault(ov.user_id, []).append((ov.start_week, ov.end_week))

    def _is_user_exempted(user_id, week_no):
        for s, e in exempt_map.get(user_id, []):
            if s <= week_no <= e:
                return True
        return False

    user_by_number = {}
    department_rows = {}
    required_submission_per_member = int(reward_settings['required_submission'] or 0)
    for department_name in selected_departments:
        users = department_user_map.get(department_name, [])
        user_by_number.update({user.number: user for user in users if user.number})
        month_rows = []
        week_map = {}
        member_count = len(users)
        for month in month_templates:
            weeks = []
            for week in month.get('weeks', []):
                wno = week['week_no']
                active_count = sum(1 for u in users if not _is_user_exempted(u.id, wno))
                week_required = active_count * required_submission_per_member
                week_row = {
                    'week_no': wno,
                    'week_label': week['week_label'],
                    'start_date': week['start_date'],
                    'end_date': week['end_date'],
                    'required_submission': week_required,
                    'effective_count': 0,
                    'error_count': 0,
                    'extra_count': 0,
                    'shortage_count': 0,
                    'missing_count': 0,
                    'missing_penalty': 0,
                    'assessment_total_score': 0.0,
                    'assessment_actual_score': 0.0,
                    'raw_score': 0.0,
                    'final_score': 0.0,
                    'assessment_tiers': [],
                    '_assessment_tier_map': {},
                    '_active_member_count': active_count
                }
                weeks.append(week_row)
                week_map[wno] = week_row
            month_summary_required = sum(w['required_submission'] for w in weeks)
            month_rows.append({
                'month_no': month['month_no'],
                'month_label': month['month_label'],
                'week_range_label': month['week_range_label'],
                'start_date': month['start_date'],
                'end_date': month['end_date'],
                'weeks': weeks,
                'summary': {
                    'required_submission': month_summary_required,
                    'effective_count': 0,
                    'error_count': 0,
                    'extra_count': 0,
                    'shortage_count': 0,
                    'missing_count': 0,
                    'missing_penalty': 0,
                    'assessment_total_score': 0.0,
                    'assessment_actual_score': 0.0,
                    'raw_score': 0.0,
                    'final_score': 0.0
                }
            })
        department_rows[department_name] = {
            'department': department_name,
            'member_count': member_count,
            'required_submission_per_member': required_submission_per_member,
            'required_submission_per_week': member_count * required_submission_per_member,
            'users': users,
            'months': month_rows,
            '_week_map': week_map
        }

    groups = latest_form_groups_for_users(list(user_by_number.keys()))
    effective_form_candidates = []
    user_week_effective = {}
    for group_data in groups:
        latest_form = group_data.get('latest_form')
        if not latest_form:
            continue
        user = user_by_number.get(latest_form.listener_number)
        if not user:
            continue
        department_name = user.department
        department_row = department_rows.get(department_name)
        if not department_row:
            continue
        week_no = get_form_effective_week_no(latest_form, reward_settings)
        if week_no not in selected_week_nos:
            continue
        week_row = department_row['_week_map'].get(week_no)
        if not week_row:
            continue

        # 跳过被豁免用户的表单
        if _is_user_exempted(user.id, week_no):
            continue

        if latest_form.status == '中心已审核':
            week_row['effective_count'] += 1
            user_week_effective[(user.id, week_no)] = user_week_effective.get((user.id, week_no), 0) + 1
            effective_form_candidates.append({
                'group_data': group_data,
                'department': department_name,
                'week_no': week_no,
            })

        if latest_form.score_record:
            for item in latest_form.score_record.items:
                score_value = float(item.personal_score or 0.0)
                if score_value <= 0:
                    continue
                week_row['assessment_total_score'] += score_value
                score_key = format_score_value_label(score_value)
                tier_entry = week_row['_assessment_tier_map'].setdefault(score_key, {
                    'score_value': score_value,
                    'score_label': score_key,
                    'item_count': 0,
                    'total_score': 0.0
                })
                tier_entry['item_count'] += 1
                tier_entry['total_score'] += score_value

    error_unique_ids = resolve_error_unique_ids([item['group_data'] for item in effective_form_candidates])
    for item in effective_form_candidates:
        if item['group_data']['unique_id'] in error_unique_ids:
            department_rows[item['department']]['_week_map'][item['week_no']]['error_count'] += 1

    for department_name, department_row in department_rows.items():
        users = department_row['users']
        for week_no, week_row in department_row['_week_map'].items():
            for user in users:
                # 跳过被豁免的用户
                if _is_user_exempted(user.id, week_no):
                    continue
                effective_count = user_week_effective.get((user.id, week_no), 0)
                diff = effective_count - department_row['required_submission_per_member']
                if diff > 0:
                    week_row['extra_count'] += diff
                elif diff < 0:
                    week_row['shortage_count'] += abs(diff)

            week_row['missing_count'] = max(
                int(week_row['required_submission']) - int(week_row['effective_count']),
                0
            )
            week_row['missing_penalty'] = week_row['missing_count'] * 3
            week_row['raw_score'] = float(week_row['extra_count']) - float(week_row['shortage_count']) * 2
            week_row['assessment_actual_score'] = max(float(week_row['assessment_total_score']) - 5, 0.0)
            week_row['final_score'] = (
                float(week_row['raw_score'])
                - float(week_row['missing_penalty'])
                - float(week_row['assessment_actual_score'])
            )
            week_row['assessment_tiers'] = sorted(
                week_row['_assessment_tier_map'].values(),
                key=lambda item: (float(item['score_value']), item['score_label'])
            )
            week_row.pop('_assessment_tier_map', None)

        for month_row in department_row['months']:
            for week_row in month_row['weeks']:
                for key in ['required_submission', 'effective_count', 'error_count', 'extra_count', 'shortage_count', 'missing_count', 'missing_penalty', 'assessment_total_score', 'assessment_actual_score', 'raw_score', 'final_score']:
                    month_row['summary'][key] += float(week_row[key])
            for int_key in ['required_submission', 'effective_count', 'error_count', 'extra_count', 'shortage_count', 'missing_count', 'missing_penalty']:
                month_row['summary'][int_key] = int(month_row['summary'][int_key])

        department_row.pop('users', None)
        department_row.pop('_week_map', None)
        payload['departments'].append({
            'department': department_row['department'],
            'member_count': department_row['member_count'],
            'required_submission_per_member': department_row['required_submission_per_member'],
            'required_submission_per_week': department_row['required_submission_per_week'],
            'months': department_row['months'],
        })

    return payload, None


def build_submission_snapshot_payload(users, start_date, end_date, sort_by, time_filter_type='created'):
    reward_settings, reward_err = get_teaching_reward_settings()
    if reward_err:
        return None, reward_err
    reward_window = compute_teaching_week_window(
        start_date,
        end_date,
        reward_settings['first_week_date'],
        reward_settings['week_start_day'],
        reward_settings.get('total_weeks'),
    )
    rows = build_submission_count_rows(
        users,
        start_date,
        end_date,
        reward_settings,
        reward_window,
        time_filter_type=time_filter_type
    )
    sort_submission_count_rows(rows, sort_by)
    return {
        'users': rows,
        'sort_by': sort_by,
        'time_filter_type': time_filter_type,
        'reward_meta': {
            'range_start': start_date.strftime('%Y-%m-%d'),
            'range_end': (end_date - timedelta(days=1)).strftime('%Y-%m-%d'),
            'week_label': reward_window['label'],
            'start_week': reward_window['start_week'],
            'end_week': reward_window['end_week'],
            'week_count': reward_window['week_count'],
            'has_full_weeks': reward_window['has_full_weeks'],
            'required_submission_per_week': reward_settings['required_submission'],
            'disclaimer': '奖励表情况为系统根据当前奖励表规则和数据自动给出，不代表真实考评情况，请认真核实，建议核实无误后保存或导出数据。'
        }
    }, None


def build_submission_snapshot_workbook(payload):
    rows = payload.get('users', [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '交表统计'
    headers = ['排名', '姓名', '编号', '部门', '小组', '交表数量', '奖励表数量', '有效表单数', '错误表单数']
    ws.append(headers)
    for row in rows:
        ws.append([
            row.get('rank'),
            row.get('name'),
            row.get('number'),
            row.get('department'),
            row.get('group'),
            row.get('submission_group_count'),
            row.get('reward_form_count'),
            row.get('reward_effective_count'),
            row.get('reward_error_count')
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    detail_ws = wb.create_sheet('奖励表详情')
    detail_headers = ['姓名', '编号', '周次', '开始日期', '结束日期', '有效表单数', '错误表单数', '每周需交表数', '奖励表数']
    detail_ws.append(detail_headers)
    for row in rows:
        for week in row.get('reward_week_details', []):
            detail_ws.append([
                row.get('name'),
                row.get('number'),
                week.get('week_label'),
                week.get('start_date'),
                week.get('end_date'),
                week.get('effective_count'),
                week.get('error_count'),
                week.get('required_submission'),
                week.get('reward_form_count')
            ])
    for cell in detail_ws[1]:
        cell.font = Font(bold=True)

    autosize_worksheet(ws)
    autosize_worksheet(detail_ws)
    return wb


def build_department_monthly_workbook(payload):
    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = '月度考评汇总'
    summary_headers = ['部门', '教学月', '周次范围', '成员数', '月多交表', '月少交表', '月错误表单扣分（原始分）', '月错误表单扣分（实际扣分）', '月总考评分']
    summary_ws.append(summary_headers)
    for department_row in payload.get('departments', []):
        for month_row in department_row.get('months', []):
            summary = month_row.get('summary', {})
            summary_ws.append([
                department_row.get('department'),
                month_row.get('month_label'),
                month_row.get('week_range_label'),
                department_row.get('member_count'),
                summary.get('extra_count'),
                summary.get('shortage_count'),
                summary.get('assessment_total_score'),
                summary.get('assessment_actual_score'),
                summary.get('final_score'),
            ])
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)

    detail_ws = wb.create_sheet('周度考评明细')
    detail_headers = ['部门', '教学月', '周次', '时间范围', '需交表数', '有效表单数', '错误表单数', '多交表（+1分/张）', '少交表（-2分/张）', '实际交表少于应交表数（-3/张）', '错误表单扣分（原始分）', '错误表单扣分（实际扣分）', '周考评分', '扣分档位汇总']
    detail_ws.append(detail_headers)
    for department_row in payload.get('departments', []):
        for month_row in department_row.get('months', []):
            for week_row in month_row.get('weeks', []):
                tier_text = '；'.join(
                    f"{tier.get('score_label')}分 x {tier.get('item_count')}项 = {tier.get('total_score')}"
                    for tier in week_row.get('assessment_tiers', [])
                )
                detail_ws.append([
                    department_row.get('department'),
                    month_row.get('month_label'),
                    week_row.get('week_label'),
                    f"{week_row.get('start_date')} ~ {week_row.get('end_date')}",
                    week_row.get('required_submission'),
                    week_row.get('effective_count'),
                    week_row.get('error_count'),
                    week_row.get('extra_count'),
                    week_row.get('shortage_count'),
                    week_row.get('missing_count'),
                    week_row.get('assessment_total_score'),
                    week_row.get('assessment_actual_score'),
                    week_row.get('final_score'),
                    tier_text
                ])
    for cell in detail_ws[1]:
        cell.font = Font(bold=True)

    autosize_worksheet(summary_ws)
    autosize_worksheet(detail_ws, max_width=60)
    return wb
