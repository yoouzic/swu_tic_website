"""Paired synthetic evidence replay on the user-provided workbook; no database writes."""
import argparse
from collections import defaultdict
from datetime import date
import hashlib
import json
import math
from pathlib import Path
import random
import sys

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from app.services.campus_buildings import load_buildings, nearby_buildings, representative_point
from app.services.lecture_site_capture import PERIODS, infer_period
from app.services.listening_assistant import ListeningAssistantService
from app.services.listening_assistant_contracts import ScheduleEntry, parse_period
from app.services.site_context_model import rank_site_candidates, split_room, _value
from app.services.teaching_calendar import TeachingCalendarConfig, date_for_teaching_weekday


def load_workbook_candidates(path):
    from openpyxl import load_workbook
    workbook=load_workbook(path,read_only=True,data_only=True)
    rows=workbook.active.iter_rows(values_only=True);header=next(rows);entries=[];skipped=0
    source_hash=hashlib.sha256(path.read_bytes()).hexdigest()
    for rownum,values in enumerate(rows,2):
        row=dict(zip(header,values))
        try:
            weekday=int(row['星期几'])
            if not 1<=weekday<=7:raise ValueError('weekday')
            main=parse_period(row.get('上课节次'));venue=parse_period(row.get('场地上课节次'))
            period=venue or main
            if not period:raise ValueError('period')
            entries.append(ScheduleEntry(entry_id=f'workbook:{rownum}',source_row=rownum,
                source_batch_id='replay-'+source_hash[:12],weekday=weekday,
                room=str(row.get('上课地点') or ''),location_raw=str(row.get('上课地点') or ''),
                period=period,period_raw=str(row.get('上课节次') or ''),
                venue_period_raw=str(row.get('场地上课节次') or ''),
                start_week_raw=str(row.get('起始周') or ''),
                venue_start_week_raw=str(row.get('场地上课起始周') or ''),
                teacher_name=str(row.get('姓名') or ''),teacher_college=str(row.get('教师所属学院') or ''),
                course_title=str(row.get('课程名称') or ''),
                student_grade_class=str(row.get('教学班组成') or ''),semester='2025-2026-2'))
        except (TypeError,ValueError):skipped+=1
    workbook.close()
    # Verified against the existing isolated debug configuration, not changed here.
    calendar=TeachingCalendarConfig(date(2026,3,2),0,20)
    service=ListeningAssistantService(schedule_loader=lambda **kwargs:entries,semester='2025-2026-2',calendar=calendar)
    daily={}
    for weekday in range(1,8):
        day=date_for_teaching_weekday(6,weekday,calendar)
        result=service.search_partial({'date':day.isoformat()},candidate_limit=None)
        daily[day.isoformat()]=[c.to_public_dict() for c in result.candidates]
    return daily,{'workbook':path.name,'sha256':source_hash,'parsed_rows':len(entries),
                  'skipped_invalid_rows':skipped,'replay_week':6,'first_week_monday':'2026-03-02'}


def spatial_single_ocr_baseline(rows,evidence):
    """Explicit benchmark baseline; not the old assistant's measured completion rate."""
    if evidence['location']:
        nearby=nearby_buildings(evidence['location'])['buildings'];distances={b['code']:b['distance_m'] for b in nearby}
    else:distances=None
    hypotheses=evidence['door_hypotheses'];top=max(hypotheses,key=lambda h:h['score'])['value'] if hypotheses else None
    number=str(int(top)) if top else None;period=infer_period(evidence['clock'])
    selected=[]
    for c in rows:
        building,room=split_room(c['room'])
        if number and room!=number:continue
        if distances is not None and building not in distances:continue
        selected.append(c)
    return sorted(selected,key=lambda c:(not(period and c['period'][0]<=period<=c['period'][1]),
                                          distances.get(split_room(c['room'])[0],0) if distances else 0,c['candidate_id']))


def evidence_for(target,variant,indexed,rng):
    building,number=split_room(target['room']);door=number.zfill(4)
    clock=PERIODS[min(target['period'][1],target['period'][0]+1)-1][0]
    h,m=map(int,clock.split(':'));minutes=h*60+m+5;clock=f'{minutes//60:02}:{minutes%60:02}'
    point=None
    if building in indexed and variant!='no_location':
        chosen=indexed[building]
        if variant=='biased_location':
            lon,lat=representative_point(chosen)
            chosen=min((f for code,f in indexed.items() if code!=building),
                       key=lambda f:sum((a-b)**2 for a,b in zip(representative_point(f),(lon,lat))))
        lon,lat=representative_point(chosen)
        point={'latitude':lat+rng.gauss(0,8)/111195,
               'longitude':lon+rng.gauss(0,8)/(111195*math.cos(math.radians(lat))),'accuracy':35}
    hypotheses=[{'value':door,'score':.99,'method':'synthetic'}]
    if variant=='ambiguous_ocr':
        wrong=door[:-1]+('8' if door[-1]!='8' else '3')
        hypotheses=[{'value':wrong,'score':.998,'method':'synthetic_whole'},
                    {'value':door,'score':.99,'method':'synthetic_region'}]
    if variant=='break_time':
        end=PERIODS[target['period'][1]-1][1];h,m=map(int,end.split(':'));minutes=h*60+m+5
        clock=f'{minutes//60:02}:{minutes%60:02}'
    return {'location':point,'door_hypotheses':hypotheses,'clock':clock}


def metrics(records):
    groups=defaultdict(list)
    for r in records:groups[r['variant']].append(r)
    result={}
    for name,rows in [('all',records),*groups.items()]:
        n=len(rows)
        result[name]={'count':n,
            'baseline_top3':sum(r['baseline_rank'] is not None and r['baseline_rank']<=3 for r in rows)/n,
            'fusion_top3':sum(r['fusion_rank'] is not None and r['fusion_rank']<=3 for r in rows)/n,
            'fusion_room_top3':sum(r['room_in_fusion_top3'] for r in rows)/n,
            'after_simulated_answers_top3':sum(r['after_questions_rank'] is not None and r['after_questions_rank']<=3 for r in rows)/n,
            'mean_simulated_questions':sum(len(r['answers']) for r in rows)/n}
    return result


def run_replay(workbook, *, count=60,seed=20261002):
    daily,source=load_workbook_candidates(workbook)
    data=load_buildings();indexed={f['properties']['code']:f for f in data['features']}
    eligible=[c for rows in daily.values() for c in rows if split_room(c['room'])[0] in indexed
              and split_room(c['room'])[1].isdigit() and c['period'][1]<=len(PERIODS) and not c['conflicts']]
    if not eligible:raise ValueError('No eligible numbered mapped candidates in replay week')
    rng=random.Random(seed);selected=rng.sample(eligible,min(count,len(eligible)));records=[]
    for target in selected:
        rows=daily[target['lecture_date']]
        for variant in ('clean','biased_location','ambiguous_ocr','no_location','break_time'):
            evidence=evidence_for(target,variant,indexed,rng)
            baseline=spatial_single_ocr_baseline(rows,evidence)
            result=rank_site_candidates(rows,**evidence)
            ranked=result['candidates'];facts={};asked=[];answers=[]
            def position(candidates):
                return next((i+1 for i,c in enumerate(candidates) if c['candidate_id']==target['candidate_id']),None)
            while result['question'] and len(answers)<4 and not (position(result['candidates']) and position(result['candidates'])<=3):
                kind=result['question']['kind'];value=_value(target,kind)
                facts[kind]=value;asked.append(kind);answers.append({'question_kind':kind,'simulated_answer':value})
                result=rank_site_candidates(rows,**evidence,confirmed_facts=facts,asked=asked)
            records.append({'variant':variant,'ground_truth':target,'input':evidence,
                'baseline_rank':position(baseline),'fusion_rank':position(ranked),
                'room_in_fusion_top3':any(c['room']==target['room'] for c in ranked[:3]),
                'after_questions_rank':position(result['candidates']),'answers':answers,
                'baseline_top3':[c['candidate_id'] for c in baseline[:3]],
                'fusion_top3':[{k:c[k] for k in ('candidate_id','room','period','course_title','match_score','match_reasons')} for c in ranked[:3]]})
    return {'scope':'Paired synthetic location/OCR/time perturbations on workbook rows; not field accuracy or human completion',
            'baseline_definition':'nearby footprint distance + highest-score exact OCR room + clock-period ordering',
            'answers_definition':'oracle simulated answers using ground truth only after a model-selected question',
            'source':source,'seed':seed,'targets':len(selected),'eligible_targets':len(eligible),
            'metrics':metrics(records),'records':records}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,default=next(ROOT.glob('*.xlsx')))
    parser.add_argument('--count',type=int,default=60);parser.add_argument('--seed',type=int,default=20261002)
    parser.add_argument('--output',type=Path,default=ROOT/'output/2026-10-02-site-context-fusion/paired-replay.json')
    args=parser.parse_args();report=run_replay(args.workbook,count=args.count,seed=args.seed)
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'source':report['source'],'targets':report['targets'],'metrics':report['metrics']},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
