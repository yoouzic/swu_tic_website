from __future__ import annotations

from dataclasses import dataclass
import hashlib
from pathlib import Path
import re
from typing import Iterable, Iterator
import unicodedata

from openpyxl import load_workbook


FEEDBACK_HEADERS = {
    '课程反馈', '课程反馈优点', '评价内容', '听课评价', '课堂教学评价',
    '不足及建议', '建议', '优质案例推荐', '异常情况反映',
}
IDENTITY_HEADER_MARKERS = (
    '信息员', '听课人', '听课者', '姓名', '教师', '老师', '见证', '签名',
    '手机号', '手机', '电话', '学号', '编号', 'qq', '宿舍', '寝室', '身份',
    '学院', '专业', '行政班', '班级', '年级',
)
PHONE_RE = re.compile(r'(?<!\d)1[3-9]\d{9}(?!\d)')
LONG_NUMBER_RE = re.compile(r'(?<!\d)\d{8,}(?!\d)')
CHINESE_RE = re.compile(r'[\u3400-\u9fff]')
KNOWN_NAME_MARKERS = ('张三', '李四', '王五', '赵六')


@dataclass(frozen=True)
class CorpusStats:
    source_files: int
    candidate_cells: int
    accepted_texts: int
    sha256: str

    def to_dict(self) -> dict[str, object]:
        return {
            'source_files': self.source_files,
            'candidate_cells': self.candidate_cells,
            'accepted_texts': self.accepted_texts,
            'sha256': self.sha256,
        }


@dataclass(frozen=True)
class CorpusResult:
    texts: tuple[str, ...]
    stats: CorpusStats

    def __iter__(self) -> Iterator[str]:
        return iter(self.texts)

    def __len__(self) -> int:
        return len(self.texts)


def _normalize_header(value: object) -> str:
    return ''.join(str(value or '').split()).strip('：:')


def _identity_header(value: object) -> bool:
    normalized = _normalize_header(value).lower()
    return any(marker.lower() in normalized for marker in IDENTITY_HEADER_MARKERS)


def _is_feedback_header(value: object) -> bool:
    normalized = _normalize_header(value)
    if _identity_header(normalized):
        return False
    return normalized in FEEDBACK_HEADERS or any(
        alias in normalized for alias in FEEDBACK_HEADERS if alias not in {'建议'}
    )


def deidentify_text(value: object, identity_values: Iterable[str] = ()) -> str:
    """Return evaluation prose with transient identity values removed."""

    text = unicodedata.normalize('NFKC', str(value or ''))
    for identity in sorted({str(item).strip() for item in identity_values if str(item).strip()}, key=len, reverse=True):
        if len(identity) >= 2:
            text = text.replace(identity, '合成对象')
    for marker in KNOWN_NAME_MARKERS:
        text = text.replace(marker, '合成对象')
    text = PHONE_RE.sub('合成联系方式', text)
    text = LONG_NUMBER_RE.sub('合成编号', text)
    # Names embedded immediately before common identity nouns are not needed as prose.
    text = re.sub(r'(?<!验收)[\u3400-\u9fff]{2,4}(?=老师)', '合成教师', text)
    text = re.sub(r'(?<!验收)[\u3400-\u9fff]{2,4}(?=同学)', '合成同学', text)
    return ' '.join(text.split())


def _acceptable_prose(value: str) -> bool:
    if not value or not 8 <= len(value) <= 600:
        return False
    return len(CHINESE_RE.findall(value)) >= 8


def _workbook_paths(source: Path) -> list[Path]:
    if source.is_file():
        return [source]
    if not source.is_dir():
        return []
    return sorted(path for path in source.iterdir() if path.suffix.lower() == '.xlsx')


def extract_evaluation_corpus(source: str | Path) -> CorpusResult:
    """Read only matched feedback columns from demo workbooks.

    No source rows, filenames, identity values, or workbook metadata are returned.
    """

    source_path = Path(source).expanduser().resolve()
    texts: set[str] = set()
    candidate_cells = 0
    source_files = 0
    for workbook_path in _workbook_paths(source_path):
        source_files += 1
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        except Exception:
            # Historical demo files can contain styles unsupported by the bundled
            # reader. Skip only that file; never expose its raw contents or error.
            continue
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = None
                header_row = -1
                for row_number, row in enumerate(rows):
                    if row_number >= 10:
                        break
                    feedback_columns = [
                        index for index, value in enumerate(row) if _is_feedback_header(value)
                    ]
                    if feedback_columns:
                        header = tuple(row)
                        header_row = row_number
                        break
                if header is None:
                    continue
                feedback_columns = [
                    index for index, value in enumerate(header) if _is_feedback_header(value)
                ]
                identity_columns = [
                    index for index, value in enumerate(header) if _identity_header(value)
                ]
                identity_values: set[str] = set()
                remaining_rows = list(rows)
                for row in remaining_rows:
                    for index in identity_columns:
                        if index < len(row) and row[index] is not None:
                            value = str(row[index]).strip()
                            if 2 <= len(value) <= 40:
                                identity_values.add(value)
                for row in remaining_rows:
                    for index in feedback_columns:
                        if index >= len(row) or not isinstance(row[index], str):
                            continue
                        candidate_cells += 1
                        text = deidentify_text(row[index], identity_values)
                        if _acceptable_prose(text):
                            texts.add(text)
        except Exception:
            # A workbook can fail while streaming rows or data validations, too.
            # Keep extraction fail-closed and continue with other demo files.
            continue
        finally:
            workbook.close()
    ordered = tuple(sorted(texts))
    digest = hashlib.sha256('\n'.join(ordered).encode('utf-8')).hexdigest()
    return CorpusResult(
        texts=ordered,
        stats=CorpusStats(
            source_files=source_files,
            candidate_cells=candidate_cells,
            accepted_texts=len(ordered),
            sha256=digest,
        ),
    )


def school_schedule_sha256(source: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(source).expanduser().resolve().open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def _parse_week_structure(value: object) -> tuple[int, ...]:
    values = []
    for start, end in re.findall(r'(\d+)\s*(?:[-~至到]\s*(\d+))?', str(value or '')):
        first = int(start)
        last = int(end or start)
        if 1 <= first <= last <= 60:
            values.extend(range(first, last + 1))
    return tuple(sorted(set(values)))


def _parse_period_structure(value: object) -> tuple[int, int] | None:
    match = re.search(r'(\d+)\s*(?:[-~至到]\s*(\d+))?', str(value or ''))
    if not match:
        return None
    start = int(match.group(1))
    end = int(match.group(2) or start)
    if not 1 <= start <= end <= 20:
        return None
    return start, end


def _parse_weekday_structure(value: object) -> int | None:
    text = str(value or '').strip()
    match = re.search(r'[1-7]', text)
    if match:
        return int(match.group(0))
    names = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '日': 7, '天': 7}
    for name, weekday in names.items():
        if name in text:
            return weekday
    return None


def extract_school_schedule_structure(source: str | Path, *, limit: int = 2000):
    """Extract only weeks, weekday, periods and synthetic location tokens."""

    from .generator import ScheduleStructure

    path = Path(source).expanduser().resolve()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ()
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = None
            header_row = -1
            positions = {}
            for row_number, row in enumerate(rows):
                if row_number >= 10:
                    break
                normalized = {_normalize_header(value): index for index, value in enumerate(row)}
                weekday_header = next((name for name in ('星期几', '星期', 'weekday') if name in normalized), None)
                weeks_header = next((name for name in ('起始周', '周次', 'weeks') if name in normalized), None)
                period_header = next((name for name in ('上课节次', '节次', 'periods') if name in normalized), None)
                location_header = next((name for name in ('上课地点', '地点', 'location') if name in normalized), None)
                if all(value is not None for value in (weekday_header, weeks_header, period_header, location_header)):
                    header = tuple(row)
                    header_row = row_number
                    positions = {
                        'weekday': normalized[weekday_header],
                        'weeks': normalized[weeks_header],
                        'periods': normalized[period_header],
                        'location': normalized[location_header],
                    }
                    break
            if header is None:
                continue
            location_tokens: dict[str, str] = {}
            structures = []
            for row in rows:
                weeks = _parse_week_structure(row[positions['weeks']] if positions['weeks'] < len(row) else None)
                periods = _parse_period_structure(row[positions['periods']] if positions['periods'] < len(row) else None)
                weekday = _parse_weekday_structure(row[positions['weekday']] if positions['weekday'] < len(row) else None)
                raw_location = str(row[positions['location']] or '').strip() if positions['location'] < len(row) else ''
                if not weeks or not periods or weekday is None:
                    continue
                if raw_location not in location_tokens:
                    location_tokens[raw_location] = f'验收地点结构{len(location_tokens) + 1:03d}'
                structures.append(ScheduleStructure(
                    weeks=weeks,
                    weekday=weekday,
                    start_period=periods[0],
                    end_period=periods[1],
                    location=location_tokens[raw_location],
                ))
                if len(structures) >= limit:
                    return tuple(structures)
            if structures:
                return tuple(structures)
    except Exception:
        return ()
    finally:
        workbook.close()
    return ()


collect_deidentified_corpus = extract_evaluation_corpus


__all__ = [
    'CorpusResult', 'CorpusStats', 'FEEDBACK_HEADERS', 'collect_deidentified_corpus',
    'deidentify_text', 'extract_evaluation_corpus', 'extract_school_schedule_structure',
    'school_schedule_sha256',
]
