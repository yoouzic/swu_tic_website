from __future__ import annotations

from dataclasses import asdict, dataclass, replace
from datetime import date, timedelta
import hashlib
import json
from pathlib import Path
import random
import re
from typing import Iterable, Mapping, Sequence

from openpyxl import Workbook

from .config import AcceptanceConfig


DEPARTMENTS = (
    '验收部门一', '验收部门二', '验收部门三', '验收部门四',
)
GROUP_SUFFIXES = ('验收A组', '验收B组')
REVIEW_MODES = ('rules_only', 'llm_only', 'combined')
ANOMALY_MARKERS = (
    'personal_schedule_exact_conflict',
    'class_schedule_approximate_conflict',
    'same_college_teacher',
    'repeated_witness_across_weeks',
    'consecutive_teacher_weeks',
    'missing_required_prefix',
    'short_or_template_feedback',
    'same_time_conflict',
    'school_schedule_mismatch',
    'insufficient_schedule_coverage',
)
DEFAULT_CORPUS = (
    '该老师讲解具体，课堂组织清晰，能够结合课程内容引导同学思考。',
    '该老师表达清楚，课堂节奏合理，重点内容安排得比较完整。',
)


@dataclass(frozen=True)
class ScheduleStructure:
    """De-identified time/location structure copied from a source schedule."""

    weeks: tuple[int, ...]
    weekday: int
    start_period: int
    end_period: int
    location: str


@dataclass(frozen=True)
class AcceptanceScheduleEntry:
    weeks: tuple[int, ...]
    weekday: int
    start_period: int
    end_period: int
    location: str
    teacher_name: str
    teacher_college: str
    course_title: str
    teaching_class: str


@dataclass(frozen=True)
class OfficerSpec:
    officer_id: str
    student_id: str
    name: str
    department: str
    group: str
    form_count: int
    coverage: str

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class FormSpec:
    synthetic_key: str
    ordinal: int
    officer_id: str
    student_id: str
    listener_name: str
    department: str
    group: str
    review_mode: str
    teaching_week: int
    lecture_date: str
    weekday: int
    start_period: int
    end_period: int
    lecture_location: str
    teacher_name: str
    teacher_college: str
    course_title: str
    student_grade_class: str
    witness_name: str
    student_signature1: str
    contact_phone1: str
    student_signature2: str
    contact_phone2: str
    course_feedback: str
    suggestions: str
    coverage: str
    normal_control: bool
    oracle_markers: tuple[str, ...]

    def to_dict(self) -> dict[str, object]:
        values = asdict(self)
        values['oracle_markers'] = list(self.oracle_markers)
        return values


@dataclass(frozen=True)
class AcceptanceManifest:
    config: AcceptanceConfig
    officers: tuple[OfficerSpec, ...]
    forms: tuple[FormSpec, ...]
    corpus_sha256: str
    sha256: str
    school_schedule_source_sha256: str = ''

    @property
    def officer_count(self) -> int:
        return len(self.officers)

    @property
    def logical_form_count(self) -> int:
        return len(self.forms)

    @property
    def department_counts(self) -> list[int]:
        return [sum(1 for officer in self.officers if officer.department == department) for department in DEPARTMENTS]

    @property
    def batch_counts(self) -> dict[str, int]:
        return {mode: sum(1 for form in self.forms if form.review_mode == mode) for mode in REVIEW_MODES}

    @property
    def llm_logical_count(self) -> int:
        return sum(form.review_mode in {'llm_only', 'combined'} for form in self.forms)

    @property
    def normal_control_count(self) -> int:
        return sum(form.normal_control for form in self.forms)

    @property
    def violation_intended_count(self) -> int:
        return sum(not form.normal_control for form in self.forms)

    def to_dict(self) -> dict[str, object]:
        return {
            'config': self.config.to_dict(),
            'corpus_sha256': self.corpus_sha256,
            'school_schedule_source_sha256': self.school_schedule_source_sha256,
            'sha256': self.sha256,
            'officer_count': self.officer_count,
            'logical_form_count': self.logical_form_count,
            'department_counts': self.department_counts,
            'batch_counts': self.batch_counts,
            'llm_logical_count': self.llm_logical_count,
            'normal_control_count': self.normal_control_count,
            'violation_intended_count': self.violation_intended_count,
            'officers': [officer.to_dict() for officer in self.officers],
            'forms': [form.to_dict() for form in self.forms],
        }

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, sort_keys=True, separators=(',', ':'))

    def oracle_rows(self) -> tuple[dict[str, object], ...]:
        return tuple({
            'synthetic_key': form.synthetic_key,
            'review_mode': form.review_mode,
            'normal_control': form.normal_control,
            'oracle_markers': list(form.oracle_markers),
            'coverage': form.coverage,
        } for form in self.forms)


def _canonical_digest_payload(
    config: AcceptanceConfig,
    officers: Sequence[OfficerSpec],
    forms: Sequence[FormSpec],
    corpus_sha256: str,
    school_schedule_source_sha256: str,
) -> str:
    config_values = config.to_dict()
    for field in ('runtime_root', 'demo_dir', 'school_schedule'):
        config_values.pop(field, None)
    payload = {
        'config': config_values,
        'corpus_sha256': corpus_sha256,
        'school_schedule_source_sha256': school_schedule_source_sha256,
        'officers': [officer.to_dict() for officer in officers],
        'forms': [form.to_dict() for form in forms],
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()


def _normal_text(value: str) -> str:
    value = value.strip()
    if value.startswith('该老师'):
        return value
    return f'该老师{value}'


_NORMAL_FEEDBACK_PADDING = (
    '\uff0c\u8bfe\u5802\u4e92\u52a8\u81ea\u7136\uff0c\u91cd\u70b9\u8bb2\u89e3\u6e05\u695a\uff0c'
    '\u5b9e\u4f8b\u5145\u5206\uff0c\u5b66\u751f\u80fd\u591f\u8ddf\u8fdb\u8bfe\u5802\u601d\u8003\u3002'
)

_FEEDBACK_VARIANTS = (
    '\uff0c\u4fa7\u91cd\u6559\u5b66\u7ec4\u7ec7\u4e0e\u8bfe\u7a0b\u8854\u63a5\u7684\u89c2\u5bdf\u3002',
    '\uff0c\u4fa7\u91cd\u91cd\u70b9\u5185\u5bb9\u4e0e\u5b9e\u4f8b\u89e3\u91ca\u7684\u8fde\u8d2f\u3002',
    '\uff0c\u4fa7\u91cd\u8bfe\u5802\u4e92\u52a8\u548c\u5b66\u751f\u7406\u89e3\u60c5\u51b5\u3002',
    '\uff0c\u4fa7\u91cd\u6559\u5b66\u65b9\u6cd5\u548c\u5b9e\u8df5\u73af\u8282\u8fde\u63a5\u3002',
    '\uff0c\u4fa7\u91cd\u677f\u4e66\u5b89\u6392\u4e0e\u8bfe\u5802\u8282\u594f\u7684\u914d\u5408\u3002',
    '\uff0c\u4fa7\u91cd\u95ee\u9898\u8bbe\u8ba1\u548c\u8ba8\u8bba\u53cd\u9988\u7684\u8bb0\u5f55\u3002',
    '\uff0c\u4fa7\u91cd\u77e5\u8bc6\u8981\u70b9\u4e0e\u5b66\u4e60\u96be\u70b9\u7684\u8bf4\u660e\u3002',
    '\uff0c\u4fa7\u91cd\u8bfe\u7a0b\u5185\u5bb9\u548c\u6559\u5b66\u6d41\u7a0b\u7684\u6574\u4f53\u8bc4\u4f30\u3002',
)


def _ensure_normal_feedback(text: str) -> str:
    while sum('\u3400' <= char <= '\u9fff' for char in text) < 50:
        text += _NORMAL_FEEDBACK_PADDING
    return text


def _violation_text(value: str, markers: Sequence[str]) -> str:
    text = value
    if 'missing_required_prefix' in markers and text.startswith('该老师'):
        text = text[3:]
    if 'short_or_template_feedback' in markers:
        text = '一般，较好。'
    if 'school_schedule_mismatch' in markers:
        text = f'{text} 课程安排与已启用课表待核对。'
    return text


def _marker_candidates(officer: OfficerSpec, form_ordinal: int, normal_control: bool) -> tuple[str, ...]:
    if normal_control:
        return ()
    if officer.coverage == 'complete':
        candidates = (
            'personal_schedule_exact_conflict',
            'same_college_teacher',
            'school_schedule_mismatch',
            'missing_required_prefix',
            'short_or_template_feedback',
        )
    elif officer.coverage == 'basic':
        candidates = (
            'class_schedule_approximate_conflict',
            'same_college_teacher',
            'school_schedule_mismatch',
            'missing_required_prefix',
            'short_or_template_feedback',
        )
    else:
        candidates = (
            'insufficient_schedule_coverage',
            'school_schedule_mismatch',
            'missing_required_prefix',
            'short_or_template_feedback',
        )
    return (candidates[form_ordinal % len(candidates)],)


def _relation_anchors(
    officers: Sequence[OfficerSpec],
    config: AcceptanceConfig,
) -> dict[str, str]:
    """Assign relation markers to distinct two-form officers."""

    eligible: list[tuple[OfficerSpec, str, int]] = []
    ordinal = 0
    for officer in officers:
        form_ordinals = tuple(range(ordinal, ordinal + officer.form_count))
        if officer.form_count == 2 and all(
            item % config.per_mode >= config.normal_per_mode
            for item in form_ordinals
        ):
            mode = REVIEW_MODES[form_ordinals[0] // config.per_mode]
            eligible.append((officer, mode, form_ordinals[0]))
        ordinal += officer.form_count
    if len(eligible) < 3:
        raise ValueError('at least three two-form officers are required for relation oracles')

    selected: list[tuple[OfficerSpec, str, int]] = []
    for desired_mode in ('llm_only', 'combined', 'llm_only'):
        unused = [
            item for item in eligible
            if item[1] == desired_mode
            and item[0] not in {chosen[0] for chosen in selected}
        ]
        candidate = next(
            (item for item in unused if item[0].department not in {chosen[0].department for chosen in selected}),
            None,
        ) or next(iter(unused), None)
        if candidate is None:
            raise ValueError(f'no eligible two-form officer for relation mode: {desired_mode}')
        selected.append(candidate)
    return {
        'repeated_witness_across_weeks': selected[0][0].officer_id,
        'consecutive_teacher_weeks': selected[1][0].officer_id,
        'same_time_conflict': selected[2][0].officer_id,
    }


def _plan_oracle_markers(
    officers: Sequence[OfficerSpec],
    config: AcceptanceConfig,
) -> tuple[dict[tuple[str, int], tuple[str, ...]], dict[str, str]]:
    planned: dict[tuple[str, int], list[str]] = {}
    ordinal = 0
    for officer in officers:
        for form_index in range(officer.form_count):
            normal_control = (ordinal % config.per_mode) < config.normal_per_mode
            candidates = _marker_candidates(officer, ordinal, normal_control)
            planned[(officer.officer_id, form_index)] = list(candidates)
            ordinal += 1

    anchors = _relation_anchors(officers, config)
    for marker, officer_id in anchors.items():
        for form_index in range(2):
            planned[(officer_id, form_index)].append(marker)
    return {
        key: tuple(dict.fromkeys(markers))
        for key, markers in planned.items()
    }, anchors


def _coverage_for(index: int, officer_count: int) -> str:
    cohort_size = officer_count // 3
    if index < cohort_size:
        return 'complete'
    if index < 2 * cohort_size:
        return 'basic'
    return 'missing'


def _department_for(index: int, counts: Sequence[int]) -> str:
    cursor = 0
    for department, count in zip(DEPARTMENTS, counts):
        cursor += count
        if index < cursor:
            return department
    return DEPARTMENTS[-1]


def _corpus_texts(corpus: Iterable[str]) -> tuple[str, ...]:
    cleaned = {
        str(text).strip()
        for text in corpus
        if isinstance(text, str) and str(text).strip()
    }
    return tuple(sorted(cleaned or DEFAULT_CORPUS))


def _usable_structures(school_structure: Sequence[ScheduleStructure] | None) -> tuple[ScheduleStructure, ...]:
    if not school_structure:
        return ()
    usable = []
    for structure in school_structure:
        weeks = tuple(sorted({int(week) for week in structure.weeks if 1 <= int(week) <= 60}))
        if not weeks:
            continue
        if not 1 <= int(structure.weekday) <= 7:
            continue
        if not 1 <= int(structure.start_period) <= int(structure.end_period) <= 20:
            continue
        location = str(structure.location or '').strip() or f'验收地点{len(usable) + 1:03d}'
        usable.append(ScheduleStructure(
            weeks=weeks,
            weekday=int(structure.weekday),
            start_period=int(structure.start_period),
            end_period=int(structure.end_period),
            location=location,
        ))
    return tuple(usable)


def _select_structure(
    structures: Sequence[ScheduleStructure],
    form_ordinal: int,
    previous_week: int | None,
) -> tuple[ScheduleStructure | None, int]:
    if not structures:
        return None, 1 + (form_ordinal % 16)
    for offset in range(len(structures)):
        structure = structures[(form_ordinal + offset) % len(structures)]
        week = structure.weeks[form_ordinal % len(structure.weeks)]
        if previous_week is None or week != previous_week:
            return structure, week
    structure = structures[form_ordinal % len(structures)]
    return structure, structure.weeks[form_ordinal % len(structure.weeks)]


def generate_manifest(
    config: AcceptanceConfig,
    corpus: Iterable[str],
    *,
    school_structure: Sequence[ScheduleStructure] | None = None,
    school_schedule_source_sha256: str = '',
) -> AcceptanceManifest:
    """Create the deterministic synthetic organization, forms, and oracle."""

    config.validate()
    texts = _corpus_texts(corpus)
    corpus_sha256 = hashlib.sha256('\n'.join(texts).encode('utf-8')).hexdigest()
    rng = random.Random(config.seed)
    structures = _usable_structures(school_structure)

    officers: list[OfficerSpec] = []
    for index in range(config.officer_count):
        department = _department_for(index, config.department_counts)
        group = GROUP_SUFFIXES[index % len(GROUP_SUFFIXES)]
        officers.append(OfficerSpec(
            officer_id=f'YA{index + 1:04d}',
            student_id=f'{260000000000 + index + 1:012d}',
            name=f'验收信息员{index + 1:04d}',
            department=department,
            group=group,
            form_count=1 if index < config.one_form_officer_count else 2,
            coverage=_coverage_for(index, config.officer_count),
        ))

    marker_plan, relation_anchors = _plan_oracle_markers(officers, config)
    forms: list[FormSpec] = []
    form_ordinal = 0
    for officer_index, officer in enumerate(officers):
        previous_week: int | None = None
        for form_index in range(officer.form_count):
            mode = REVIEW_MODES[form_ordinal // config.per_mode]
            mode_ordinal = form_ordinal % config.per_mode
            normal_control = mode_ordinal < config.normal_per_mode
            markers = marker_plan[(officer.officer_id, form_index)]
            source_text = texts[rng.randrange(len(texts))]
            normal_text = _normal_text(source_text)
            if normal_control or 'short_or_template_feedback' not in markers:
                normal_text = _ensure_normal_feedback(normal_text)
            normal_text += _FEEDBACK_VARIANTS[form_ordinal % len(_FEEDBACK_VARIANTS)]
            feedback = _violation_text(normal_text, markers) if not normal_control else normal_text
            structure, week = _select_structure(structures, form_ordinal, previous_week)
            if structure is None:
                weekday = 1 + ((form_ordinal + officer_index) % 5)
                start_period = 1 + ((form_ordinal * 2) % 10)
                end_period = start_period + 1
                location = f'验收地点{1 + (form_ordinal % 20):02d}'
            else:
                weekday = structure.weekday
                start_period = structure.start_period
                end_period = structure.end_period
                location = structure.location
            previous_week = week
            lecture_date = date.fromisoformat(config.semester_monday) + timedelta(days=(week - 1) * 7 + weekday - 1)
            teacher_index = 1 + (form_ordinal % 250)
            department_index = DEPARTMENTS.index(officer.department) + 1
            teacher_college = (
                officer.department
                if 'same_college_teacher' in markers
                else f'楠屾敹瀛﹂櫌{department_index}'
            )
            witness_name = f'楠屾敹瑙佽瘉{1 + (form_ordinal % 60):03d}'
            contact_phone = f'{13900000000 + form_ordinal:011d}'
            witness_name2 = f'楠屾敹瑙佽瘉{61 + (form_ordinal % 60):03d}'
            contact_phone2 = f'{13800000000 + form_ordinal:011d}'
            if 'school_schedule_mismatch' in markers:
                location = '验收偏离地点'
            forms.append(FormSpec(
                synthetic_key=f'AC-{form_ordinal + 1:04d}',
                ordinal=form_ordinal,
                officer_id=officer.officer_id,
                student_id=officer.student_id,
                listener_name=officer.name,
                department=officer.department,
                group=officer.group,
                review_mode=mode,
                teaching_week=week,
                lecture_date=lecture_date.isoformat(),
                weekday=weekday,
                start_period=start_period,
                end_period=end_period,
                lecture_location=location,
                teacher_name=f'验收教师{teacher_index:03d}',
                teacher_college=f'验收学院{department_index}',
                course_title=f'验收课程{1 + (form_ordinal % 300):03d}',
                student_grade_class=f'验收行政班{department_index}{1 + (officer_index % 12):02d}',
                witness_name=f'验收见证{1 + (form_ordinal % 60):03d}',
                student_signature1=witness_name,
                contact_phone1=contact_phone,
                student_signature2=witness_name2,
                contact_phone2=contact_phone2,
                course_feedback=feedback,
                suggestions='无' if normal_control else '建议结合证据进一步核对。',
                coverage=officer.coverage,
                normal_control=normal_control,
                oracle_markers=markers,
            ))
            form_ordinal += 1

    form_indexes = {
        form.synthetic_key: index
        for index, form in enumerate(forms)
    }

    def update_form(form: FormSpec, **changes) -> FormSpec:
        index = form_indexes[form.synthetic_key]
        updated = replace(form, **changes)
        forms[index] = updated
        return updated

    def update_week(form: FormSpec, week: int) -> FormSpec:
        lecture_date = date.fromisoformat(config.semester_monday) + timedelta(
            days=(week - 1) * 7 + form.weekday - 1,
        )
        return update_form(form, teaching_week=week, lecture_date=lecture_date.isoformat())

    for form in tuple(forms):
        update_form(
            form,
            teacher_name=f'\u9a8c\u6536\u6559\u5e08{form.ordinal + 1:04d}',
            course_title=f'\u9a8c\u6536\u8bfe\u7a0b{form.ordinal + 1:04d}',
        )

    for form in tuple(forms):
        if 'same_college_teacher' in form.oracle_markers:
            update_form(form, teacher_college=form.department)

    repeated_officer = relation_anchors['repeated_witness_across_weeks']
    repeated_forms = [form for form in forms if form.officer_id == repeated_officer]
    repeated_phone = f'{13700000001:011d}'
    repeated_name = '楠屾敹閲嶅瑙佽瘉001'
    for form in repeated_forms:
        update_form(
            form,
            witness_name=repeated_name,
            student_signature1=repeated_name,
            contact_phone1=repeated_phone,
        )

    consecutive_officer = relation_anchors['consecutive_teacher_weeks']
    consecutive_forms = [form for form in forms if form.officer_id == consecutive_officer]
    if len(consecutive_forms) != 2:
        raise ValueError('consecutive-teacher anchor must have two forms')
    first, second = consecutive_forms
    base_week = min(first.teaching_week, 59)
    first = update_week(first, base_week)
    second = update_week(second, base_week + 1)
    update_form(
        second,
        teacher_name=first.teacher_name,
        course_title=first.course_title,
        teacher_college=first.teacher_college,
        weekday=first.weekday,
        start_period=first.start_period,
        end_period=first.end_period,
        lecture_location=first.lecture_location,
    )

    same_time_officer = relation_anchors['same_time_conflict']
    same_time_forms = [form for form in forms if form.officer_id == same_time_officer]
    if len(same_time_forms) != 2:
        raise ValueError('same-time anchor must have two forms')
    first, second = same_time_forms
    update_form(
        second,
        teaching_week=first.teaching_week,
        lecture_date=first.lecture_date,
        weekday=first.weekday,
        start_period=first.start_period,
        end_period=first.end_period,
    )

    # Keep the cross-week contract for every two-form officer except the
    # deliberately colliding same-time oracle pair.
    for officer in officers:
        officer_forms = [form for form in forms if form.officer_id == officer.officer_id]
        if len(officer_forms) != 2:
            continue
        marker_union = set().union(*(set(form.oracle_markers) for form in officer_forms))
        if 'same_time_conflict' in marker_union:
            continue
        if officer_forms[0].teaching_week == officer_forms[1].teaching_week:
            update_week(officer_forms[1], min(officer_forms[0].teaching_week + 1, 60))

    forms = [forms[index] for index in range(len(forms))]

    if len(forms) != config.logical_form_count:
        raise ValueError('generated form count did not close')
    digest = _canonical_digest_payload(
        config,
        officers,
        forms,
        corpus_sha256,
        school_schedule_source_sha256,
    )
    return AcceptanceManifest(
        config=config,
        officers=tuple(officers),
        forms=tuple(forms),
        corpus_sha256=corpus_sha256,
        sha256=digest,
        school_schedule_source_sha256=school_schedule_source_sha256,
    )


def assert_real_identity_patterns_absent(serialized: str) -> None:
    """Fail closed on phone/long-number/name patterns in an output string."""

    # The ignored runtime manifest contains one generated login secret. It is
    # intentionally excluded from the identity-pattern scan without exposing it.
    serialized = re.sub(r'("run_password"\s*:\s*")[^"]*(")', r'\1<redacted>\2', serialized)
    serialized = re.sub(r'("(?:sha256|corpus_sha256)"\s*:\s*")[^"]*(")', r'\1<redacted>\2', serialized)
    serialized = re.sub(r'("contact_phone[12]"\s*:\s*")1[3-9]\d{9}(")', r'\1<synthetic-phone>\2', serialized)
    serialized = serialized.replace('合成教师老师', '<synthetic-teacher>')
    serialized = serialized.replace('合成同学同学', '<synthetic-student>')
    serialized = serialized.replace('合成对象老师', '<synthetic-teacher>')
    serialized = serialized.replace('合成对象同学', '<synthetic-student>')
    if re.search(r'(?<!\d)1[3-9]\d{9}(?!\d)', serialized):
        raise AssertionError('raw phone-like identity pattern found')
    for match in re.finditer(r'(?<!\d)\d{8,}(?!\d)', serialized):
        value = match.group(0)
        if len(value) == 12 and value.startswith('26000000'):
            continue
        raise AssertionError('raw long numeric identity pattern found')
    for marker in ('张三', '李四', '王五', '赵六', '张三老师', '李四老师', '王五老师', '赵六老师'):
        if marker in serialized:
            raise AssertionError('known example identity found')


def _listener_admin_class(officer_id: str) -> str:
    """Return a synthetic administrative class owned by one information officer."""

    return f'acceptance-listener-class-{officer_id}'


def _slot_overlaps_form(
    week: int,
    weekday: int,
    start_period: int,
    end_period: int,
    form: FormSpec,
) -> bool:
    return (
        week == form.teaching_week
        and weekday == form.weekday
        and start_period <= form.end_period
        and form.start_period <= end_period
    )


def _first_non_overlapping_slot(forms: Sequence[FormSpec]) -> tuple[int, int, int, int]:
    for week in range(1, 61):
        for weekday in range(1, 8):
            for start_period in range(1, 20):
                end_period = start_period + 1
                if all(
                    not _slot_overlaps_form(week, weekday, start_period, end_period, form)
                    for form in forms
                ):
                    return week, weekday, start_period, end_period
    raise ValueError('could not place a synthetic non-overlapping schedule slot')


def build_acceptance_school_schedule(manifest: AcceptanceManifest) -> tuple[AcceptanceScheduleEntry, ...]:
    """Build synthetic rows for form matching and independent class mappings.

    ``student_grade_class`` identifies the class being observed, not the
    information officer's own administrative class.  Basic officers therefore
    receive a separate synthetic class token.  Its base slot is deliberately
    disjoint from all of that officer's forms; only forms carrying the class
    conflict oracle receive an overlapping approximate slot.
    """

    entries: list[AcceptanceScheduleEntry] = []
    seen: set[tuple[object, ...]] = set()

    def add_entry(entry: AcceptanceScheduleEntry) -> None:
        key = (
            entry.weeks, entry.weekday, entry.start_period, entry.end_period,
            entry.location, entry.teacher_name, entry.teacher_college,
            entry.course_title, entry.teaching_class,
        )
        if key not in seen:
            seen.add(key)
            entries.append(entry)

    forms_by_officer: dict[str, list[FormSpec]] = {}
    for form in manifest.forms:
        forms_by_officer.setdefault(form.officer_id, []).append(form)
        schedule_location = (
            'acceptance-school-location-mismatch'
            if 'school_schedule_mismatch' in form.oracle_markers
            else form.lecture_location
        )
        add_entry(AcceptanceScheduleEntry(
            weeks=(form.teaching_week,),
            weekday=form.weekday,
            start_period=form.start_period,
            end_period=form.end_period,
            location=schedule_location,
            teacher_name=form.teacher_name,
            teacher_college=form.teacher_college,
            course_title=form.course_title,
            teaching_class=form.student_grade_class,
        ))

    for officer in manifest.officers:
        if officer.coverage != 'basic':
            continue
        forms = forms_by_officer.get(officer.officer_id, [])
        if not forms:
            raise ValueError(f'basic officer has no form: {officer.officer_id}')
        admin_class = _listener_admin_class(officer.officer_id)
        week, weekday, start_period, end_period = _first_non_overlapping_slot(forms)
        add_entry(AcceptanceScheduleEntry(
            weeks=(week,),
            weekday=weekday,
            start_period=start_period,
            end_period=end_period,
            location='acceptance-listener-location',
            teacher_name=f'\u9a8c\u6536\u6559\u5e08-{officer.officer_id}',
            teacher_college='acceptance-listener-college',
            course_title=f'\u9a8c\u6536\u8bfe\u7a0b-{officer.officer_id}',
            teaching_class=admin_class,
        ))
        for form in forms:
            if 'class_schedule_approximate_conflict' not in form.oracle_markers:
                continue
            add_entry(AcceptanceScheduleEntry(
                weeks=(form.teaching_week,),
                weekday=form.weekday,
                start_period=form.start_period,
                end_period=form.end_period,
                location=form.lecture_location,
                teacher_name=form.teacher_name,
                teacher_college=form.teacher_college,
                course_title=form.course_title,
                teaching_class=admin_class,
            ))
    return tuple(entries)


def write_acceptance_school_schedule(manifest: AcceptanceManifest, output_path: str | Path) -> Path:
    """Write the de-identified school schedule consumed by acceptance import."""

    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = 'school_schedule'
    sheet.append([
        'semester', 'teacher_name', 'teacher_college', 'course_title',
        'teaching_class', 'weeks', 'weekday', 'periods', 'location',
    ])
    for entry in build_acceptance_school_schedule(manifest):
        weeks = ','.join(str(week) for week in entry.weeks)
        sheet.append([
            manifest.config.semester,
            entry.teacher_name,
            entry.teacher_college,
            entry.course_title,
            entry.teaching_class,
            weeks,
            entry.weekday,
            f'{entry.start_period}-{entry.end_period}',
            entry.location,
        ])
    book.save(path)
    book.close()
    return path


def write_acceptance_workbooks(manifest: AcceptanceManifest, output_dir: str | Path) -> dict[str, Path]:
    """Write only synthetic schedule datasets beneath an acceptance output root."""

    output = Path(output_dir).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    forms_by_officer: dict[str, list[FormSpec]] = {}
    for form in manifest.forms:
        forms_by_officer.setdefault(form.officer_id, []).append(form)

    mapping_path = output / 'class-mapping.xlsx'
    mapping_book = Workbook()
    mapping_sheet = mapping_book.active
    mapping_sheet.title = 'class_mapping'
    mapping_sheet.append(['listener_number', 'student_id', 'admin_class'])
    for officer in manifest.officers:
        if officer.coverage != 'basic':
            continue
        mapping_sheet.append([
            officer.officer_id,
            officer.student_id,
            _listener_admin_class(officer.officer_id),
        ])
    mapping_book.save(mapping_path)
    mapping_book.close()

    personal_path = output / 'personal-schedule.xlsx'
    personal_book = Workbook()
    personal_sheet = personal_book.active
    personal_sheet.title = 'personal_schedule'
    personal_sheet.append([
        'semester', 'listener_number', 'student_id', 'course_title',
        'weeks', 'weekday', 'periods',
    ])
    seen_officers: set[str] = set()
    for officer in manifest.officers:
        if officer.coverage != 'complete' or officer.officer_id in seen_officers:
            continue
        seen_officers.add(officer.officer_id)
        forms = forms_by_officer.get(officer.officer_id, [])
        if not forms:
            raise ValueError(f'complete officer has no form: {officer.officer_id}')
        exact_conflict_form = next(
            (form for form in forms if 'personal_schedule_exact_conflict' in form.oracle_markers),
            None,
        )
        if exact_conflict_form is not None:
            week = exact_conflict_form.teaching_week
            weekday = exact_conflict_form.weekday
            start_period = exact_conflict_form.start_period
            end_period = exact_conflict_form.end_period
            course_title = exact_conflict_form.course_title
        else:
            # Complete coverage comes from this personal dataset.  Keep its
            # slot separate from every form unless the oracle explicitly asks
            # for the exact personal-schedule conflict.
            week, weekday, start_period, end_period = _first_non_overlapping_slot(forms)
            course_title = f'楠屾敹涓汉璇剧▼{officer.officer_id}'
        personal_sheet.append([
            manifest.config.semester,
            officer.officer_id,
            officer.student_id,
            course_title,
            str(week),
            weekday,
            f'{start_period}-{end_period}',
        ])
    personal_book.save(personal_path)
    personal_book.close()
    return {'class_mapping': mapping_path, 'personal': personal_path}


__all__ = [
    'ANOMALY_MARKERS', 'AcceptanceManifest', 'AcceptanceScheduleEntry', 'FormSpec',
    'OfficerSpec', 'ScheduleStructure',
    'assert_real_identity_patterns_absent', 'generate_manifest',
    'build_acceptance_school_schedule', 'write_acceptance_school_schedule',
    'write_acceptance_workbooks',
]
